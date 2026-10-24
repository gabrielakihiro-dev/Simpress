"""Equipment export: fill the fixed export model with the active inventory's aliases and IdCliente.

The export model (``modelos/EquipamentosDetalhes-exportacao.xlsx``) defines the
format and rows; the values come from the active EquipamentosDetalhes, the
Endpoint hostnames and the manual aliases. ``count_export_changes`` keeps the
last count per file signature and records whether exporting is currently
possible (``get_export_status``).

``main.py`` wires the file locations and readers in ``EquipmentExportServices``
with late-bound callables, so tests that patch ``main`` keep affecting the
export. Errors are domain errors; ``external_errors`` lists the exception types
raised by those callbacks that must pass through unchanged.
"""

import json
import logging
import threading
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, ContextManager, Dict, List, Optional, Tuple, Type
from urllib.parse import quote

import openpyxl

from core.errors import DomainConflictError, DomainError
from core.integrity import (
    FileBusyError,
    FileTransactionError,
    WorkbookSafetyError,
    load_workbook_safely,
    serialize_workbook_validated,
    workbook_snapshot,
)
from core.xlsx_safety import WORKBOOK_READ_ERRORS
from domains.effective_alias import endpoint_alias_review_items, merge_export_aliases
from domains.inventory_comparison import normalize_reference_value
from domains.uploaded_workbooks import (
    DETAILS_ALIAS_HEADER_ALIASES,
    DETAILS_SERIAL_HEADER_ALIASES,
    find_workbook_header_column,
)


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class EquipmentExportServices:
    export_model_file: Callable[[], Path]
    endpoint_upload_file: Callable[[], Path]
    modified_aliases_file: Callable[[], Path]
    transaction_root: Callable[[], Path]
    read_lock_timeout: Callable[[], float]
    get_parque_file: Callable[[], Optional[Path]]
    get_details_files: Callable[[], List[Path]]
    find_details_sheet: Callable[[Path], Optional[str]]
    read_endpoint_alias_index: Callable[[], Dict[str, Any]]
    load_modified_apelidos: Callable[[], Dict[str, str]]
    latest_comparison_update_entry: Callable[[], Optional[Dict[str, Any]]]
    validate_xlsx: Callable[[Path], Any]
    lock_paths: Callable[..., ContextManager[Any]]
    logger: logging.Logger
    external_errors: Tuple[Type[BaseException], ...] = ()
    # Hostnames the user discarded from the Endpoint duplicates; they decide which one is exported.
    removed_hostnames_file: Optional[Callable[[], Path]] = None


def _removed_hostnames_paths(services: EquipmentExportServices) -> List[Path]:
    return [services.removed_hostnames_file()] if services.removed_hostnames_file is not None else []


@dataclass(frozen=True)
class EquipmentExport:
    """The exported workbook and the response headers that describe it."""

    content: bytes
    headers: Dict[str, str]


def load_export_model_aliases(services: EquipmentExportServices) -> Dict[str, str]:
    """Lê os apelidos do modelo fixo usado como base de exportação."""
    if not services.export_model_file().is_file():
        raise DomainError(f"O modelo de exportação '{services.export_model_file().name}' não foi encontrado.", status_code=404)

    workbook = None
    try:
        services.validate_xlsx(services.export_model_file())
        workbook = openpyxl.load_workbook(services.export_model_file(), read_only=True, data_only=True)
        for worksheet in workbook.worksheets:
            header_row = [cell.value for cell in next(worksheet.iter_rows(max_row=1), ())]
            serial_col = find_workbook_header_column(header_row, DETAILS_SERIAL_HEADER_ALIASES)
            alias_col = find_workbook_header_column(header_row, ("APELIDO",))
            if serial_col is None or alias_col is None:
                continue

            aliases_by_serial: Dict[str, List[str]] = {}
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                serial = row[serial_col] if serial_col < len(row) else None
                if serial is None or not str(serial).strip():
                    continue
                alias = row[alias_col] if alias_col < len(row) else None
                serial_key = normalize_reference_value(serial)
                if serial_key:
                    aliases_by_serial.setdefault(serial_key, []).append("" if alias is None else str(alias).strip())
            return {
                serial: values[0]
                for serial, values in aliases_by_serial.items()
                if len(values) == 1
            }

        raise DomainError(f"O modelo '{services.export_model_file().name}' não contém as colunas NumeroSerie e Apelido.")
    except DomainError:
        raise
    except WORKBOOK_READ_ERRORS as error:
        services.logger.warning(
            "Falha ao ler modelo (arquivo=%s, papel=exportacao, erro=%s).",
            services.export_model_file().name, type(error).__name__,
        )
        raise DomainError("Não foi possível ler o modelo de exportação.") from None
    finally:
        if workbook is not None:
            workbook.close()


def get_export_aliases(
    services: EquipmentExportServices,
    endpoint_index: Optional[Dict[str, Any]] = None,
) -> Dict[str, str]:
    """Combina hostnames da InvAll com apelidos manuais, que têm prioridade."""
    endpoint_aliases = (
        endpoint_index.get("unique", {})
        if endpoint_index is not None
        else services.read_endpoint_alias_index()["unique"]
    )
    return merge_export_aliases(endpoint_aliases, services.load_modified_apelidos())


def _normalized_record(record: Dict[str, Any]) -> Dict[str, Any]:
    return {
        normalize_reference_value(header): value
        for header, value in record.items()
        if not str(header).startswith("__") and normalize_reference_value(header)
    }


def _record_value(record: Dict[str, Any], aliases: Tuple[str, ...]) -> Any:
    values = _normalized_record(record)
    return next((values[key] for key in map(normalize_reference_value, aliases) if key in values), None)


def _same_workbook_value(left: Any, right: Any) -> bool:
    if isinstance(left, str) and isinstance(right, str):
        return left.strip().casefold() == right.strip().casefold()
    return left == right


def _model_row_matches_record(headers: List[Any], cells: Any, record: Dict[str, Any]) -> bool:
    """Confirm a same-position replacement by comparing every shared, stable field."""
    source = _normalized_record(record)
    mutable_headers = {
        *(normalize_reference_value(value) for value in DETAILS_SERIAL_HEADER_ALIASES),
        *(normalize_reference_value(value) for value in DETAILS_ALIAS_HEADER_ALIASES),
        normalize_reference_value("IDCLIENTE"),
    }
    compared = 0
    for index, header in enumerate(headers):
        key = normalize_reference_value(header)
        if not key or key in mutable_headers or key not in source:
            continue
        current_value = source[key]
        if current_value is None or current_value == "":
            continue
        cell = cells[index] if index < len(cells) else None
        if cell is None or cell.data_type == "f" or not _same_workbook_value(cell.value, current_value):
            return False
        compared += 1
    return compared >= 2


def _model_row_matches_source(headers: List[Any], cells: Any, source: Dict[str, Any]) -> bool:
    record = {header: value for header, value in source.get("values", {}).items()}
    return _model_row_matches_record(headers, cells, record)


def _current_inventory_rows(
    worksheet: Any,
    headers: List[Any],
    serial_column: int,
    alias_column: int,
    id_column: Optional[int],
) -> Tuple[Dict[str, List[Dict[str, Any]]], Dict[int, Dict[str, Any]]]:
    by_serial: Dict[str, List[Dict[str, Any]]] = {}
    by_row: Dict[int, Dict[str, Any]] = {}
    for row_number, cells in enumerate(worksheet.iter_rows(min_row=2), start=2):
        serial = cells[serial_column].value if serial_column < len(cells) else None
        serial_key = normalize_reference_value(serial)
        if not serial_key:
            continue
        row = {
            "serial": serial,
            "serial_key": serial_key,
            "apelido": cells[alias_column].value if alias_column < len(cells) else None,
            "id_cliente": cells[id_column].value if id_column is not None and id_column < len(cells) else None,
            "values": {
                normalize_reference_value(header): cells[index].value
                for index, header in enumerate(headers)
                if index < len(cells) and normalize_reference_value(header)
            },
        }
        by_row[row_number] = row
        by_serial.setdefault(serial_key, []).append(row)
    return by_serial, by_row


def _match_current_row(
    headers: List[Any],
    cells: Any,
    serial_key: str,
    row_number: int,
    current_by_serial: Dict[str, List[Dict[str, Any]]],
    current_by_row: Dict[int, Dict[str, Any]],
) -> Optional[Dict[str, Any]]:
    candidates = current_by_serial.get(serial_key, [])
    if len(candidates) == 1:
        return candidates[0]
    if candidates:
        return None
    row_candidate = current_by_row.get(row_number)
    if (
        row_candidate is not None
        and len(current_by_serial.get(row_candidate["serial_key"], [])) == 1
        and _model_row_matches_source(headers, cells, row_candidate)
    ):
        return row_candidate
    return None


def prepare_export_model_alias_sync(
    services: EquipmentExportServices,
    details_records: List[Dict[str, Any]],
    endpoint_index: Dict[str, Any],
) -> Optional[bytes]:
    """Update only serials and aliases in the imported model, preserving its layout."""
    model_file = services.export_model_file()
    if not model_file.is_file():
        return None

    workbook = load_workbook_safely(model_file, allow_threaded_comments=True)
    try:
        sheet_name = next(
            (name for name in workbook.sheetnames if "parque" in name.lower() or "equipamento" in name.lower()),
            workbook.sheetnames[0],
        )
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in next(worksheet.iter_rows(max_row=1), ())]
        serial_column = find_workbook_header_column(headers, DETAILS_SERIAL_HEADER_ALIASES)
        alias_column = find_workbook_header_column(headers, ("APELIDO",))
        if serial_column is None or alias_column is None:
            return None

        records_by_serial: Dict[str, List[Dict[str, Any]]] = {}
        records_by_row: Dict[int, Dict[str, Any]] = {}
        for record in details_records:
            serial = _record_value(record, DETAILS_SERIAL_HEADER_ALIASES)
            serial_key = normalize_reference_value(serial)
            if not serial_key:
                continue
            records_by_serial.setdefault(serial_key, []).append(record)
            row_number = record.get("__excel_row_number__")
            if isinstance(row_number, int) and not isinstance(row_number, bool):
                records_by_row[row_number] = record

        aliases = get_export_aliases(services, endpoint_index)
        expected_values: Dict[tuple, Any] = {}
        for row_number in range(2, worksheet.max_row + 1):
            serial_cell = worksheet.cell(row_number, serial_column + 1)
            alias_cell = worksheet.cell(row_number, alias_column + 1)
            if serial_cell.data_type == "f" or alias_cell.data_type == "f":
                continue
            serial_key = normalize_reference_value(serial_cell.value)
            candidates = records_by_serial.get(serial_key, [])
            matched_record: Optional[Dict[str, Any]] = candidates[0] if len(candidates) == 1 else None
            if not candidates:
                row_record = records_by_row.get(row_number)
                if row_record is not None:
                    row_serial = normalize_reference_value(
                        _record_value(row_record, DETAILS_SERIAL_HEADER_ALIASES)
                    )
                    if len(records_by_serial.get(row_serial, [])) == 1 and _model_row_matches_record(
                        headers,
                        tuple(worksheet.cell(row_number, column + 1) for column in range(len(headers))),
                        row_record,
                    ):
                        matched_record = row_record
            if matched_record is None:
                continue

            current_serial = _record_value(matched_record, DETAILS_SERIAL_HEADER_ALIASES)
            current_serial_key = normalize_reference_value(current_serial)
            current_alias = _record_value(matched_record, DETAILS_ALIAS_HEADER_ALIASES)
            updated_alias = aliases.get(current_serial_key, current_alias)
            if serial_key != current_serial_key:
                expected_values[(worksheet.title, serial_cell.coordinate)] = current_serial
            if updated_alias is None:
                updated_alias = ""
            if alias_cell.value != updated_alias:
                expected_values[(worksheet.title, alias_cell.coordinate)] = updated_alias

        if not expected_values:
            return None
        snapshot = workbook_snapshot(workbook)
        return serialize_workbook_validated(
            model_file,
            workbook,
            snapshot,
            expected_values,
            allow_threaded_comments=True,
        )
    finally:
        workbook.close()


@dataclass(frozen=True)
class ExportStatus:
    """Whether the export can run, with the reason shown when it cannot."""

    available: bool
    reason: str
    retryable: bool = False

    def as_dict(self) -> Dict[str, Any]:
        return {"available": self.available, "reason": self.reason}

    def readiness_dict(self) -> Dict[str, Any]:
        return {
            **self.as_dict(),
            "state": "ready" if self.available else "unavailable" if self.retryable else "invalid",
        }


@dataclass(frozen=True)
class CachedExportCounts:
    """Counts and status computed for one signature of the export files."""

    counts: Dict[str, int]
    status: ExportStatus


_EXPORT_STATUS_LOCK = threading.Lock()
_EXPORT_COUNTS_LOCK = threading.Lock()
_export_status = ExportStatus(available=False, reason="A exportação ainda não foi conferida.")
_EXPORT_COUNTS_CACHE: Dict[Any, CachedExportCounts] = {}


def _set_export_status(available: bool, reason: str) -> None:
    global _export_status
    with _EXPORT_STATUS_LOCK:
        _export_status = ExportStatus(available=available, reason=reason)


def _current_export_status() -> ExportStatus:
    with _EXPORT_STATUS_LOCK:
        return _export_status


def get_export_status() -> Dict[str, Any]:
    return _current_export_status().as_dict()


def export_files_signature(services: EquipmentExportServices) -> tuple:
    signature: List[Any] = []
    try:
        active_file = services.get_parque_file()
    except FileBusyError:
        active_file = None
    for path in (
        active_file,
        services.export_model_file(),
        services.endpoint_upload_file(),
        services.modified_aliases_file(),
        *_removed_hostnames_paths(services),
    ):
        if path is None:
            signature.append(None)
            continue
        try:
            stat = path.stat()
            signature.append((str(path.resolve(strict=False)), stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns))
        except OSError:
            signature.append((str(path.resolve(strict=False)), None))
    return tuple(signature)


def count_export_changes_with_status(
    services: EquipmentExportServices,
    locked: bool = False,
    endpoint_index: Optional[Dict[str, Any]] = None,
) -> Tuple[Dict[str, int], Dict[str, Any]]:
    """Count export changes and return the status from the same check."""
    empty_counts = {"apelidos": 0, "id_clientes": 0, "total": 0}

    def count() -> Tuple[Dict[str, int], ExportStatus, bool]:
        source_file = services.get_parque_file()
        if source_file is None or not source_file.is_file() or not services.export_model_file().is_file():
            return empty_counts.copy(), ExportStatus(False, "A base ativa ou o modelo de exportação está ausente."), True

        model_wb = None
        current_wb = None
        try:
            services.validate_xlsx(services.export_model_file())
            model_wb = openpyxl.load_workbook(services.export_model_file(), read_only=True, data_only=False)
            model_sheet_name = next(
                (name for name in model_wb.sheetnames if "parque" in name.lower() or "equipamento" in name.lower()),
                model_wb.sheetnames[0],
            )
            model_ws = model_wb[model_sheet_name]
            model_headers = [cell.value for cell in next(model_ws.iter_rows(max_row=1), ())]
            model_serial_col = find_workbook_header_column(model_headers, DETAILS_SERIAL_HEADER_ALIASES)
            model_alias_col = find_workbook_header_column(model_headers, ("APELIDO",))
            model_id_col = find_workbook_header_column(model_headers, ("IDCLIENTE",))
            if model_serial_col is None or model_alias_col is None:
                return empty_counts.copy(), ExportStatus(False, "O modelo de exportação não contém as colunas obrigatórias."), True

            current_sheet_name = services.find_details_sheet(source_file)
            if current_sheet_name is None:
                return empty_counts.copy(), ExportStatus(False, "A base ativa não contém uma aba de inventário compatível."), True
            current_wb = openpyxl.load_workbook(source_file, read_only=True, data_only=False)
            current_ws = current_wb[current_sheet_name]
            current_headers = [cell.value for cell in next(current_ws.iter_rows(max_row=1), ())]
            current_serial_col = find_workbook_header_column(current_headers, DETAILS_SERIAL_HEADER_ALIASES)
            current_alias_col = find_workbook_header_column(current_headers, DETAILS_ALIAS_HEADER_ALIASES)
            current_id_col = find_workbook_header_column(current_headers, ("IDCLIENTE",))
            if current_serial_col is None or current_alias_col is None:
                return empty_counts.copy(), ExportStatus(False, "A base ativa não contém as colunas de série e apelido."), True

            current_by_serial, current_by_row = _current_inventory_rows(
                current_ws,
                current_headers,
                current_serial_col,
                current_alias_col,
                current_id_col,
            )
            if any(len(rows) > 1 for rows in current_by_serial.values()):
                return empty_counts.copy(), ExportStatus(False, "A base ativa contém séries duplicadas; a exportação exige revisão."), True

            model_rows = []
            model_serial_counts: Dict[str, int] = {}
            for row_number, cells in enumerate(model_ws.iter_rows(min_row=2), start=2):
                serial_key = normalize_reference_value(
                    cells[model_serial_col].value if model_serial_col < len(cells) else None
                )
                if not serial_key:
                    continue
                model_serial_counts[serial_key] = model_serial_counts.get(serial_key, 0) + 1
                model_rows.append((serial_key, row_number, cells))

            effective_endpoint_index = endpoint_index if endpoint_index is not None else services.read_endpoint_alias_index()
            aliases = get_export_aliases(services, effective_endpoint_index)
            alias_count = 0
            id_cliente_count = 0
            for serial_key, row_number, cells in model_rows:
                if model_serial_counts.get(serial_key) != 1:
                    continue
                current = _match_current_row(
                    model_headers,
                    cells,
                    serial_key,
                    row_number,
                    current_by_serial,
                    current_by_row,
                )
                if current is None:
                    continue
                alias_cell = cells[model_alias_col] if model_alias_col < len(cells) else None
                if alias_cell is not None and alias_cell.data_type != "f":
                    updated_alias = aliases.get(current["serial_key"], current["apelido"])
                    if updated_alias is None:
                        updated_alias = ""
                    if alias_cell.value != updated_alias:
                        alias_count += 1

                if model_id_col is not None and current_id_col is not None:
                    id_cell = cells[model_id_col] if model_id_col < len(cells) else None
                    if id_cell is not None and id_cell.data_type != "f":
                        updated_id = current["id_cliente"]
                        if updated_id is None:
                            updated_id = ""
                        if id_cell.value != updated_id:
                            id_cliente_count += 1

            return {
                "apelidos": alias_count,
                "id_clientes": id_cliente_count,
                "total": alias_count + id_cliente_count,
            }, ExportStatus(True, "OK"), True
        except Exception as error:
            services.logger.warning("Contagem de exportação indisponível (%s).", type(error).__name__)
            return empty_counts.copy(), ExportStatus(
                False,
                "Não foi possível conferir o modelo e a base de exportação.",
                retryable=True,
            ), False
        finally:
            if model_wb is not None:
                model_wb.close()
            if current_wb is not None:
                current_wb.close()

    def count_with_cache() -> Tuple[Dict[str, int], ExportStatus]:
        if endpoint_index is not None:
            values, status, _ = count()
            _set_export_status(status.available, status.reason)
            return values, status
        key = export_files_signature(services)
        with _EXPORT_COUNTS_LOCK:
            cached = _EXPORT_COUNTS_CACHE.get(key)
            if cached is not None:
                _set_export_status(cached.status.available, cached.status.reason)
                return {name: cached.counts[name] for name in empty_counts}, cached.status
        values, status, cacheable = count()
        _set_export_status(status.available, status.reason)
        if cacheable:
            with _EXPORT_COUNTS_LOCK:
                _EXPORT_COUNTS_CACHE.clear()
                _EXPORT_COUNTS_CACHE[key] = CachedExportCounts(counts=dict(values), status=status)
        return values, status

    if locked:
        values, status = count_with_cache()
    else:
        with services.lock_paths([services.transaction_root() / ".atomic-transactions"], timeout=services.read_lock_timeout()):
            detail_files = services.get_details_files()
            with services.lock_paths([
                *detail_files,
                services.endpoint_upload_file(),
                services.modified_aliases_file(),
                *_removed_hostnames_paths(services),
                services.export_model_file(),
            ], timeout=services.read_lock_timeout()):
                values, status = count_with_cache()
    return values, status.readiness_dict()


def count_export_changes(
    services: EquipmentExportServices,
    locked: bool = False,
    endpoint_index: Optional[Dict[str, Any]] = None,
) -> Dict[str, int]:
    """Count writable alias and IdCliente cells that export would change."""
    counts, _ = count_export_changes_with_status(services, locked, endpoint_index)
    return counts


def build_equipment_export(services: EquipmentExportServices) -> EquipmentExport:
    """Exporta a apresentação de referência com os registros da base ativa."""
    source_file = services.export_model_file()
    current_file = services.get_parque_file()
    if not source_file.is_file():
        raise DomainError(f"O modelo de exportação '{source_file.name}' não foi encontrado.", status_code=404)
    if current_file is None or not current_file.is_file():
        raise DomainError("A base atual da Simpress não foi encontrada.", status_code=404)

    wb = None
    current_wb = None
    passthrough_errors: Tuple[Type[BaseException], ...] = (DomainError, *services.external_errors)
    try:
        with services.lock_paths([
            source_file,
            current_file,
            services.endpoint_upload_file(),
            services.modified_aliases_file(),
            *_removed_hostnames_paths(services),
        ]):
            services.validate_xlsx(source_file)
            wb = load_workbook_safely(source_file, allow_threaded_comments=True)
            original = workbook_snapshot(wb)
            sheet_target = next(
                (name for name in wb.sheetnames if "parque" in name.lower() or "equipamento" in name.lower()),
                wb.sheetnames[0],
            )
            ws = wb[sheet_target]
            header_row = [cell.value for cell in next(ws.iter_rows(max_row=1), ())]
            serial_col_index = find_workbook_header_column(header_row, DETAILS_SERIAL_HEADER_ALIASES)
            apelido_col_index = find_workbook_header_column(header_row, ("APELIDO",))
            id_cliente_col_index = find_workbook_header_column(header_row, ("IDCLIENTE",))
            if serial_col_index is None or apelido_col_index is None:
                raise DomainError(f"A planilha '{source_file.name}' não contém as colunas NumeroSerie e Apelido esperadas.")
            serial_col = serial_col_index + 1
            apelido_col = apelido_col_index + 1
            id_cliente_col = id_cliente_col_index + 1 if id_cliente_col_index is not None else None
            header_style_updates = {}
            default_header_style_id = ws.cell(row=1, column=serial_col).style_id
            for column in (apelido_col, id_cliente_col):
                if column is None:
                    continue
                header_cell = ws.cell(row=1, column=column)
                if header_cell.style_id != default_header_style_id:
                    header_style_updates[(ws.title, header_cell.coordinate)] = default_header_style_id

            current_sheet_name = services.find_details_sheet(current_file)
            if current_sheet_name is None:
                raise DomainConflictError("A base atual da Simpress não tem os cabeçalhos esperados.")
            current_wb = openpyxl.load_workbook(current_file, read_only=True, data_only=False)
            current_ws = current_wb[current_sheet_name]
            current_headers = [cell.value for cell in next(current_ws.iter_rows(max_row=1), ())]
            current_serial_col_index = find_workbook_header_column(current_headers, DETAILS_SERIAL_HEADER_ALIASES)
            if current_serial_col_index is None:
                raise DomainConflictError("A base atual não contém NumeroSerie.")
            current_alias_col_index = find_workbook_header_column(current_headers, DETAILS_ALIAS_HEADER_ALIASES)
            if current_alias_col_index is None:
                raise DomainConflictError("A base atual não contém Apelido ou Hostname.")
            current_serial_col = current_serial_col_index + 1
            current_alias_col = current_alias_col_index + 1
            current_id_cliente_col_index = find_workbook_header_column(current_headers, ("IDCLIENTE",))
            current_id_cliente_col = current_id_cliente_col_index + 1 if current_id_cliente_col_index is not None else None
            current_by_serial, current_by_row = _current_inventory_rows(
                current_ws,
                current_headers,
                current_serial_col_index,
                current_alias_col_index,
                current_id_cliente_col_index,
            )
            if any(len(rows) > 1 for rows in current_by_serial.values()):
                raise DomainConflictError("A base atual contém séries duplicadas; confira antes de exportar.")

            endpoint_index = services.read_endpoint_alias_index()
            aliases = get_export_aliases(services, endpoint_index)
            serial_counts: Dict[str, int] = {}
            for row in range(2, ws.max_row + 1):
                serial_key = normalize_reference_value(ws.cell(row=row, column=serial_col).value)
                if serial_key:
                    serial_counts[serial_key] = serial_counts.get(serial_key, 0) + 1

            expected_values = {}
            alias_update_count = 0
            id_cliente_update_count = 0
            skipped_duplicate_rows = set()
            skipped_formula_rows = 0
            skipped_formula_serials = set()
            skipped_unmatched_rows = set()
            for row in range(2, ws.max_row + 1):
                serial_key = normalize_reference_value(ws.cell(row=row, column=serial_col).value)
                if not serial_key:
                    continue
                if serial_counts.get(serial_key) != 1:
                    skipped_duplicate_rows.add(serial_key)
                    continue

                current = _match_current_row(
                    headers=header_row,
                    cells=tuple(ws.cell(row=row, column=column + 1) for column in range(len(header_row))),
                    serial_key=serial_key,
                    row_number=row,
                    current_by_serial=current_by_serial,
                    current_by_row=current_by_row,
                )
                if current is None:
                    skipped_unmatched_rows.add(serial_key)
                    continue

                formula_skipped = False
                serial_cell = ws.cell(row=row, column=serial_col)
                if (
                    serial_cell.data_type != "f"
                    and serial_key != current["serial_key"]
                ):
                    serial_cell.value = current["serial"]
                    expected_values[(ws.title, serial_cell.coordinate)] = current["serial"]
                alias_cell = ws.cell(row=row, column=apelido_col)
                if alias_cell.data_type == "f":
                    formula_skipped = True
                else:
                    updated_alias = aliases.get(
                        current["serial_key"],
                        current["apelido"],
                    )
                    if updated_alias is None:
                        updated_alias = ""
                    if alias_cell.value != updated_alias:
                        alias_cell.value = updated_alias
                        expected_values[(ws.title, alias_cell.coordinate)] = updated_alias
                        alias_update_count += 1

                if id_cliente_col and current_id_cliente_col:
                    id_cliente_cell = ws.cell(row=row, column=id_cliente_col)
                    if id_cliente_cell.data_type == "f":
                        formula_skipped = True
                    else:
                        updated_id_cliente = current["id_cliente"]
                        if updated_id_cliente is None:
                            updated_id_cliente = ""
                        if id_cliente_cell.value != updated_id_cliente:
                            id_cliente_cell.value = updated_id_cliente
                            expected_values[(ws.title, id_cliente_cell.coordinate)] = updated_id_cliente
                            id_cliente_update_count += 1

                if formula_skipped:
                    skipped_formula_rows += 1
                    skipped_formula_serials.add(serial_key)

            review_items = endpoint_alias_review_items(endpoint_index, set(current_by_serial))
            pending_review_count = (
                len(review_items) + len(skipped_duplicate_rows) + skipped_formula_rows
                + len(skipped_unmatched_rows)
            )
            review_summary = [
                {
                    "motivo": "Registros do Endpoint que exigem conferência",
                    "total": len(review_items),
                    "series": [
                        str(item.get("serie") or "")[:64]
                        for item in review_items if item.get("serie")
                    ][:10],
                },
                {
                    "motivo": "Séries duplicadas no modelo de exportação",
                    "total": len(skipped_duplicate_rows),
                    "series": sorted(skipped_duplicate_rows)[:10],
                },
                {
                    "motivo": "Linhas protegidas por fórmulas",
                    "total": skipped_formula_rows,
                    "series": sorted(skipped_formula_serials)[:10],
                },
                {
                    "motivo": "Linhas sem correspondência segura na base atual",
                    "total": len(skipped_unmatched_rows),
                    "series": sorted(skipped_unmatched_rows)[:10],
                },
            ]
            review_summary = [item for item in review_summary if item["total"]]
            output_bytes = serialize_workbook_validated(
                source_file, wb, original, expected_values,
                expected_cell_styles=header_style_updates,
                allow_threaded_comments=True,
            )

        now_str = datetime.now().strftime("%d-%m-%Y_%H%M%S")
        export_filename = f"EquipamentosDetalhes-{now_str}.xlsx"
        latest_history = None
        try:
            latest_history = services.latest_comparison_update_entry()
        except services.external_errors:
            pass
        history_alias_count = 0
        history_timestamp = ""
        if isinstance(latest_history, dict):
            history_comparison = latest_history.get("comparacao")
            history_comparison = history_comparison if isinstance(history_comparison, dict) else {}
            history_alias_changes = history_comparison.get("apelidos_atualizados")
            if isinstance(history_alias_changes, list):
                history_alias_count = len(history_alias_changes)
            else:
                history_summary = latest_history.get("resumo")
                if isinstance(history_summary, dict):
                    try:
                        history_alias_count = max(0, int(history_summary.get("apelidos_atualizados", 0)))
                    except (TypeError, ValueError):
                        history_alias_count = 0
            history_timestamp = str(
                latest_history.get("horario")
                or history_comparison.get("horario_comparacao")
                or ""
            )

        headers = {
            "Content-Disposition": f'attachment; filename="{export_filename}"',
            "Access-Control-Expose-Headers": (
                "Content-Disposition, X-Total-Alterados, X-Total-IdCliente-Alterados, "
                "X-Total-Alteracoes-Exportadas, X-Total-Conferencia, X-Conferencia-Resumo, "
                "X-Historico-Apelidos-Atualizados, X-Historico-Horario"
            ),
            "X-Total-Alterados": str(alias_update_count),
            "X-Total-IdCliente-Alterados": str(id_cliente_update_count),
            "X-Total-Alteracoes-Exportadas": str(alias_update_count + id_cliente_update_count),
            "X-Total-Conferencia": str(pending_review_count),
            "X-Conferencia-Resumo": quote(json.dumps(review_summary, ensure_ascii=False, separators=(",", ":"))),
            "X-Historico-Apelidos-Atualizados": str(history_alias_count),
            "X-Historico-Horario": history_timestamp,
        }

        return EquipmentExport(content=output_bytes, headers=headers)
    except passthrough_errors:
        raise
    except FileBusyError as error:
        raise DomainConflictError(str(error)) from error
    except WorkbookSafetyError as error:
        raise DomainConflictError(f"Exportação cancelada para preservar a planilha: {error}") from error
    except FileTransactionError as error:
        raise DomainError(f"Não foi possível validar a exportação: {error}") from error
    except Exception:
        services.logger.exception("Falha ao exportar a base local de equipamentos.")
        raise DomainError("Erro ao exportar a base de equipamentos.") from None
    finally:
        if wb is not None:
            wb.close()
        if current_wb is not None:
            current_wb.close()
