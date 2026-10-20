"""Excel presentation helpers for inventory comparison history."""

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple
from core.xlsx_safety import set_cell_value_as_text
from domains.history_annotations import annotation_export_text
from domains.history_restore import comparison_event_type
from domains.inventory_comparison import _comparison_source_text_for_display

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet


_EXPORT_NAVY = "10213A"
_EXPORT_NAVY_DARK = "0B1729"
_EXPORT_CYAN = "00A3E0"
_EXPORT_TEXT = "26364A"
_EXPORT_MUTED = "63748A"
_EXPORT_PALE = "F2F6FA"
_EXPORT_BORDER = "D9E3ED"


def _export_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        value = "; ".join(f"{key}: {item}" for key, item in value.items())
    elif isinstance(value, (list, tuple, set)):
        value = "; ".join(str(item) for item in value if item is not None)
    return str(value)[:32767]


def _history_display_files(value: Any) -> Dict[str, str]:
    if not isinstance(value, dict):
        return {}
    return {
        key: _comparison_source_text_for_display(_export_text(name))
        for key, name in value.items() if isinstance(key, str)
    }


def _history_section_text(value: Any, *, bare_sources: bool = False) -> str:
    text = _comparison_source_text_for_display(_export_text(value))
    if bare_sources:
        text = re.sub(r"\bno endpoint\b", "Endpoint", text, flags=re.IGNORECASE)
        return re.sub(r"\bna Simpress\b", "Simpress", text, flags=re.IGNORECASE)
    return re.sub(r"\bno endpoint\b", "no Endpoint", text, flags=re.IGNORECASE)


_USER_CHANGE_LABELS = {
    "novo_endpoint": "Usuário novo no Endpoint",
    "alterado_endpoint": "Usuário alterado no Endpoint",
    "removido_endpoint": "Usuário removido do Endpoint",
    "manual_voltou": "Manual voltou a valer",
}


def _export_user(user: Any, origin: Any) -> str:
    text = _export_text(user).strip()
    return f"{text} (M)" if text and origin == "manual" else text


def _export_datetime(value: Any) -> Optional[datetime]:
    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.fromisoformat(str(value or "").replace("Z", "+00:00"))
        except (TypeError, ValueError, OverflowError):
            return None
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone().replace(tzinfo=None)
    return parsed


def _export_event_timestamp(event: Dict[str, Any]) -> float:
    parsed = _export_datetime(event.get("horario"))
    return parsed.timestamp() if parsed is not None else float("-inf")


def _export_rows(comparison: Dict[str, Any], key: str) -> list:
    rows = comparison.get(key)
    return [item for item in rows if isinstance(item, dict)] if isinstance(rows, list) else []


def _export_note(
    comparison: Dict[str, Any],
    category: str,
    index: int,
    item: Optional[Dict[str, Any]] = None,
) -> str:
    if isinstance(item, dict):
        if item.get("anotacao_conflito"):
            return _export_text(annotation_export_text(item))
        if isinstance(item.get("anotacao"), str):
            return _export_text(item["anotacao"])
        if item.get("anotacao_id"):
            return ""
    annotations = comparison.get("anotacoes")
    if not isinstance(annotations, dict):
        return ""
    return _export_text(annotations.get(f"{category}:{index}"))


def _style_export_title(ws: Worksheet, title: str, subtitle: str, last_column: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.fill = PatternFill("solid", fgColor=_EXPORT_NAVY_DARK)
    title_cell.font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    subtitle_cell = ws.cell(row=2, column=1, value=subtitle)
    subtitle_cell.fill = PatternFill("solid", fgColor=_EXPORT_NAVY)
    subtitle_cell.font = Font(name="Aptos", size=10, color="D7E5F3")
    subtitle_cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 10


def _safe_export_cell(cell: Cell, value: Any) -> None:
    if isinstance(value, str):
        set_cell_value_as_text(cell, value[:32767])
        cell.number_format = "@"
    else:
        cell.value = value
        if isinstance(value, datetime):
            cell.number_format = "dd/mm/yyyy hh:mm"
        elif isinstance(value, int) and not isinstance(value, bool):
            cell.number_format = "#,##0"
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _estimated_row_height(values: Sequence[Any], widths: Sequence[Optional[float]]) -> float:
    lines = 1
    for value, width in zip(values, widths):
        if value is None:
            continue
        effective_width = max(10, min(float(width or 20), 60))
        for line in str(value).splitlines() or [""]:
            lines = max(lines, (len(line) // int(effective_width * 1.35)) + 1)
    return min(90, max(25, lines * 15))


def _write_export_table_sheet(
    workbook: openpyxl.Workbook,
    sheet_name: str,
    title: str,
    subtitle: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    table_name: str,
    widths: Sequence[float],
) -> None:
    ws = workbook.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = _EXPORT_CYAN
    _style_export_title(ws, title, subtitle, len(headers))

    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor=_EXPORT_NAVY)
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=_EXPORT_CYAN))
        ws.column_dimensions[get_column_letter(column)].width = widths[column - 1]
    ws.row_dimensions[4].height = 34
    ws.freeze_panes = "A5"

    for row_index, values in enumerate(rows, start=5):
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=column)
            _safe_export_cell(cell, value)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=_EXPORT_PALE)
            cell.font = Font(name="Aptos", size=10, color=_EXPORT_TEXT)
            cell.border = Border(bottom=Side(style="hair", color=_EXPORT_BORDER))
        ws.row_dimensions[row_index].height = _estimated_row_height(values, widths)

    if rows:
        table = Table(
            displayName=table_name,
            ref=f"A4:{get_column_letter(len(headers))}{4 + len(rows)}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    else:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}4"

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"


def _write_export_summary_sheet(
    workbook: openpyxl.Workbook,
    generated_at: datetime,
    selected_date: date,
    events: Sequence[Dict[str, Any]],
    event_counts: Mapping[str, int],
    category_counts: Mapping[str, int],
    detail_total: int,
) -> None:
    ws = workbook.active
    ws.title = "Resumo"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 95
    ws.sheet_properties.tabColor = _EXPORT_CYAN
    _style_export_title(
        ws,
        "Histórico de planilhas",
        f"Data selecionada: {selected_date:%d/%m/%Y} · Exportado em {generated_at:%d/%m/%Y às %H:%M} · {len(events):,} eventos".replace(",", "."),
        8,
    )
    for column in range(1, 9):
        ws.column_dimensions[get_column_letter(column)].width = 18
    for start_column, label, value in (
        (1, "EVENTOS", len(events)),
        (3, "ATUALIZAÇÕES", event_counts["atualizacao"]),
        (5, "RESTAURAÇÕES", event_counts["restauracao"]),
        (7, "LINHAS DE DETALHE", detail_total),
    ):
        end_column = start_column + 1
        ws.merge_cells(start_row=4, start_column=start_column, end_row=4, end_column=end_column)
        ws.merge_cells(start_row=5, start_column=start_column, end_row=5, end_column=end_column)
        label_cell = ws.cell(row=4, column=start_column, value=label)
        label_cell.fill = PatternFill("solid", fgColor="E7F3FA")
        label_cell.font = Font(name="Aptos", size=9, bold=True, color=_EXPORT_MUTED)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell = ws.cell(row=5, column=start_column, value=value)
        value_cell.fill = PatternFill("solid", fgColor="F4F8FC")
        value_cell.font = Font(name="Aptos Display", size=22, bold=True, color=_EXPORT_NAVY)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in (4, 5):
            for column in range(start_column, end_column + 1):
                ws.cell(row=row, column=column).border = Border(
                    left=Side(style="thin", color=_EXPORT_BORDER),
                    right=Side(style="thin", color=_EXPORT_BORDER),
                    top=Side(style="thin", color=_EXPORT_BORDER),
                    bottom=Side(style="thin", color=_EXPORT_BORDER),
                )
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 42

    ws.merge_cells("A7:B7")
    section = ws["A7"]
    section.value = "Registros por categoria"
    section.fill = PatternFill("solid", fgColor=_EXPORT_NAVY)
    section.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
    section.alignment = Alignment(vertical="center")
    ws.row_dimensions[7].height = 25
    for column, header in enumerate(("Categoria", "Linhas"), start=1):
        cell = ws.cell(row=8, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor="DDEAF4")
        cell.font = Font(name="Aptos", size=10, bold=True, color=_EXPORT_NAVY)
        cell.alignment = Alignment(vertical="center")
    for row_index, (label, count) in enumerate(category_counts.items(), start=9):
        for column, cell_value in enumerate((label, count), start=1):
            cell = ws.cell(row=row_index, column=column, value=cell_value)
            cell.font = Font(name="Aptos", size=10, color=_EXPORT_TEXT)
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=Side(style="hair", color=_EXPORT_BORDER))
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=_EXPORT_PALE)
        ws.cell(row=row_index, column=2).number_format = "#,##0"
        ws.row_dimensions[row_index].height = 23

    note_row = 10 + len(category_counts)
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    note = ws.cell(row=note_row, column=1)
    note.value = "As abas seguintes mostram os eventos e os registros detalhados, incluindo as anotações salvas em cada comparação."
    note.font = Font(name="Aptos", size=9, italic=True, color=_EXPORT_MUTED)
    note.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[note_row].height = 30


def build_history_export_workbook(
    display_events: List[Dict[str, Any]],
    selected_date: date,
) -> Tuple[BytesIO, str]:
    """Build the history workbook for one day; returns the rewound stream and its filename."""
    display_events = sorted(display_events, key=_export_event_timestamp)

    event_counts = {"atualizacao": 0, "restauracao": 0}
    category_counts = {
        "Apelidos atualizados": 0,
        "Apelidos para conferência": 0,
        "Divergências manuais": 0,
        "Sem verificação há 7+ dias": 0,
        "Ausentes no endpoint": 0,
        "Ausentes na Simpress": 0,
        "Divergências de localidade": 0,
        "Duplicidades e ambiguidades": 0,
        "Séries sem local definido": 0,
        "Mudanças de usuário": 0,
    }
    history_rows = []
    alias_rows = []
    absence_rows = []
    no_access_rows = []
    location_rows = []
    ambiguity_rows = []
    user_rows = []

    for order, event in enumerate(display_events, start=1):
        comparison = event.get("comparacao")
        comparison = comparison if isinstance(comparison, dict) else {}
        summary = event.get("resumo")
        summary = summary if isinstance(summary, dict) else {}
        event_type = comparison_event_type(event)
        event_type_label = "Restauração" if event_type == "restauracao" else "Atualização"
        event_counts[event_type] += 1
        timestamp = _export_datetime(event.get("horario") or comparison.get("horario_comparacao"))
        event_id = _export_text(event.get("id") or comparison.get("evento_id"))
        alias_updates = _export_rows(comparison, "apelidos_atualizados")
        alias_reviews = _export_rows(comparison, "apelidos_sem_atualizacao")
        manual_alias_conflicts = _export_rows(comparison, "divergencias_apelidos_manuais")
        absent_endpoint = _export_rows(comparison, "ausentes_no_endpoint")
        absent_simpress = _export_rows(comparison, "ausentes_em_detalhes")
        no_recent_verification = _export_rows(comparison, "sem_acesso_7_dias")
        local_divergences = _export_rows(comparison, "divergencias_local")
        unmapped_series = _export_rows(comparison, "series_sem_mapeamento")
        ambiguities = _export_rows(comparison, "correspondencias_ambiguas")
        ambiguous_records = _export_rows(comparison, "registros_ambiguos")

        manual_pending_count = sum(
            1 for item in manual_alias_conflicts
            if item.get("resolucao", "pendente") == "pendente"
        )
        ambiguity_count = len(ambiguities) + len(ambiguous_records)
        category_counts["Apelidos atualizados"] += len(alias_updates)
        category_counts["Apelidos para conferência"] += len(alias_reviews)
        category_counts["Divergências manuais"] += len(manual_alias_conflicts)
        category_counts["Sem verificação há 7+ dias"] += len(no_recent_verification)
        category_counts["Ausentes no endpoint"] += len(absent_endpoint)
        category_counts["Ausentes na Simpress"] += len(absent_simpress)
        category_counts["Divergências de localidade"] += len(local_divergences)
        category_counts["Duplicidades e ambiguidades"] += ambiguity_count
        category_counts["Séries sem local definido"] += len(unmapped_series)
        user_changes = _export_rows(comparison, "mudancas_usuarios")
        category_counts["Mudanças de usuário"] += len(user_changes)
        for index, item in enumerate(user_changes):
            user_rows.append((
                timestamp, event_type_label,
                _USER_CHANGE_LABELS.get(str(item.get("tipo") or ""), _export_text(item.get("tipo"))),
                item.get("serie"), item.get("computador"),
                _export_user(item.get("usuario_anterior"), item.get("origem_anterior")),
                _export_user(item.get("usuario_novo"), item.get("origem_nova")),
                _history_section_text(item.get("motivo")),
                _export_note(comparison, "mudancas_usuarios", index, item),
            ))

        files_uploaded = event.get("arquivos_enviados")
        if not isinstance(files_uploaded, dict):
            files_uploaded = comparison.get("arquivos_enviados")
        files_uploaded = files_uploaded if isinstance(files_uploaded, dict) else {}
        files_restored = comparison.get("arquivos_restaurados")
        files_restored = files_restored if isinstance(files_restored, dict) else {}
        endpoint_file = files_restored.get("endpoint") if event_type == "restauracao" else files_uploaded.get("endpoint")
        simpress_file = files_restored.get("detalhes") if event_type == "restauracao" else files_uploaded.get("detalhes")
        history_rows.append((
            order,
            timestamp,
            event_type_label,
            event_id,
            _comparison_source_text_for_display(_export_text(endpoint_file)),
            _comparison_source_text_for_display(_export_text(simpress_file)),
            comparison.get("total_endpoint", summary.get("total_endpoint")),
            comparison.get("total_detalhes", summary.get("total_detalhes")),
            comparison.get("total_equipamentos_carregados", summary.get("total_equipamentos_carregados")),
            len(alias_updates),
            len(alias_reviews),
            manual_pending_count,
            len(absent_endpoint),
            len(absent_simpress),
            len(no_recent_verification),
            len(local_divergences),
            ambiguity_count,
            len(unmapped_series),
            _export_datetime((comparison.get("restauracao_de") or {}).get("horario"))
            if isinstance(comparison.get("restauracao_de"), dict) else None,
            _export_text(comparison.get("backup_diretorio")),
        ))

        for item in alias_updates:
            alias_rows.append((
                timestamp, event_type_label, "Apelido atualizado",
                item.get("serie"), "", item.get("apelido_anterior"),
                item.get("apelido_equipamentos_detalhes"),
                item.get("apelido_novo") or item.get("apelido_invall"),
                _export_datetime(item.get("horario")), "Concluído", "",
                _comparison_source_text_for_display(_export_text(item.get("arquivos_origem"))), "",
            ))
        for index, item in enumerate(alias_reviews):
            alias_rows.append((
                timestamp, event_type_label, "Para conferência",
                item.get("serie"), item.get("computador"), "", "", "", None,
                "Pendente", _history_section_text(item.get("motivo") or "Conferir associação."), "",
                _export_note(comparison, "apelidos_sem_atualizacao", index, item),
            ))
        manual_status = {
            "manter": "Alteração manual mantida",
            "excluir": "Alteração manual removida",
            "substituida": "Atualizada por nova planilha",
            "pendente": "Pendente",
        }
        for item in manual_alias_conflicts:
            resolution = item.get("resolucao") or "pendente"
            alias_rows.append((
                timestamp, event_type_label, "Divergência manual",
                item.get("serie"), "", item.get("apelido_manual"),
                item.get("apelido_equipamentos_detalhes"), item.get("apelido_invall"),
                _export_datetime(item.get("horario_resolucao")),
                manual_status.get(resolution, _export_text(resolution)),
                "Diferença entre apelido manual e as planilhas de origem.", "", "",
            ))

        for category, items, note_category in (
            ("Ausente no endpoint", absent_endpoint, "ausentes_no_endpoint"),
            ("Ausente na Simpress", absent_simpress, "ausentes_em_detalhes"),
        ):
            missing_source = "no endpoint" if category == "Ausente no endpoint" else "na Simpress"
            for index, item in enumerate(items):
                absence_rows.append((
                    timestamp, category,
                    ", ".join(
                        _history_section_text(source, bare_sources=True)
                        for source in (
                            item["ausente_em"] if isinstance(item.get("ausente_em"), list)
                            else [item.get("ausente_em") or item.get("fonte_ausente") or missing_source]
                        ) if source
                    ),
                    item.get("computador"), item.get("serie"), item.get("local_previsto"),
                    _export_note(comparison, note_category, index, item),
                ))

        for index, item in enumerate(no_recent_verification):
            no_access_rows.append((
                timestamp, event_type_label,
                (
                    f"{item.get('computador')} (hostname sumiu do Endpoint)"
                    if item.get("hostname_sumiu_endpoint") else item.get("computador")
                ),
                item.get("serie"),
                _export_datetime(item.get("ultima_verificacao")), item.get("dias_sem_verificacao"),
                item.get("linha_endpoint"), _export_note(comparison, "sem_acesso_7_dias", index, item),
            ))

        for index, item in enumerate(local_divergences):
            location_rows.append((
                timestamp, "Localidade divergente", item.get("computador"),
                item.get("serie"), item.get("local_previsto"), item.get("local_detalhes"),
                None, "", _export_note(comparison, "divergencias_local", index, item),
            ))
        for index, item in enumerate(unmapped_series):
            location_rows.append((
                timestamp, "Série sem local definido", "", item.get("serie"),
                "", "", item.get("quantidade"), _export_text(item.get("exemplos")),
                _export_note(comparison, "series_sem_mapeamento", index, item),
            ))

        for index, item in enumerate(ambiguities):
            sources = []
            for source in item.get("fontes") or []:
                if not isinstance(source, dict):
                    continue
                label = _history_section_text(source.get("fonte"), bare_sources=True)
                if source.get("quantidade"):
                    label += f" ({source['quantidade']} registros)"
                elif source.get("serie"):
                    label += f" — série {source['serie']}"
                sources.append(label)
            note = _export_note(comparison, "ambiguidades", index, item)
            ambiguity_rows.append((
                timestamp, item.get("tipo") or "Ambiguidade",
                item.get("chave"), "; ".join(sources),
                _export_datetime(item.get("ultima_verificacao")),
                _history_section_text(item.get("mensagem"), bare_sources=True),
                note,
            ))
        for index, item in enumerate(ambiguous_records, start=len(ambiguities)):
            ambiguous_sources = item.get("fontes") or {}
            source_summary = "; ".join(
                _history_section_text(source, bare_sources=True) for source, status in ambiguous_sources.items()
                if status == "ambíguo"
            ) if isinstance(ambiguous_sources, dict) else ""
            record_key = f"{item.get('computador') or 'Sem hostname'} · {item.get('serie') or 'Sem série'}"
            notes = [_export_note(comparison, "ambiguidades", index, item)]
            notes.extend(
                _export_text(annotation_export_text(associated))
                for associated in item.get("anotacoes_associadas", [])
                if isinstance(associated, dict)
            )
            note = "\n".join(value for value in notes if value)
            ambiguity_rows.append((
                timestamp, "Presença ambígua", record_key,
                source_summary, _export_datetime(item.get("ultima_verificacao")),
                _history_section_text(item.get("motivo"), bare_sources=True), note,
            ))

    detail_total = sum(category_counts.values())
    workbook = openpyxl.Workbook()
    workbook.properties.title = "Histórico de atualizações de planilhas"
    workbook.properties.subject = "Comparações, registros e anotações do histórico de planilhas"
    workbook.properties.creator = "Simpress"
    generated_at = datetime.now().astimezone().replace(tzinfo=None)
    _write_export_summary_sheet(workbook, generated_at, selected_date, display_events, event_counts, category_counts, detail_total)

    _write_export_table_sheet(
        workbook, "Histórico", "Eventos de planilhas",
        "Uma linha por atualização ou restauração, em ordem cronológica.",
        ["Ordem", "Data e hora", "Operação", "ID do evento", "Arquivo no endpoint", "Arquivo na Simpress",
         "Equipamentos no endpoint", "Equipamentos na Simpress", "Equipamentos carregados", "Apelidos atualizados",
         "Apelidos para conferência", "Divergências manuais pendentes", "Ausentes no endpoint",
         "Ausentes na Simpress", "Sem verificação há 7+ dias", "Divergências de local", "Duplicidades e ambiguidades",
         "Séries sem local", "Restauração de", "Diretório de backup"],
        history_rows, "tblHistoricoPlanilhas",
        [10, 20, 16, 34, 32, 42, 18, 18, 20, 17, 20, 22, 18, 18, 23, 18, 22, 17, 20, 36],
    )
    _write_export_table_sheet(
        workbook, "Apelidos", "Apelidos e hostnames",
        "Atualizações, linhas que precisam de conferência e divergências manuais registradas.",
        ["Data e hora", "Operação", "Categoria", "Série", "Hostname no Endpoint",
         "Apelido anterior/manual", "Apelido na Simpress", "Apelido no endpoint/novo", "Horário da alteração",
         "Situação", "Motivo", "Arquivos de origem", "Anotação"],
        alias_rows, "tblHistoricoApelidos",
        [20, 16, 24, 18, 28, 30, 30, 32, 20, 29, 56, 36, 42],
    )
    _write_export_table_sheet(
        workbook, "Ausências", "Equipamentos ausentes por fonte",
        "Registros sem correspondência segura no endpoint ou na Simpress.",
        ["Data e hora", "Categoria", "Ausente em", "Hostname/Apelido", "Série",
         "Local previsto", "Anotação"],
        absence_rows, "tblHistoricoAusencias",
        [20, 24, 22, 32, 18, 28, 42],
    )
    _write_export_table_sheet(
        workbook, "Sem acesso", "Notebooks sem verificação recente",
        "Notebooks cuja última verificação com êxito ocorreu há pelo menos 7 dias.",
        ["Data e hora", "Operação", "Hostname/Apelido", "Série", "Última verificação com êxito",
         "Dias sem verificação", "Linha no endpoint", "Anotação"],
        no_access_rows, "tblHistoricoSemAcesso",
        [20, 16, 32, 20, 26, 20, 16, 42],
    )
    _write_export_table_sheet(
        workbook, "Localidades", "Divergências e séries sem local",
        "Diferenças de localidade e prefixos sem regra cadastrada.",
        ["Data e hora", "Categoria", "Computador", "Série", "Local previsto",
         "Local na Simpress", "Quantidade", "Exemplos", "Anotação"],
        location_rows, "tblHistoricoLocalidades",
        [20, 28, 30, 18, 28, 28, 14, 54, 42],
    )
    _write_export_table_sheet(
        workbook, "Ambiguidades", "Duplicidades e correspondências ambíguas",
        "Presenças ambíguas e correspondências sem identificação única.",
        ["Data e hora", "Ocorrência", "Série ou hostname/Apelido",
         "Fontes", "Última verificação com êxito", "Tratamento", "Anotação"],
        ambiguity_rows, "tblHistoricoAmbiguidades",
        [20, 27, 48, 34, 26, 66, 42],
    )

    _write_export_table_sheet(
        workbook, "Usuários", "Mudanças de usuário logado",
        "Usuários novos, alterados ou removidos no Endpoint e usuários manuais que voltaram a valer.",
        ["Data e hora", "Operação", "Mudança", "Série", "Hostname/Apelido",
         "Usuário anterior", "Usuário novo", "Motivo", "Anotação"],
        user_rows, "tblHistoricoUsuarios",
        [20, 16, 30, 18, 28, 28, 28, 66, 42],
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"Historico_Planilhas_{selected_date:%Y%m%d}_{generated_at:%H%M%S}.xlsx"
    return output, filename
