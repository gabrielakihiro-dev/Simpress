"""Reading and validating the inventory workbooks (Simpress and Endpoint).

Header aliases, sheet discovery and upload table selection live here. Messages
and HTTP status codes are part of the API contract; keep them unchanged.
"""

import io
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional

import openpyxl

from core.errors import DomainError, DomainValidationError
from core.integrity import FileBusyError
from core.xlsx_safety import WORKBOOK_READ_ERRORS, XlsxLimitExceeded, XlsxPackageError, validate_xlsx
from domains.inventory_comparison import _record_value, normalize_reference_value


DETAILS_SERIAL_HEADER_ALIASES = ("NUMEROSERIE", "NUMERODESERIE", "SERIAL")
DETAILS_ALIAS_HEADER_ALIASES = (
    "APELIDO", "HOSTNAME", "NOMEDOCOMPUTADOR", "COMPUTERNAME"
)
DETAILS_TYPE_HEADER_ALIASES = ("Tipo", "ClasseEquipamento")
DETAILS_INSTALLATION_CNPJ_HEADER_ALIASES = (
    "CNPJ Instalacao",
    "CNPJ de Instalacao",
    "CNPJ de Instalação",
    "CnpjInstalacao",
    "CnpjDeInstalacao",
    "CNPJ Instalacao Filial",
    "CNPJ de Instalacao Filial",
    "CNPJ da Instalacao Filial",
    "CNPJ da Instalacao",
    "CNPJ da Filial",
)
DETAILS_REQUIRED_HEADERS = frozenset({"MODELO", "CIDADE"})
DETAILS_COMPLETE_SOURCE_COLUMN_COUNT = 50


@dataclass(frozen=True)
class WorkbookReaderServices:
    validate_xlsx: Callable[[Any], Any]
    validate_xlsx_package: Callable[[bytes], None]
    ensure_workbook_readable: Callable[[Path], None]
    logger: logging.Logger


def normalize_workbook_header(value: Any) -> str:
    """Normaliza cabeçalhos com a mesma regra usada pelas tabelas enviadas."""
    return normalize_reference_value(value)


class WorkbookHeaderError(ValueError):
    """An uploaded workbook has duplicate or semantically ambiguous columns."""


def find_workbook_header_column(headers: Iterable[Any], aliases: Iterable[str]) -> Optional[int]:
    """Return one unambiguous column whose normalized header matches an accepted alias."""
    normalized_aliases = {normalize_workbook_header(alias) for alias in aliases}
    matches = [
        index
        for index, value in enumerate(headers)
        if normalize_workbook_header(value) in normalized_aliases
    ]
    if len(matches) > 1:
        columns = ", ".join(str(index + 1) for index in matches)
        raise ValueError(f"Mais de uma coluna corresponde ao mesmo campo (colunas {columns}).")
    return matches[0] if matches else None


def validate_unambiguous_workbook_headers(headers: Iterable[Any]) -> None:
    """Reject duplicate normalized headers before mapping rows into dictionaries."""
    normalized_headers = [normalize_workbook_header(header) for header in headers]
    positions: Dict[str, List[int]] = {}
    for index, normalized in enumerate(normalized_headers):
        if normalized:
            positions.setdefault(normalized, []).append(index)
    duplicates = [indexes for indexes in positions.values() if len(indexes) > 1]
    if duplicates:
        columns = ", ".join(
            "/".join(str(index + 1) for index in indexes)
            for indexes in duplicates
        )
        raise WorkbookHeaderError(f"A planilha contém cabeçalhos repetidos nas colunas {columns}.")

    alias_groups = (
        DETAILS_SERIAL_HEADER_ALIASES,
        DETAILS_ALIAS_HEADER_ALIASES,
        ("NOMEDOCOMPUTADOR", "COMPUTERNAME", "HOSTNAME"),
        ("ROTULODESERVICONUMERODESERIE", "NUMEROSERIE", "NUMERODESERIE", "SERIAL", "SERVICETAG"),
    )
    for aliases in alias_groups:
        accepted = {normalize_workbook_header(alias) for alias in aliases}
        matches = [index for index, value in enumerate(normalized_headers) if value in accepted]
        if len(matches) > 1:
            columns = ", ".join(str(index + 1) for index in matches)
            raise WorkbookHeaderError(f"A planilha contém colunas candidatas conflitantes nas posições {columns}.")


def read_details_sheet_records(services: WorkbookReaderServices, file_path: Path, sheet_name: str):
    """Read inventory rows with openpyxl, preserving the cells used by the application."""
    workbook = None
    try:
        services.validate_xlsx(file_path)
        workbook = openpyxl.load_workbook(
            file_path, read_only=True, data_only=True, keep_links=False
        )
        worksheet = workbook[sheet_name]
        raw_headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        validate_unambiguous_workbook_headers(raw_headers)
        headers = []
        used_headers = set()
        for index, value in enumerate(raw_headers):
            base = str(value).strip() if value is not None else f"Unnamed: {index}"
            header = base
            suffix = 1
            while header in used_headers:
                header = f"{base}.{suffix}"
                suffix += 1
            used_headers.add(header)
            headers.append(header)

        records = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None and value != "" for value in values):
                continue
            records.append({
                header: values[index] if index < len(values) and values[index] is not None else ""
                for index, header in enumerate(headers)
            })
    finally:
        if workbook is not None:
            workbook.close()

    return headers, records


def ensure_workbook_readable(file_path: Path) -> None:
    """Distingue arquivo indisponível (travado, sincronizando) de arquivo inválido."""
    try:
        with open(file_path, "rb"):
            return
    except FileNotFoundError:
        return
    except OSError as error:
        raise FileBusyError(
            f"Não foi possível ler '{Path(file_path).name}' agora; verifique se o arquivo "
            "está sendo sincronizado ou usado por outro programa e tente novamente."
        ) from error


def find_details_sheet(services: WorkbookReaderServices, file_path: Path) -> Optional[str]:
    """Retorna uma aba EquipamentosDetalhes com os campos usados pelo site.

    Levanta FileBusyError quando o arquivo existe mas não pode ser lido; ``None``
    significa somente que o arquivo não tem uma aba de inventário válida.
    """
    services.ensure_workbook_readable(file_path)
    workbook = None
    try:
        services.validate_xlsx(file_path)
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        matching_sheets = []
        for worksheet in workbook.worksheets:
            headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            validate_unambiguous_workbook_headers(headers)
            normalized_headers = {
                normalize_workbook_header(value)
                for value in headers
                if value is not None
            }

            has_serial = bool(set(DETAILS_SERIAL_HEADER_ALIASES) & normalized_headers)
            has_alias = bool(set(DETAILS_ALIAS_HEADER_ALIASES) & normalized_headers)
            if has_serial and has_alias and DETAILS_REQUIRED_HEADERS.issubset(normalized_headers):
                priority = int("parque" in worksheet.title.lower() or "equipamento" in worksheet.title.lower())
                matching_sheets.append((priority, worksheet.title))

        if matching_sheets:
            matching_sheets.sort(key=lambda item: item[0], reverse=True)
            return matching_sheets[0][1]
    except WORKBOOK_READ_ERRORS as error:
        services.logger.warning(
            "Falha ao localizar aba de inventário (arquivo=%s, papel=detalhes, erro=%s).",
            file_path.name, type(error).__name__,
        )
        return None
    finally:
        if workbook is not None:
            workbook.close()
    return None


def validate_xlsx_package(file_bytes: bytes, validate: Callable[[Any], Any] = validate_xlsx) -> None:
    """Rejects non-XLSX, malformed, or excessively large ZIP packages before parsing."""
    try:
        validate(file_bytes)
    except XlsxLimitExceeded as error:
        raise DomainValidationError(str(error), status_code=413) from None
    except XlsxPackageError:
        raise DomainValidationError(
            detail="O arquivo enviado não é uma planilha .xlsx válida.",
        ) from None
    except DomainError:
        raise
    except (OSError, ValueError):
        raise DomainValidationError(
            detail="O arquivo enviado não é uma planilha .xlsx válida.",
        ) from None


def read_workbook_tables(services: WorkbookReaderServices, file_bytes: bytes) -> List[Dict[str, Any]]:
    """Lê abas não vazias de um XLSX e retorna cabeçalhos normalizados e linhas."""
    workbook = None
    try:
        services.validate_xlsx_package(file_bytes)
        workbook = openpyxl.load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        tables = []
        for worksheet in workbook.worksheets:
            iterator = worksheet.iter_rows(values_only=True)
            raw_headers = next(iterator, ())
            headers = [normalize_workbook_header(value) for value in raw_headers]
            if not any(headers):
                continue
            validate_unambiguous_workbook_headers(raw_headers)

            records = []
            for excel_row_number, values in enumerate(iterator, start=2):
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                record = {
                    headers[index]: values[index]
                    for index in range(min(len(headers), len(values)))
                    if headers[index]
                }
                record["__excel_row_number__"] = excel_row_number
                records.append(record)
            tables.append({
                "sheet": worksheet.title,
                "headers": set(headers),
                "column_count": sum(bool(header) for header in headers),
                "records": records,
            })
        return tables
    except DomainError:
        raise
    except WorkbookHeaderError as error:
        raise DomainValidationError(f"Cabeçalhos inválidos: {error}") from None
    except WORKBOOK_READ_ERRORS as error:
        services.logger.warning("Falha ao ler tabelas da planilha (papel=upload, erro=%s).", type(error).__name__)
        raise DomainValidationError(
            detail="O arquivo enviado não contém uma planilha .xlsx válida ou legível.",
        ) from None
    finally:
        if workbook is not None:
            workbook.close()


def select_upload_table(tables: List[Dict[str, Any]], role: str) -> Dict[str, Any]:
    """Validates the required columns and refuses to guess between matching sheets."""
    serial_headers = set(DETAILS_SERIAL_HEADER_ALIASES)
    alias_headers = set(DETAILS_ALIAS_HEADER_ALIASES)

    if role == "detalhes":
        required = set(DETAILS_REQUIRED_HEADERS)
        basic_candidates = [
            table for table in tables
            if required.issubset(table["headers"])
            and table["headers"] & serial_headers
            and table["headers"] & alias_headers
        ]
        installation_cnpj_headers = {
            normalize_workbook_header(alias)
            for alias in DETAILS_INSTALLATION_CNPJ_HEADER_ALIASES
        }
        candidates = [
            table for table in basic_candidates
            if table.get("column_count", len(table["headers"] - {""})) == DETAILS_COMPLETE_SOURCE_COLUMN_COUNT
            and table["headers"] & installation_cnpj_headers
        ]
        description = (
            "a base completa de EquipamentosDetalhes com exatamente "
            f"{DETAILS_COMPLETE_SOURCE_COLUMN_COUNT} colunas, incluindo número de série, "
            "apelido/hostname, modelo, cidade e CNPJ de instalação"
        )
    elif role == "endpoint":
        name_headers = {"NOMEDOCOMPUTADOR", "COMPUTERNAME", "HOSTNAME"}
        service_tag_headers = {
            "ROTULODESERVICONUMERODESERIE", "NUMEROSERIE", "NUMERODESERIE", "SERIAL", "SERVICETAG"
        }
        candidates = [
            table for table in tables
            if table["headers"] & name_headers and table["headers"] & service_tag_headers
        ]
        description = "nome do computador e rótulo de serviço/número de série"
    else:
        raise ValueError(f"Papel de planilha desconhecido: {role}")

    candidates = [table for table in candidates if table["records"]]
    if not candidates:
        display_role = {
            "endpoint": "Endpoint",
            "detalhes": "Simpress",
        }[role]
        if role == "detalhes" and any(table["records"] for table in basic_candidates):
            actual_column_counts = sorted({
                table.get("column_count", len(table["headers"] - {""}))
                for table in basic_candidates
            })
            schema_problems = []
            if any(count != DETAILS_COMPLETE_SOURCE_COLUMN_COUNT for count in actual_column_counts):
                schema_problems.append(
                    "possui " + ", ".join(str(count) for count in actual_column_counts) + " colunas"
                )
            if any(not (table["headers"] & installation_cnpj_headers) for table in basic_candidates):
                schema_problems.append("não contém o cabeçalho CNPJ de instalação")
            problem_summary = " e ".join(schema_problems) or "não corresponde ao formato completo"
            raise DomainValidationError(
                detail=(
                    f"A planilha enviada como Simpress foi recusada: {problem_summary}. "
                    "A base de entrada precisa ter exatamente "
                    f"{DETAILS_COMPLETE_SOURCE_COLUMN_COUNT} colunas e o cabeçalho CNPJ de instalação. "
                    "A planilha reduzida do Exportar Excel não pode ser usada em Atualizar Planilhas."
                ),
            )
        raise DomainValidationError(
            detail=f"A planilha enviada como {display_role} não contém uma aba com os cabeçalhos obrigatórios: {description}.",
        )

    if len(candidates) > 1:
        sheet_names = ", ".join(f"'{table['sheet']}'" for table in candidates)
        display_role = {
            "endpoint": "Endpoint",
            "detalhes": "Simpress",
        }[role]
        raise DomainValidationError(
            detail=(
                f"O arquivo selecionado para {display_role} tem mais de uma aba compatível "
                f"({sheet_names}). A seleção automática foi cancelada; envie um arquivo com "
                "apenas a aba correta para evitar comparar a tabela errada."
            ),
        )

    table = candidates[0]
    if role == "endpoint":
        name_column = next(key for key in ("NOMEDOCOMPUTADOR", "COMPUTERNAME", "HOSTNAME") if key in table["headers"])
        populated_rows = sum(1 for row in table["records"] if str(row.get(name_column) or "").strip())
    else:
        populated_rows = sum(
            1 for row in table["records"]
            if _record_value(row, "NumeroSerie", "Número de série", "Serial", "ServiceTag")
        )

    if populated_rows == 0:
        raise DomainValidationError(f"A aba '{table['sheet']}' não tem linhas de dados utilizáveis.")

    if role == "detalhes":
        installation_cnpj_column = next(
            header for header in table["headers"] if header in installation_cnpj_headers
        )
        missing_installation_cnpj_rows = sum(
            1
            for row in table["records"]
            if _record_value(row, "NumeroSerie", "Número de série", "Serial", "ServiceTag")
            and not str(row.get(installation_cnpj_column) or "").strip()
        )
        if missing_installation_cnpj_rows:
            raise DomainValidationError(
                detail=(
                    "A base completa da Simpress contém "
                    f"{missing_installation_cnpj_rows} série(s) sem CNPJ de instalação. "
                    "Preencha esse campo para todas as séries antes de atualizar as planilhas."
                ),
            )

    table["populated_rows"] = populated_rows
    return table
