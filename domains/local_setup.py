"""Validation for first-use imports, independent of application storage."""

import io
import json
import logging
import lzma
import os
import re
import unicodedata
import zipfile
import zlib
from pathlib import Path, PurePosixPath

from fastapi import HTTPException
from openpyxl.cell.cell import MergedCell

from core.config import format_size_limit
from core.integrity import atomic_write_json, load_workbook_safely
from core.openpyxl_compat import stored_cells


logger = logging.getLogger(__name__)


TRANSFER_CELLS = (
    "B5", "J5", "B8", "F8", "B17", "F17", "B19", "F19", "J19",
    "B21", "J21", "B23", "F23", "J23", "B25", "J25", "B29", "F29",
    "B31", "F31", "J31", "B33", "J33", "B35", "F35", "J35", "B37", "J37",
)


def unpack_simpress_zip(data: bytes, limit: int) -> tuple[bytes, str]:
    """Read one bounded XLSX member without extracting anything to disk."""
    size_label = format_size_limit(limit)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(item.file_size for item in members) > 2 * limit:
                raise HTTPException(413, "O ZIP excede os limites de descompactação.")
            candidates = []
            for item in members:
                name = item.filename.replace("\\", "/")
                path = PurePosixPath(name)
                if path.is_absolute() or ".." in path.parts or ":" in name or "\x00" in name:
                    raise HTTPException(400, "O ZIP contém um caminho inválido.")
                if item.flag_bits & 1:
                    raise HTTPException(400, "Envie um ZIP sem senha.")
                if item.is_dir():
                    continue
                if item.file_size > limit:
                    raise HTTPException(413, f"Um arquivo no ZIP excede o limite de {size_label}.")
                if path.suffix.lower() == ".xlsx" and not path.name.startswith("~$"):
                    candidates.append(item)
            if len(candidates) != 1:
                raise HTTPException(400, "O ZIP da Simpress precisa conter exatamente uma planilha .xlsx.")
            member = candidates[0]
            with archive.open(member) as stream:
                workbook = stream.read(limit + 1)
            if len(workbook) > limit:
                raise HTTPException(413, f"A planilha descompactada excede o limite de {size_label}.")
            if not workbook:
                raise HTTPException(400, "A planilha dentro do ZIP está vazia.")
            return workbook, member.filename
    except HTTPException:
        raise
    except (OSError, ValueError, RuntimeError, NotImplementedError, zipfile.BadZipFile, zlib.error, lzma.LZMAError) as error:
        raise HTTPException(400, "Não foi possível ler o ZIP da Simpress.") from error


def _normalize(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"[^A-Z0-9]", "", text.encode("ascii", "ignore").decode().upper())


def validate_local_model(source, kind: str) -> None:
    """Use the same selected sheet and writable fields as the generators."""
    wb = load_workbook_safely(
        source, allow_threaded_comments=True,
        allow_legacy_comments=kind == "transferencia",
        allow_custom_xml_passthrough=kind == "transferencia",
    )
    try:
        if kind == "transferencia":
            name = next((name for name in wb.sheetnames if "transfer" in name.lower()
                         and "lote" not in name.lower()), wb.sheetnames[0])
            ws = wb[name]
            stored_addresses = {cell.coordinate for cell in stored_cells(ws)}
            for address, expected in (("B2", "TRANSFERENCIADEEQUIPAMENTO"),
                                      ("B15", "DADOSDERETIRADA"), ("B27", "DADOSDEENTREGA")):
                if _normalize(ws[address].value) != expected:
                    raise ValueError("O modelo precisa usar o layout oficial de transferência da Simpress.")
            if any(isinstance(ws[address], MergedCell) for address in TRANSFER_CELLS):
                raise ValueError("O modelo contém campos de transferência que não podem ser preenchidos.")
            if not set(TRANSFER_CELLS).issubset(stored_addresses):
                raise ValueError("O modelo não contém todas as células de preenchimento da transferência.")
        elif kind == "exportacao":
            name = next((name for name in wb.sheetnames if "parque" in name.lower()
                         or "equipamento" in name.lower()), wb.sheetnames[0])
            ws = wb[name]
            stored_addresses = {cell.coordinate for cell in stored_cells(ws)}
            headers = [_normalize(cell.value) for cell in next(ws.iter_rows(max_row=1), ())]
            serial_columns = [index for index, header in enumerate(headers)
                              if header in {"NUMEROSERIE", "NUMERODESERIE", "SERIAL"}]
            if len(serial_columns) != 1 or headers.count("APELIDO") != 1:
                raise ValueError("O modelo de exportação precisa ter colunas únicas de número de série e apelido.")
            serial_col = serial_columns[0]
            alias_col = headers.index("APELIDO")
            id_col = headers.index("IDCLIENTE") if "IDCLIENTE" in headers else None
            has_equipment = False
            for row in ws.iter_rows(min_row=2):
                if not _normalize(row[serial_col].value):
                    continue
                has_equipment = True
                for column in (alias_col, id_col):
                    if column is None:
                        continue
                    cell = row[column]
                    if isinstance(cell, MergedCell) or cell.coordinate not in stored_addresses:
                        raise ValueError("O modelo de exportação contém campos ausentes ou mesclados que não podem ser preenchidos.")
            if not has_equipment:
                raise ValueError("O modelo de exportação não contém equipamentos.")
        else:
            raise ValueError("Tipo de modelo inválido.")
    finally:
        wb.close()


# Payload field, environment variable, key in config/contato_origem.json and length limit.
ORIGIN_CONTACT_FIELDS = (
    ("retirada_contato", "SIMPRESS_TRANSFER_ORIGIN_CONTACT", "contato", 160),
    ("retirada_email", "SIMPRESS_TRANSFER_ORIGIN_EMAIL", "email", 254),
    ("retirada_telefone", "SIMPRESS_TRANSFER_ORIGIN_PHONE", "telefone", 64),
)


def _read_origin_contact_file(path: Path) -> tuple[dict, bool]:
    """Return the stored values and whether the file exists but could not be used."""
    try:
        stored = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        return {}, False
    except (OSError, ValueError):
        logger.warning("Arquivo de contato de origem ilegível; os campos ficam vazios.")
        return {}, True
    if not isinstance(stored, dict):
        logger.warning("Arquivo de contato de origem sem objeto JSON; os campos ficam vazios.")
        return {}, True
    return stored, False


def load_origin_contact(path: Path, environment=None) -> dict:
    """Resolve each origin contact field: environment variable, then local file, then empty."""
    env = os.environ if environment is None else environment
    stored, invalid_file = _read_origin_contact_file(path)
    values, sources = {}, {}
    for field, variable, key, limit in ORIGIN_CONTACT_FIELDS:
        from_env = str(env.get(variable) or "").strip()
        from_file = stored.get(key)
        from_file = from_file.strip() if isinstance(from_file, str) else ""
        if from_env:
            values[field], sources[field] = from_env[:limit], "ambiente"
        elif from_file:
            values[field], sources[field] = from_file[:limit], "arquivo"
        else:
            values[field], sources[field] = "", ""
    return {
        **values,
        "contato_origem_configurado": all(values.values()),
        "contato_origem_fontes": sources,
        "contato_origem_arquivo_invalido": invalid_file,
    }


def save_origin_contact(path: Path, contato: str, email: str, telefone: str) -> None:
    """Write the already validated contact to the git-ignored local file."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_json(target, {"contato": contato, "email": email, "telefone": telefone})
