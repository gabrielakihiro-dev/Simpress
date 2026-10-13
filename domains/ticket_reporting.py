"""History, attention queue, filtering, and export for ticket reports."""

from __future__ import annotations

from collections import Counter
from copy import deepcopy
from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import json
import re
import unicodedata
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from domains.history_annotations import annotation_export_text


MAX_TICKET_SNAPSHOTS = 50
TICKET_SNAPSHOT_FIELDS = ("atualizado_em", "fontes", "resumo", "chamados")
ORACLE_CHANGE_FIELDS = (
    ("solicitacao", "Solicitação Oracle"),
    ("ordem_servico", "OS Oracle"),
    ("logradouro", "Logradouro Oracle"),
    ("numero_endereco", "Número do endereço Oracle"),
    ("bairro", "Bairro Oracle"),
    ("cidade", "Cidade Oracle"),
    ("atendido_pela_inibicao", "Atendido pela inibição Oracle"),
    ("nome_filial", "Nome da filial Oracle"),
    ("status", "Status Oracle"),
    ("sub_status", "Sub-status Oracle"),
    ("fila", "Fila Oracle"),
    ("fase_atendimento", "Fase Oracle"),
    ("status_entregas", "Status das entregas"),
    ("status_atividades", "Status das atividades"),
    ("assunto", "Assunto Oracle"),
    ("descricao", "Descrição do chamado Oracle"),
    ("data_criacao", "Criação Oracle"),
    ("ultima_atualizacao", "Última atualização Oracle"),
    ("data_conclusao", "Conclusão Oracle"),
)
SIMPRESS_CHANGE_FIELDS = (
    ("pedidos", "Pedidos Simpress"),
    ("ordens_venda", "Ordem de venda (OV) Simpress"),
    ("ordens_servico", "OS Simpress"),
    ("chamados_cliente", "Chamados do cliente"),
    ("status", "Status Simpress"),
    ("datas_pedido", "Datas dos pedidos"),
    ("modelos", "Modelos Simpress"),
    ("descricoes_produto", "Descrição do produto Simpress"),
    ("previsoes_entrega", "Previsão de entrega Simpress"),
    ("cidades_entrega", "Cidade de entrega Simpress"),
    ("enderecos_entrega", "Endereço de entrega Simpress"),
)
# "Data do pedido" and "Previsão de entrega" are already compared through the lists above.
SIMPRESS_STAGE_CHANGE_FIELDS = (
    ("liberacao_pedido", "Liberação do pedido Simpress"),
    ("verificado_estoque", "Verificado e estoque Simpress"),
    ("pedido_separacao", "Pedido em separação Simpress"),
    ("nota_fiscal_emitida", "Nota fiscal emitida Simpress"),
)


def _stable_serial(value: Any) -> str:
    plain = unicodedata.normalize("NFKD", str(value or ""))
    ascii_value = plain.encode("ascii", "ignore").decode("ascii").upper()
    return re.sub(r"[^A-Z0-9]+", "", ascii_value)


def _normalized_text(value: Any) -> str:
    plain = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join(plain.encode("ascii", "ignore").decode("ascii").casefold().split())


def _order_service_descending_key(value: Any) -> Tuple[bool, int, str]:
    text = str(value or "").strip()
    ascii_value = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    digits = "" if re.search(r"[A-Za-z]", ascii_value) else re.sub(r"\D", "", ascii_value)
    number = int(digits) if digits else None
    return number is not None, number if number is not None else -1, _normalized_text(text)


def _ticket_order_service_sort_key(ticket: Dict[str, Any]) -> Tuple[int, int, str]:
    """Return the highest OS for a ticket, with missing OS values sorting last."""
    simpress_value = ticket.get("simpress")
    simpress: Dict[str, Any] = simpress_value if isinstance(simpress_value, dict) else {}
    oracle_value = ticket.get("oracle")
    oracle: List[Any] = oracle_value if isinstance(oracle_value, list) else []
    values: List[Any] = []
    for field_values in (simpress.get("ordens_servico"), ticket.get("ordens_servico_conferidas")):
        if isinstance(field_values, (list, tuple, set)):
            values.extend(field_values)
        elif field_values:
            values.append(field_values)
    values.extend(
        item.get("ordem_servico")
        for item in oracle
        if isinstance(item, dict)
    )

    missing_labels = {"nao informado", "nao informada", "sem os", "n/a", "na"}
    candidates = [
        str(value).strip()
        for value in values
        if value is not None and str(value).strip()
        and _normalized_text(value) not in missing_labels
    ]
    if not candidates:
        return 0, -1, ""

    highest = max(candidates, key=_order_service_descending_key)
    numeric, number, normalized = _order_service_descending_key(highest)
    return (2 if numeric else 1), number, normalized


def _ticket_snapshot(report: Dict[str, Any], snapshot_id: Optional[str] = None) -> Optional[Dict[str, Any]]:
    if not isinstance(report, dict) or not isinstance(report.get("chamados"), list):
        return None
    snapshot = {field: deepcopy(report.get(field)) for field in TICKET_SNAPSHOT_FIELDS}
    snapshot["atualizado_em"] = str(snapshot.get("atualizado_em") or "")
    existing_id = snapshot_id or report.get("snapshot_id")
    if not existing_id:
        fingerprint = json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        existing_id = "legado-" + sha256(fingerprint.encode("utf-8")).hexdigest()[:20]
    snapshot["id"] = str(existing_id)
    return snapshot


def ticket_report_snapshots(report: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Return ordered snapshots, including the latest report when migrating legacy data."""
    snapshots: List[Dict[str, Any]] = []
    raw_history = report.get("historico") if isinstance(report, dict) else None
    if isinstance(raw_history, list):
        for item in raw_history:
            snapshot = _ticket_snapshot(item, item.get("id") if isinstance(item, dict) else None)
            if snapshot:
                snapshots.append(snapshot)

    has_imported_current = bool(
        isinstance(report, dict)
        and (report.get("atualizado_em") or report.get("fontes") or report.get("snapshot_id") or report.get("chamados"))
    )
    current = _ticket_snapshot(report) if has_imported_current else None
    if current and not any(item["id"] == current["id"] for item in snapshots):
        snapshots.append(current)

    # Keep the newest occurrence of an ID while preserving chronological order.
    unique: Dict[str, Dict[str, Any]] = {}
    for snapshot in snapshots:
        unique[snapshot["id"]] = snapshot
    ordered = list(unique.values())
    return ordered[-MAX_TICKET_SNAPSHOTS:]


def _order_service_key(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper()


def _numeric_order_service(value: Any) -> Optional[int]:
    normalized = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    if re.search(r"[A-Za-z]", normalized):
        return None
    digits = re.sub(r"\D", "", normalized)
    return int(digits) if digits else None


def _latest_simpress_identifier(values: Sequence[Any]) -> str:
    candidates = [
        item.strip()
        for value in values
        if value is not None
        for item in re.split(r"[,;\n|]+", str(value))
        if item.strip()
    ]
    if not candidates:
        return ""
    numeric = [(item, _numeric_order_service(item)) for item in candidates]
    if all(number is not None for _, number in numeric):
        return max(
            ((item, number) for item, number in numeric if number is not None),
            key=lambda entry: entry[1],
        )[0]
    return candidates[-1]


def normalize_ticket_snapshot(snapshot: Dict[str, Any]) -> Dict[str, Any]:
    """Project legacy reports with current Simpress pedido and OS IDs without changing stored data."""
    normalized = deepcopy(snapshot)
    tickets = normalized.get("chamados")
    if not isinstance(tickets, list):
        return normalized

    changed = False
    for ticket in tickets:
        if not isinstance(ticket, dict):
            continue
        simpress = ticket.get("simpress")
        if not isinstance(simpress, dict):
            simpress = {}
            ticket["simpress"] = simpress
        orders = simpress.get("ordens_servico")
        orders_changed = False
        if isinstance(orders, list):
            latest_os = _latest_simpress_identifier(orders)
            latest_orders = [latest_os] if latest_os else []
            orders_changed = orders != latest_orders
            if orders_changed:
                changed = True
                simpress["ordens_servico"] = latest_orders

        pedidos = simpress.get("pedidos")
        if isinstance(pedidos, list):
            latest_pedido = _latest_simpress_identifier(pedidos)
            latest_pedidos = [latest_pedido] if latest_pedido else []
            if pedidos != latest_pedidos:
                changed = True
                simpress["pedidos"] = latest_pedidos

        if orders_changed:
            oracle_value = ticket.get("oracle")
            oracle: List[Any] = oracle_value if isinstance(oracle_value, list) else []
            oracle_orders = [
                str(item.get("ordem_servico") or "").strip()
                for item in oracle
                if isinstance(item, dict) and str(item.get("ordem_servico") or "").strip()
            ]
            matched = (
                [latest_os]
                if latest_os and any(_order_service_key(item) == _order_service_key(latest_os) for item in oracle_orders)
                else []
            )
            ticket["ordens_servico_conferidas"] = matched

            simpress_present = bool(simpress.get("serie") or simpress.get("pedidos") or latest_os)
            if simpress_present and oracle:
                if matched:
                    ticket["relacao"] = "conferida"
                elif latest_os and oracle_orders:
                    ticket["relacao"] = "divergente"
                else:
                    ticket["relacao"] = "sem_os"
            elif simpress_present:
                ticket["relacao"] = "somente_simpress"
            elif oracle:
                ticket["relacao"] = "somente_oracle"
            else:
                ticket["relacao"] = "sem_os"

        ticket["fontes_relacao"] = ticket_relation_provenance(ticket)
        ticket["alertas_relacao"] = ticket_relation_alerts(ticket)
        ticket["revisao_relacao"] = bool(ticket["alertas_relacao"])

    if changed and isinstance(normalized.get("resumo"), dict):
        relation_counts = Counter(
            str(ticket.get("relacao") or "")
            for ticket in tickets
            if isinstance(ticket, dict)
        )
        summary = dict(normalized["resumo"])
        summary.update({
            "series_total": len(tickets),
            "conferidas": relation_counts["conferida"],
            "divergentes": relation_counts["divergente"],
            "sem_os": relation_counts["sem_os"],
            "somente_simpress": relation_counts["somente_simpress"],
            "somente_oracle": relation_counts["somente_oracle"],
        })
        normalized["resumo"] = summary
    return normalized


def ticket_relation_provenance(ticket: Dict[str, Any]) -> Dict[str, Any]:
    """Show which source supplied each identifier without changing stored rows."""
    simpress_value = ticket.get("simpress")
    simpress = simpress_value if isinstance(simpress_value, dict) else {}
    oracle_value = ticket.get("oracle")
    oracle = [item for item in oracle_value if isinstance(item, dict)] if isinstance(oracle_value, list) else []
    relation = str(ticket.get("relacao") or "")

    def values(value: Any) -> List[str]:
        if not isinstance(value, (list, tuple, set)):
            value = [value] if value else []
        return sorted({str(item).strip() for item in value if str(item or "").strip()}, key=_normalized_text)

    simpress_series = bool(
        simpress.get("serie") or simpress.get("pedidos") or simpress.get("ordens_servico") or simpress.get("ordens_venda")
    ) or relation == "somente_simpress"
    oracle_series = bool(oracle) or relation == "somente_oracle"
    return {
        "serie": {"simpress": simpress_series, "oracle": oracle_series},
        "pedido": {
            "simpress": values(simpress.get("pedidos")),
            "oracle_solicitacao": values([item.get("solicitacao") for item in oracle]),
        },
        "ordem_servico": {
            "simpress": sorted(values(simpress.get("ordens_servico")), key=_order_service_descending_key, reverse=True),
            "oracle": sorted(values([item.get("ordem_servico") for item in oracle]), key=_order_service_descending_key, reverse=True),
        },
    }


def ticket_relation_alerts(ticket: Dict[str, Any]) -> List[Dict[str, str]]:
    """Return actionable relation warnings, preserving parser-detected duplicates."""
    provenance = ticket_relation_provenance(ticket)
    simpress_os = provenance["ordem_servico"]["simpress"]
    oracle_os = provenance["ordem_servico"]["oracle"]
    existing = ticket.get("alertas_relacao")
    alerts = [dict(item) for item in existing if isinstance(item, dict) and item.get("mensagem")] if isinstance(existing, list) else []
    known = {str(item.get("tipo") or "") for item in alerts}

    def add(kind: str, message: str) -> None:
        if kind not in known:
            alerts.append({"tipo": kind, "mensagem": message})
            known.add(kind)

    relation = str(ticket.get("relacao") or "")
    series_sources = provenance["serie"]
    if relation == "divergente":
        add("os_divergente", f"OS diferentes entre as fontes: Simpress {', '.join(simpress_os) or 'sem OS'}; Oracle {', '.join(oracle_os) or 'sem OS'}. Conferir.")
    elif relation == "sem_os":
        add("sem_os", "A série aparece nas fontes, mas não há uma OS para confirmar a relação.")
    elif relation == "somente_simpress" or (series_sources["simpress"] and not series_sources["oracle"]):
        add("fonte_oracle_ausente", "Série encontrada na Simpress, sem registro correspondente no Oracle.")
    elif relation == "somente_oracle" or (series_sources["oracle"] and not series_sources["simpress"]):
        add("fonte_simpress_ausente", "Série encontrada no Oracle, sem pedido correspondente na Simpress.")
    return alerts


def ticket_snapshot_metadata(snapshot: Dict[str, Any]) -> Dict[str, Any]:
    sources = snapshot.get("fontes") or {}
    summary = snapshot.get("resumo") or {}
    return {
        "id": snapshot.get("id", ""),
        "atualizado_em": snapshot.get("atualizado_em", ""),
        "fontes": {
            "simpress": deepcopy((sources.get("simpress") or {}).get("arquivo", "")),
            "oracle": deepcopy((sources.get("oracle") or {}).get("arquivo", "")),
        },
        "resumo": deepcopy(summary),
    }


TICKET_TIMELINE_ANNOTATION_CATEGORY = "linha_do_tempo_chamados"
TICKET_TIMELINE_STATE_FIELDS = (
    ("relacao", "Relação das fontes"),
    ("pedidos_simpress", "Pedido Simpress"),
    ("solicitacoes_oracle", "Solicitação Oracle"),
    ("ordens_venda_simpress", "OV Simpress"),
    ("ordens_servico_simpress", "OS Simpress"),
    ("ordens_servico_oracle", "OS Oracle"),
    ("status_simpress", "Status Simpress"),
    ("status_oracle", "Status Oracle"),
    ("filas_oracle", "Fila Oracle"),
    ("cidades", "Cidade"),
    ("atualizacoes_oracle", "Última atualização Oracle"),
)


def ticket_service_providers(ticket: Dict[str, Any]) -> List[str]:
    """Classify Oracle service branches into triage-friendly provider groups."""
    providers = set()
    oracle = ticket.get("oracle")
    for item in oracle if isinstance(oracle, list) else []:
        if not isinstance(item, dict):
            continue
        branch = _normalized_text(item.get("nome_filial"))
        if "parceiro" in branch:
            providers.add("parceiro")
        elif "simpress" in branch:
            providers.add("simpress")
        else:
            providers.add("outro" if branch else "nao informado")
    return sorted(providers)


def ticket_city_values(ticket: Dict[str, Any]) -> List[str]:
    """Return distinct service and delivery cities used by daily triage."""
    values: Dict[str, str] = {}
    simpress = ticket.get("simpress")
    simpress = simpress if isinstance(simpress, dict) else {}
    delivery_cities = simpress.get("cidades_entrega", [])
    delivery_cities = delivery_cities if isinstance(delivery_cities, (list, tuple, set)) else [delivery_cities]
    for value in delivery_cities:
        text = str(value or "").strip()
        if text:
            values.setdefault(_normalized_text(text), text)
    oracle = ticket.get("oracle")
    for item in oracle if isinstance(oracle, list) else []:
        if not isinstance(item, dict):
            continue
        text = str(item.get("cidade") or "").strip()
        if text:
            values.setdefault(_normalized_text(text), text)
    return sorted(values.values(), key=_normalized_text)


def ticket_timeline_item_id(snapshot_id: Any, serial: Any) -> str:
    """Build an event-specific stable annotation ID for a ticket in a snapshot."""
    material = f"{snapshot_id or ''}\0{_stable_serial(serial)}".encode("utf-8")
    return "ticket-" + sha256(material).hexdigest()[:48]


def ticket_timeline_state(ticket: Dict[str, Any]) -> Dict[str, Any]:
    """Project the identifiers and service state that matter to ticket follow-up."""
    simpress_value = ticket.get("simpress")
    simpress = simpress_value if isinstance(simpress_value, dict) else {}
    oracle_value = ticket.get("oracle")
    oracle = [item for item in oracle_value if isinstance(item, dict)] if isinstance(oracle_value, list) else []

    def simpress_values(field: str) -> List[str]:
        values = simpress.get(field)
        if not isinstance(values, (list, tuple, set)):
            values = [values] if values else []
        return sorted({str(value).strip() for value in values if str(value or "").strip()}, key=_normalized_text)

    def oracle_values(field: str) -> List[str]:
        return sorted({
            str(item.get(field) or "").strip()
            for item in oracle
            if str(item.get(field) or "").strip()
        }, key=_normalized_text)

    status_oracle = sorted({
        " · ".join(value for value in (
            str(item.get("status") or "").strip(),
            str(item.get("sub_status") or "").strip(),
        ) if value)
        for item in oracle
        if str(item.get("status") or "").strip() or str(item.get("sub_status") or "").strip()
    }, key=_normalized_text)
    return {
        "relacao": str(ticket.get("relacao") or ""),
        "pedidos_simpress": simpress_values("pedidos"),
        "solicitacoes_oracle": oracle_values("solicitacao"),
        "ordens_venda_simpress": simpress_values("ordens_venda"),
        "ordens_servico_simpress": simpress_values("ordens_servico"),
        "ordens_servico_oracle": oracle_values("ordem_servico"),
        "status_simpress": simpress_values("status"),
        "status_oracle": status_oracle,
        "filas_oracle": oracle_values("fila"),
        "cidades": ticket_city_values(ticket),
        "atualizacoes_oracle": oracle_values("ultima_atualizacao"),
    }


@dataclass(frozen=True)
class TicketTimelineIndex:
    """Operation-scoped timeline projection, retaining rows for duplicate series."""

    snapshots: Tuple[Dict[str, Any], ...]
    events_by_serial: Mapping[str, Tuple[Dict[str, Any], ...]]
    events_by_item_id: Mapping[str, Dict[str, Any]]
    events_by_snapshot_serial: Mapping[Tuple[int, str], Dict[str, Any]]
    rows_by_snapshot_serial: Mapping[Tuple[int, str], Tuple[Dict[str, Any], ...]]

    def event_for_snapshot(self, snapshot: Dict[str, Any], serial: Any) -> Optional[Dict[str, Any]]:
        return self.events_by_snapshot_serial.get((id(snapshot), _stable_serial(serial)))


def build_ticket_timeline_index(snapshots: Sequence[Dict[str, Any]]) -> TicketTimelineIndex:
    """Index selected immutable snapshots once, preserving first-row timeline semantics."""
    snapshot_values = tuple(snapshots)
    events_by_serial: Dict[str, List[Dict[str, Any]]] = {}
    events_by_item_id: Dict[str, Dict[str, Any]] = {}
    events_by_snapshot_serial: Dict[Tuple[int, str], Dict[str, Any]] = {}
    rows_by_snapshot_serial: Dict[Tuple[int, str], Tuple[Dict[str, Any], ...]] = {}
    previous_state_by_serial: Dict[str, Dict[str, Any]] = {}

    for snapshot in snapshot_values:
        if not isinstance(snapshot, dict):
            continue
        rows_by_key: Dict[str, List[Dict[str, Any]]] = {}
        for ticket in snapshot.get("chamados", []):
            if isinstance(ticket, dict):
                rows_by_key.setdefault(_stable_serial(ticket.get("serie")), []).append(ticket)

        snapshot_identity = id(snapshot)
        snapshot_id = str(snapshot.get("id") or "")
        for serial_key, matching_rows in rows_by_key.items():
            # Keep every duplicate row available for audit; the historical projection
            # intentionally uses the first matching row, as the former `next(...)` did.
            rows_by_snapshot_serial[(snapshot_identity, serial_key)] = tuple(matching_rows)
            ticket = matching_rows[0]
            state = ticket_timeline_state(ticket)
            previous_state = previous_state_by_serial.get(serial_key)
            changes = []
            for field, label in TICKET_TIMELINE_STATE_FIELDS:
                current_value = state.get(field)
                previous_value = previous_state.get(field) if previous_state is not None else None
                if previous_state is None:
                    if current_value in (None, "", []):
                        continue
                    changes.append({"campo": label, "anterior": [], "atual": current_value})
                elif current_value != previous_value:
                    changes.append({"campo": label, "anterior": previous_value, "atual": current_value})
            event: Dict[str, Any] = {
                "snapshot_id": snapshot_id,
                "item_id": ticket_timeline_item_id(snapshot_id, ticket.get("serie") or ""),
                "serie": str(ticket.get("serie") or ""),
                "atualizado_em": str(snapshot.get("atualizado_em") or ""),
                "estado": state,
                "mudancas": changes,
                "tipo": "Importação inicial" if previous_state is None else "Dados alterados" if changes else "Sem mudanças",
                "anotacao": "",
            }
            events_by_serial.setdefault(serial_key, []).append(event)
            events_by_item_id.setdefault(event["item_id"], event)
            events_by_snapshot_serial[(snapshot_identity, serial_key)] = event
            previous_state_by_serial[serial_key] = state

    return TicketTimelineIndex(
        snapshots=snapshot_values,
        events_by_serial={key: tuple(events) for key, events in events_by_serial.items()},
        events_by_item_id=events_by_item_id,
        events_by_snapshot_serial=events_by_snapshot_serial,
        rows_by_snapshot_serial=rows_by_snapshot_serial,
    )


def build_ticket_timeline(
    snapshots: Sequence[Dict[str, Any]], serial: str,
    *, index: Optional[TicketTimelineIndex] = None,
) -> List[Dict[str, Any]]:
    """List one auditable state checkpoint per import, with only changed fields highlighted."""
    timeline_index = index if index is not None else build_ticket_timeline_index(snapshots)
    events = deepcopy(list(timeline_index.events_by_serial.get(_stable_serial(serial), ())))
    for event in events:
        if not event.get("serie") and serial:
            event["serie"] = str(serial)
    return events


def _parse_date(value: Any) -> Optional[date]:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for pattern in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _is_closed(ticket: Dict[str, Any]) -> bool:
    text = _normalized_text(f"{ticket.get('status', '')} {ticket.get('sub_status', '')}")
    return any(token in text for token in (
        "conclu", "encerr", "finaliz", "cancel", "resolvid", "fechad", "closed", "complete",
    ))


def ticket_attention(
    ticket: Dict[str, Any],
    *,
    stale_days: int = 3,
    today: Optional[date] = None,
) -> Dict[str, Any]:
    """Flag active Oracle calls with no update date or no update within the threshold."""
    oracle_value = ticket.get("oracle")
    oracle: List[Any] = oracle_value if isinstance(oracle_value, list) else []
    active: List[Dict[str, Any]] = [
        item for item in oracle if isinstance(item, dict) and not _is_closed(item)
    ]
    update_dates: List[date] = []
    for item in active:
        update_date = _parse_date(item.get("ultima_atualizacao"))
        if update_date is not None:
            update_dates.append(update_date)
    latest_date = max(update_dates) if update_dates else None
    reference = today or datetime.now().astimezone().date()
    age = max(0, (reference - latest_date).days) if latest_date else None
    is_stale = bool(active and ((age is not None and age >= max(0, stale_days)) or latest_date is None))

    priority = 0 if is_stale else 1
    reason = (
        "Sem data de atualização" if age is None else f"Sem atualização há {age} dias"
    ) if is_stale else "Em acompanhamento"

    return {
        "prioridade": priority,
        "em_atencao": is_stale,
        "desatualizado": is_stale,
        "dias_sem_atualizacao": age,
        "ultima_atualizacao": latest_date.isoformat() if latest_date else "",
        "motivo": reason,
    }


def _walk_values(value: Any) -> Iterable[str]:
    if isinstance(value, dict):
        for nested in value.values():
            yield from _walk_values(nested)
    elif isinstance(value, list):
        for nested in value:
            yield from _walk_values(nested)
    elif value is not None:
        yield str(value)


def filter_ticket_records(
    tickets: Sequence[Dict[str, Any]],
    *,
    relation: str = "",
    status: str = "",
    queue: str = "",
    city: str = "",
    provider: str = "",
    search: str = "",
    stale_days: int = 3,
    age_filter: str = "",
    attention_only: bool = False,
    today: Optional[date] = None,
) -> List[Dict[str, Any]]:
    relation_filter = _normalized_text(relation)
    status_filter = _normalized_text(status)
    queue_filter = _normalized_text(queue)
    city_filter = _normalized_text(city)
    provider_filter = _normalized_text(provider)
    query = _normalized_text(search)
    minimum_age = int(age_filter) if str(age_filter).isdigit() else None
    missing_age = str(age_filter) == "sem_data"
    matched: List[Tuple[Tuple[int, int, str], Dict[str, Any]]] = []

    for ticket in tickets:
        if not isinstance(ticket, dict):
            continue
        oracle = [item for item in ticket.get("oracle", []) if isinstance(item, dict)]
        attention = ticket_attention(ticket, stale_days=stale_days, today=today)
        if relation_filter and _normalized_text(ticket.get("relacao")) != relation_filter:
            continue
        if city_filter and city_filter not in {_normalized_text(value) for value in ticket_city_values(ticket)}:
            continue
        if provider_filter and provider_filter not in {
            _normalized_text(value) for value in ticket_service_providers(ticket)
        }:
            continue
        if (status_filter or queue_filter) and not any(
            (not status_filter or _normalized_text(item.get("status")) == status_filter)
            and (not queue_filter or _normalized_text(item.get("fila")) == queue_filter)
            for item in oracle
        ):
            continue
        if attention_only and not attention["em_atencao"]:
            continue
        active_oracle = [item for item in oracle if not _is_closed(item)]
        if missing_age and not (active_oracle and attention["dias_sem_atualizacao"] is None):
            continue
        if minimum_age is not None and not (
            attention["dias_sem_atualizacao"] is not None
            and attention["dias_sem_atualizacao"] >= minimum_age
        ):
            continue
        if query and query not in _normalized_text(" ".join(_walk_values(ticket))):
            continue
        age = attention["dias_sem_atualizacao"] or 0
        matched.append((
            (attention["prioridade"], -age, _normalized_text(ticket.get("serie"))),
            ticket,
        ))

    # Preserve attention ordering for equal OS values, while ordering the full
    # result by the newest available OS and leaving records without OS at the end.
    matched.sort(key=lambda item: item[0])
    matched.sort(key=lambda item: _ticket_order_service_sort_key(item[1]), reverse=True)
    return [ticket for _, ticket in matched]


def _ticket_field_values(ticket: Dict[str, Any], field: str, source: str) -> List[str]:
    if source == "simpress":
        values = (ticket.get("simpress") or {}).get(field, [])
        return sorted({str(item) for item in values if str(item)}, key=_normalized_text)
    values = [item.get(field, "") for item in ticket.get("oracle", []) if isinstance(item, dict)]
    return sorted({str(item) for item in values if str(item)}, key=_normalized_text)


def ticket_stage_date(ticket: Dict[str, Any], stage_key: str) -> str:
    """Date of one Simpress order stage, empty when the stage has no date."""
    stages = (ticket.get("simpress") or {}).get("etapas_pedido")
    for stage in stages if isinstance(stages, list) else []:
        if isinstance(stage, dict) and stage.get("chave") == stage_key:
            return str(stage.get("data") or "").strip()
    return ""


def _oracle_ticket_rows(ticket: Dict[str, Any]) -> List[Dict[str, str]]:
    rows = []
    for oracle in ticket.get("oracle", []):
        if not isinstance(oracle, dict):
            continue
        rows.append({field: str(oracle.get(field, "") or "") for field, _label in ORACLE_CHANGE_FIELDS})
    return sorted(rows, key=lambda row: json.dumps(row, ensure_ascii=False, sort_keys=True))


def compare_ticket_snapshots(previous: Dict[str, Any], current: Dict[str, Any]) -> Dict[str, Any]:
    previous_by_serial = {
        _stable_serial(item.get("serie")): item
        for item in previous.get("chamados", [])
        if isinstance(item, dict) and _stable_serial(item.get("serie"))
    }
    current_by_serial = {
        _stable_serial(item.get("serie")): item
        for item in current.get("chamados", [])
        if isinstance(item, dict) and _stable_serial(item.get("serie"))
    }
    before_keys, after_keys = set(previous_by_serial), set(current_by_serial)
    new_rows = [current_by_serial[key].get("serie", key) for key in sorted(after_keys - before_keys)]
    absent_rows = [previous_by_serial[key].get("serie", key) for key in sorted(before_keys - after_keys)]
    unchanged: List[str] = []
    changed: List[Dict[str, Any]] = []

    for key in sorted(before_keys & after_keys):
        old, new = previous_by_serial[key], current_by_serial[key]
        fields: List[Dict[str, Any]] = []
        if old.get("relacao") != new.get("relacao"):
            fields.append({"campo": "Relação das OS", "anterior": old.get("relacao", ""), "atual": new.get("relacao", "")})
        for field, label in SIMPRESS_CHANGE_FIELDS:
            old_values = _ticket_field_values(old, field, "simpress")
            new_values = _ticket_field_values(new, field, "simpress")
            if old_values != new_values:
                fields.append({"campo": label, "anterior": old_values, "atual": new_values})
        for stage_key, label in SIMPRESS_STAGE_CHANGE_FIELDS:
            old_date, new_date = ticket_stage_date(old, stage_key), ticket_stage_date(new, stage_key)
            if old_date != new_date:
                fields.append({
                    "campo": label,
                    "anterior": [old_date] if old_date else [],
                    "atual": [new_date] if new_date else [],
                })
        for field, label in ORACLE_CHANGE_FIELDS:
            old_values = _ticket_field_values(old, field, "oracle")
            new_values = _ticket_field_values(new, field, "oracle")
            if old_values != new_values:
                fields.append({"campo": label, "anterior": old_values, "atual": new_values})
        old_oracle_rows = _oracle_ticket_rows(old)
        new_oracle_rows = _oracle_ticket_rows(new)
        if old_oracle_rows != new_oracle_rows and not fields:
            fields.append({
                "campo": "Vínculo entre solicitações Oracle",
                "anterior": old_oracle_rows,
                "atual": new_oracle_rows,
            })
        if fields:
            changed.append({"serie": new.get("serie", old.get("serie", key)), "campos": fields})
        else:
            unchanged.append(new.get("serie", key))

    return {
        "anterior": ticket_snapshot_metadata(previous),
        "atual": ticket_snapshot_metadata(current),
        "contagens": {
            "novos": len(new_rows),
            "alterados": len(changed),
            "iguais": len(unchanged),
            "ausentes_arquivo": len(absent_rows),
        },
        "mudancas": {
            "novos": new_rows,
            "alterados": changed,
            "iguais": unchanged,
            "ausentes_arquivo": absent_rows,
        },
        "aviso_ausentes": "Ausência na planilha atual não confirma que o chamado foi resolvido.",
    }


def _safe_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool, date, datetime)):
        result = str(value) if isinstance(value, (str, int, float, bool)) else value.isoformat()
    else:
        result = ", ".join(str(item) for item in value) if isinstance(value, list) else str(value)
    if isinstance(result, str) and result.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + result
    return result


def build_ticket_workbook(
    snapshot: Dict[str, Any],
    tickets: Sequence[Dict[str, Any]],
    filters: Dict[str, Any],
    *,
    timelines: Optional[Mapping[str, Sequence[Dict[str, Any]]]] = None,
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"

    def append_summary_row(label: str, value: Any) -> None:
        # Escape external text without changing native numbers or dates.
        summary.append([label, _safe_cell(value) if isinstance(value, str) else value])

    summary.append(["Acompanhamento de Chamados"])
    append_summary_row("Versão importada", snapshot.get("atualizado_em", ""))
    sources = snapshot.get("fontes") or {}
    append_summary_row("Arquivo Simpress", (sources.get("simpress") or {}).get("arquivo", ""))
    append_summary_row("Arquivo Oracle", (sources.get("oracle") or {}).get("arquivo", ""))
    append_summary_row("Séries no resultado", len(tickets))
    for label, key in (("Relação", "relation"), ("Status Oracle", "status"), ("Fila Oracle", "queue"), ("Cidade", "city"), ("Atendimento", "provider"), ("Atualização", "age_filter"), ("Busca", "search"), ("Somente atenção (3 dias)", "attention_only")):
        value = filters.get(key, "")
        if value not in (None, "", False):
            append_summary_row(label, str(value))
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 48

    sheet = workbook.create_sheet("Chamados")
    headers = [
        "Série", "Atenção", "Relação", "Pedidos Simpress", "OV Simpress", "OS Simpress", "OS Oracle", "Chamado do cliente", "Status Simpress", "Data do pedido",
        "Liberação do pedido", "Verificado e estoque", "Pedido em separação", "Nota fiscal emitida", "Etapa atual", "Modelo",
        "Descrição do produto", "Cidade de entrega", "Endereço de entrega", "Previsão de entrega",
        "Solicitação Oracle", "Logradouro Oracle", "Número do endereço Oracle", "Bairro Oracle", "Cidade Oracle",
        "Atendido pela inibição Oracle", "Nome da filial Oracle", "Status Oracle", "Sub-status Oracle", "Fila Oracle",
        "Fase Oracle", "Última atualização Oracle", "Status entregas", "Status atividades", "Assunto",
        "Descrição do chamado Oracle", "Série na Simpress", "Série no Oracle", "Origem do pedido", "Origem da OS", "Revisão da relação", "Motivo da revisão",
    ]
    sheet.append(headers)
    for ticket in tickets:
        oracle_rows = sorted(
            (item for item in ticket.get("oracle", []) if isinstance(item, dict)),
            key=lambda item: _order_service_descending_key(item.get("ordem_servico", "")),
            reverse=True,
        ) or [{}]
        simpress = ticket.get("simpress") or {}
        simpress_os = sorted(
            (value for value in simpress.get("ordens_servico", []) if value),
            key=_order_service_descending_key,
            reverse=True,
        )
        attention = ticket_attention(ticket, stale_days=3)
        provenance = ticket_relation_provenance(ticket)
        alerts = ticket_relation_alerts(ticket)
        for oracle in oracle_rows:
            values = [
                ticket.get("serie", ""), attention["motivo"], ticket.get("relacao", ""),
                simpress.get("pedidos", []), simpress.get("ordens_venda", []), simpress_os, oracle.get("ordem_servico", ""),
                simpress.get("chamados_cliente", []),
                simpress.get("status", []), simpress.get("datas_pedido", []),
                ticket_stage_date(ticket, "liberacao_pedido"), ticket_stage_date(ticket, "verificado_estoque"),
                ticket_stage_date(ticket, "pedido_separacao"), ticket_stage_date(ticket, "nota_fiscal_emitida"),
                (simpress.get("etapa_atual") or {}).get("rotulo", "") if isinstance(simpress.get("etapa_atual"), dict) else "",
                simpress.get("modelos", []),
                simpress.get("descricoes_produto", []), simpress.get("cidades_entrega", []),
                simpress.get("enderecos_entrega", []) if simpress.get("pedido_em_andamento") else [],
                simpress.get("previsoes_entrega", []),
                oracle.get("solicitacao", ""), oracle.get("logradouro", ""),
                oracle.get("numero_endereco", ""), oracle.get("bairro", ""), oracle.get("cidade", ""),
                oracle.get("atendido_pela_inibicao", ""), oracle.get("nome_filial", ""), oracle.get("status", ""),
                oracle.get("sub_status", ""), oracle.get("fila", ""), oracle.get("fase_atendimento", ""),
                oracle.get("ultima_atualizacao", ""), oracle.get("status_entregas", ""),
                oracle.get("status_atividades", ""), oracle.get("assunto", ""), oracle.get("descricao", ""),
                "Simpress" if provenance["serie"]["simpress"] else "",
                "Oracle" if provenance["serie"]["oracle"] else "",
                "Simpress: pedido · Oracle: solicitação",
                "Simpress / Oracle" if provenance["ordem_servico"]["oracle"] else "Simpress" if simpress_os else "Oracle",
                "Sim" if alerts else "Não",
                " · ".join(item["mensagem"] for item in alerts),
            ]
            sheet.append([_safe_cell(value) for value in values])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="166534")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate((18, 36, 20, 24, 22, 20, 18, 24, 24, 22, 22, 22, 22, 22, 24, 28, 42, 24, 48, 20, 22, 30, 18, 24, 24, 30, 30, 24, 24, 24, 24, 22, 22, 40, 64, 64, 16, 16, 36, 24, 18, 72), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = sheet.sheet_format.defaultRowHeight
    description_column = headers.index("Descrição do chamado Oracle") + 1
    for row in sheet.iter_rows(min_row=2, min_col=description_column, max_col=description_column):
        row[0].alignment = Alignment(vertical="top", wrap_text=False)

    review_sheet = workbook.create_sheet("Conferência das relações")
    review_headers = [
        "Série", "Relação", "Série na Simpress", "Série no Oracle", "Pedido Simpress",
        "Solicitação Oracle", "OS Simpress", "OS Oracle", "Revisão necessária", "Motivo da revisão",
    ]
    review_sheet.append(review_headers)
    for ticket in tickets:
        provenance = ticket_relation_provenance(ticket)
        simpress_pedidos = provenance["pedido"]["simpress"]
        oracle_solicitations = provenance["pedido"]["oracle_solicitacao"]
        simpress_orders = provenance["ordem_servico"]["simpress"]
        oracle_orders = provenance["ordem_servico"]["oracle"]
        alerts = ticket_relation_alerts(ticket)
        review_sheet.append([_safe_cell(value) for value in (
            ticket.get("serie", ""), ticket.get("relacao", ""),
            "Simpress" if provenance["serie"]["simpress"] else "",
            "Oracle" if provenance["serie"]["oracle"] else "",
            simpress_pedidos, oracle_solicitations, simpress_orders, oracle_orders,
            "Sim" if alerts else "Não", " · ".join(item["mensagem"] for item in alerts),
        )])
    review_sheet.freeze_panes = "A2"
    review_sheet.auto_filter.ref = review_sheet.dimensions
    review_sheet.row_dimensions[1].height = 30
    for cell in review_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="9A3412")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate((20, 22, 18, 18, 28, 32, 24, 24, 20, 90), start=1):
        review_sheet.column_dimensions[get_column_letter(index)].width = width

    timeline_sheet = workbook.create_sheet("Linha do tempo")
    timeline_headers = [
        "Série", "Importação", "Evento", "Pedido Simpress", "Solicitação Oracle", "OV Simpress", "OS Simpress",
        "OS Oracle", "Status Simpress", "Status Oracle", "Fila Oracle", "Última atualização Oracle", "Anotação",
    ]
    timeline_sheet.append(timeline_headers)
    for ticket in tickets:
        serial = str(ticket.get("serie") or "")
        for event in (timelines or {}).get(serial, []):
            state = event.get("estado") or {}
            timeline_sheet.append([_safe_cell(value) for value in (
                serial, event.get("atualizado_em", ""), event.get("tipo", ""),
                state.get("pedidos_simpress", []), state.get("solicitacoes_oracle", []), state.get("ordens_venda_simpress", []),
                state.get("ordens_servico_simpress", []), state.get("ordens_servico_oracle", []),
                state.get("status_simpress", []), state.get("status_oracle", []),
                state.get("filas_oracle", []), state.get("atualizacoes_oracle", []), annotation_export_text(event),
            )])
    timeline_sheet.freeze_panes = "A2"
    timeline_sheet.auto_filter.ref = timeline_sheet.dimensions
    timeline_sheet.row_dimensions[1].height = 30
    for cell in timeline_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155E75")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate((20, 24, 24, 28, 32, 24, 24, 24, 36, 36, 32, 32, 64), start=1):
        timeline_sheet.column_dimensions[get_column_letter(index)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
