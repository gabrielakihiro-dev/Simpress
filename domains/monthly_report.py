"""Monthly report for the contract manager, built only from data the application already keeps.

See docs/operacao/relatorio-mensal.md for where each number comes from.
"""

from __future__ import annotations

import calendar
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
import re
from typing import Any, Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.config import SUPPORTED_CONTRACT, SUPPORTED_CONTRACT_LEGAL_NAME
from core.errors import DomainValidationError
from domains.history_insights import is_restoration, PINNED_BASELINE_ID
from domains.inventory_comparison import (
    filter_curitiba_hostname_location_mismatches,
    filter_sebrae_owned_unmatched_items,
)
from domains.ticket_reporting import (
    _is_closed,
    _parse_date,
    normalize_ticket_snapshot,
    ticket_report_snapshots,
    ticket_stage_date,
)
from domains.usage_metrics import METRICS, summarize

MONTH_NAMES = (
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
)
USAGE_LABELS = {
    "transferencias": "Transferências geradas",
    "comparacoes": "Comparações feitas",
    "exportacoes": "Exportações",
    "os_conciliadas": "OS conciliadas",
    "buscas": "Buscas",
}
SECTION_ENDPOINT = "Sem correspondência — no endpoint"
SECTION_SIMPRESS = "Sem correspondência — na Simpress"
SECTION_LOCATION = "Localidade divergente"
OPENED = "Aberta no mês"
CLOSED = "Concluída no mês"
PENDING = "Pendente no fim do mês"
NO_RECORDS = "Sem registros no mês"
_SIMPRESS_FINISHED_TOKENS = ("entreg", "cancel", "conclu", "finaliz", "encerr")


class InvalidReportMonthError(DomainValidationError):
    default_status_code = 422


@dataclass(frozen=True)
class ReportMonth:
    year: int
    month: int

    @property
    def start(self) -> date:
        return date(self.year, self.month, 1)

    @property
    def end(self) -> date:
        return date(self.year, self.month, calendar.monthrange(self.year, self.month)[1])

    @property
    def key(self) -> str:
        return f"{self.year:04d}-{self.month:02d}"

    @property
    def label(self) -> str:
        return f"{MONTH_NAMES[self.month - 1]} de {self.year}"

    def contains(self, value: Optional[date]) -> bool:
        return value is not None and self.start <= value <= self.end


def parse_report_month(value: Any, today: date) -> ReportMonth:
    """Accept AAAA-MM up to the current month."""
    match = re.fullmatch(r"(\d{4})-(\d{2})", str(value or "").strip())
    if not match or not 1 <= int(match.group(2)) <= 12 or int(match.group(1)) < 2000:
        raise InvalidReportMonthError("Selecione um mês válido no formato AAAA-MM.")
    month = ReportMonth(int(match.group(1)), int(match.group(2)))
    if month.start > today.replace(day=1):
        raise InvalidReportMonthError("Selecione o mês atual ou um mês anterior.")
    return month


# Inventory ---------------------------------------------------------------

def inventory_section(equipment: Iterable[Mapping[str, Any]], stats: Mapping[str, Any]) -> Dict[str, Any]:
    """Counts from the official EquipamentosDetalhes cache (one entry per unique serial)."""
    items = list(equipment)
    by_type = Counter(str(item.get("tipo") or "").strip() or "Sem tipo" for item in items)
    by_unit = Counter(
        (str(item.get("retirada_cidade") or "").strip() or "Sem cidade",
         str(item.get("endereco_instalacao") or "").strip())
        for item in items
    )
    raw_duplicates = stats.get("series_duplicadas")
    duplicates = raw_duplicates if isinstance(raw_duplicates, dict) else {}
    empty_rows = stats.get("linhas_sem_serie")
    return {
        "total": len(items),
        "unidades": len(by_unit),
        "por_tipo": sorted(by_type.items(), key=lambda entry: (-entry[1], entry[0].casefold())),
        "por_unidade": sorted(
            ((city, address, count) for (city, address), count in by_unit.items()),
            key=lambda entry: (-entry[2], entry[0].casefold(), entry[1].casefold()),
        ),
        "linhas_sem_serie": empty_rows if isinstance(empty_rows, int) else 0,
        "series_duplicadas": sorted(str(serial) for serial in duplicates),
        "arquivo_origem": str(stats.get("arquivo_origem") or ""),
        "carregada_em": str(stats.get("ultima_atualizacao") or ""),
    }


# Comparison --------------------------------------------------------------

def _event_local_date(event: Mapping[str, Any]) -> Optional[datetime]:
    try:
        parsed = datetime.fromisoformat(str(event.get("horario") or "").replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return None
    return parsed.astimezone().replace(tzinfo=None) if parsed.tzinfo else parsed


def comparison_section(
    history: Sequence[Any],
    month: ReportMonth,
    display: Callable[[Dict[str, Any]], Dict[str, Any]],
) -> Dict[str, Any]:
    """Latest inventory update up to the end of the month, with the current history rules."""
    updates: List[Tuple[datetime, Dict[str, Any]]] = []
    for event in history:
        if not isinstance(event, dict) or is_restoration(event) or event.get("id") == PINNED_BASELINE_ID:
            continue
        moment = _event_local_date(event)
        if moment is not None and moment.date() <= month.end:
            updates.append((moment, event))
    in_month = sum(1 for moment, _event in updates if month.contains(moment.date()))
    if not updates:
        return {"horario": None, "atualizacoes_no_mes": 0, "linhas": [], "contagens": {
            SECTION_ENDPOINT: 0, SECTION_SIMPRESS: 0, SECTION_LOCATION: 0,
        }}
    moment, latest = max(updates, key=lambda entry: entry[0])
    view = display(latest).get("comparacao") or {}

    def rows(key: str) -> List[Dict[str, Any]]:
        value = view.get(key)
        return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []

    # The display view already applies these rules; repeating them keeps the report safe on its own.
    sections = (
        (SECTION_ENDPOINT, filter_sebrae_owned_unmatched_items(rows("ausentes_no_endpoint"))),
        (SECTION_SIMPRESS, filter_sebrae_owned_unmatched_items(rows("ausentes_em_detalhes"))),
        (SECTION_LOCATION, filter_curitiba_hostname_location_mismatches(rows("divergencias_local"))),
    )
    lines = [
        {
            "secao": section,
            "hostname": str(item.get("computador") or ""),
            "serie": str(item.get("serie") or ""),
            "local_previsto": str(item.get("local_previsto") or ""),
            "local_simpress": str(item.get("local_detalhes") or ""),
        }
        for section, items in sections
        for item in items
    ]
    return {
        "horario": moment,
        "atualizacoes_no_mes": in_month,
        "linhas": lines,
        "contagens": {section: len(items) for section, items in sections},
    }


# Tickets -----------------------------------------------------------------

def _key(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper()


def _latest_date(values: Any) -> Optional[date]:
    candidates = values if isinstance(values, list) else [values]
    parsed = [found for found in (_parse_date(value) for value in candidates) if found is not None]
    return max(parsed) if parsed else None


def _month_situation(month: ReportMonth, opened: Optional[date], closed: Optional[date], finished: bool) -> List[str]:
    situation = []
    if month.contains(opened):
        situation.append(OPENED)
    if month.contains(closed):
        situation.append(CLOSED)
    open_at_end = opened is not None and opened <= month.end and (
        (closed is None and not finished) or (closed is not None and closed > month.end)
    )
    if open_at_end:
        situation.append(PENDING)
    return situation


def tickets_section(report: Mapping[str, Any], month: ReportMonth) -> Dict[str, Any]:
    """Oracle OS and Simpress orders by date; the newest import wins for a repeated identifier."""
    oracle: Dict[str, Dict[str, Any]] = {}
    simpress: Dict[str, Dict[str, Any]] = {}
    snapshots = [normalize_ticket_snapshot(item) for item in ticket_report_snapshots(dict(report))]
    for snapshot in snapshots:
        for ticket in snapshot.get("chamados") or []:
            if not isinstance(ticket, dict):
                continue
            serial = str(ticket.get("serie") or "")
            for item in ticket.get("oracle") or []:
                if not isinstance(item, dict):
                    continue
                identifier = str(item.get("ordem_servico") or item.get("solicitacao") or "").strip()
                key = _key(identifier) or f"{_key(serial)}|{item.get('data_criacao') or ''}"
                oracle[key] = {
                    "fonte": "Oracle",
                    "identificador": identifier,
                    "serie": str(item.get("serie") or serial),
                    "status": " / ".join(
                        part for part in (str(item.get("status") or ""), str(item.get("sub_status") or "")) if part
                    ),
                    "criacao": _parse_date(item.get("data_criacao")),
                    "conclusao": _parse_date(item.get("data_conclusao")),
                    "encerrada": _is_closed(item),
                }
            data = ticket.get("simpress")
            if not isinstance(data, dict) or not (data.get("pedidos") or data.get("datas_pedido")):
                continue
            orders = data.get("pedidos") if isinstance(data.get("pedidos"), list) else []
            identifier = str(orders[-1] if orders else "").strip()
            opened = _parse_date(ticket_stage_date(ticket, "data_pedido")) or _latest_date(data.get("datas_pedido"))
            key = _key(identifier) or f"{_key(serial)}|{opened or ''}"
            statuses = [str(value) for value in data.get("status") or [] if value]
            status_text = " ".join(statuses).casefold()
            simpress[key] = {
                "fonte": "Simpress",
                "identificador": identifier,
                "serie": serial,
                "status": ", ".join(statuses),
                "criacao": opened,
                "conclusao": _parse_date(ticket_stage_date(ticket, "entregue")),
                "encerrada": any(token in status_text for token in _SIMPRESS_FINISHED_TOKENS),
            }

    def count(records: Mapping[str, Dict[str, Any]]) -> Dict[str, int]:
        counts = {"abertas": 0, "concluidas": 0, "pendentes": 0, "encerradas_sem_data": 0}
        for record in records.values():
            record["situacao"] = _month_situation(month, record["criacao"], record["conclusao"], record["encerrada"])
            counts["abertas"] += OPENED in record["situacao"]
            counts["concluidas"] += CLOSED in record["situacao"]
            counts["pendentes"] += PENDING in record["situacao"]
            counts["encerradas_sem_data"] += bool(
                record["encerrada"] and record["conclusao"] is None
                and record["criacao"] is not None and record["criacao"] <= month.end
            )
        return counts

    oracle_counts = count(oracle)
    simpress_counts = count(simpress)
    detail = sorted(
        (record for record in [*oracle.values(), *simpress.values()] if record["situacao"]),
        key=lambda record: (record["fonte"], record["criacao"] or date.min, record["identificador"]),
    )
    return {
        "oracle": oracle_counts,
        "simpress": simpress_counts,
        "linhas": detail,
        "importacoes": len(snapshots),
        "ultima_importacao": str(report.get("atualizado_em") or ""),
    }


# Usage -------------------------------------------------------------------

def usage_section(store: Mapping[str, Any], month: ReportMonth) -> Dict[str, Any]:
    summary = summarize(store, month.end, (month.end - month.start).days + 1)
    return {
        "totais": summary["totais"],
        "por_dia": summary["por_dia"],
        "registro_desde": summary["registro_desde"],
        "tempos_configurados": summary["tempos_configurados"],
        "horas_economizadas": summary["horas_economizadas"],
    }


# Report ------------------------------------------------------------------

def build_monthly_report(
    month: ReportMonth,
    *,
    equipment: Iterable[Mapping[str, Any]],
    stats: Mapping[str, Any],
    comparison_history: Sequence[Any],
    comparison_display: Callable[[Dict[str, Any]], Dict[str, Any]],
    ticket_report: Mapping[str, Any],
    usage_store: Mapping[str, Any],
    generated_at: datetime,
) -> Dict[str, Any]:
    return {
        "mes": month,
        "gerado_em": generated_at,
        "inventario": inventory_section(equipment, stats),
        "comparacao": comparison_section(comparison_history, month, comparison_display),
        "os": tickets_section(ticket_report, month),
        "uso": usage_section(usage_store, month),
    }


def _date_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, str) and value:
        parsed = _parse_date(value)
        return parsed.strftime("%d/%m/%Y") if parsed else value
    return ""


def summary_rows(report: Mapping[str, Any]) -> List[Tuple[str, str, Any, str]]:
    """(bloco, indicador, valor, origem) shown on the Resumo sheet."""
    inventory = report["inventario"]
    comparison = report["comparacao"]
    tickets = report["os"]
    usage = report["uso"]
    base = f"Base oficial EquipamentosDetalhes ({inventory['arquivo_origem'] or 'arquivo não informado'})"
    oracle, simpress = tickets["oracle"], tickets["simpress"]
    rows: List[Tuple[str, str, Any, str]] = [
        ("Inventário", "Equipamentos na base", inventory["total"], f"{base}: séries únicas."),
        ("Inventário", "Unidades com equipamentos", inventory["unidades"], "Cidade + endereço de instalação distintos."),
        ("Inventário", "Linhas sem número de série", inventory["linhas_sem_serie"], "Não entram no total."),
        ("Inventário", "Séries duplicadas", len(inventory["series_duplicadas"]),
         ", ".join(inventory["series_duplicadas"]) or "Nenhuma; cada série conta uma vez."),
        ("Última comparação", "Data da comparação", _date_text(comparison["horario"]) or NO_RECORDS,
         "Atualização de inventário mais recente até o fim do mês."),
        ("Última comparação", SECTION_ENDPOINT, comparison["contagens"][SECTION_ENDPOINT], "Histórico de planilhas."),
        ("Última comparação", SECTION_SIMPRESS, comparison["contagens"][SECTION_SIMPRESS],
         "Histórico de planilhas, sem os notebooks próprios do Sebrae."),
        ("Última comparação", SECTION_LOCATION, comparison["contagens"][SECTION_LOCATION],
         "Histórico de planilhas; PRNOTESJP, PRNOTERL e PRNOTEARC contam como Curitiba."),
        ("Última comparação", "Atualizações registradas no mês", comparison["atualizacoes_no_mes"], "Histórico de planilhas."),
        ("OS do mês", "OS Oracle abertas no mês", oracle["abertas"], "Data de criação no Oracle."),
        ("OS do mês", "OS Oracle concluídas no mês", oracle["concluidas"], "Data de conclusão no Oracle."),
        ("OS do mês", "OS Oracle pendentes no fim do mês", oracle["pendentes"], "Criadas até o fim do mês e sem conclusão até ele."),
        ("OS do mês", "Pedidos Simpress abertos no mês", simpress["abertas"], "Data do pedido na Simpress."),
        ("OS do mês", "Pedidos Simpress entregues no mês", simpress["concluidas"], "Etapa Pedido entregue na Simpress."),
        ("OS do mês", "Pedidos Simpress em andamento no fim do mês", simpress["pendentes"], "Feitos até o fim do mês e sem entrega até ele."),
        ("OS do mês", "Encerrados sem data", oracle["encerradas_sem_data"] + simpress["encerradas_sem_data"],
         "Status encerrado sem data de conclusão ou entrega; fora das contagens do mês."),
    ]
    rows.extend(
        ("Uso no mês", USAGE_LABELS[metric], usage["totais"][metric], "data/state/metricas_uso.json")
        for metric in METRICS
    )
    rows.append((
        "Uso no mês", "Horas economizadas (estimativa)",
        usage["horas_economizadas"] if usage["tempos_configurados"] else "Tempos manuais não informados",
        "Σ(total × minutos informados) / 60.",
    ))
    rows.append(("Uso no mês", "Registrado desde", _date_text(usage["registro_desde"]) or NO_RECORDS, "Primeiro dia com registro."))
    return rows


_HEADER_FILL = PatternFill("solid", fgColor="10213A")
_BLOCK_FILL = PatternFill("solid", fgColor="E6F4FA")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _safe(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value)[:32767]
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def _table(sheet: Any, headers: Sequence[str], rows: Iterable[Sequence[Any]], empty: str = NO_RECORDS) -> None:
    sheet.append(list(headers))
    for cell in sheet[sheet.max_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    count = 0
    for row in rows:
        sheet.append([_safe(value) for value in row])
        count += 1
    if not count:
        sheet.append([empty])
    sheet.freeze_panes = "A2"
    _fit(sheet)


def _fit(sheet: Any) -> None:
    for column in sheet.columns:
        width = max((len(str(cell.value or "")) for cell in column), default=8)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 2, 10), 70)


def report_filename(month: ReportMonth) -> str:
    return f"Relatorio-mensal-{month.key}.xlsx"


def build_monthly_report_workbook(report: Mapping[str, Any]) -> Tuple[BytesIO, str]:
    month: ReportMonth = report["mes"]
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary.append([f"Relatório mensal — {month.label}"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Contrato {SUPPORTED_CONTRACT} - {SUPPORTED_CONTRACT_LEGAL_NAME}"])
    summary.append([f"Gerado em {_date_text(report['gerado_em'])}"])
    inventory = report["inventario"]
    summary.append([f"Inventário: base atual ({inventory['arquivo_origem'] or 'sem arquivo'}, carregada em {inventory['carregada_em'] or '-'})."])
    tickets = report["os"]
    summary.append([
        f"OS: {tickets['importacoes']} importação(ões) no histórico; última em {_date_text(tickets['ultima_importacao']) or '-'}."
    ])
    summary.append([])
    summary.append(["Indicador", "Valor", "Origem"])
    for cell in summary[summary.max_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    current_block = None
    for block, label, value, origin in summary_rows(report):
        if block != current_block:
            summary.append([block])
            summary[summary.max_row][0].font = Font(bold=True)
            for cell in summary[summary.max_row][:3]:
                cell.fill = _BLOCK_FILL
            current_block = block
        summary.append([_safe(label), _safe(value), _safe(origin)])
        summary[summary.max_row][1].alignment = Alignment(horizontal="right")
    summary.column_dimensions["A"].width = 46
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 80

    _table(workbook.create_sheet("Inventário por tipo"), ("Tipo", "Equipamentos"), inventory["por_tipo"])
    _table(workbook.create_sheet("Inventário por unidade"), ("Cidade", "Endereço", "Equipamentos"), inventory["por_unidade"])
    _table(
        workbook.create_sheet("Comparação"),
        ("Seção", "Hostname", "Série", "Local previsto", "Local na Simpress"),
        ((line["secao"], line["hostname"], line["serie"], line["local_previsto"], line["local_simpress"])
         for line in report["comparacao"]["linhas"]),
        "Sem pendências na última comparação.",
    )
    _table(
        workbook.create_sheet("OS do mês"),
        ("Fonte", "OS/Pedido", "Série", "Status", "Criação", "Conclusão", "Situação no mês"),
        ((record["fonte"], record["identificador"], record["serie"], record["status"],
          _date_text(record["criacao"]), _date_text(record["conclusao"]), "; ".join(record["situacao"]))
         for record in tickets["linhas"]),
    )
    _table(
        workbook.create_sheet("Uso por dia"),
        ("Data", *(USAGE_LABELS[metric] for metric in METRICS)),
        ((_date_text(date.fromisoformat(row["data"])), *(row[metric] for metric in METRICS))
         for row in report["uso"]["por_dia"]),
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output, report_filename(month)
