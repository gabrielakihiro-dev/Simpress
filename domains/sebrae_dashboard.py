"""Pure rules of the Dashboard-import flow (docs/operacao/sincronizar-chamados-sebrae.md, "Modelo atual").

The user imports the Dashboard ``.xlsx``; column C tells which SoftExpert tickets are already logged.
New tickets become rows the user pastes back: nothing here writes to Excel. No browser, no network.
Ticket content is data: nothing read from a ticket is ever interpreted as an instruction.
"""

from __future__ import annotations

import io
import re
import unicodedata
import warnings
import zipfile
from dataclasses import dataclass
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from core.config import MAX_XLSX_ARCHIVE_ENTRIES, MAX_XLSX_UNCOMPRESSED_BYTES
from domains.equipment_search import normalize_search_term
from domains.inventory_comparison import SEBRAE_OWNED_NOTEBOOK_SERIALS
from domains.sebrae_chamados import (
    ChamadoSoftExpert,
    PENDENCIAS_DEFINIDAS,
    UF,
    UNIDADE,
    FORMATO_DATA,
    detalhes_atendimento_do_chamado,
    interpretar_data,
    mascarar_cpfs,
    numeros_do_dashboard,
    triagem,
)

ABA = "CHAMADOS"
LINHAS_VAZIAS_PARA_PARAR = 5_000  # the real table has ~480 formula-only rows after the last ticket

# Header of each column of Tabela3 (A:W) since column G (Título) was added on 28/10/2026.
CABECALHOS: Dict[str, str] = {
    "A": "RESPONSAVEL", "B": "CATEGORIA", "C": "CHAMADO", "D": "STATUS CHAMADO", "E": "OS",
    "F": "ATENDIMENTO REALIZADO", "G": "TITULO", "H": "DESCRICAO CHAMADO", "I": "NOME USUARIO",
    "J": "TIPO DO CHAMADO", "K": "SERIE", "L": "MODELO", "M": "UNIDADE", "N": "CIDADE", "O": "UF",
    "P": "DATA ABERTURA", "Q": "MES", "R": "DT ENCERRAMENTO", "S": "TIPO SLA", "T": "DIAS ABERTOS",
    "U": "AGING", "V": "HR CONSUMIDA", "W": "PRODUZINDO",
}
MANUAIS = ("A", "B", "D", "J", "R", "S")
FORMULAS = ("Q", "T", "U", "V")  # the site never writes them: a paste would erase the formula
INTOCADAS = ("W",)

MOTIVO_ALERTA_ENDPOINT = "Alerta automático do Endpoint; fica pendente até o usuário classificar."
_ALERTA_ENDPOINT = re.compile(r"endpoint\s+central\s*-", re.IGNORECASE)

CRUZ_UNICO = "unico"
CRUZ_AUSENTE = "ausente"
CRUZ_AMBIGUO = "ambiguo"


class DashboardInvalido(ValueError):
    """The workbook is not the expected Dashboard; the message is meant for the user."""


# ---------------------------------------------------------------------------- import

def _sem_acentos_maiusculas(texto: Any) -> str:
    decomposto = unicodedata.normalize("NFKD", str(texto or ""))
    return "".join(c for c in decomposto if not unicodedata.combining(c)).strip().upper()


def conferir_cabecalhos(cabecalho: Sequence[Any]) -> None:
    """Refuses a Dashboard whose A:V headers differ from ``CABECALHOS`` (for example an export without column G)."""
    diferentes = []
    for indice, (letra, esperado) in enumerate(CABECALHOS.items()):
        if letra == "W":
            continue  # PRODUZINDO is optional in older exports
        lido = _sem_acentos_maiusculas(cabecalho[indice] if indice < len(cabecalho) else "")
        if lido != esperado:
            diferentes.append(f"{letra}: esperado '{esperado.title()}', veio '{lido.title()}'")
    if diferentes:
        raise DashboardInvalido(
            "A aba CHAMADOS não tem as colunas esperadas (" + "; ".join(diferentes[:3]) + "). "
            "Exporte o Dashboard atual, com a coluna Título em G."
        )


def _validar_pacote(arquivo: zipfile.ZipFile) -> None:
    """ZIP-level checks only. The per-sheet row limit of ``validate_xlsx_archive`` would reject the real
    Dashboard (other tabs have formatted empty cells down to row 1,048,576); the sheet is read as a stream."""
    entradas = arquivo.infolist()
    nomes = {entrada.filename for entrada in entradas}
    if not entradas or len(entradas) > MAX_XLSX_ARCHIVE_ENTRIES or len(nomes) != len(entradas):
        raise DashboardInvalido("O arquivo XLSX tem uma estrutura ZIP inválida ou grande demais.")
    if sum(entrada.file_size for entrada in entradas) > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise DashboardInvalido("O conteúdo descompactado da planilha excede o limite permitido.")
    if "[Content_Types].xml" not in nomes:
        raise DashboardInvalido("O arquivo enviado não é uma planilha .xlsx válida.")


def ler_coluna_c(conteudo: bytes) -> frozenset[str]:
    """Ticket numbers of column C of the ``CHAMADOS`` sheet (hidden rows included), headers checked first."""
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as arquivo:
            _validar_pacote(arquivo)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # openpyxl warns about the Dashboard's validation extension
            livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError, ValueError) as erro:
        if isinstance(erro, DashboardInvalido):
            raise
        raise DashboardInvalido("O arquivo enviado não é uma planilha .xlsx válida.") from erro
    try:
        if ABA not in livro.sheetnames:
            raise DashboardInvalido(f"A planilha não tem a aba {ABA}.")
        linhas = livro[ABA].iter_rows(min_row=1, max_col=len(CABECALHOS), values_only=True)
        conferir_cabecalhos(next(linhas, ()))
        valores: List[Any] = []
        vazias = 0
        for linha in linhas:
            valor = linha[2] if len(linha) > 2 else None
            if valor in (None, ""):
                vazias += 1
                if vazias > LINHAS_VAZIAS_PARA_PARAR:
                    break  # the table ends long before the formatted-but-empty rows of the sheet
                continue
            vazias = 0
            valores.append(valor)
        numeros = numeros_do_dashboard(valores)
    finally:
        livro.close()
    if not numeros:
        raise DashboardInvalido("A coluna C não trouxe nenhum número UTIC; a planilha parece vazia ou incorreta.")
    return numeros


# ---------------------------------------------------------------------------- triage

@dataclass(frozen=True)
class NovosChamados:
    novos: Tuple[str, ...]
    pendentes: Tuple[Tuple[str, str], ...]
    ignorados: Tuple[str, ...]


def triar_novos(titulos_softexpert: Mapping[str, str], existentes: Iterable[str]) -> NovosChamados:
    """SoftExpert tickets minus column C; known Endpoint alerts (by list or by title) stay pending."""
    resultado = triagem(titulos_softexpert.keys(), existentes)
    novos: List[str] = []
    pendentes = list(resultado.pendentes)
    for numero in resultado.a_extrair:
        if _ALERTA_ENDPOINT.search(titulos_softexpert.get(numero, "")):
            pendentes.append((numero, MOTIVO_ALERTA_ENDPOINT))
        else:
            novos.append(numero)
    return NovosChamados(tuple(novos), tuple(pendentes), resultado.ignorados)


# ---------------------------------------------------------------------------- cross-reference by name

@dataclass(frozen=True)
class Cruzamento:
    status: str
    serie: str = ""
    modelo: str = ""
    cidade: str = ""
    candidatos: int = 0


def _chave_serie(serie: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", normalize_search_term(serie))


def _mesma_serie(a: str, b: str, macbook: bool) -> bool:
    """Equal series; for MacBooks the Endpoint series has no leading ``S`` that EquipamentosDetalhes adds."""
    return a == b or (macbook and (a == "S" + b or b == "S" + a))


def cruzar_por_nome(nome: str, registros: Iterable[Mapping[str, Any]]) -> Cruzamento:
    """Series/model/city of the person's one equipment; absent or ambiguous never picks a candidate.

    ``registros`` come from EquipamentosDetalhes, InvAll and Entra already normalised to the keys
    ``nome``, ``serie``, ``modelo`` and ``cidade``. Sebrae-owned series are ignored. A MacBook series
    without the leading ``S`` equals the same series with it, and counts as one candidate.
    """
    alvo = normalize_search_term(" ".join(str(nome or "").split()))
    if not alvo:
        return Cruzamento(CRUZ_AUSENTE)
    achados: Dict[str, Dict[str, str]] = {}
    for registro in registros:
        if normalize_search_term(" ".join(str(registro.get("nome") or "").split())) != alvo:
            continue
        chave = _chave_serie(registro.get("serie"))
        if not chave or chave in SEBRAE_OWNED_NOTEBOOK_SERIALS:
            continue
        mac = "MACBOOK" in normalize_search_term(registro.get("modelo"))
        equivalente = next((k for k in achados if _mesma_serie(k, chave, mac or bool(achados[k]["mac"]))), chave)
        atual = achados.setdefault(equivalente, {"serie": "", "modelo": "", "cidade": "", "mac": ""})
        if mac:
            atual["mac"] = "MACBOOK"
        for campo in ("serie", "modelo", "cidade"):  # first non-empty value wins; sources are walked in priority order
            if not atual[campo] and str(registro.get(campo) or "").strip():
                atual[campo] = str(registro[campo]).strip()
    if not achados:
        return Cruzamento(CRUZ_AUSENTE)
    if len(achados) > 1:
        return Cruzamento(CRUZ_AMBIGUO, candidatos=len(achados))
    unico = next(iter(achados.values()))
    return Cruzamento(CRUZ_UNICO, unico["serie"], unico["modelo"], unico["cidade"], 1)


# ---------------------------------------------------------------------------- row

OMITIDO = "[DADO SENSÍVEL OMITIDO]"  # same wording the Dashboard already uses for passwords
_EMAIL = re.compile(r"^[\w.+-]+@[\w-]+(?:\.[\w-]+)+$")
_SENHA_PADRAO = re.compile(r"(?i)\bsebrae@\d+\b")  # initial password pattern seen in real tickets
_SENHA_ROTULADA = re.compile(r"(?i)\b(senha|password|pwd)(\s*[:=-]\s*)(\S+)")


def omitir_credenciais(texto: Any) -> str:
    """Hides passwords: the value after ``senha:`` and the lone token on the line after a lone e-mail.

    shortcut: a password written next to other words on its own line is not detected; widen the
    patterns if a ticket leaks one that way.
    """
    saida: List[str] = []
    apos_email = False
    for linha in _SENHA_PADRAO.sub(OMITIDO, str(texto or "")).splitlines():
        palavra = linha.strip()
        if not palavra:
            saida.append(linha)
            continue
        if apos_email and len(palavra.split()) == 1 and len(palavra) >= 6 and not _EMAIL.match(palavra):
            saida.append(OMITIDO)
        else:
            saida.append(_SENHA_ROTULADA.sub(lambda m: m.group(1) + m.group(2) + OMITIDO, linha))
        apos_email = bool(_EMAIL.match(palavra)) or (apos_email and not palavra)
    return "\n".join(saida)


def _texto_de_celula(valor: Any) -> str:
    """One line, CPF and passwords masked; a leading ``= + - @`` gets an apostrophe so Excel never reads it as a formula."""
    texto = mascarar_cpfs(" ".join(omitir_credenciais(valor).split()))
    return "'" + texto if texto and texto[0] in "=+-@" else texto


def montar_linha(chamado: ChamadoSoftExpert, cruzamento: Optional[Cruzamento] = None, os_da_serie: str = "") -> Dict[str, str]:
    """Automatic columns of the Dashboard row; manual, formula and untouched columns are not in the result."""
    cruzamento = cruzamento or Cruzamento(CRUZ_AUSENTE)
    data = interpretar_data(chamado.habilitada_em)
    unico = cruzamento.status == CRUZ_UNICO
    return {
        "C": chamado.numero,
        "E": _texto_de_celula(os_da_serie) if unico else "",
        "F": _texto_de_celula(detalhes_atendimento_do_chamado(chamado)),
        "G": _texto_de_celula(chamado.titulo),
        "H": _texto_de_celula(chamado.descricao),
        "I": _texto_de_celula(chamado.solicitante),
        "K": _texto_de_celula(cruzamento.serie) if unico else "",
        "L": _texto_de_celula(cruzamento.modelo) if unico else "",
        "M": UNIDADE,
        "N": _texto_de_celula(cruzamento.cidade) if unico else "",
        "O": UF,
        "P": data.strftime(FORMATO_DATA) if data else "",
    }


def linhas_para_colar(linhas: Sequence[Mapping[str, str]]) -> Tuple[str, str]:
    """Two tab-separated blocks, one row per line: ``A:P`` and ``R:S``.

    Q (Mês) is a formula, so the paste skips it; the manual R and S go in a second block.
    """
    def bloco(letras: Sequence[str]) -> str:
        return "\n".join("\t".join(str(linha.get(letra, "")) for letra in letras) for linha in linhas)

    return bloco("ABCDEFGHIJKLMNOP"), bloco("RS")


__all__ = [
    "ABA", "CABECALHOS", "CRUZ_AMBIGUO", "CRUZ_AUSENTE", "CRUZ_UNICO", "Cruzamento", "DashboardInvalido",
    "FORMULAS", "INTOCADAS", "MANUAIS", "NovosChamados", "PENDENCIAS_DEFINIDAS", "conferir_cabecalhos",
    "cruzar_por_nome", "ler_coluna_c", "linhas_para_colar", "montar_linha", "omitir_credenciais", "triar_novos",
]
