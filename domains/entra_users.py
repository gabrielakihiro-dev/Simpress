"""Local Entra name directory; the Endpoint login remains the equipment identity."""

from __future__ import annotations

import csv
import io
import json
import logging
import os
import re
import tempfile
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from domains.inventory_comparison import normalize_reference_value


logger = logging.getLogger(__name__)

DOMAIN = "pr.sebrae.com.br"
HEADERS = ("id", "displayName", "userPrincipalName", "userType", "onPremisesSyncEnabled", "identities", "companyName", "creationType")
MAX_CSV_BYTES = 60 * 1024 * 1024
MAX_ROWS = 100_000
MAX_NAME_LENGTH = 300
_LOGIN = re.compile(r"[a-z0-9][a-z0-9._-]*\$?", re.IGNORECASE)
_LOCK = threading.RLock()


class EntraImportError(ValueError):
    pass


def login_key(value: Any) -> str:
    """Accept one Endpoint login, optionally qualified by domain or PR UPN."""
    text = str(value or "").strip()
    if "\\" in text:
        parts = text.split("\\")
        if len(parts) != 2:
            return ""
        text = parts[1]
    if "@" in text:
        local, separator, domain = text.partition("@")
        if separator and domain.casefold() != DOMAIN:
            return ""
        text = local
    key = text.casefold()
    return key if _LOGIN.fullmatch(key) else ""


def _csv_row(outer: list[str], row_number: int) -> list[str]:
    if len(outer) == 1:
        line = outer[0]
        if line.endswith(";"):
            line = line[:-1]
        try:
            row = next(csv.reader([line]))
        except csv.Error as error:
            raise EntraImportError(f"CSV inválido na linha {row_number}.") from error
    else:
        row = outer
    if len(row) != len(HEADERS):
        raise EntraImportError(f"CSV inválido na linha {row_number}: esperadas 8 colunas.")
    row[-1] = row[-1].removesuffix(";")
    return [value.strip() for value in row]


def parse_csv(content: bytes) -> dict[str, Any]:
    if not content or len(content) > MAX_CSV_BYTES:
        raise EntraImportError("O CSV está vazio ou excede 60 MB.")
    try:
        stream = io.StringIO(content.decode("utf-8-sig"), newline="")
        reader = csv.reader(stream)
        header = _csv_row(next(reader), 1)
        if tuple(header) != HEADERS:
            raise EntraImportError("Cabeçalhos do CSV do Entra inválidos.")
        users: dict[str, str] = {}
        conflicts: set[str] = set()
        source_rows = 0
        filtered_rows = 0
        invalid_rows = 0
        for outer in reader:
            if not any(value.strip() for value in outer):
                continue
            source_rows += 1
            if source_rows > MAX_ROWS:
                raise EntraImportError("O CSV excede o limite de 100.000 linhas.")
            # Some tenant-wide rows contain broken quoting. Only PR rows are relevant;
            # never let an unrelated account prevent the filtered directory update.
            if f"@{DOMAIN}" not in " ".join(outer).casefold():
                continue
            row = _csv_row(outer, reader.line_num)
            name = row[1].strip().strip('"').strip()
            upn = row[2].strip().strip('"').strip()
            local, separator, domain = upn.rpartition("@")
            if not separator or domain.casefold() != DOMAIN:
                continue
            filtered_rows += 1
            key = login_key(upn)
            if not key or not name or len(name) > MAX_NAME_LENGTH:
                # An invalid login has no key to track: count the row itself.
                invalid_rows += 1
                continue
            previous = users.get(key)
            if previous is not None and previous != name:
                conflicts.add(key)
            users[key] = name
    except (UnicodeError, csv.Error, StopIteration) as error:
        raise EntraImportError("Não foi possível ler o CSV UTF-8 do Entra.") from error
    if not filtered_rows:
        raise EntraImportError("O CSV não contém contas @pr.sebrae.com.br.")
    # A login seen with different names is ambiguous: leave it out rather than pick one.
    for key in conflicts:
        del users[key]
    if not users:
        raise EntraImportError("Nenhum login PR do CSV é válido: todos têm nomes conflitantes ou dados inválidos.")
    return {
        "version": 1,
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "source_rows": source_rows,
        "filtered_rows": filtered_rows,
        "logins_ignorados": len(conflicts) + invalid_rows,
        "users": users,
    }


def _read_json(path: Path, default: dict[str, Any]) -> dict[str, Any]:
    if not path.exists():
        return default
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError(f"Estado inválido: {path.name}")
    return data


def read_directory(path: Path) -> dict[str, Any]:
    with _LOCK:
        data = _read_json(path, {"version": 1, "users": {}, "source_rows": 0, "filtered_rows": 0, "updated_at": ""})
        if not isinstance(data.get("users"), dict):
            raise ValueError("Cadastro do Entra inválido.")
        return data


def read_manual_names(path: Path) -> dict[str, Any]:
    with _LOCK:
        data = _read_json(path, {})
        if any(not isinstance(value, dict) for value in data.values()):
            raise ValueError("Cadastro de nomes manuais inválido.")
        return data


def read_names_safe(directory_path: Path, manual_path: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    """Read both name sources for the alias index; a bad file only drops the names."""
    try:
        directory = read_directory(directory_path)
    except (OSError, ValueError) as error:
        logger.warning("Cadastro do Entra ignorado no índice do Endpoint: %s", error)
        directory = {}
    try:
        manual = read_manual_names(manual_path)
    except (OSError, ValueError) as error:
        logger.warning("Nomes manuais ignorados no índice do Endpoint: %s", error)
        manual = {}
    return directory, manual


def _atomic_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(handle, "w", encoding="utf-8") as stream:
            json.dump(data, stream, ensure_ascii=False, separators=(",", ":"))
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def import_csv(content: bytes, path: Path) -> dict[str, Any]:
    parsed = parse_csv(content)
    with _LOCK:
        _atomic_json(path, parsed)
    return {key: parsed[key] for key in ("updated_at", "source_rows", "filtered_rows", "logins_ignorados")} | {"users": len(parsed["users"])}


def save_manual_name(path: Path, serial: str, login: str, name: str) -> None:
    key = normalize_reference_value(serial)
    clean = name.strip()
    owner = login_key(login)
    if not key or len(clean) > MAX_NAME_LENGTH or any(ord(char) < 32 for char in clean):
        raise ValueError("Nome ou série inválidos.")
    if clean and not owner:
        raise ValueError("Não há login válido para vincular o nome.")
    with _LOCK:
        data = read_manual_names(path)
        if clean:
            data[key] = {"login": owner, "nome": clean}
        else:
            data.pop(key, None)
        _atomic_json(path, data)


def resolve_name(user: Any, serial: Any, directory: dict[str, Any], manual_names: dict[str, Any]) -> tuple[str, str]:
    key = login_key(user)
    if not key:
        return "", ""
    name = directory.get("users", {}).get(key)
    if isinstance(name, str) and name:
        return name, "entra"
    entry = manual_names.get(normalize_reference_value(serial), {})
    if isinstance(entry, dict) and str(entry.get("login") or "") == key:
        name = str(entry.get("nome") or "").strip()
        if name:
            return name, "manual"
    return "", ""


def export_xlsx(directory: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Nomes PR"
    sheet.append(("Login", "Display Name"))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26
    for login, name in sorted(directory.get("users", {}).items()):
        sheet.append((login, name))
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"  # A displayName beginning with '=' is never a formula.
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 52
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
