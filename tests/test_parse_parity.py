"""Paridade do parse de planilhas antes de otimizá-lo.

Protege o resultado da comparação de inventário contra mudanças na leitura das
planilhas (opção (a) de docs/arquitetura/medicao_parse_2026-10.md: passada única,
só a aba e as colunas usadas). As planilhas são sintéticas, geradas aqui com
openpyxl; nenhum arquivo de data/ é lido.

Os dois caminhos reais de leitura são exercitados:

- upload: read_workbook_tables + select_upload_table;
- arquivo local: find_details_sheet + read_details_sheet_records.

O resultado completo (tabelas, __excel_row_number__, headers, column_count e a
comparação inteira) é comparado com tests/fixtures/parse_parity_snapshot.json.
Uma diferença indica mudança de comportamento. Se ela for intencional, regenere
o snapshot com SIMPRESS_UPDATE_PARSE_SNAPSHOT=1 e revise o diff antes do commit.
"""

import io
import json
import logging
import os
import tempfile
import unittest
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List

import openpyxl

from core.xlsx_safety import validate_xlsx
from domains import uploaded_workbooks
from domains.inventory_comparison import compare_uploaded_workbooks, normalize_reference_value


SNAPSHOT_PATH = Path(__file__).resolve().parent / "fixtures" / "parse_parity_snapshot.json"
UPDATE_SNAPSHOT_ENV = "SIMPRESS_UPDATE_PARSE_SNAPSHOT"
AS_OF = datetime(2026, 10, 15, 12, 0, tzinfo=timezone(timedelta(hours=-3)))
UPLOADED_NAMES = {
    "endpoint": "InvAllComputersSummary.xlsx",
    "detalhes": "EquipamentosDetalhes.xlsx",
}

# Aliases propositais: "Número de Série" (NUMERODESERIE), " Cidade " com espaços
# e "CNPJ de Instalação" com acento. Completam as 50 colunas da base completa.
DETAILS_HEADERS = [
    "Número de Série",
    "Apelido",
    "Modelo",
    " Cidade ",
    "CNPJ de Instalação",
    "Tipo",
] + [f"Campo {index:02d}" for index in range(7, 51)]

ENDPOINT_HEADERS = [
    "Nome do computador",
    "Rótulo de serviço/Número de série",
    "Usuários logados",
    "Última verificação com êxito",
    "Sistema operacional",
]

CNPJ = "08.123.456/0001-90"

# (série, apelido, modelo, cidade, CNPJ, tipo)
DETAILS_ROWS: List[Any] = [
    ("BRJ100AAA1", "PRNOTELON01", "ProBook 440", "Londrina", CNPJ, "Notebook"),
    # Localidade divergente real: hostname de Londrina, cadastro em Maringá.
    ("BRJ100AAA2", "PRNOTELON02", "ProBook 440", "Maringá", CNPJ, "Notebook"),
    # Exceções de Curitiba (PRNOTESJP/PRNOTERL/PRNOTEARC) com outra cidade cadastrada.
    ("BRJ200SJP1", "PRNOTESJP01", "ProBook 440", "São José dos Pinhais", CNPJ, "Notebook"),
    ("BRJ200RL01", "PRNOTERL01", "ProBook 440", "Colombo", CNPJ, "Notebook"),
    ("BRJ200ARC1", "PRNOTEARC01", "ProBook 440", "Araucária", CNPJ, "Notebook"),
    # MacBooks: Endpoint sem S inicial, Simpress com S.
    ("SCMVXKV7YCM", "PRMACMTZ01", "MacBook Air", "Curitiba", CNPJ, "Notebook"),
    ("SC02XYZ123", "PRMACMTZ02", "MacBook Pro", "Curitiba", CNPJ, "Notebook"),
    # MacBook ambíguo: duas linhas com S para a mesma série do Endpoint.
    ("SFVFAMB001", "PRMACAMB01", "MacBook Air", "Curitiba", CNPJ, "Notebook"),
    ("SFVFAMB001", "PRMACAMB02", "MacBook Air", "Curitiba", CNPJ, "Notebook"),
    # Série com S que não é MacBook no Endpoint: o S deve ser preservado.
    ("SXA9988776", "PRNOTEMAR01", "ThinkPad", "Maringá", CNPJ, "Notebook"),
    # Série duplicada no cadastro.
    ("BRJ300DUP1", "PRNOTECAS01", "ProBook 440", "Cascavel", CNPJ, "Notebook"),
    ("BRJ300DUP1", "PRNOTECAS02", "ProBook 440", "Cascavel", CNPJ, "Notebook"),
    # Série vazia com apelido (sem CNPJ, permitido sem série).
    (None, "PRNOTEFOZ01", "ProBook 440", "Foz do Iguaçu", None, "Notebook"),
    # Série vazia, apelido com espaços (casado pelo hostname do Endpoint).
    ("", "  PRNOTEPG01  ", "ProBook 440", "Ponta Grossa", "", "Notebook"),
    # Só no cadastro.
    ("BRJ400ONLY", "PRNOTETOL01", "ProBook 440", "Toledo", CNPJ, "Notebook"),
    # Série numérica (célula int) e série com hífen/minúsculas.
    (123456789, "PRNOTEGPV01", "ProBook 440", "Guarapuava", CNPJ, "Notebook"),
    ("brj-500-low", "PRNOTEIRA01", "ProBook 440", "Irati", CNPJ, "Notebook"),
    # Equipamento que não é notebook, sem hostname PRNOTE.
    ("CN12345678", "IMPRESSORA-MTZ", "LaserJet", "Curitiba", CNPJ, "Impressora"),
    # Hostname do cadastro igual a dois no Endpoint com séries divergentes.
    ("BRJ600HOST", "PRNOTEUMU01", "ProBook 440", "Umuarama", CNPJ, "Notebook"),
]

# (hostname, série, usuário, última verificação, sistema)
ENDPOINT_ROWS: List[Any] = [
    ("PRNOTELON01", "BRJ100AAA1", "sebrae\\ana", datetime(2026, 10, 14, 9, 30), "Windows 11"),
    ("PRNOTELON02", "BRJ100AAA2", "sebrae\\bia", datetime(2026, 10, 1, 8, 0), "Windows 11"),
    ("PRNOTESJP01", "BRJ200SJP1", "", "Oct 1, 2026 10:00 AM", "Windows 11"),
    ("PRNOTERL01", "BRJ200RL01", None, "08/10/2026 14:00", "Windows 11"),
    ("PRNOTEARC01", "BRJ200ARC1", "sebrae\\caio", "2026-10-14T10:00:00Z", "Windows 10"),
    # MacBooks sem S inicial.
    ("PRMACMTZ01", "CMVXKV7YCM", "sebrae\\dani", date(2026, 10, 13), "macOS"),
    ("PRMACMTZ02", "C02XYZ123", "sebrae\\edu", "12 out 2026 09:00", "macOS"),
    ("PRMACAMB01", "FVFAMB001", "sebrae\\fabi", None, "macOS"),
    # Notebooks próprios do Sebrae: ignorados nos ausentes em Simpress.
    ("PRNOTEMTZ90", "BRJ135P5XS", "sebrae\\gil", datetime(2026, 10, 14), "Windows 11"),
    ("PRNOTEMTZ91", "BRJ142QKBQ", "", None, "Windows 11"),
    ("PRNOTEMTZ92", "brj9435vmr", "", None, "Windows 11"),
    # Duplicidade de série no Endpoint.
    ("PRNOTECAS01", "BRJ300DUP1", "", None, "Windows 11"),
    # Casa por hostname, sem série no Endpoint.
    ("PRNOTEFOZ01", "", "", None, "Windows 11"),
    ("PRNOTEPG01", None, "", None, "Windows 11"),
    # Só no Endpoint e com série sem mapeamento de local.
    ("PRNOTEXYZ01", "BRJ700ENDP", "", datetime(2026, 9, 1, 7, 15), "Windows 11"),
    ("PRNOTEBKPLON1", "BRJ700ENDQ", "", None, "Windows 11"),
    # Hostname duplicado no Endpoint com séries divergentes.
    ("PRNOTEUMU01", "BRJ600HOST", "", None, "Windows 11"),
    ("PRNOTEUMU01", "BRJ600HOSX", "", None, "Windows 11"),
    # Série numérica e série normalizada.
    ("PRNOTEGPV01", 123456789, "", None, "Windows 11"),
    ("PRNOTEIRA01", "BRJ500LOW", "", None, "Windows 11"),
    # Fora do escopo da comparação (não começa com PRNOTE/PRMAC).
    ("PRDESKMTZ01", "BRJ800DESK", "", None, "Windows 11"),
]


def _generic_value(row_index: int, column_index: int) -> Any:
    """Valores variados nas colunas não usadas: texto, números, datas e vazios."""
    kind = (row_index + column_index) % 5
    if kind == 0:
        return None
    if kind == 1:
        return f"valor {row_index}-{column_index}"
    if kind == 2:
        return row_index * 100 + column_index
    if kind == 3:
        return round(row_index + column_index / 10, 1)
    return datetime(2026, 1 + row_index % 12, 1 + column_index % 28)


def build_details_workbook() -> bytes:
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary.append(["Indicador", "Valor"])
    summary.append(["Total", len(DETAILS_ROWS)])
    summary.append(["Gerado em", datetime(2026, 10, 15, 8, 0)])

    workbook.create_sheet("Vazia")

    sheet = workbook.create_sheet("EquipamentosDetalhes")
    sheet.append(DETAILS_HEADERS)
    for row_index, row in enumerate(DETAILS_ROWS):
        generic = [_generic_value(row_index, column) for column in range(7, 51)]
        sheet.append(list(row) + generic)
        if row_index == 4:
            # Linha totalmente vazia no meio: gera lacuna em __excel_row_number__.
            sheet.append([])
        if row_index == 9:
            # Linha só com espaços numa coluna genérica (o upload a descarta).
            sheet.append([None] * 10 + ["   "])
    # Linha final com dados apenas em colunas genéricas, sem série nem apelido.
    sheet.append([None] * 6 + ["sobra"] + [None] * 43)
    return _workbook_bytes(workbook)


def build_endpoint_workbook() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "InvAllComputersSummary"
    sheet.append(ENDPOINT_HEADERS)
    for row_index, row in enumerate(ENDPOINT_ROWS):
        sheet.append(list(row))
        if row_index == 2:
            sheet.append([])
    return _workbook_bytes(workbook)


def _workbook_bytes(workbook: Any) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _to_json(value: Any) -> Any:
    """Serialização determinística que preserva o tipo dos valores das células."""
    if isinstance(value, dict):
        return {str(key): _to_json(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_to_json(item) for item in value]
    if isinstance(value, (set, frozenset)):
        return sorted(_to_json(item) for item in value)
    if isinstance(value, datetime):
        return {"__datetime__": value.isoformat()}
    if isinstance(value, date):
        return {"__date__": value.isoformat()}
    if isinstance(value, time):
        return {"__time__": value.isoformat()}
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    raise TypeError(f"Tipo inesperado no snapshot: {type(value).__name__}")


def _local_details_table(sheet_name: str, headers: List[str], records: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Adaptador só de teste: põe os registros do caminho local no formato do upload."""
    return {
        "sheet": sheet_name,
        "headers": {normalize_reference_value(header) for header in headers},
        "records": [
            {normalize_reference_value(key): value for key, value in record.items()}
            for record in records
        ],
    }


class ParseParityTests(unittest.TestCase):
    maxDiff = None

    @classmethod
    def setUpClass(cls) -> None:
        cls.services = uploaded_workbooks.WorkbookReaderServices(
            validate_xlsx=validate_xlsx,
            validate_xlsx_package=uploaded_workbooks.validate_xlsx_package,
            ensure_workbook_readable=uploaded_workbooks.ensure_workbook_readable,
            logger=logging.getLogger("tests.parse_parity"),
        )
        cls.details_bytes = build_details_workbook()
        cls.endpoint_bytes = build_endpoint_workbook()
        cls.actual = cls._build_actual()

    @classmethod
    def _build_actual(cls) -> Dict[str, Any]:
        details_tables = uploaded_workbooks.read_workbook_tables(cls.services, cls.details_bytes)
        endpoint_tables = uploaded_workbooks.read_workbook_tables(cls.services, cls.endpoint_bytes)
        details_table = uploaded_workbooks.select_upload_table(details_tables, "detalhes")
        endpoint_table = uploaded_workbooks.select_upload_table(endpoint_tables, "endpoint")
        upload_comparison = compare_uploaded_workbooks(
            endpoint_table, details_table, dict(UPLOADED_NAMES), as_of=AS_OF,
        )

        with tempfile.TemporaryDirectory() as directory:
            details_path = Path(directory) / "EquipamentosDetalhes_teste.xlsx"
            details_path.write_bytes(cls.details_bytes)
            sheet_name = uploaded_workbooks.find_details_sheet(cls.services, details_path)
            if sheet_name is None:
                raise AssertionError("find_details_sheet não encontrou a aba sintética.")
            local_headers, local_records = uploaded_workbooks.read_details_sheet_records(
                cls.services, details_path, sheet_name,
            )
        local_comparison = compare_uploaded_workbooks(
            endpoint_table,
            _local_details_table(sheet_name, local_headers, local_records),
            dict(UPLOADED_NAMES),
            as_of=AS_OF,
        )

        return _to_json({
            "upload": {
                "tabelas_detalhes": details_tables,
                "tabelas_endpoint": endpoint_tables,
                "selecionada_detalhes": details_table,
                "selecionada_endpoint": endpoint_table,
                "comparacao": upload_comparison,
            },
            "arquivo_local": {
                "aba": sheet_name,
                "headers": local_headers,
                "records": local_records,
                "comparacao": local_comparison,
            },
        })

    def _expected(self) -> Dict[str, Any]:
        if os.environ.get(UPDATE_SNAPSHOT_ENV) == "1":
            SNAPSHOT_PATH.write_text(
                json.dumps(self.actual, ensure_ascii=False, indent=1, sort_keys=True) + "\n",
                encoding="utf-8",
            )
        expected: Dict[str, Any] = json.loads(SNAPSHOT_PATH.read_text(encoding="utf-8"))
        return expected

    def test_synthetic_workbooks_cover_the_complete_details_schema(self) -> None:
        self.assertEqual(len(DETAILS_HEADERS), uploaded_workbooks.DETAILS_COMPLETE_SOURCE_COLUMN_COUNT)
        selected = self.actual["upload"]["selecionada_detalhes"]
        self.assertEqual(selected["sheet"], "EquipamentosDetalhes")
        self.assertEqual(selected["column_count"], 50)
        self.assertEqual(self.actual["arquivo_local"]["aba"], "EquipamentosDetalhes")
        self.assertEqual(len(self.actual["arquivo_local"]["headers"]), 50)

    def test_upload_tables_match_snapshot(self) -> None:
        expected = self._expected()["upload"]
        actual = self.actual["upload"]
        for key in ("tabelas_detalhes", "tabelas_endpoint", "selecionada_detalhes", "selecionada_endpoint"):
            with self.subTest(parte=key):
                self.assertEqual(actual[key], expected[key])

    def test_upload_comparison_matches_snapshot(self) -> None:
        self.assertEqual(
            self.actual["upload"]["comparacao"],
            self._expected()["upload"]["comparacao"],
        )

    def test_local_file_records_match_snapshot(self) -> None:
        expected = self._expected()["arquivo_local"]
        actual = self.actual["arquivo_local"]
        for key in ("aba", "headers", "records"):
            with self.subTest(parte=key):
                self.assertEqual(actual[key], expected[key])

    def test_local_file_comparison_matches_snapshot(self) -> None:
        self.assertEqual(
            self.actual["arquivo_local"]["comparacao"],
            self._expected()["arquivo_local"]["comparacao"],
        )

    def test_snapshot_keeps_the_business_rules_visible(self) -> None:
        """Âncoras legíveis das regras do AGENTS.md, além do snapshot bruto."""
        comparison = self.actual["upload"]["comparacao"]
        missing_in_details = {item["serie"] for item in comparison["ausentes_em_detalhes"]}
        for serial in ("BRJ135P5XS", "BRJ142QKBQ", "brj9435vmr"):
            self.assertNotIn(serial, missing_in_details)
        self.assertIn("BRJ700ENDP", missing_in_details)

        mismatched_hosts = {item["computador"] for item in comparison["divergencias_local"]}
        self.assertIn("PRNOTELON02", mismatched_hosts)
        for hostname in ("PRNOTESJP01", "PRNOTERL01", "PRNOTEARC01"):
            self.assertNotIn(hostname, mismatched_hosts)

        missing_in_endpoint = {item["serie"] for item in comparison["ausentes_no_endpoint"]}
        self.assertNotIn("SCMVXKV7YCM", missing_in_endpoint)
        self.assertNotIn("CMVXKV7YCM", missing_in_details)

        details_rows = [
            record["__excel_row_number__"]
            for record in self.actual["upload"]["selecionada_detalhes"]["records"]
        ]
        self.assertEqual(details_rows[:6], [2, 3, 4, 5, 6, 8])

        # Diferença atual entre os caminhos: o arquivo local mantém a linha só com
        # espaços, o upload a descarta. O resto da comparação é idêntico.
        local = self.actual["arquivo_local"]["comparacao"]
        self.assertEqual(
            [key for key in comparison if comparison[key] != local.get(key)],
            ["total_detalhes"],
        )
        self.assertEqual(local["total_detalhes"], comparison["total_detalhes"] + 1)


if __name__ == "__main__":
    unittest.main()
