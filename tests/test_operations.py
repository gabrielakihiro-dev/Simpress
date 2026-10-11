"""Tests for the operational history and ticket workflows."""

from io import BytesIO
from datetime import date, timedelta
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
import openpyxl
from openpyxl import Workbook
from starlette.datastructures import UploadFile

from domains.history_insights import compare_updates, latest_update, pending_rows
from domains.ticket_reporting import build_ticket_workbook
from routers.operations import TicketTimelineAnnotationPayload, create_operations_router

PROJECT_ROOT = Path(__file__).resolve().parents[1]


def route_handler(router, path, method):
    return next(
        route.endpoint for route in router.routes
        if getattr(route, "path", None) == path and method in getattr(route, "methods", set())
    )


class HistoryInsightTests(unittest.TestCase):
    def test_pending_center_uses_latest_update_and_keeps_annotation_identity(self):
        old = {"id": "old", "horario": "2026-01-01", "comparacao": {"sem_acesso_14_dias": []}}
        current = {
            "id": "current", "horario": "2026-02-01", "comparacao": {
                "sem_acesso_14_dias": [{"serie": "ABC-1", "anotacao_id": "row-1", "anotacao": "Ver chamado"}],
                "divergencias_local": [{"serie": "XYZ-2"}],
            },
        }
        restoration = {"id": "restore", "horario": "2026-03-01", "tipo": "restauracao", "comparacao": {}}

        self.assertEqual(latest_update([old, current, restoration]), current)
        rows = pending_rows(current)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["item"]["anotacao_id"], "row-1")
        self.assertEqual(rows[0]["item"]["anotacao"], "Ver chamado")

    def test_update_evolution_labels_new_continued_and_resolved_rows(self):
        previous = {"id": "before", "horario": "2026-01-01", "comparacao": {
            "divergencias_local": [{"serie": "A-1"}, {"serie": "B-2"}],
        }}
        current = {"id": "after", "horario": "2026-02-01", "comparacao": {
            "divergencias_local": [{"serie": "A1"}, {"serie": "C-3"}],
        }}

        result = compare_updates(previous, current)
        self.assertEqual(result["contagens"], {"novos": 1, "continuam": 1, "resolvidos": 1})
        self.assertEqual(result["mudancas"]["novos"][0]["evento_id"], "after")
        self.assertEqual(result["mudancas"]["resolvidos"][0]["evento_id"], "before")

    def test_operational_read_routes_serve_pending_and_evolution_views(self):
        history = [
            {"id": "before", "horario": "2026-01-01", "comparacao": {"divergencias_local": [{"serie": "A-1"}]}},
            {"id": "after", "horario": "2026-02-01", "comparacao": {
                "divergencias_local": [{"serie": "A1", "anotacao_id": "note-1"}],
                "sem_acesso_14_dias": [{"serie": "B-2", "computador": "PRNOTE2", "anotacao_id": "note-2"}],
            }},
            {"id": "restore", "horario": "2026-03-01", "tipo": "restauracao", "comparacao": {}},
        ]
        router = create_operations_router(
            load_comparison_history=lambda: history,
            comparison_history_for_display=lambda event: event,
            ticket_report_path=lambda: Path(tempfile.gettempdir()) / "simpress-unused-tickets.json",
            transaction_root=lambda: Path(tempfile.gettempdir()),
        )

        pending = route_handler(router, "/api/operacoes/pendencias", "GET")()
        evolution = route_handler(router, "/api/operacoes/evolucao", "GET")(anterior="before", atual="after")

        self.assertEqual(pending["evento"]["id"], "after")
        self.assertEqual(len(pending["pendencias"]), 2)
        self.assertEqual(evolution["contagens"], {"novos": 1, "continuam": 1, "resolvidos": 0})


class TicketImportRouteTests(unittest.TestCase):
    @staticmethod
    def workbook_bytes(rows):
        workbook = Workbook()
        worksheet = workbook.active
        for row in rows:
            worksheet.append(row)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_import_reconciles_sources_and_persists_latest_report(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            path = Path(directory) / "state" / "tickets.json"
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: path,
                transaction_root=lambda: Path(directory),
            )
            simpress = self.workbook_bytes([
                ["NumeroPedido", "NumeroSerie", "OrdemServico", "NumeroChamadoCliente", "StatusAtual"],
                ["PED-1", "SN-1", 1001, "INC-1", "Em atendimento"],
                ["PED-1", "SN-1", 1001, "INC-1", "Em atendimento"],
                ["PED-2", "SN-2", 1002, None, "Aguardando"],
                ["PED-3", "SN-3", 1003, None, "Aguardando"],
            ])
            oracle = self.workbook_bytes([
                ["Gestão de Solicitações"],
                ["Solicitação", "ID OS Corporate", "Número de Série", "Logradouro", "Número", "Bairro", "Cidade", "Atendido pela Inibição", "Nome da Filial", "Fila", "Status", "Sub-status", "Descrição"],
                ["REQ-1", 1001, "sn 1", "R Caeté", 150, "Prado Velho", "Curitiba", "Não", "SIMPRESS - PR", "Fila A", "Em execução", "Pendente", "Solicito troca do teclado e mouse."],
                ["REQ-2", 2002, "SN-2", "R Vitória", 2564, "Centro", "Cascavel", "Sim", "PARCEIRO", "Fila B", "Aguardando", "", "Aguardando diagnóstico técnico."],
                ["REQ-4", 9004, "SN-4", "R Alfredo Bufren", 365, "Centro", "Irati", "Não", "SIMPRESS - PR", "Fila C", "Aberto", "", "Verificar equipamento."],
            ])

            endpoint = route_handler(router, "/api/operacoes/chamados/importar", "POST")
            report = endpoint(
                simpress_file=UploadFile(filename="pedidos.xlsx", file=BytesIO(simpress)),
                oracle_file=UploadFile(filename="solicitacoes.xlsx", file=BytesIO(oracle)),
            )
            by_serial = {ticket["serie"]: ticket for ticket in report["chamados"]}
            self.assertEqual(report["resumo"], {
                "series_total": 4, "conferidas": 1, "divergentes": 1,
                "sem_os": 0, "somente_simpress": 1, "somente_oracle": 1,
            })
            self.assertEqual(by_serial["SN-1"]["simpress"]["ordens_servico"], ["1001"])
            self.assertEqual(by_serial["SN-1"]["oracle"][0]["solicitacao"], "REQ-1")
            self.assertEqual(by_serial["SN-1"]["oracle"][0]["descricao"], "Solicito troca do teclado e mouse.")
            self.assertEqual(by_serial["SN-1"]["oracle"][0]["logradouro"], "R Caeté")
            self.assertEqual(by_serial["SN-1"]["oracle"][0]["numero_endereco"], "150")
            self.assertEqual(by_serial["SN-1"]["oracle"][0]["bairro"], "Prado Velho")
            self.assertEqual(by_serial["SN-1"]["oracle"][0]["cidade"], "Curitiba")
            self.assertEqual(by_serial["SN-1"]["oracle"][0]["atendido_pela_inibicao"], "Não")
            self.assertEqual(by_serial["SN-1"]["oracle"][0]["nome_filial"], "SIMPRESS - PR")
            self.assertEqual(by_serial["SN-2"]["oracle"][0]["nome_filial"], "PARCEIRO")
            self.assertEqual(by_serial["SN-1"]["relacao"], "conferida")
            self.assertEqual(by_serial["SN-2"]["relacao"], "divergente")
            self.assertEqual(report["fontes"]["simpress"]["linhas"], 4)
            persisted = route_handler(router, "/api/operacoes/chamados", "GET")()
            self.assertEqual(persisted["chamados"], report["chamados"])
            self.assertEqual(persisted["atualizado_em"], report["atualizado_em"])

    def test_import_keeps_highest_simpress_os_even_when_its_order_date_is_older(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: Path(directory) / "tickets.json",
                transaction_root=lambda: Path(directory),
            )
            simpress = self.workbook_bytes([
                ["NumeroPedido", "NumeroSerie", "OrdemServico", "StatusAtual", "DataPedido"],
                ["PED-OLD", "SN-MULTI", 7124895, "Pedido antigo", "10/10/2026"],
                ["PED-NEW", "SN-MULTI", 7134662, "Pedido ativo", "09/10/2026"],
            ])
            oracle = self.workbook_bytes([
                ["Solicitação", "ID OS Corporate", "Número de Série", "Data de Criação", "Status"],
                ["REQ-OLD", 7124895, "SN-MULTI", "2026-10-09T10:00:00", "Concluído"],
                ["REQ-NEW", 7134662, "SN-MULTI", "2026-10-10T10:00:00", "Em execução"],
            ])

            report = route_handler(router, "/api/operacoes/chamados/importar", "POST")(
                simpress_file=UploadFile(filename="pedidos.xlsx", file=BytesIO(simpress)),
                oracle_file=UploadFile(filename="solicitacoes.xlsx", file=BytesIO(oracle)),
            )

            ticket = report["chamados"][0]
            self.assertEqual(ticket["simpress"]["pedidos"], ["PED-NEW"])
            self.assertEqual(ticket["simpress"]["ordens_servico"], ["7134662"])
            self.assertEqual(ticket["simpress"]["status"], ["Pedido ativo"])
            self.assertEqual(len(ticket["oracle"]), 1)
            self.assertEqual(ticket["oracle"][0]["ordem_servico"], "7134662")
            self.assertEqual(ticket["oracle"][0]["solicitacao"], "REQ-NEW")
            self.assertEqual(ticket["ordens_servico_conferidas"], ["7134662"])

    def test_import_keeps_highest_pedido_attached_to_the_latest_simpress_os(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: Path(directory) / "tickets.json",
                transaction_root=lambda: Path(directory),
            )
            simpress = self.workbook_bytes([
                ["NumeroPedido", "NumeroSerie", "OrdemServico", "StatusAtual", "DataPedido"],
                [46735940, "SN-PEDIDO", 7134662, "Pedido antigo", "10/10/2026"],
                [48076976, "SN-PEDIDO", 7134662, "Pedido atual", "09/10/2026"],
                [46000000, "SN-PEDIDO", 7124895, "OS antiga", "11/10/2026"],
            ])
            oracle = self.workbook_bytes([
                ["Solicitação", "ID OS Corporate", "Número de Série"],
                ["REQ-ATUAL", 7134662, "SN-PEDIDO"],
            ])

            report = route_handler(router, "/api/operacoes/chamados/importar", "POST")(
                simpress_file=UploadFile(filename="pedidos.xlsx", file=BytesIO(simpress)),
                oracle_file=UploadFile(filename="solicitacoes.xlsx", file=BytesIO(oracle)),
            )

            ticket = report["chamados"][0]
            self.assertEqual(ticket["simpress"]["pedidos"], ["48076976"])
            self.assertEqual(ticket["simpress"]["ordens_servico"], ["7134662"])
            self.assertEqual(ticket["simpress"]["status"], ["Pedido atual"])

    def test_import_includes_simpress_product_city_and_only_active_delivery_addresses(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: Path(directory) / "tickets.json",
                transaction_root=lambda: Path(directory),
            )
            simpress = self.workbook_bytes([
                [
                    "NumeroPedido", "NumeroSerie", "OrdemServico", "OrdemVenda", "StatusAtual", "DescricaoProduto",
                    "PrevisaoEntrega", "CidadeEntrega", "EnderecoEntrega", "NumeroEntrega", "ComplementoEntrega",
                    "BairroEntrega", "EstadoEntrega", "CepEntrega",
                ],
                [
                    48076976, "SN-ATIVO", 7232035, "OV_15799764", "Verificação de estoque", "Notebook HP ProBook",
                    "28/09/2026", "Curitiba", "Rua das Flores", 123, "Sala 4", "Centro", "PR", "80000-000",
                ],
                [
                    48076977, "SN-ENTREGUE", 7232036, "OV_15799772", "Pedido entregue", "Monitor HP",
                    "15/09/2026", "Maringá", "Avenida Central", 500, "Bloco B", "Zona 1", "PR", "87000-000",
                ],
                [
                    48076978, "SN-CANCELADO", 7232037, "OV_15785398", "Cancelado", "Teclado sem fio",
                    "", "Londrina", "Rua Norte", 25, "", "Centro", "PR", "86000-000",
                ],
            ])
            oracle = self.workbook_bytes([
                ["Solicitação", "ID OS Corporate", "Número de Série"],
                ["REQ-ATIVO", 7232035, "SN-ATIVO"],
            ])

            report = route_handler(router, "/api/operacoes/chamados/importar", "POST")(
                simpress_file=UploadFile(filename="pedidos.xlsx", file=BytesIO(simpress)),
                oracle_file=UploadFile(filename="solicitacoes.xlsx", file=BytesIO(oracle)),
            )

            by_serial = {ticket["serie"]: ticket["simpress"] for ticket in report["chamados"]}
            self.assertEqual(by_serial["SN-ATIVO"]["descricoes_produto"], ["Notebook HP ProBook"])
            self.assertEqual(by_serial["SN-ATIVO"]["ordens_venda"], ["OV_15799764"])
            self.assertEqual(by_serial["SN-ATIVO"]["previsoes_entrega"], ["28/09/2026"])
            self.assertEqual(by_serial["SN-ATIVO"]["cidades_entrega"], ["Curitiba"])
            self.assertTrue(by_serial["SN-ATIVO"]["pedido_em_andamento"])
            self.assertEqual(
                by_serial["SN-ATIVO"]["enderecos_entrega"],
                ["Rua das Flores, 123, Sala 4, Centro, Curitiba - PR, CEP 80000-000"],
            )
            for serial in ("SN-ENTREGUE", "SN-CANCELADO"):
                self.assertFalse(by_serial[serial]["pedido_em_andamento"])
                self.assertEqual(by_serial[serial]["enderecos_entrega"], [])
                self.assertEqual(len(by_serial[serial]["descricoes_produto"]), 1)
                self.assertEqual(len(by_serial[serial]["cidades_entrega"]), 1)
            self.assertEqual(by_serial["SN-ENTREGUE"]["previsoes_entrega"], ["15/09/2026"])
            self.assertEqual(by_serial["SN-CANCELADO"]["previsoes_entrega"], [])

    def test_import_keeps_the_sales_order_attached_to_the_latest_simpress_order(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: Path(directory) / "tickets.json",
                transaction_root=lambda: Path(directory),
            )
            simpress = self.workbook_bytes([
                ["NumeroPedido", "NumeroSerie", "OrdemServico", "OrdemVenda"],
                ["PED-OLD", "SN-OV", 7124895, "OV_15799763"],
                ["PED-NEW", "SN-OV", 7134662, "OV_15799764"],
            ])
            oracle = self.workbook_bytes([
                ["Solicitação", "ID OS Corporate", "Número de Série"],
                ["REQ-NEW", 7134662, "SN-OV"],
            ])

            report = route_handler(router, "/api/operacoes/chamados/importar", "POST")(
                simpress_file=UploadFile(filename="pedidos.xlsx", file=BytesIO(simpress)),
                oracle_file=UploadFile(filename="solicitacoes.xlsx", file=BytesIO(oracle)),
            )

            ticket = report["chamados"][0]
            self.assertEqual(ticket["simpress"]["pedidos"], ["PED-NEW"])
            self.assertEqual(ticket["simpress"]["ordens_servico"], ["7134662"])
            self.assertEqual(ticket["simpress"]["ordens_venda"], ["OV_15799764"])

    def test_import_uses_highest_numeric_order_service_when_dates_are_missing(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: Path(directory) / "tickets.json",
                transaction_root=lambda: Path(directory),
            )
            simpress = self.workbook_bytes([
                ["NumeroPedido", "NumeroSerie", "OrdemServico"],
                ["PED-OLD", "SN-MULTI", 7056941],
                ["PED-NEW", "SN-MULTI", 7232035],
            ])
            oracle = self.workbook_bytes([
                ["Solicitação", "ID OS Corporate", "Número de Série"],
                ["REQ-OLD", 7056941, "SN-MULTI"],
                ["REQ-NEW", 7232035, "SN-MULTI"],
            ])

            report = route_handler(router, "/api/operacoes/chamados/importar", "POST")(
                simpress_file=UploadFile(filename="pedidos.xlsx", file=BytesIO(simpress)),
                oracle_file=UploadFile(filename="solicitacoes.xlsx", file=BytesIO(oracle)),
            )

            ticket = report["chamados"][0]
            self.assertEqual(ticket["simpress"]["pedidos"], ["PED-NEW"])
            self.assertEqual(ticket["simpress"]["ordens_servico"], ["7232035"])
            self.assertEqual([row["ordem_servico"] for row in ticket["oracle"]], ["7232035"])
            self.assertEqual(ticket["ordens_servico_conferidas"], ["7232035"])

    def test_ticket_timeline_annotations_follow_exact_rows_with_independent_import_ids(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            root = Path(directory)
            shared = root / "shared-annotations"
            outbox = root / "annotation-outbox"
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: root / "tickets.json",
                transaction_root=lambda: root,
                history_annotations_directory=lambda: shared,
                local_annotation_outbox_directory=lambda: outbox,
                history_annotations_sync_enabled=lambda: False,
            )
            simpress = self.workbook_bytes([
                ["NumeroPedido", "NumeroSerie", "OrdemServico", "StatusAtual", "CidadeEntrega"],
                [48076976, "SN-TIMELINE", 7232035, "Em separação", "Curitiba"],
            ])
            oracle = self.workbook_bytes([
                ["Solicitação", "ID OS Corporate", "Número de Série", "Data da Última Atualização", "Nome da Filial", "Fila", "Status"],
                ["REQ-TIMELINE", 7232035, "SN-TIMELINE", "2026-10-10 09:00", "PARCEIRO", "Fila A", "Em execução"],
            ])
            import_route = route_handler(router, "/api/operacoes/chamados/importar", "POST")

            def import_same_rows():
                return import_route(
                    simpress_file=UploadFile(filename="pedidos.xlsx", file=BytesIO(simpress)),
                    oracle_file=UploadFile(filename="oracle.xlsx", file=BytesIO(oracle)),
                )

            first = import_same_rows()
            first_event = route_handler(router, "/api/operacoes/chamados/linha-do-tempo/{serial}", "GET")(
                "SN-TIMELINE", snapshot_id=first["snapshot_id"],
            )["eventos"][0]
            save_annotation = route_handler(
                router, "/api/operacoes/chamados/linha-do-tempo/{event_id}/{item_id}", "PUT",
            )
            saved = save_annotation(
                event_id=first_event["snapshot_id"], item_id=first_event["item_id"],
                payload=TicketTimelineAnnotationPayload(anotacao="Cobrar o parceiro", base_revisoes=first_event["anotacao_revisoes"]),
            )
            self.assertEqual(saved["sincronizacao"], "local")

            second = import_same_rows()
            events = route_handler(router, "/api/operacoes/chamados/linha-do-tempo/{serial}", "GET")(
                "SN-TIMELINE",
            )["eventos"]
            self.assertEqual(len(events), 2)
            self.assertEqual(events[-1]["anotacao"], "Cobrar o parceiro")
            self.assertNotEqual(events[0]["item_id"], events[1]["item_id"])

            save_annotation(
                event_id=second["snapshot_id"], item_id=events[-1]["item_id"],
                payload=TicketTimelineAnnotationPayload(anotacao="", base_revisoes=events[-1]["anotacao_revisoes"]),
            )
            third = import_same_rows()
            third_events = route_handler(router, "/api/operacoes/chamados/linha-do-tempo/{serial}", "GET")(
                "SN-TIMELINE",
            )["eventos"]
            self.assertEqual(third_events[-1]["snapshot_id"], third["snapshot_id"])
            self.assertEqual(third_events[-1]["anotacao"], "")

            fourth = import_same_rows()
            fourth_events = route_handler(router, "/api/operacoes/chamados/linha-do-tempo/{serial}", "GET")(
                "SN-TIMELINE",
            )["eventos"]
            self.assertEqual(fourth_events[-1]["snapshot_id"], fourth["snapshot_id"])
            self.assertEqual(fourth_events[-1]["anotacao"], "")

    def test_ticket_import_flags_duplicate_rows_and_ambiguous_source_relations(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: Path(directory) / "tickets.json",
                transaction_root=lambda: Path(directory),
            )
            simpress = self.workbook_bytes([
                ["NumeroPedido", "NumeroSerie", "OrdemServico", "OrdemVenda"],
                [48076976, "SN-REVISAO", 7279287, "OV_15785398"],
                [48076976, "SN-REVISAO", 7279287, "OV_15785398"],
            ])
            oracle = self.workbook_bytes([
                ["Solicitação", "ID OS Corporate", "Número de Série"],
                ["REQ-1", 7279287, "SN-REVISAO"],
                ["REQ-1", 7279287, "SN-REVISAO"],
                ["REQ-2", 7279287, "SN-REVISAO"],
            ])

            report = route_handler(router, "/api/operacoes/chamados/importar", "POST")(
                simpress_file=UploadFile(filename="simpress.xlsx", file=BytesIO(simpress)),
                oracle_file=UploadFile(filename="oracle.xlsx", file=BytesIO(oracle)),
            )
            ticket = report["chamados"][0]
            alert_types = {alert["tipo"] for alert in ticket["alertas_relacao"]}

            self.assertEqual(ticket["fontes_relacao"]["serie"], {"simpress": True, "oracle": True})
            self.assertEqual(ticket["fontes_relacao"]["pedido"]["simpress"], ["48076976"])
            self.assertEqual(ticket["fontes_relacao"]["pedido"]["oracle_solicitacao"], ["REQ-1", "REQ-2"])
            self.assertIn("duplicidade_simpress", alert_types)
            self.assertIn("duplicidade_oracle", alert_types)
            self.assertIn("solicitacoes_ambiguas_oracle", alert_types)
            self.assertTrue(ticket["revisao_relacao"])

    def test_export_keeps_attention_rule_fixed_at_three_days(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            path = Path(directory) / "tickets.json"
            today = date.today()
            tickets = []
            for serial, order, age in (("SN-FRESH", "7279287", 2), ("SN-STALE", "7279280", 3)):
                updated = (today - timedelta(days=age)).isoformat()
                tickets.append({
                    "serie": serial, "relacao": "conferida", "ordens_servico_conferidas": [order],
                    "simpress": {"serie": serial, "pedidos": ["48076976"], "ordens_servico": [order]},
                    "oracle": [{"ordem_servico": order, "status": "Em execução", "ultima_atualizacao": updated}],
                })
            path.write_text(json.dumps({
                "snapshot_id": "attention-version", "atualizado_em": today.isoformat(),
                "fontes": {}, "resumo": {"series_total": 2}, "chamados": tickets,
            }), encoding="utf-8")
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: path,
                transaction_root=lambda: Path(directory),
            )

            captured = {}

            def capture_export(snapshot, filtered, filters, **kwargs):
                captured["serials"] = [ticket["serie"] for ticket in filtered]
                captured["stale_days"] = filters["stale_days"]
                return b"xlsx"

            with patch("routers.operations.build_ticket_workbook", side_effect=capture_export):
                response = route_handler(router, "/api/operacoes/chamados/exportar", "GET")(
                    snapshot_id="attention-version", relation="", status="", queue="", city="", provider="",
                    search="", stale_days=30, age_filter="", attention_only=True,
                )
            self.assertEqual(captured["serials"], ["SN-STALE"])
            self.assertEqual(captured["stale_days"], 3)
            self.assertEqual(response.headers["content-type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_import_rejects_non_xlsx_and_missing_source_columns(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: Path(directory) / "tickets.json",
                transaction_root=lambda: Path(directory),
            )
            endpoint = route_handler(router, "/api/operacoes/chamados/importar", "POST")
            simpress = self.workbook_bytes([["NumeroPedido", "NumeroSerie", "OrdemServico"], ["1", "A1", "100"]])
            with self.assertRaises(HTTPException) as invalid_extension:
                endpoint(
                    simpress_file=UploadFile(filename="pedidos.csv", file=BytesIO(simpress)),
                    oracle_file=UploadFile(filename="oracle.xlsx", file=BytesIO(simpress)),
                )
            self.assertEqual(invalid_extension.exception.status_code, 400)

            oracle = self.workbook_bytes([
                ["Solicitação", "Número de Série"],
                ["REQ-1", "A1"],
            ])
            with self.assertRaises(HTTPException) as missing_columns:
                endpoint(
                    simpress_file=UploadFile(filename="pedidos.xlsx", file=BytesIO(simpress)),
                    oracle_file=UploadFile(filename="oracle.xlsx", file=BytesIO(oracle)),
                )
            self.assertEqual(missing_columns.exception.status_code, 422)

    def test_import_history_snapshot_comparison_and_filtered_excel_export(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            path = Path(directory) / "state" / "tickets.json"
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: path,
                transaction_root=lambda: Path(directory),
            )
            import_endpoint = route_handler(router, "/api/operacoes/chamados/importar", "POST")

            def import_version(simpress_rows, oracle_rows):
                return import_endpoint(
                    simpress_file=UploadFile(filename="pedidos.xlsx", file=BytesIO(self.workbook_bytes(simpress_rows))),
                    oracle_file=UploadFile(filename="solicitacoes.xlsx", file=BytesIO(self.workbook_bytes(oracle_rows))),
                )

            first = import_version(
                [["NumeroPedido", "NumeroSerie", "OrdemServico", "StatusAtual"], ["PED-1", "SN-1", 1001, "Aberto"]],
                [["Solicitação", "ID OS Corporate", "Número de Série", "Status", "Fila"], ["REQ-1", 1001, "SN-1", "Em execução", "Fila A"]],
            )
            current = import_version(
                [["NumeroPedido", "NumeroSerie", "OrdemServico", "StatusAtual"], ["PED-1", "SN-1", 2002, "Aberto"], ["PED-2", "SN-2", 3003, "Aguardando"]],
                [["Solicitação", "ID OS Corporate", "Número de Série", "Status", "Fila"], ["REQ-1", 2002, "SN-1", "Em espera", "Fila B"]],
            )

            history = route_handler(router, "/api/operacoes/chamados/historico", "GET")()
            self.assertEqual(history["total"], 2)
            self.assertEqual(history["eventos"][0]["id"], current["snapshot_id"])
            self.assertEqual(history["eventos"][1]["id"], first["snapshot_id"])

            snapshot = route_handler(router, "/api/operacoes/chamados/snapshot/{snapshot_id}", "GET")(first["snapshot_id"])
            self.assertEqual(len(snapshot["chamados"]), 1)
            evolution = route_handler(router, "/api/operacoes/chamados/evolucao", "GET")(
                anterior=first["snapshot_id"], atual=current["snapshot_id"]
            )
            self.assertEqual(evolution["contagens"], {"novos": 1, "alterados": 1, "iguais": 0, "ausentes_arquivo": 0})
            self.assertIn("Status Oracle", [field["campo"] for field in evolution["mudancas"]["alterados"][0]["campos"]])

            response = route_handler(router, "/api/operacoes/chamados/exportar", "GET")(
                snapshot_id=current["snapshot_id"], relation="conferida", status="", queue="",
                search="", stale_days=7, age_filter="", attention_only=False,
            )
            self.assertEqual(response.headers["content-type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.assertIn("Acompanhamento_de_Chamados_", response.headers["content-disposition"])

    def test_legacy_reports_show_only_latest_simpress_pedido_and_os_without_rewriting_raw_data(self):
        with tempfile.TemporaryDirectory(dir=PROJECT_ROOT) as directory:
            path = Path(directory) / "tickets.json"
            router = create_operations_router(
                load_comparison_history=lambda: [],
                comparison_history_for_display=lambda event: event,
                ticket_report_path=lambda: path,
                transaction_root=lambda: Path(directory),
            )
            legacy_ticket = {
                "serie": "BRJ5512DK7",
                "relacao": "conferida",
                "ordens_servico_conferidas": ["7124895"],
                "simpress": {
                    "serie": "BRJ5512DK7", "pedidos": ["48076976", "46735940"],
                    "ordens_servico": ["7134662", "7124895"], "chamados_cliente": [],
                    "status": ["Pedido entregue", "Cancelado"], "datas_pedido": [], "modelos": [],
                },
                "oracle": [{"ordem_servico": "7124895", "solicitacao": "REQ-1"}],
            }
            previous = {
                "id": "previous", "atualizado_em": "2026-10-09T10:00:00",
                "fontes": {}, "resumo": {"series_total": 1, "conferidas": 1},
                "chamados": [{
                    **legacy_ticket,
                    "simpress": {**legacy_ticket["simpress"], "ordens_servico": ["7124895"]},
                }],
            }
            current = {
                "id": "current", "atualizado_em": "2026-10-10T10:00:00",
                "fontes": {}, "resumo": {"series_total": 1, "conferidas": 1},
                "chamados": [legacy_ticket],
            }
            report = {
                "snapshot_id": "current", "atualizado_em": current["atualizado_em"],
                "fontes": current["fontes"], "resumo": current["resumo"],
                "chamados": current["chamados"], "historico": [previous, current],
            }
            path.write_text(json.dumps(report, ensure_ascii=False), encoding="utf-8")

            displayed = route_handler(router, "/api/operacoes/chamados", "GET")()
            self.assertEqual(displayed["chamados"][0]["simpress"]["pedidos"], ["48076976"])
            self.assertEqual(displayed["chamados"][0]["simpress"]["ordens_servico"], ["7134662"])
            self.assertEqual(displayed["chamados"][0]["ordens_servico_conferidas"], [])
            self.assertEqual(displayed["chamados"][0]["relacao"], "divergente")
            self.assertEqual(displayed["resumo"]["divergentes"], 1)
            history = route_handler(router, "/api/operacoes/chamados/historico", "GET")()
            self.assertEqual(history["eventos"][0]["resumo"]["divergentes"], 1)

            snapshot = route_handler(router, "/api/operacoes/chamados/snapshot/{snapshot_id}", "GET")("current")
            self.assertEqual(snapshot["chamados"][0]["simpress"]["pedidos"], ["48076976"])
            self.assertEqual(snapshot["chamados"][0]["simpress"]["ordens_servico"], ["7134662"])
            evolution = route_handler(router, "/api/operacoes/chamados/evolucao", "GET")(
                anterior="previous", atual="current"
            )
            self.assertIn("OS Simpress", [field["campo"] for field in evolution["mudancas"]["alterados"][0]["campos"]])

            exported_workbook = {}

            def build_and_capture(snapshot, tickets, filters, **kwargs):
                exported_workbook["bytes"] = build_ticket_workbook(snapshot, tickets, filters, **kwargs)
                return exported_workbook["bytes"]

            with patch("routers.operations.build_ticket_workbook", side_effect=build_and_capture):
                export = route_handler(router, "/api/operacoes/chamados/exportar", "GET")(
                    snapshot_id="current", relation="", status="", queue="", search="",
                    stale_days=7, age_filter="", attention_only=False,
                )
            workbook_bytes = exported_workbook["bytes"]
            self.assertEqual(export.headers["content-type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            workbook = openpyxl.load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
            try:
                sheet = workbook["Chamados"]
                headers = next(sheet.iter_rows(min_row=1, max_row=1))
                columns = {cell.value: index for index, cell in enumerate(headers, start=1)}
                self.assertEqual(sheet.cell(2, columns["Pedidos Simpress"]).value, "48076976")
                self.assertEqual(sheet.cell(2, columns["OS Simpress"]).value, "7134662")
            finally:
                workbook.close()

            raw_ticket = json.loads(path.read_text(encoding="utf-8"))["chamados"][0]
            self.assertEqual(raw_ticket["simpress"]["pedidos"], ["48076976", "46735940"])
            self.assertEqual(raw_ticket["simpress"]["ordens_servico"], ["7134662", "7124895"])

if __name__ == "__main__":
    unittest.main()
