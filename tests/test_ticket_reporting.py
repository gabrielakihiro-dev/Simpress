"""Tests for ticket filters, attention, snapshots, and spreadsheet export."""

from datetime import date
from io import BytesIO
import unittest

import openpyxl

from domains.ticket_reporting import (
    TICKET_TIMELINE_ANNOTATION_CATEGORY,
    build_ticket_timeline,
    build_ticket_workbook,
    compare_ticket_snapshots,
    filter_ticket_records,
    ticket_attention,
    ticket_relation_alerts,
    ticket_relation_provenance,
    ticket_report_snapshots,
)


def _ticket(
    serial="SN-1",
    *,
    relation="conferida",
    status="Em execução",
    queue="Fila A",
    last_update="2026-10-01",
    order="OS-1",
):
    oracle = [] if status is None else [{
        "solicitacao": "REQ-1", "ordem_servico": order, "status": status,
        "sub_status": "Pendente", "fila": queue, "fase_atendimento": "Execução",
        "ultima_atualizacao": last_update, "status_entregas": "", "status_atividades": "",
        "assunto": "Atendimento", "descricao": "Solicito suporte para o equipamento.",
    }]
    return {
        "serie": serial,
        "relacao": relation,
        "ordens_servico_conferidas": [order] if relation == "conferida" else [],
        "simpress": {"pedidos": ["PED-1"], "ordens_servico": [order], "chamados_cliente": [], "status": ["Em atendimento"]},
        "oracle": oracle,
    }


class TicketReportingTests(unittest.TestCase):
    def test_filters_match_relation_status_queue_search_and_age(self):
        tickets = [
            _ticket(),
            _ticket("SN-2", relation="divergente", queue="Fila B", last_update="2026-10-09"),
            _ticket("SN-3", relation="somente_simpress", status=None),
        ]
        result = filter_ticket_records(
            tickets, relation="conferida", status="Em execução", queue="Fila A",
            search="sn-1", age_filter="7", today=date(2026, 10, 10),
        )
        self.assertEqual([row["serie"] for row in result], ["SN-1"])

        split_match = _ticket("SN-4")
        split_match["oracle"].append({"status": "Aguardando", "fila": "Fila B"})
        self.assertEqual(
            filter_ticket_records([split_match], status="Em execução", queue="Fila B", today=date(2026, 10, 10)),
            [],
        )

    def test_triage_filters_by_city_and_simpress_or_partner_provider(self):
        curitiba = _ticket("SN-CURITIBA")
        parceiro = _ticket("SN-PARCEIRO")
        parceiro["oracle"][0]["nome_filial"] = "PARCEIRO"
        parceiro["oracle"][0]["cidade"] = "Maringá"
        parceiro["simpress"]["cidades_entrega"] = ["Curitiba"]

        rows = filter_ticket_records(
            [curitiba, parceiro], city="Maringá", provider="parceiro", queue="Fila A",
            today=date(2026, 10, 10),
        )

        self.assertEqual([row["serie"] for row in rows], ["SN-PARCEIRO"])

    def test_attention_queue_only_flags_active_tickets_without_update_for_three_days(self):
        divergent_but_recent = _ticket("SN-DIVERGENTE", relation="divergente", last_update="2026-10-09")
        no_order_but_recent = _ticket("SN-SEM-OS", relation="sem_os", last_update="2026-10-09", order="")
        only_one_source_but_recent = _ticket("SN-SOMENTE-SIMPRESS", relation="somente_simpress", status=None)
        stale_at_threshold = _ticket("SN-3-DIAS", last_update="2026-10-07")
        stale = _ticket("SN-ANTIGA", last_update="2026-09-20")
        closed = _ticket("SN-FECHADA", status="Concluído", last_update="2026-01-01")

        rows = filter_ticket_records(
            [closed, only_one_source_but_recent, no_order_but_recent, divergent_but_recent, stale, stale_at_threshold],
            attention_only=True, today=date(2026, 10, 10),
        )

        self.assertEqual([row["serie"] for row in rows], ["SN-ANTIGA", "SN-3-DIAS"])
        self.assertFalse(ticket_attention(divergent_but_recent, today=date(2026, 10, 10))["em_atencao"])
        self.assertFalse(ticket_attention(no_order_but_recent, today=date(2026, 10, 10))["em_atencao"])
        self.assertFalse(ticket_attention(only_one_source_but_recent, today=date(2026, 10, 10))["em_atencao"])
        self.assertFalse(ticket_attention(closed, today=date(2026, 10, 10))["em_atencao"])
        self.assertIn("Sem atualização há 3 dias", ticket_attention(stale_at_threshold, today=date(2026, 10, 10))["motivo"])
        self.assertIn("Sem atualização há 20 dias", ticket_attention(stale, today=date(2026, 10, 10))["motivo"])

    def test_ticket_results_sort_by_highest_os_and_leave_missing_os_last(self):
        high = _ticket("SN-7279287", order="7279287")
        middle = _ticket("SN-7279280", order="7279280")
        low = _ticket("SN-7265402", order="7265402")
        no_order = _ticket("SN-SEM-OS", status=None, relation="sem_os", order="")
        no_order["simpress"]["ordens_servico"] = []
        no_order["ordens_servico_conferidas"] = []

        rows = filter_ticket_records([low, no_order, middle, high], today=date(2026, 10, 10))

        self.assertEqual(
            [row["serie"] for row in rows],
            ["SN-7279287", "SN-7279280", "SN-7265402", "SN-SEM-OS"],
        )

    def test_missing_active_update_is_a_filterable_attention_item(self):
        missing = _ticket("SN-4", last_update="")
        attention = ticket_attention(missing, today=date(2026, 10, 10))
        self.assertTrue(attention["em_atencao"])
        self.assertIsNone(attention["dias_sem_atualizacao"])
        self.assertEqual(
            [row["serie"] for row in filter_ticket_records([missing], age_filter="sem_data", today=date(2026, 10, 10))],
            ["SN-4"],
        )

    def test_ticket_timeline_tracks_order_status_and_update_changes_per_import(self):
        first_ticket = _ticket("SN-LINHA", order="7232035", last_update="2026-10-01")
        first_ticket["simpress"]["pedidos"] = ["48076976"]
        first_ticket["simpress"]["status"] = ["Em separação"]
        second_ticket = _ticket("SN-LINHA", order="7279287", last_update="2026-10-09")
        second_ticket["simpress"]["pedidos"] = ["48076980"]
        second_ticket["simpress"]["status"] = ["Pedido em transporte"]
        second_ticket["oracle"][0]["solicitacao"] = "REQ-2"
        snapshots = [
            {"id": "version-1", "atualizado_em": "2026-10-01T09:00:00", "chamados": [first_ticket]},
            {"id": "version-2", "atualizado_em": "2026-10-10T09:00:00", "chamados": [second_ticket]},
        ]

        events = build_ticket_timeline(snapshots, "sn linha")
        changed_fields = {change["campo"] for change in events[1]["mudancas"]}

        self.assertEqual(len(events), 2)
        self.assertEqual(events[0]["item_id"], build_ticket_timeline(snapshots, "SN-LINHA")[0]["item_id"])
        self.assertNotEqual(events[0]["item_id"], events[1]["item_id"])
        self.assertTrue({"Pedido Simpress", "Solicitação Oracle", "OS Oracle", "Status Simpress", "Última atualização Oracle"}.issubset(changed_fields))
        self.assertEqual(events[1]["estado"]["pedidos_simpress"], ["48076980"])

    def test_provenance_and_uncertain_relationships_are_exported_for_review(self):
        ticket = _ticket("SN-REVISAR", relation="divergente", order="7232035")
        ticket["simpress"]["pedidos"] = ["48076976"]
        ticket["oracle"][0]["ordem_servico"] = "7124895"
        ticket["oracle"][0]["solicitacao"] = "REQ-9"
        provenance = ticket_relation_provenance(ticket)
        alerts = ticket_relation_alerts(ticket)
        timeline = build_ticket_timeline([{"id": "version-1", "atualizado_em": "2026-10-10", "chamados": [ticket]}], "SN-REVISAR")
        timeline[0]["anotacao"] = "Cobrar retorno do parceiro"

        workbook = openpyxl.load_workbook(BytesIO(build_ticket_workbook(
            {"atualizado_em": "2026-10-10"}, [ticket], {}, timelines={"SN-REVISAR": timeline},
        )), data_only=True)
        review = workbook["Conferência das relações"]
        review_columns = {cell.value: index for index, cell in enumerate(review[1], start=1)}
        self.assertTrue(provenance["serie"]["simpress"] and provenance["serie"]["oracle"])
        self.assertEqual(provenance["pedido"]["simpress"], ["48076976"])
        self.assertEqual(provenance["pedido"]["oracle_solicitacao"], ["REQ-9"])
        self.assertIn("os_divergente", {item["tipo"] for item in alerts})
        self.assertEqual(review.cell(2, review_columns["Série na Simpress"]).value, "Simpress")
        self.assertEqual(review.cell(2, review_columns["Solicitação Oracle"]).value, "REQ-9")
        self.assertEqual(review.cell(2, review_columns["Revisão necessária"]).value, "Sim")
        exported_timeline = workbook["Linha do tempo"]
        self.assertEqual(exported_timeline.cell(2, 13).value, "Cobrar retorno do parceiro")
        self.assertEqual(TICKET_TIMELINE_ANNOTATION_CATEGORY, "linha_do_tempo_chamados")
        workbook.close()

    def test_snapshot_comparison_reports_changes_new_and_absent_without_calling_absent_resolved(self):
        previous = {"id": "old", "atualizado_em": "2026-10-01", "fontes": {}, "resumo": {}, "chamados": [
            _ticket("ABC-1", order="100"), _ticket("OLD-2", status=None, relation="somente_simpress"),
        ]}
        current = {"id": "new", "atualizado_em": "2026-10-10", "fontes": {}, "resumo": {}, "chamados": [
            _ticket("abc1", order="200"), _ticket("NEW-3", status=None, relation="somente_simpress"),
        ]}

        result = compare_ticket_snapshots(previous, current)
        self.assertEqual(result["contagens"], {"novos": 1, "alterados": 1, "iguais": 0, "ausentes_arquivo": 1})
        self.assertEqual(result["mudancas"]["ausentes_arquivo"], ["OLD-2"])
        self.assertNotIn("resolvidos", result["contagens"])
        self.assertIn("não confirma", result["aviso_ausentes"])

    def test_simpress_delivery_details_are_compared_and_exported_only_while_in_progress(self):
        previous_ticket = _ticket("SN-DELIVERY")
        previous_ticket["simpress"].update({
            "descricoes_produto": ["Notebook modelo antigo"],
            "previsoes_entrega": ["28/09/2026"],
            "cidades_entrega": ["Curitiba"],
            "enderecos_entrega": ["Rua antiga, 10, Curitiba - PR"],
            "pedido_em_andamento": True,
        })
        current_ticket = _ticket("SN-DELIVERY")
        current_ticket["simpress"].update({
            "ordens_venda": ["OV_15799764"],
            "descricoes_produto": ["Notebook modelo novo"],
            "previsoes_entrega": ["15/09/2026"],
            "cidades_entrega": ["Maringá"],
            "enderecos_entrega": [],
            "pedido_em_andamento": False,
        })
        comparison = compare_ticket_snapshots(
            {"chamados": [previous_ticket]}, {"chamados": [current_ticket]},
        )
        changed_fields = {
            field["campo"]
            for row in comparison["mudancas"]["alterados"]
            for field in row["campos"]
        }
        self.assertTrue({
            "Descrição do produto Simpress", "Previsão de entrega Simpress",
            "Cidade de entrega Simpress", "Endereço de entrega Simpress",
        }.issubset(changed_fields))

        current_ticket["simpress"].update({
            "descricoes_produto": ["Notebook modelo novo"],
            "previsoes_entrega": ["15/09/2026"],
            "cidades_entrega": ["Maringá"],
            "enderecos_entrega": ["Avenida Central, 500, Maringá - PR"],
            "pedido_em_andamento": True,
        })
        content = build_ticket_workbook({"atualizado_em": "2026-10-10"}, [current_ticket], {})
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
        sheet = workbook["Chamados"]
        columns = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        self.assertEqual(sheet.cell(2, columns["OV Simpress"]).value, "OV_15799764")
        self.assertEqual(sheet.cell(2, columns["Descrição do produto"]).value, "Notebook modelo novo")
        self.assertEqual(sheet.cell(2, columns["Cidade de entrega"]).value, "Maringá")
        self.assertEqual(sheet.cell(2, columns["Endereço de entrega"]).value, "Avenida Central, 500, Maringá - PR")
        self.assertEqual(sheet.cell(2, columns["Previsão de entrega"]).value, "15/09/2026")
        workbook.close()

        current_ticket["simpress"]["pedido_em_andamento"] = False
        hidden_address = openpyxl.load_workbook(
            BytesIO(build_ticket_workbook({"atualizado_em": "2026-10-10"}, [current_ticket], {})),
            data_only=True,
        )
        hidden_columns = {cell.value: index for index, cell in enumerate(hidden_address["Chamados"][1], start=1)}
        self.assertEqual(
            hidden_address["Chamados"].cell(2, hidden_columns["Endereço de entrega"]).value,
            None,
        )
        hidden_address.close()

    def test_oracle_description_is_compared_and_exported(self):
        previous = _ticket("SN-DESCRIPTION")
        previous["oracle"][0]["descricao"] = "Solicito diagnóstico do equipamento."
        current = _ticket("SN-DESCRIPTION")
        long_description = "Solicito a troca do teclado e do mouse. " * 20
        current["oracle"][0]["descricao"] = long_description
        current["oracle"][0].update({
            "logradouro": "R Caeté", "numero_endereco": "150", "bairro": "Prado Velho",
            "cidade": "Curitiba", "atendido_pela_inibicao": "Não", "nome_filial": "SIMPRESS - PR",
        })
        without_description = _ticket("SN-SEM-DESCRICAO")
        without_description["oracle"][0]["descricao"] = ""

        comparison = compare_ticket_snapshots(
            {"chamados": [previous]}, {"chamados": [current]},
        )
        changed_fields = [
            field["campo"]
            for row in comparison["mudancas"]["alterados"]
            for field in row["campos"]
        ]
        self.assertIn("Descrição do chamado Oracle", changed_fields)
        self.assertTrue({
            "Logradouro Oracle", "Número do endereço Oracle", "Bairro Oracle", "Cidade Oracle",
            "Atendido pela inibição Oracle", "Nome da filial Oracle",
        }.issubset(changed_fields))

        workbook = openpyxl.load_workbook(
            BytesIO(build_ticket_workbook({"atualizado_em": "2026-10-10"}, [current, without_description], {})),
            data_only=True,
        )
        sheet = workbook["Chamados"]
        description_column = next(
            index for index, cell in enumerate(sheet[1], start=1)
            if cell.value == "Descrição do chamado Oracle"
        )
        self.assertEqual(sheet.cell(2, description_column).value, long_description)
        self.assertFalse(sheet.cell(2, description_column).alignment.wrap_text)
        self.assertEqual(sheet.row_dimensions[2].height, sheet.row_dimensions[3].height)
        exported_columns = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        self.assertEqual(sheet.cell(2, exported_columns["Logradouro Oracle"]).value, "R Caeté")
        self.assertEqual(sheet.cell(2, exported_columns["Número do endereço Oracle"]).value, "150")
        self.assertEqual(sheet.cell(2, exported_columns["Bairro Oracle"]).value, "Prado Velho")
        self.assertEqual(sheet.cell(2, exported_columns["Cidade Oracle"]).value, "Curitiba")
        self.assertEqual(sheet.cell(2, exported_columns["Atendido pela inibição Oracle"]).value, "Não")
        self.assertEqual(sheet.cell(2, exported_columns["Nome da filial Oracle"]).value, "SIMPRESS - PR")
        workbook.close()

    def test_oracle_service_orders_export_in_descending_numeric_order(self):
        ticket = _ticket("SN-OS-ORDER")
        ticket["simpress"]["ordens_servico"] = ["7124895", "7134662"]
        ticket["oracle"] = [
            {"solicitacao": "REQ-OLD", "ordem_servico": "7124895", "descricao": "OS anterior"},
            {"solicitacao": "REQ-NEW", "ordem_servico": "7134662", "descricao": "OS atual"},
        ]

        workbook = openpyxl.load_workbook(
            BytesIO(build_ticket_workbook({"atualizado_em": "2026-10-10"}, [ticket], {})),
            data_only=True,
        )
        sheet = workbook["Chamados"]
        order_column = next(index for index, cell in enumerate(sheet[1], start=1) if cell.value == "OS Oracle")
        simpress_order_column = next(index for index, cell in enumerate(sheet[1], start=1) if cell.value == "OS Simpress")
        self.assertEqual(order_column, simpress_order_column + 1)
        self.assertEqual([sheet.cell(row, order_column).value for row in (2, 3)], ["7134662", "7124895"])
        self.assertEqual(sheet.cell(2, simpress_order_column).value, "7134662, 7124895")
        workbook.close()

    def test_sales_order_is_included_in_simpress_history_comparison(self):
        previous = _ticket("SN-OV")
        current = _ticket("SN-OV")
        previous["simpress"]["ordens_venda"] = ["OV_15799763"]
        current["simpress"]["ordens_venda"] = ["OV_15799764"]

        comparison = compare_ticket_snapshots(
            {"chamados": [previous]}, {"chamados": [current]},
        )

        changed_fields = {
            field["campo"]
            for row in comparison["mudancas"]["alterados"]
            for field in row["campos"]
        }
        self.assertIn("Ordem de venda (OV) Simpress", changed_fields)

    def test_legacy_report_gets_a_stable_snapshot_id_and_export_is_a_valid_filtered_xlsx(self):
        legacy = {"atualizado_em": "2026-10-09", "fontes": {}, "resumo": {}, "chamados": [_ticket()]}
        first = ticket_report_snapshots(legacy)
        second = ticket_report_snapshots(legacy)
        self.assertEqual(first[0]["id"], second[0]["id"])

        filtered = filter_ticket_records(legacy["chamados"], relation="conferida")
        content = build_ticket_workbook(first[0], filtered, {"relation": "conferida"})
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Resumo", "Chamados", "Conferência das relações", "Linha do tempo"])
        sheet = workbook["Chamados"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 1).value, "SN-1")
        self.assertEqual(sheet.cell(1, 1).value, "Série")

        unsafe = _ticket("=1+1")
        unsafe_content = build_ticket_workbook(first[0], [unsafe], {})
        unsafe_workbook = openpyxl.load_workbook(BytesIO(unsafe_content), data_only=False)
        self.assertEqual(unsafe_workbook["Chamados"].cell(2, 1).value, "'=1+1")
        self.assertEqual(unsafe_workbook["Chamados"].cell(2, 1).data_type, "s")

    def test_snapshot_retention_keeps_only_the_50_most_recent_versions(self):
        history = [
            {"id": f"version-{index}", "atualizado_em": f"2026-10-{(index % 28) + 1:02d}", "chamados": []}
            for index in range(60)
        ]
        report = {**history[-1], "snapshot_id": history[-1]["id"], "historico": history}
        snapshots = ticket_report_snapshots(report)
        self.assertEqual(len(snapshots), 50)
        self.assertEqual(snapshots[0]["id"], "version-10")
        self.assertEqual(snapshots[-1]["id"], "version-59")


if __name__ == "__main__":
    unittest.main()
