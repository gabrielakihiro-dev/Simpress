"""User text that looks like a formula is written as literal text in every export."""

import io
from pathlib import Path
import tempfile
import unittest
import zipfile

import openpyxl
from openpyxl import Workbook

from core import integrity
from core.xlsx_safety import set_cell_value_as_text
from domains.history_export import _safe_export_cell

PROJECT_ROOT = Path(__file__).resolve().parents[1]
PAYLOADS = (
    '=HYPERLINK("http://example.invalid","abrir")',
    "=1+1",
    "+cmd",
    "@SUM(1)",
    "-2+3",
    "\t=1+1",
    " =1+1",
)


def sheet_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("xl/worksheets/sheet1.xml").decode("utf-8")


class FormulaTextCellTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(self.enterContext(tempfile.TemporaryDirectory(prefix=".formula-test-", dir=tempfile.gettempdir())))
        self.source = self.root / "modelo.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plan"
        sheet.append(["Serie", "Apelido"])
        sheet.append(["SN-FICTICIA", "antigo"])
        workbook.save(self.source)

    def assert_literal(self, data: bytes, expected: str) -> None:
        self.assertNotIn("<f>", sheet_xml(data))
        workbook = openpyxl.load_workbook(io.BytesIO(data))
        try:
            cell = workbook["Plan"]["B2"]
            self.assertEqual((cell.value, cell.data_type), (expected, "s"))
        finally:
            workbook.close()

    def test_helper_keeps_non_text_values_and_plain_text_untouched(self):
        sheet = Workbook().active
        set_cell_value_as_text(sheet["A1"], 12)
        set_cell_value_as_text(sheet["A2"], "texto")
        set_cell_value_as_text(sheet["A3"], "=1+1")
        self.assertEqual((sheet["A1"].value, sheet["A1"].data_type), (12, "n"))
        self.assertEqual((sheet["A2"].value, sheet["A2"].data_type), ("texto", "s"))
        self.assertEqual((sheet["A3"].value, sheet["A3"].data_type), ("=1+1", "s"))

    def test_validated_template_download_writes_formula_markers_as_text(self):
        for payload in PAYLOADS:
            with self.subTest(payload=payload):
                workbook = integrity.load_workbook_safely(self.source)
                try:
                    snapshot = integrity.workbook_snapshot(workbook)
                    set_cell_value_as_text(workbook["Plan"]["B2"], payload)
                    data = integrity.serialize_workbook_validated(
                        self.source, workbook, snapshot, {("Plan", "B2"): payload},
                    )
                finally:
                    workbook.close()
                self.assert_literal(data, payload)

    def test_transactional_workbook_edit_writes_formula_markers_as_text(self):
        for payload in PAYLOADS:
            with self.subTest(payload=payload):
                target = self.root / "editado.xlsx"
                target.write_bytes(self.source.read_bytes())

                def update(workbook):
                    set_cell_value_as_text(workbook["Plan"]["B2"], payload)
                    return {("Plan", "B2"): payload}

                with integrity.AtomicFileTransaction(self.root, [target]) as transaction:
                    self.assertTrue(integrity.stage_workbook_update(transaction, target, update))
                    transaction.commit()
                self.assert_literal(target.read_bytes(), payload)

    def test_history_export_cells_keep_formula_markers_as_text(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plan"
        for row, payload in enumerate(PAYLOADS, start=1):
            _safe_export_cell(sheet.cell(row=row, column=1), payload)
        buffer = io.BytesIO()
        workbook.save(buffer)
        self.assertNotIn("<f>", sheet_xml(buffer.getvalue()))
        reloaded = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
        values = [reloaded["Plan"].cell(row=row, column=1).value for row in range(1, len(PAYLOADS) + 1)]
        self.assertEqual(values, list(PAYLOADS))


if __name__ == "__main__":
    unittest.main()
