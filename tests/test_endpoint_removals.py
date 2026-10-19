"""Remoção do hostname mais antigo nas duplicidades do Endpoint."""

import copy
import json
import tempfile
import unittest
from contextlib import ExitStack
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

import openpyxl

import main
from core.errors import DomainConflictError, DomainError
from domains import endpoint_removals, inventory_state
from domains.effective_alias import build_endpoint_alias_index
from domains.history_annotations import history_annotation_row_fingerprint
from domains.inventory_comparison import compare_uploaded_workbooks
import test_integrity_workflows as base
from test_integrity_workflows import make_details, make_model

REFERENCE = datetime(2026, 10, 19, 12, 0, tzinfo=timezone(timedelta(hours=-3)))


def endpoint_record(serial, hostname, verified=None):
    record = {"SERIAL": serial, "NOMEDOCOMPUTADOR": hostname}
    if verified is not None:
        record["ULTIMAVERIFICACAOCOMEXITO"] = verified.isoformat()
    return record


def source_table(records):
    return {"sheet": "Inventario", "headers": set(), "records": records}


def make_endpoint(path: Path, rows):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "InvAllComputersSummary"
    worksheet.append(["ROTULODESERVICONUMERODESERIE", "NOMEDOCOMPUTADOR", "ULTIMAVERIFICACAOCOMEXITO"])
    for row in rows:
        worksheet.append(list(row))
    workbook.save(path)
    workbook.close()


class IndexTests(unittest.TestCase):
    def test_removed_pair_makes_the_remaining_hostname_unique(self):
        records = [endpoint_record("SN-1", "HOST-OLD"), endpoint_record("SN-1", "HOST-NEW")]
        self.assertIn("SN1", build_endpoint_alias_index(records)["duplicates"])

        index = build_endpoint_alias_index(records, endpoint_removals.removal_keys([
            {"serie": "sn-1", "hostname": "host old"},
        ]))

        self.assertEqual(index["unique"], {"SN1": "HOST-NEW"})
        self.assertEqual(index["duplicates"], {})

    def test_malformed_entries_never_match(self):
        keys = endpoint_removals.removal_keys([None, "x", {"serie": "SN-1"}, {"hostname": "H"}, {}])
        self.assertEqual(keys, frozenset())


class ComparisonTests(unittest.TestCase):
    def _records(self):
        return [
            endpoint_record("SN-1", "PRNOTE-OLD", REFERENCE - timedelta(days=30)),
            endpoint_record("SN-1", "PRNOTE-NEW", REFERENCE - timedelta(hours=2)),
        ]

    def test_ambiguous_record_flags_only_the_single_oldest_hostname(self):
        result = compare_uploaded_workbooks(source_table(self._records()), source_table([]), {}, as_of=REFERENCE)

        rows = result["registros_ambiguos"][0]["registros_endpoint"]
        self.assertEqual({row["computador"]: row["mais_antigo"] for row in rows}, {"PRNOTE-OLD": True, "PRNOTE-NEW": False})
        self.assertTrue(all(row["ultima_verificacao"] for row in rows))

    def test_no_button_when_dates_are_missing_tied_or_hostnames_repeat(self):
        same = REFERENCE - timedelta(days=3)
        cases = {
            "missing": [endpoint_record("SN-1", "PRNOTE-A", same), endpoint_record("SN-1", "PRNOTE-B")],
            "tied": [endpoint_record("SN-1", "PRNOTE-A", same), endpoint_record("SN-1", "PRNOTE-B", same)],
            "same host": [
                endpoint_record("SN-1", "PRNOTE-A", same - timedelta(days=1)),
                endpoint_record("SN-1", "prnote-a", same),
            ],
        }
        for name, records in cases.items():
            with self.subTest(name):
                result = compare_uploaded_workbooks(source_table(records), source_table([]), {}, as_of=REFERENCE)
                rows = result["registros_ambiguos"][0]["registros_endpoint"]
                self.assertFalse(any(row["mais_antigo"] for row in rows))

    def test_removed_hostname_leaves_the_comparison_but_not_the_total(self):
        removed = endpoint_removals.removal_keys([{"serie": "SN-1", "hostname": "PRNOTE-OLD"}])

        result = compare_uploaded_workbooks(
            source_table(self._records()), source_table([]), {}, as_of=REFERENCE,
            removed_endpoint_hostnames=removed,
        )

        self.assertEqual(result["total_endpoint"], 2)
        self.assertEqual(result["registros_ambiguos"], [])
        self.assertEqual(result["sem_acesso_7_dias"], [])

    def test_annotation_fingerprint_ignores_the_volatile_ambiguity_fields(self):
        base_row = {"computador": "H", "serie": "S", "motivo": "m"}
        enriched = {
            **base_row,
            "ultima_verificacao": "2026-10-19T10:00",
            "registros_endpoint": [{"computador": "H", "mais_antigo": True}],
        }
        self.assertEqual(
            history_annotation_row_fingerprint("registros_ambiguos", base_row),
            history_annotation_row_fingerprint("registros_ambiguos", enriched),
        )
        self.assertNotEqual(
            history_annotation_row_fingerprint("sem_acesso_7_dias", base_row),
            history_annotation_row_fingerprint("sem_acesso_7_dias", enriched),
        )


class RemovalWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp_context = tempfile.TemporaryDirectory(prefix=".endpoint-removal-test-", dir=base.PROJECT_ROOT)
        self.root = Path(self.temp_context.name)
        self.input_dir = self.root / "planilhas_entrada"
        self.input_dir.mkdir()
        self.state_dir = self.root / "state"
        self.state_dir.mkdir()
        self.details = self.input_dir / "EquipamentosDetalhes_atual.xlsx"
        self.endpoint = self.input_dir / "InvAllComputersSummary_atual.xlsx"
        self.model = self.input_dir / "modelo-ficticio.xlsx"
        self.removed = self.root / "hostnames_removidos.json"
        self.stack = ExitStack()
        for name, value in (
            ("BASE_DIR", self.root), ("STATE_DIR", self.state_dir), ("INPUT_DIR", self.input_dir),
            ("DETAILS_UPLOAD_FILE", self.details), ("ENDPOINT_UPLOAD_FILE", self.endpoint),
            ("EXPORT_MODEL_FILE", self.model), ("UPLOAD_BACKUPS_DIR", self.input_dir / "_backups"),
            ("MODIFIED_APELIDOS_FILE", self.root / "apelidos_modificados.json"),
            ("MANUAL_USERS_FILE", self.root / "usuarios_manuais.json"),
            ("REMOVED_ENDPOINT_HOSTNAMES_FILE", self.removed),
            ("ALIAS_HISTORY_FILE", self.root / "historico_alteracoes.json"),
            ("COMPARISON_HISTORY_FILE", self.root / "historico_comparacoes.json"),
            ("CUSTOM_UNIDADES_FILE", self.root / "unidades_custom.json"), ("TRANSACTION_ROOT", self.root),
        ):
            self.stack.enter_context(patch.object(main, name, value))
        self.old_state = (
            copy.deepcopy(inventory_state.equipamentos_cache),
            copy.deepcopy(inventory_state.unidades_cache),
            copy.deepcopy(inventory_state.stats_cache),
        )
        inventory_state.equipamentos_cache = {}
        inventory_state.unidades_cache = []
        inventory_state.stats_cache = {}
        (self.root / "apelidos_modificados.json").write_text("{}", encoding="utf-8")
        (self.root / "unidades_custom.json").write_text("[]", encoding="utf-8")
        (self.root / "historico_alteracoes.json").write_text("[]", encoding="utf-8")
        (self.root / "historico_comparacoes.json").write_text("[]", encoding="utf-8")
        make_details(self.details, [("SN-1", "ALIAS-1", "Modelo A", "Maringá", "Notebook", "87000000")])
        make_endpoint(self.endpoint, [
            ("SN-1", "HOST-OLD", "2026-09-01T10:00:00"),
            ("SN-1", "HOST-NEW", "2026-10-19T10:00:00"),
        ])
        make_model(self.model)
        self.assertTrue(main.load_database())
        self.services = main._endpoint_removal_services

    def tearDown(self):
        inventory_state.equipamentos_cache, inventory_state.unidades_cache, inventory_state.stats_cache = self.old_state
        self.stack.close()
        self.temp_context.cleanup()

    def test_removed_hostname_is_the_one_exported_and_can_be_restored(self):
        self.assertEqual(main.load_endpoint_aliases(), {})
        self.assertEqual(inventory_state.equipamentos_cache["SN-1"]["apelido"], "ALIAS-1")

        result = endpoint_removals.remove_endpoint_hostname(self.services, "SN-1", "HOST-OLD")

        self.assertEqual(result["mantidos"], ["HOST-NEW"])
        self.assertEqual(main.load_endpoint_aliases(), {"SN1": "HOST-NEW"})
        self.assertEqual(main.equipment_export.get_export_aliases(main._equipment_export_services, main.read_endpoint_alias_index())["SN1"], "HOST-NEW")
        self.assertEqual(inventory_state.equipamentos_cache["SN-1"]["apelido"], "HOST-NEW")
        stored = json.loads(self.removed.read_text(encoding="utf-8"))
        self.assertEqual([(item["serie"], item["hostname"]) for item in stored], [("SN-1", "HOST-OLD")])
        # A planilha de origem continua com as duas linhas.
        self.assertEqual(openpyxl.load_workbook(self.endpoint).active.max_row, 3)

        # Repetir a remoção (por exemplo, um duplo clique) não grava uma segunda entrada.
        with self.assertRaises(DomainConflictError):
            endpoint_removals.remove_endpoint_hostname(self.services, "SN-1", "host-old")
        self.assertEqual(len(json.loads(self.removed.read_text(encoding="utf-8"))), 1)

        endpoint_removals.restore_endpoint_hostname(self.services, "SN-1", "HOST-OLD")
        self.assertEqual(main.load_endpoint_aliases(), {})
        self.assertEqual(inventory_state.equipamentos_cache["SN-1"]["apelido"], "ALIAS-1")
        self.assertEqual(json.loads(self.removed.read_text(encoding="utf-8")), [])

    def test_the_last_remaining_hostname_and_unknown_pairs_are_refused(self):
        endpoint_removals.remove_endpoint_hostname(self.services, "SN-1", "HOST-OLD")

        with self.assertRaises(DomainConflictError):
            endpoint_removals.remove_endpoint_hostname(self.services, "SN-1", "HOST-NEW")
        with self.assertRaises(DomainConflictError):
            endpoint_removals.remove_endpoint_hostname(self.services, "SN-9", "HOST-X")
        with self.assertRaises(DomainError):
            endpoint_removals.remove_endpoint_hostname(self.services, "SN-1", "   ")
        self.assertEqual(
            [item["hostname"] for item in json.loads(self.removed.read_text(encoding="utf-8"))],
            ["HOST-OLD"],
        )

    def test_listing_returns_the_stored_pairs(self):
        endpoint_removals.remove_endpoint_hostname(self.services, "SN-1", "HOST-OLD")
        listed = endpoint_removals.list_removals(self.services)["removidos"]
        self.assertEqual([(item["serie"], item["hostname"]) for item in listed], [("SN-1", "HOST-OLD")])


if __name__ == "__main__":
    unittest.main()
