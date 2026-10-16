import asyncio
import copy
from concurrent.futures import ThreadPoolExecutor
import io
import json
import os
import shutil
import tempfile
import threading
import time
import unittest
import warnings
import zipfile
from datetime import datetime, timezone
from contextlib import ExitStack
from dataclasses import replace
from pathlib import Path
from urllib.parse import unquote
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
from domains.history_annotations import resolve_history_annotations
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from pydantic import ValidationError
from starlette.datastructures import UploadFile
from domains import inventory_update

with patch("openpyxl.load_workbook", side_effect=RuntimeError("testes não acessam planilhas locais")):
    import main


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PHONE_FORMAT_CASES = json.loads(
    (Path(__file__).parent / "fixtures" / "phone_format_cases.json").read_text(encoding="utf-8")
)


def make_details(path: Path, rows, headers=None):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "EquipamentosDetalhes"
    worksheet.append(headers or ["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP"])
    for row in rows:
        worksheet.append(list(row))
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["B"].width = 31
    worksheet["A1"].font = Font(bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    metadata = workbook.create_sheet("Metadados")
    metadata["A1"] = "Fórmula e estilo devem permanecer"
    metadata["B1"] = "=1+1"
    metadata["B1"].number_format = "0.00"
    metadata.row_dimensions[1].height = 29
    workbook.save(path)
    workbook.close()


def complete_details_headers(headers, include_installation_cnpj=True):
    """Build a 50-column Simpress upload header while preserving the test columns."""
    complete = list(headers)
    normalized = {main.normalize_workbook_header(header) for header in complete}
    installation_cnpj_headers = {
        main.normalize_workbook_header(header)
        for header in main.DETAILS_INSTALLATION_CNPJ_HEADER_ALIASES
    }
    if include_installation_cnpj and not (normalized & installation_cnpj_headers):
        complete.append("CNPJ Instalacao")
        normalized.add(main.normalize_workbook_header("CNPJ Instalacao"))

    index = 1
    while len(complete) < main.DETAILS_COMPLETE_SOURCE_COLUMN_COUNT:
        header = f"Detalhe Complementar {index:02d}"
        index += 1
        if main.normalize_workbook_header(header) in normalized:
            continue
        complete.append(header)
        normalized.add(main.normalize_workbook_header(header))
    return complete


def add_structured_table(path: Path, reference="A1:F2", name="EquipmentInventory"):
    workbook = openpyxl.load_workbook(path)
    worksheet = workbook["EquipamentosDetalhes"]
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    worksheet.add_table(table)
    workbook.save(path)
    workbook.close()


def make_endpoint(path: Path, rows):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "InvAllComputersSummary"
    worksheet.append(["ROTULODESERVICONUMERODESERIE", "NOMEDOCOMPUTADOR"])
    for row in rows:
        worksheet.append(list(row))
    workbook.save(path)
    workbook.close()


def make_model(path: Path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Parque"
    worksheet.append(["NumeroSerie", "Apelido", "Observacao", "Formula"])
    values = [
        ("S-1", "ANTIGO", "manter-a", "=1+1"),
        ("DUP", "DUP-ANTIGO", "manter-b", "=2+2"),
        ("SAME", "HOST-SAME", "manter-c", "=3+3"),
        ("FORMULA", "=CONCAT(\"HOST\",\"-F\")", "manter-d", "=4+4"),
        ("OUTDUP", "DUP-OUT-A", "manter-e", "=5+5"),
        ("OUTDUP", "DUP-OUT-B", "manter-f", "=6+6"),
    ]
    for row in values:
        worksheet.append(list(row))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:D7"
    worksheet.column_dimensions["B"].width = 37
    worksheet.row_dimensions[2].height = 27
    worksheet["A1"].font = Font(bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    other = workbook.create_sheet("Outra aba")
    other["A1"] = "Não alterar"
    other["B1"] = "=SUM(1,2)"
    other["B1"].font = Font(italic=True)
    other.column_dimensions["A"].width = 24
    workbook.save(path)
    workbook.close()


def make_transfer_template(path: Path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Transferência"
    for address in (
        "B5", "J5", "B8", "F8", "B17", "F17", "B19", "F19", "J19",
        "B21", "J21", "B23", "F23", "J23", "B25", "J25", "B29", "F29",
        "B31", "F31", "J31", "B33", "J33", "B35", "F35", "J35", "B37", "J37",
    ):
        worksheet[address] = "MARCADOR FICTICIO"
        worksheet[address].fill = PatternFill("solid", fgColor="EAF2F8")
    worksheet.merge_cells("B8:E8")
    worksheet.freeze_panes = "B5"
    workbook.create_sheet("Instruções")
    workbook.save(path)
    workbook.close()
    return _add_threaded_comment_parts(path)


def make_transfer_payload(**overrides):
    values = {
        "serial": "SERIE-FICTICIA",
        "descricao_equipamento": "EQUIPAMENTO FICTICIO",
        "contrato": 16467,
        "retirada_cnpj": "00.000.000/0000-00",
        "retirada_endereco": "ORIGEM FICTICIA",
        "retirada_contato": "REMETENTE FICTICIO",
        "retirada_email": "remetente@example.invalid",
        "retirada_telefone": "41900000000",
        "retirada_obs": "OBSERVACAO FICTICIA",
        "entrega_cnpj": "00.000.000/0000-00",
        "entrega_endereco": "DESTINO FICTICIO",
        "entrega_contato": "DESTINATARIO FICTICIO",
        "entrega_email": "destinatario@example.invalid",
        "entrega_telefone": "41911112222",
        "entrega_obs": "OBSERVACAO FICTICIA",
        "entrega_almoco": "12:00:00",
        "entrega_pausa_duracao": "1h",
    }
    values.update(overrides)
    return main.TransferenciaPayload(**values)


def _add_threaded_comment_parts(source: Path) -> Path:
    """Add synthetic threaded-comment package parts without touching source fixtures."""
    destination = source.with_name(f"{source.stem}-comentarios-ficticios.xlsx")
    comments_path = "xl/threadedComments/threadedComment1.xml"
    persons_path = "xl/persons/person.xml"
    sheet_path = "xl/worksheets/sheet1.xml"
    sheet_relationships_path = "xl/worksheets/_rels/sheet1.xml.rels"
    comment_id = "{00000000-0000-0000-0000-000000000001}"
    with zipfile.ZipFile(source, "r") as original, zipfile.ZipFile(destination, "w") as output:
        output.comment = original.comment
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "[Content_Types].xml":
                content = content.replace(
                    b"</Types>",
                    b'<Default Extension="threadedcomments" ContentType="application/vnd.ms-excel.threadedcomments+xml"/>'
                    b'<Override PartName="/xl/persons/person.xml" ContentType="application/vnd.ms-excel.person+xml"/>'
                    b"</Types>",
                )
            elif item.filename == sheet_path:
                content = content.replace(
                    b"</worksheet>",
                    b'<extLst><ext uri="{00000000-0000-0000-0000-000000000002}">'
                    b'<threadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments" '
                    b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
                    b'r:id="rIdThreadedComments"/></ext></extLst></worksheet>',
                )
            elif item.filename == sheet_relationships_path:
                content = content.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdThreadedComments" '
                    b'Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
                    b'Target="../threadedComments/threadedComment1.xml"/></Relationships>',
                )
            output.writestr(item, content)

        if sheet_relationships_path not in original.namelist():
            output.writestr(
                sheet_relationships_path,
                b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                b'<Relationship Id="rIdThreadedComments" '
                b'Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
                b'Target="../threadedComments/threadedComment1.xml"/></Relationships>',
            )
        output.writestr(
            comments_path,
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments">'
            b'<threadedComment ref="A1" dT="2026-01-01T00:00:00Z" personId="'
            + comment_id.encode("ascii")
            + b'"><text><r><t>Comentario ficticio</t></r></text></threadedComment></ThreadedComments>',
        )
        output.writestr(
            persons_path,
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<personList xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments">'
            b'<person displayName="Pessoa Ficticia" id="'
            + comment_id.encode("ascii")
            + b'" userId="pessoa@example.invalid" providerId="None"/></personList>',
        )
    return destination


class WorkflowIntegrityTests(unittest.TestCase):
    def setUp(self):
        self.temp_context = tempfile.TemporaryDirectory(prefix=".workflow-test-", dir=PROJECT_ROOT)
        self.root = Path(self.temp_context.name)
        self.input_dir = self.root / "planilhas_entrada"
        self.input_dir.mkdir()
        self.state_dir = self.root / "state"
        self.state_dir.mkdir()
        self.details = self.input_dir / "EquipamentosDetalhes_atual.xlsx"
        self.endpoint = self.input_dir / "InvAllComputersSummary_atual.xlsx"
        self.model = self.input_dir / "modelo-ficticio.xlsx"
        self.backups = self.input_dir / "_backups"
        self.overrides = self.root / "apelidos_modificados.json"
        self.manual_history = self.root / "historico_alteracoes.json"
        self.comparison_history = self.root / "historico_comparacoes.json"
        self.shared_annotations = self.root / "shared_annotations"
        self.pending_annotations = self.root / "pending_annotations"
        self.custom_units = self.root / "unidades_custom.json"
        self.stack = ExitStack()
        self.stack.enter_context(patch.object(main, "BASE_DIR", self.root))
        self.stack.enter_context(patch.object(main, "STATE_DIR", self.state_dir))
        self.stack.enter_context(patch.object(main, "INPUT_DIR", self.input_dir))
        self.stack.enter_context(patch.object(main, "DETAILS_UPLOAD_FILE", self.details))
        self.stack.enter_context(patch.object(main, "ENDPOINT_UPLOAD_FILE", self.endpoint))
        self.stack.enter_context(patch.object(main, "EXPORT_MODEL_FILE", self.model))
        self.stack.enter_context(patch.object(main, "UPLOAD_BACKUPS_DIR", self.backups))
        self.stack.enter_context(patch.object(main, "MODIFIED_APELIDOS_FILE", self.overrides))
        self.stack.enter_context(patch.object(main, "ALIAS_HISTORY_FILE", self.manual_history))
        self.stack.enter_context(patch.object(main, "COMPARISON_HISTORY_FILE", self.comparison_history))
        self.stack.enter_context(patch.object(main, "HISTORY_ANNOTATIONS_DIR", self.shared_annotations))
        self.stack.enter_context(patch.object(main, "LOCAL_HISTORY_ANNOTATION_OUTBOX_DIR", self.pending_annotations))
        self.stack.enter_context(patch.object(main, "HISTORY_ANNOTATIONS_SYNC_ENABLED", True))
        self.stack.enter_context(patch.object(main, "CUSTOM_UNIDADES_FILE", self.custom_units))
        self.stack.enter_context(patch.object(main, "TRANSACTION_ROOT", self.root))
        self.old_state = (
            copy.deepcopy(main.equipamentos_cache),
            copy.deepcopy(main.unidades_cache),
            copy.deepcopy(main.stats_cache),
        )
        main.equipamentos_cache = {}
        main.unidades_cache = []
        main.stats_cache = {}
        self.overrides.write_text("{}", encoding="utf-8")
        self.custom_units.write_text("[]", encoding="utf-8")

    def tearDown(self):
        main.equipamentos_cache, main.unidades_cache, main.stats_cache = self.old_state
        self.stack.close()
        self.temp_context.cleanup()

    def _submit_inventory_update(
        self,
        endpoint_bytes,
        details_bytes,
        key,
        *,
        endpoint_name="InvAll_idempotente.xlsx",
        details_name="Detalhes_idempotentes.xlsx",
        services=None,
    ):
        endpoint_upload = UploadFile(filename=endpoint_name, file=io.BytesIO(endpoint_bytes))
        details_upload = UploadFile(filename=details_name, file=io.BytesIO(details_bytes))
        return asyncio.run(main.update_inventory_files(
            endpoint_upload,
            details_upload,
            services=services or main._inventory_update_services,
            idempotency_key=key,
        ))

    def _prepare_inventory_upload_fixture(self):
        make_details(self.details, [
            ("SN-IDEMPOTENT", "HOST-OLD", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-IDEMPOTENT", "HOST-OLD")])
        make_model(self.model)
        self.comparison_history.write_text("[]", encoding="utf-8")
        self.manual_history.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())

    def test_partial_equipment_search_does_not_return_an_arbitrary_exact_lookup(self):
        main.equipamentos_cache = {
            "SERIE-ABC-001": {"serial": "SERIE-ABC-001", "modelo": "MODELO A"},
            "SERIE-ABC-002": {"serial": "SERIE-ABC-002", "modelo": "MODELO B"},
        }

        with self.assertRaises(main.HTTPException) as raised:
            main.get_equipamento("ABC")

        self.assertEqual(raised.exception.status_code, 404)
        self.assertEqual(
            [item["serial"] for item in main._buscar_equipamentos_por_termo("ABC")],
            ["SERIE-ABC-001", "SERIE-ABC-002"],
        )
        self.assertEqual(main.get_equipamento("serie-abc-002")["serial"], "SERIE-ABC-002")

    def test_alias_edit_payload_matches_equipment_serial_and_alias_limits(self):
        for serial in ("S", "S" * 128):
            with self.subTest(serial_length=len(serial)):
                equipment = main.EquipamentoUpdatePayload(serial=serial)
                alias = main.ApelidoUpdatePayload(serial=serial, apelido="A" * 160)
                self.assertEqual(alias.serial, equipment.serial)
                self.assertEqual(len(alias.apelido), 160)

        for serial in ("", "   ", "---", "S" * 129):
            with self.subTest(invalid_serial=serial[:12], serial_length=len(serial)):
                with self.assertRaises(ValidationError):
                    main.EquipamentoUpdatePayload(serial=serial)
                with self.assertRaises(ValidationError):
                    main.ApelidoUpdatePayload(serial=serial, apelido="HOST")

        with self.assertRaises(ValidationError):
            main.ApelidoUpdatePayload(serial="S-1", apelido="A" * 161)
        with self.assertRaises(ValidationError):
            main.EquipamentoUpdatePayload(serial="S-1", apelido="A" * 161)

        # Empty aliases are the existing contract for clearing a manual alias.
        self.assertEqual(main.ApelidoUpdatePayload(serial="S-1", apelido="").apelido, "")
        self.assertEqual(main.ApelidoUpdatePayload(serial="S-1", apelido="   ").apelido, "   ")

    def test_unit_ids_survive_sorting_edits_additions_and_database_reload(self):
        make_details(self.details, [
            ("SERIE-LONDRINA", "HOST-LONDRINA", "MODELO A", "Londrina", "Notebook", "86000000"),
            ("SERIE-MARINGA", "HOST-MARINGA", "MODELO B", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [
            ("SERIE-LONDRINA", "HOST-LONDRINA"),
            ("SERIE-MARINGA", "HOST-MARINGA"),
        ])
        saved_zeta = {
            "id": 47,
            "cidade": "Zeta",
            "endereco": "Rua Fictícia, 47",
            "uf": "PR",
            "cep": "85000-000",
            "cnpj": "47.000.000/0000-47",
            "is_sede": False,
            "endereco_completo": "Rua Fictícia, 47 - Zeta/PR - CEP 85000-000",
            "label": "Zeta - Rua Fictícia, 47 (47.000.000/0000-47)",
        }
        self.custom_units.write_text(json.dumps([saved_zeta]), encoding="utf-8")

        self.assertTrue(main.load_database())
        zeta = next(unit for unit in main.unidades_cache if unit["cidade"] == "Zeta")
        self.assertEqual(zeta["id"], 47)
        ids_before = {unit["cidade"]: unit["id"] for unit in main.unidades_cache}
        self.assertEqual(len(set(ids_before.values())), len(ids_before))

        main.salvar_unidade(main.UnidadePayload(
            cidade="Arapongas",
            endereco="Rua Nova, 8",
            uf="PR",
            cep="86700-000",
            cnpj="08.000.000/0000-08",
        ))
        ids_after_add = {unit["cidade"]: unit["id"] for unit in main.unidades_cache}
        for city, unit_id in ids_before.items():
            self.assertEqual(ids_after_add[city], unit_id)
        self.assertIn("Arapongas", ids_after_add)
        self.assertEqual(len(set(ids_after_add.values())), len(ids_after_add))

        main.salvar_unidade(main.UnidadePayload(
            id=47,
            cidade="Zeta",
            endereco="Rua Fictícia, 48",
            uf="PR",
            cep="85000-001",
            cnpj="47.000.000/0000-47",
        ))
        saved_units = json.loads(self.custom_units.read_text(encoding="utf-8"))
        self.assertEqual(len(saved_units), 2)
        edited_zeta = next(unit for unit in saved_units if unit["id"] == 47)
        self.assertEqual(edited_zeta["endereco"], "Rua Fictícia, 48")
        self.assertEqual(
            next(unit for unit in main.unidades_cache if unit["cidade"] == "Zeta")["id"],
            47,
        )

        before_stale_edit = self.custom_units.read_bytes()
        with self.assertRaises(main.HTTPException) as raised:
            main.salvar_unidade(main.UnidadePayload(
                id=999,
                cidade="Duplicata",
                endereco="Rua que não deve ser incluída",
                uf="PR",
                cep="80000-000",
                cnpj="99.000.000/0000-99",
            ))
        self.assertEqual(raised.exception.status_code, 409)
        self.assertEqual(self.custom_units.read_bytes(), before_stale_edit)

    def test_inventory_unit_edits_keep_source_identity_when_city_or_cnpj_changes(self):
        headers = ["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP", "CNPJ Instalacao"]
        make_details(
            self.details,
            [
                ("SN-A", "HOST-A", "Modelo A", "Cidade A", "Notebook", "80000000", "11.111.111/0001-11"),
                ("SN-B", "HOST-B", "Modelo B", "Cidade B", "Notebook", "80000001", "22.222.222/0001-22"),
                ("SN-C", "HOST-C", "Modelo C", "Cidade C", "Notebook", "80000002", "33.333.333/0001-33"),
            ],
            headers=headers,
        )
        make_endpoint(self.endpoint, [("SN-A", "HOST-A"), ("SN-B", "HOST-B"), ("SN-C", "HOST-C")])
        self.custom_units.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())

        edits = (
            ("Cidade A", "Cidade A Editada", "11.111.111/0001-11"),
            ("Cidade B", "Cidade B", "22.222.222/0001-99"),
            ("Cidade C", "Cidade C Editada", "33.333.333/0001-99"),
        )
        for index, (original_city, updated_city, updated_cnpj) in enumerate(edits, start=1):
            with self.subTest(city=original_city):
                source_unit = next(unit for unit in main.unidades_cache if unit["cidade"] == original_city)
                source_key = source_unit["_source_key"]
                source_id = source_unit["id"]
                saved = main.salvar_unidade(main.UnidadePayload(
                    id=source_id,
                    cidade=updated_city,
                    endereco=f"Rua Fictícia, {index}",
                    uf="PR",
                    cep=f"80000-00{index}",
                    cnpj=updated_cnpj,
                ))

                linked_units = [
                    unit for unit in main.unidades_cache
                    if unit.get("_source_key") == source_key
                ]
                self.assertEqual(len(linked_units), 1)
                self.assertEqual(linked_units[0]["id"], source_id)
                self.assertEqual(linked_units[0]["cidade"], updated_city)
                self.assertEqual(linked_units[0]["cnpj"], updated_cnpj)
                self.assertTrue(all("_source_key" not in unit for unit in saved["unidades"]))

        stored = json.loads(self.custom_units.read_text(encoding="utf-8"))
        self.assertEqual(len(stored), 3)
        self.assertEqual(len({unit["_source_key"] for unit in stored}), 3)
        self.assertTrue(all("_source_key" not in unit for unit in main.get_unidades()))

    def test_concurrent_unit_saves_preserve_both_catalog_updates(self):
        make_details(
            self.details,
            [("SN-1", "HOST-1", "Modelo A", "Maringá", "Notebook", "87000000")],
        )
        make_endpoint(self.endpoint, [("SN-1", "HOST-1")])
        self.custom_units.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())

        start_together = threading.Barrier(3)
        results = []

        def save_unit(city, cep):
            start_together.wait(timeout=5)
            try:
                results.append(main.salvar_unidade(main.UnidadePayload(
                    cidade=city,
                    endereco="Rua Fictícia, 10",
                    uf="PR",
                    cep=cep,
                    cnpj="10.000.000/0000-00",
                )))
            except Exception as error:
                results.append(error)

        workers = [
            threading.Thread(target=save_unit, args=("Unidade Concorrente A", "80000-001")),
            threading.Thread(target=save_unit, args=("Unidade Concorrente B", "80000-002")),
        ]
        for worker in workers:
            worker.start()
        start_together.wait(timeout=5)
        for worker in workers:
            worker.join(timeout=10)

        self.assertTrue(all(not worker.is_alive() for worker in workers))
        self.assertEqual(len(results), 2)
        self.assertFalse([result for result in results if isinstance(result, Exception)])
        persisted = json.loads(self.custom_units.read_text(encoding="utf-8"))
        saved_cities = {unit["cidade"] for unit in persisted}
        self.assertEqual(saved_cities, {"Unidade Concorrente A", "Unidade Concorrente B"})
        saved_ids = [unit["id"] for unit in persisted]
        self.assertEqual(len(saved_ids), len(set(saved_ids)))
        self.assertTrue({"Unidade Concorrente A", "Unidade Concorrente B"}.issubset(
            {unit["cidade"] for unit in main.unidades_cache}
        ))

    def test_unit_save_transaction_failure_preserves_catalog_and_cache(self):
        make_details(
            self.details,
            [("SN-1", "HOST-1", "Modelo A", "Maringá", "Notebook", "87000000")],
        )
        make_endpoint(self.endpoint, [("SN-1", "HOST-1")])
        self.custom_units.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())
        catalog_before = self.custom_units.read_bytes()
        units_before = copy.deepcopy(main.unidades_cache)

        with patch.object(
            main.AtomicFileTransaction,
            "update_json",
            side_effect=main.FileTransactionError("falha fictícia de gravação"),
        ):
            with self.assertRaises(main.HTTPException) as raised:
                main.salvar_unidade(main.UnidadePayload(
                    cidade="Unidade Não Persistida",
                    endereco="Rua Fictícia, 11",
                    uf="PR",
                    cep="80000-003",
                    cnpj="11.000.000/0000-00",
                ))

        self.assertEqual(raised.exception.status_code, 500)
        self.assertEqual(self.custom_units.read_bytes(), catalog_before)
        self.assertEqual(main.unidades_cache, units_before)

    def test_unit_save_rolls_back_when_resulting_inventory_cannot_be_reloaded(self):
        make_details(self.details, [
            ("SN-1", "HOST-1", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-1", "HOST-1")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        self.custom_units.write_text("[]", encoding="utf-8")
        units_before = copy.deepcopy(main.unidades_cache)
        payload = main.UnidadePayload(
            cidade="Unidade Não Persistida",
            endereco="Rua Fictícia, 12",
            uf="PR",
            cep="80000-004",
            cnpj="12.000.000/0000-00",
        )

        with patch.object(main, "load_database", return_value=False):
            with self.assertRaises(main.HTTPException) as raised:
                main.salvar_unidade(payload)

        self.assertEqual(raised.exception.status_code, 500)
        self.assertEqual(self.custom_units.read_text(encoding="utf-8"), "[]")
        self.assertEqual(main.unidades_cache, units_before)

    def test_repeating_unit_creation_is_idempotent(self):
        make_details(self.details, [
            ("SN-1", "HOST-1", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-1", "HOST-1")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        self.custom_units.write_text("[]", encoding="utf-8")
        payload = main.UnidadePayload(
            cidade="Unidade Repetida",
            endereco="Rua Fictícia, 13",
            uf="PR",
            cep="80000-005",
            cnpj="13.000.000/0000-00",
        )

        first = main.salvar_unidade(payload)
        second = main.salvar_unidade(payload)

        saved = json.loads(self.custom_units.read_text(encoding="utf-8"))
        matches = [unit for unit in saved if unit["cidade"] == "Unidade Repetida"]
        self.assertEqual(len(matches), 1)
        returned_ids = [
            next(unit["id"] for unit in result["unidades"] if unit["cidade"] == "Unidade Repetida")
            for result in (first, second)
        ]
        self.assertEqual(returned_ids[0], returned_ids[1])

    def test_detail_upload_headers_and_active_reader_share_normalization(self):
        variants = ("NumeroSerie", "numero_serie", "NÚMERO-SERIE", "Número de Série")
        for index, serial_header in enumerate(variants):
            with self.subTest(serial_header=serial_header):
                path = self.input_dir / f"EquipamentosDetalhes_variante_{index}.xlsx"
                make_details(
                    path,
                    [(
                        "SERIE-FICTICIA", "ALIAS-FICTICIO", "Modelo A", "Maringá",
                        "Notebook", "87000000", "00.000.000/0000-00",
                    )],
                    headers=complete_details_headers([
                        serial_header, "HOSTNAME", "Modelo", "Cidade", "Tipo", "CEP",
                    ]),
                )

                uploaded = main.validate_uploaded_workbook(path.read_bytes(), "detalhes")

                self.assertEqual(uploaded["populated_rows"], 1)
                self.assertEqual(main.find_details_sheet(path), "EquipamentosDetalhes")

    def test_inventory_update_rejects_reduced_export_and_50_columns_without_installation_cnpj(self):
        reduced_export = self.root / "EquipamentosDetalhes-exportado.xlsx"
        make_details(
            reduced_export,
            [("SERIE-FICTICIA", "ALIAS-FICTICIO", "Modelo A", "Maringá", "Notebook", "87000000")],
        )
        reduced_bytes = reduced_export.read_bytes()

        with self.assertRaises(main.HTTPException) as reduced_error:
            main.validate_uploaded_workbook(reduced_bytes, "detalhes")

        self.assertEqual(reduced_error.exception.status_code, 400)
        self.assertIn("50 colunas", reduced_error.exception.detail)

        export_headers = [
            "CodigoContrato", "ItemDeContrato", "RazaoSocial", "NumeroSerie", "Apelido",
            "IdCliente", "Modelo", "Endereco", "Cidade", "UF", "CEP", "Ip", "Tipo",
            "MaquinaBackup", "Monitoramento", "DataPrimeiraInstalacao", "DataUltimaInstalacao",
        ]
        export_row = [
            "CONTRATO-FICTICIO", "1", "EMPRESA FICTICIA", "SERIE-FICTICIA", "ALIAS-FICTICIO",
            "CLIENTE-FICTICIO", "Modelo A", "Rua Ficticia", "Maringa", "PR", "87000000",
            "192.0.2.1", "Notebook", "Nao", "Nao", "2020-01-01", "2020-01-01",
        ]
        make_details(reduced_export, [export_row], headers=export_headers)
        with self.assertRaises(main.HTTPException) as seventeen_column_error:
            main.validate_uploaded_workbook(reduced_export.read_bytes(), "detalhes")

        self.assertEqual(seventeen_column_error.exception.status_code, 400)
        self.assertIn("17 colunas", seventeen_column_error.exception.detail)
        self.assertIn("não contém o cabeçalho CNPJ de instalação", seventeen_column_error.exception.detail)
        self.assertIn("50 colunas", seventeen_column_error.exception.detail)

        headers_without_cnpj = complete_details_headers(
            ["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP"],
            include_installation_cnpj=False,
        )
        make_details(
            reduced_export,
            [("SERIE-FICTICIA", "ALIAS-FICTICIO", "Modelo A", "Maringá", "Notebook", "87000000")],
            headers=headers_without_cnpj,
        )
        with self.assertRaises(main.HTTPException) as missing_cnpj_error:
            main.validate_uploaded_workbook(reduced_export.read_bytes(), "detalhes")

        self.assertEqual(missing_cnpj_error.exception.status_code, 400)
        self.assertIn("CNPJ de instalação", missing_cnpj_error.exception.detail)

        headers_with_empty_cnpj = complete_details_headers(
            ["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP"]
        )
        make_details(
            reduced_export,
            [("SERIE-FICTICIA", "ALIAS-FICTICIO", "Modelo A", "Maringá", "Notebook", "87000000")],
            headers=headers_with_empty_cnpj,
        )
        with self.assertRaises(main.HTTPException) as empty_cnpj_error:
            main.validate_uploaded_workbook(reduced_export.read_bytes(), "detalhes")

        self.assertEqual(empty_cnpj_error.exception.status_code, 400)
        self.assertIn("série(s) sem CNPJ de instalação", empty_cnpj_error.exception.detail)

    def test_update_route_rejects_export_format_before_changing_active_files(self):
        make_details(self.details, [("SERIE-ATUAL", "ALIAS-ATUAL", "Modelo A", "Maringá", "Notebook", "87000000")])
        make_endpoint(self.endpoint, [("SERIE-ATUAL", "ALIAS-ATUAL")])
        before = {path: path.read_bytes() for path in (self.details, self.endpoint)}
        with tempfile.TemporaryDirectory(prefix=".reduced-export-test-", dir=PROJECT_ROOT) as directory:
            reduced_path = Path(directory) / "exportacao-reduzida.xlsx"
            make_details(
                reduced_path,
                [("SERIE-NOVA", "ALIAS-NOVO", "Modelo B", "Londrina", "Notebook", "86000000")],
            )
            reduced_bytes = reduced_path.read_bytes()

        with self.assertRaises(main.HTTPException) as raised:
            asyncio.run(main.atualizar_planilhas(
                UploadFile(filename="InvAll_nova.xlsx", file=io.BytesIO(_endpoint_bytes([("SERIE-NOVA", "ALIAS-NOVO")]))),
                UploadFile(filename="EquipamentosDetalhes-exportado.xlsx", file=io.BytesIO(reduced_bytes)),
            ))

        self.assertEqual(raised.exception.status_code, 400)
        self.assertIn("Atualizar Planilhas", raised.exception.detail)
        self.assertEqual({path: path.read_bytes() for path in before}, before)

    def test_uploaded_details_preserves_text_identifiers_and_na_values(self):
        headers = [
            "nÚmero_de_Série", "hOsT-name", "mÓde-lo", "Cidade", "Tipo", "CEP",
            "IdCliente", "CNPJInstalacao",
        ]
        complete_headers = complete_details_headers(headers)
        make_details(
            self.details,
            [
                ("001234", "HOST-TEXTO", "Modelo A", "Maringá", "Notebook", "87000000", "101", "00123456000199"),
                (None, "HOST-SEM-SERIE", "Modelo B", "Londrina", "Notebook", "86000000", "102", "00123456000199"),
                (1234, "HOST-NUMERICO", "Modelo C", "Ponta Grossa", "Notebook", "84000000", "103", "00123456000199"),
                ("NA", "HOST-NA", "Modelo D", "Cascavel", "Notebook", "85800000", "NA", "00987654000188"),
            ],
            headers=complete_headers,
        )
        make_endpoint(self.endpoint, [("OUTRA-SERIE", "HOST-OUTRO")])

        uploaded = main.validate_uploaded_workbook(self.details.read_bytes(), "detalhes")
        self.assertEqual(uploaded["populated_rows"], 3)
        read_headers, records = main._read_details_sheet_records(self.details, "EquipamentosDetalhes")
        self.assertEqual(records[0][read_headers[0]], "001234")
        self.assertEqual(records[2][read_headers[0]], 1234)
        self.assertEqual(records[3][read_headers[0]], "NA")
        self.assertEqual(records[3]["IdCliente"], "NA")
        self.assertEqual(records[0]["CNPJInstalacao"], "00123456000199")
        self.assertEqual(records[3]["CNPJInstalacao"], "00987654000188")

        self.assertTrue(main.load_database())
        self.assertIn("001234", main.equipamentos_cache)
        self.assertIn("1234", main.equipamentos_cache)
        self.assertIn("NA", main.equipamentos_cache)
        self.assertNotIn("HOST-SEM-SERIE", main.equipamentos_cache)
        self.assertEqual(main.equipamentos_cache["001234"]["apelido"], "HOST-TEXTO")
        self.assertEqual(main.equipamentos_cache["1234"]["apelido"], "HOST-NUMERICO")
        self.assertEqual(main.equipamentos_cache["NA"]["retirada_cnpj"], "00987654000188")

    def test_installation_cnpj_header_with_branch_suffix_populates_equipment_and_unit(self):
        cnpj = "12.345.678/0001-90"
        make_details(
            self.details,
            [(
                "SERIE-CNPJ-FICTICIA",
                "HOST-CNPJ-FICTICIO",
                "Modelo A",
                "Foz do Iguaçu",
                "Notebook",
                "85852-010",
                cnpj,
            )],
            headers=[
                "NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP",
                "CNPJ de Instalação da Unidade / Filial SEBRAE",
            ],
        )
        make_endpoint(self.endpoint, [("SERIE-CNPJ-FICTICIA", "HOST-CNPJ-FICTICIO")])

        self.assertTrue(main.load_database())
        equipment = main.get_equipamento("SERIE-CNPJ-FICTICIA")
        delivery_unit = next(unit for unit in main.unidades_cache if unit["cidade"] == "Foz do Iguaçu")

        self.assertEqual(equipment["retirada_cnpj"], cnpj)
        self.assertEqual(delivery_unit["cnpj"], cnpj)

    def test_upload_rejects_duplicate_normalized_or_semantic_headers(self):
        make_details(
            self.details,
            [("SN-1", "ALIAS-A", "ALIAS-B", "Modelo", "Maringá", "Notebook")],
            headers=["NumeroSerie", "Apelido", "apelido", "Modelo", "Cidade", "Tipo"],
        )

        with self.assertRaises(main.HTTPException) as raised:
            main.validate_uploaded_workbook(self.details.read_bytes(), "detalhes")

        self.assertEqual(raised.exception.status_code, 400)
        self.assertIn("cabeçalhos", raised.exception.detail.lower())

    def test_unreadable_newest_source_does_not_fall_back_to_older_inventory(self):
        archived = self.input_dir / "EquipamentosDetalhes_arquivo_antigo.xlsx"
        make_details(self.details, [
            ("SN-NOVO", "ALIAS-NOVO", "Modelo novo", "Maringá", "Notebook", "87000000"),
        ])
        make_details(archived, [
            ("SN-ANTIGO", "ALIAS-ANTIGO", "Modelo antigo", "Ponta Grossa", "Notebook", "84000000"),
        ])
        make_endpoint(self.endpoint, [("OUTRA-SERIE", "HOST-OUTRO")])
        make_model(self.model)
        timestamp = time.time()
        os.utime(archived, (timestamp - 60, timestamp - 60))
        os.utime(self.details, (timestamp, timestamp))
        self.assertTrue(main.load_database())
        self.assertIn("SN-NOVO", main.equipamentos_cache)

        real_open = open
        locked = self.details.resolve()

        def open_with_locked_newest(file, *args, **kwargs):
            if Path(file).resolve() == locked:
                raise PermissionError(13, "arquivo em uso")
            return real_open(file, *args, **kwargs)

        with patch("builtins.open", open_with_locked_newest):
            with self.assertRaises(main.FileBusyError):
                main.find_details_sheet(self.details)
            with self.assertRaises(main.FileBusyError):
                main.get_primary_equipment_file()
            self.assertFalse(main.load_database())

        self.assertIn("SN-NOVO", main.equipamentos_cache)
        self.assertNotIn("SN-ANTIGO", main.equipamentos_cache)
        self.assertEqual(main.get_primary_equipment_file().resolve(), self.details.resolve())

    def test_equipment_edit_changes_only_loaded_source_and_keeps_other_active_records(self):
        archived = self.input_dir / "EquipamentosDetalhes_arquivo_antigo.xlsx"
        make_details(self.details, [
            ("SN-EDITAR", "ALIAS-EDITAR", "Modelo antes", "Maringá", "Notebook", "87000000"),
            ("SN-MANTER", "ALIAS-MANTER", "Modelo manter", "Londrina", "Notebook", "86000000"),
        ])
        make_details(archived, [
            ("SN-ARQUIVO", "ALIAS-ARQUIVO", "Modelo arquivo", "Ponta Grossa", "Notebook", "84000000"),
        ])
        make_endpoint(self.endpoint, [("OUTRA-SERIE", "HOST-OUTRO")])
        make_model(self.model)
        timestamp = time.time()
        os.utime(archived, (timestamp - 60, timestamp - 60))
        os.utime(self.details, (timestamp, timestamp))
        archived_before = archived.read_bytes()

        self.assertTrue(main.load_database())
        self.assertEqual(main.get_primary_equipment_file().resolve(), self.details.resolve())
        selected = main.get_equipamento("SN-EDITAR")
        self.assertEqual(selected["modelo"], "Modelo antes")

        result = main.salvar_equipamento(main.EquipamentoUpdatePayload(
            serial=selected["serial"],
            modelo="Modelo atualizado",
            tipo="Notebook",
            idCliente="CLIENTE-ATUALIZADO",
            revisao=selected["revisao"],
        ))

        self.assertTrue(result["sucesso"])
        self.assertEqual(archived.read_bytes(), archived_before)
        self.assertEqual(main.get_primary_equipment_file().resolve(), self.details.resolve())
        self.assertTrue(main.load_database())
        self.assertIn("SN-EDITAR", main.equipamentos_cache)
        self.assertIn("SN-MANTER", main.equipamentos_cache)
        self.assertEqual(main.equipamentos_cache["SN-EDITAR"]["modelo"], "Modelo atualizado")
        self.assertEqual(main.equipamentos_cache["SN-MANTER"]["modelo"], "Modelo manter")

    def test_explicit_null_clears_the_same_field_in_excel_and_cache(self):
        make_details(self.details, [
            ("SN-CLEAR", "ALIAS-CLEAR", "Modelo original", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-OUTRO", "HOST-OUTRO")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        selected = main.get_equipamento("SN-CLEAR")

        main.salvar_equipamento(main.EquipamentoUpdatePayload(
            serial="SN-CLEAR",
            revisao=selected["revisao"],
            modelo=None,
        ))

        workbook = openpyxl.load_workbook(self.details, data_only=False)
        self.assertIsNone(workbook["EquipamentosDetalhes"]["C2"].value)
        workbook.close()
        self.assertEqual(main.get_equipamento("SN-CLEAR")["modelo"], "")

    def test_cache_reloads_after_another_process_replaces_the_active_files(self):
        make_details(self.details, [
            ("SN-EXTERNO", "ALIAS-EXTERNO", "Modelo antes", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-OUTRO", "HOST-OUTRO")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        self.assertFalse(main.reload_database_if_changed_externally())

        # An in-process edit publishes its own state and must not trigger a reload.
        main.aplicar_alteracao_apelido("SN-EXTERNO", "MANUAL-LOCAL")
        with patch.object(main, "load_database", side_effect=AssertionError("recarga desnecessária")):
            self.assertFalse(main.reload_database_if_changed_externally())

        # The inbox script running in another process rewrites the workbook.
        make_details(self.details, [
            ("SN-EXTERNO", "ALIAS-EXTERNO", "Modelo depois", "Maringá", "Notebook", "87000000"),
            ("SN-NOVO-EXTERNO", "ALIAS-NOVO", "Modelo novo", "Maringá", "Notebook", "87000000"),
        ])
        later = time.time() + 5
        os.utime(self.details, (later, later))
        self.assertEqual(main.equipamentos_cache["SN-EXTERNO"]["modelo"], "Modelo antes")

        self.assertTrue(main.reload_database_if_changed_externally())

        self.assertEqual(main.equipamentos_cache["SN-EXTERNO"]["modelo"], "Modelo depois")
        self.assertIn("SN-NOVO-EXTERNO", main.equipamentos_cache)
        self.assertEqual(main.equipamentos_cache["SN-EXTERNO"]["apelido"], "MANUAL-LOCAL")
        self.assertFalse(main.reload_database_if_changed_externally())

    def test_equipment_text_starting_with_formula_marker_is_saved_as_literal_text(self):
        make_details(self.details, [
            ("SN-FORMULA", "ALIAS-FORMULA", "Modelo original", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-OUTRO", "HOST-OUTRO")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        selected = main.get_equipamento("SN-FORMULA")
        payload_text = '=HYPERLINK("http://example.invalid","abrir")'

        result = main.salvar_equipamento(main.EquipamentoUpdatePayload(
            serial="SN-FORMULA",
            revisao=selected["revisao"],
            modelo=payload_text,
        ))

        self.assertTrue(result["sucesso"])
        with zipfile.ZipFile(self.details) as archive:
            sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertNotIn("<f>", sheet_xml)
        workbook = openpyxl.load_workbook(self.details, data_only=False)
        cell = workbook["EquipamentosDetalhes"]["C2"]
        self.assertEqual((cell.value, cell.data_type), (payload_text, "s"))
        workbook.close()

    def test_equipment_edit_preserves_the_supported_16467_contract_metadata(self):
        headers = [
            "NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP",
            "CodigoContrato", "RazaoSocial",
        ]
        make_details(
            self.details,
            [("SN-16467", "HOST-16467", "Modelo original", "Maringá", "Notebook", "87000000", 16467, "RAZAO SOCIAL FICTICIA")],
            headers=headers,
        )
        make_endpoint(self.endpoint, [("SN-16467", "HOST-16467")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        selected = main.get_equipamento("SN-16467")
        self.assertEqual(selected["contrato"], "16467")
        self.assertEqual(selected["razao_social"], "RAZAO SOCIAL FICTICIA")

        main.salvar_equipamento(main.EquipamentoUpdatePayload(
            serial="SN-16467",
            revisao=selected["revisao"],
            modelo="Modelo atualizado",
        ))
        updated = main.get_equipamento("SN-16467")
        self.assertEqual(updated["contrato"], "16467")
        self.assertEqual(updated["razao_social"], "RAZAO SOCIAL FICTICIA")

        previous_cache = copy.deepcopy(main.equipamentos_cache)
        make_details(
            self.details,
            [("SN-OUTRO-CONTRATO", "HOST-OUTRO", "Modelo", "Maringá", "Notebook", "87000000", 777, "RAZAO OUTRA")],
            headers=headers,
        )
        self.assertFalse(main.load_database())
        self.assertEqual(main.equipamentos_cache, previous_cache)

    def test_equipment_serial_that_normalizes_to_empty_is_rejected(self):
        make_details(self.details, [("SN-BASE", "HOST-BASE", "Modelo", "Maringá", "Notebook", "87000000")])
        make_endpoint(self.endpoint, [("SN-BASE", "HOST-BASE")])
        make_model(self.model)
        self.assertTrue(main.load_database())

        with self.assertRaises(ValueError):
            main.EquipamentoUpdatePayload(serial="---", novo=True)
        self.assertNotIn("---", main.equipamentos_cache)

    def test_unit_payload_rejects_invalid_cep_uf_and_cnpj(self):
        base = {"cidade": "Maringá", "endereco": "Rua Fictícia", "cnpj": "11.000.000/0000-00"}
        for invalid in (
            {"cep": "abc"},
            {"cep": "123"},
            {"uf": "XX"},
            {"cnpj": "123"},
        ):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                main.UnidadePayload(**{**base, **invalid})

        formatted = main.UnidadePayload(**base, cep="87000-000", uf="pr")
        self.assertEqual(formatted.cep, "87000-000")
        self.assertEqual(formatted.uf, "PR")

    def test_equipment_edit_refuses_duplicate_serial_without_writing(self):
        make_details(self.details, [
            ("SN-DUP", "ALIAS-PRIMEIRO", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("OUTRA-SERIE", "HOST-OUTRO")])
        make_model(self.model)
        self.assertTrue(main.load_database())

        workbook = openpyxl.load_workbook(self.details)
        workbook["EquipamentosDetalhes"].append(
            ["SN-DUP", "ALIAS-SEGUNDO", "Modelo B", "Londrina", "Notebook", "86000000"]
        )
        workbook.save(self.details)
        workbook.close()
        before = self.details.read_bytes()
        cache_before = copy.deepcopy(main.equipamentos_cache)

        with self.assertRaises(main.HTTPException) as raised:
            main.salvar_equipamento(main.EquipamentoUpdatePayload(
                serial="SN-DUP", modelo="Modelo incorreto", tipo="Notebook",
                revisao=main.get_equipamento("SN-DUP")["revisao"],
            ))

        self.assertEqual(raised.exception.status_code, 409)
        self.assertIn("mais de uma vez", raised.exception.detail)
        self.assertEqual(self.details.read_bytes(), before)
        self.assertEqual(main.equipamentos_cache, cache_before)

        self.assertTrue(main.load_database())
        cache_before = copy.deepcopy(main.equipamentos_cache)
        with self.assertRaises(main.HTTPException) as raised:
            main.salvar_equipamento(main.EquipamentoUpdatePayload(
                serial="SN-DUP", modelo="Modelo incorreto", tipo="Notebook",
                revisao=main.get_equipamento("SN-DUP")["revisao"],
            ))
        self.assertEqual(raised.exception.status_code, 409)
        self.assertIn("duplicada", raised.exception.detail)
        self.assertEqual(self.details.read_bytes(), before)
        self.assertEqual(main.equipamentos_cache, cache_before)

    def test_equipment_edit_refuses_a_new_active_source_after_cache_load(self):
        make_details(self.details, [
            ("SN-CACHE", "ALIAS-CACHE", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("OUTRA-SERIE", "HOST-OUTRO")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        cache_before = copy.deepcopy(main.equipamentos_cache)
        source_before = self.details.read_bytes()

        newer_source = self.input_dir / "EquipamentosDetalhes_mais_novo.xlsx"
        make_details(newer_source, [
            ("SN-NOVO", "ALIAS-NOVO", "Modelo B", "Londrina", "Notebook", "86000000"),
        ])
        timestamp = time.time() + 60
        os.utime(newer_source, (timestamp, timestamp))
        newer_before = newer_source.read_bytes()

        with self.assertRaises(main.HTTPException) as raised:
            main.salvar_equipamento(main.EquipamentoUpdatePayload(
                serial="SN-CACHE", modelo="Modelo alterado", tipo="Notebook",
                revisao=main.get_equipamento("SN-CACHE")["revisao"],
            ))

        self.assertEqual(raised.exception.status_code, 409)
        self.assertIn("base ativa mudou", raised.exception.detail)
        self.assertEqual(self.details.read_bytes(), source_before)
        self.assertEqual(newer_source.read_bytes(), newer_before)
        self.assertEqual(main.equipamentos_cache, cache_before)

    def test_equipment_edit_failure_preserves_workbook_and_cache(self):
        make_details(self.details, [
            ("SN-FALHA", "ALIAS-ORIGINAL", "Modelo original", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("OUTRA-SERIE", "HOST-OUTRO")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        before = self.details.read_bytes()
        cache_before = copy.deepcopy(main.equipamentos_cache)

        with patch.object(main, "stage_workbook_update", side_effect=main.FileTransactionError("falha fictícia")):
            with self.assertRaises(main.HTTPException) as raised:
                main.salvar_equipamento(main.EquipamentoUpdatePayload(
                    serial="SN-FALHA", modelo="Modelo alterado", tipo="Notebook",
                    revisao=main.get_equipamento("SN-FALHA")["revisao"],
                ))

        self.assertEqual(raised.exception.status_code, 500)
        self.assertEqual(self.details.read_bytes(), before)
        self.assertEqual(main.equipamentos_cache, cache_before)

    def test_equipment_edit_prepares_endpoint_aliases_before_committing_workbook(self):
        make_details(self.details, [
            ("SN-PREPARE", "ALIAS-ORIGINAL", "Modelo antes", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-PREPARE", "HOST-ENDPOINT")])
        self.assertTrue(main.load_database())
        selected = main.get_equipamento("SN-PREPARE")
        workbook_before = self.details.read_bytes()
        cache_before = copy.deepcopy(main.equipamentos_cache)

        with patch.object(main, "read_endpoint_alias_index", side_effect=RuntimeError("falha de leitura simulada")):
            with self.assertRaisesRegex(RuntimeError, "falha de leitura simulada"):
                main.salvar_equipamento(main.EquipamentoUpdatePayload(
                    serial="SN-PREPARE",
                    revisao=selected["revisao"],
                    modelo="Modelo depois",
                ))

        self.assertEqual(self.details.read_bytes(), workbook_before)
        self.assertEqual(main.equipamentos_cache, cache_before)
        self.assertEqual(main.get_equipamento("SN-PREPARE")["modelo"], "Modelo antes")

    def test_equipment_edit_prepares_search_index_before_committing_workbook(self):
        make_details(self.details, [
            ("SN-INDEX-PREPARE", "ALIAS-ORIGINAL", "Modelo antes", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-INDEX-PREPARE", "HOST-ENDPOINT")])
        self.assertTrue(main.load_database())
        selected = main.get_equipamento("SN-INDEX-PREPARE")
        workbook_before = self.details.read_bytes()
        cache_before = copy.deepcopy(main.equipamentos_cache)

        with patch.object(main, "build_search_index", side_effect=RuntimeError("falha de índice simulada")):
            with self.assertRaisesRegex(RuntimeError, "falha de índice simulada"):
                main.salvar_equipamento(main.EquipamentoUpdatePayload(
                    serial="SN-INDEX-PREPARE",
                    revisao=selected["revisao"],
                    modelo="Modelo depois",
                ))

        self.assertEqual(self.details.read_bytes(), workbook_before)
        self.assertEqual(main.equipamentos_cache, cache_before)
        self.assertEqual(main.get_equipamento("SN-INDEX-PREPARE")["modelo"], "Modelo antes")

    def test_equipment_edit_refuses_to_recreate_a_record_removed_after_load(self):
        make_details(self.details, [
            ("SN-REMOVIDO", "ALIAS-ORIGINAL", "Modelo original", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("OUTRA-SERIE", "HOST-OUTRO")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        cache_before = copy.deepcopy(main.equipamentos_cache)
        equipment_revision = main.get_equipamento("SN-REMOVIDO")["revisao"]

        make_details(self.details, [
            ("SN-ATUAL", "ALIAS-ATUAL", "Modelo atual", "Londrina", "Notebook", "86000000"),
        ])
        source_before = self.details.read_bytes()

        with self.assertRaises(main.HTTPException) as raised:
            main.salvar_equipamento(main.EquipamentoUpdatePayload(
                serial="SN-REMOVIDO", modelo="Modelo alterado", tipo="Notebook",
                revisao=equipment_revision,
            ))

        self.assertEqual(raised.exception.status_code, 409)
        self.assertIn("não está mais", raised.exception.detail)
        self.assertEqual(self.details.read_bytes(), source_before)
        self.assertEqual(main.equipamentos_cache, cache_before)

    def test_new_equipment_is_appended_only_when_marked_as_new(self):
        make_details(self.details, [
            ("SN-BASE", "ALIAS-BASE", "Modelo base", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("OUTRA-SERIE", "HOST-OUTRO")])
        make_model(self.model)
        self.assertTrue(main.load_database())

        result = main.salvar_equipamento(main.EquipamentoUpdatePayload(
            serial="SN-NOVO", novo=True, modelo="Modelo novo", tipo="Notebook",
            apelido="ALIAS-NOVO", idCliente="CLIENTE-NOVO",
        ))

        self.assertTrue(result["sucesso"])
        self.assertIn("SN-NOVO", main.equipamentos_cache)
        self.assertTrue(main.load_database())
        self.assertIn("SN-BASE", main.equipamentos_cache)
        self.assertIn("SN-NOVO", main.equipamentos_cache)

    def test_invall_is_joined_by_normalized_serial_for_site_and_search_without_rewriting_details(self):
        make_details(self.details, [
            ("S-1", "ALIAS-LOCAL", "Modelo A", "Maringá", "Notebook", "87000000"),
            ("DUP", "ALIAS-DUP", "Modelo B", "Londrina", "Notebook", "86000000"),
            ("NO-SOURCE", "ALIAS-KEEP", "Modelo C", "Ponta Grossa", "Notebook", "84000000"),
            ("DETAIL-DUP", "ALIAS-DETAIL-DUP-A", "Modelo D", "Pato Branco", "Notebook", "85500000"),
            ("detail dup", "ALIAS-DETAIL-DUP-B", "Modelo E", "Pato Branco", "Notebook", "85500000"),
        ])
        make_endpoint(self.endpoint, [
            (" s 1 ", "HOST-UNICO"),
            ("DUP", "HOST-DUP-A"),
            ("dup", "HOST-DUP-B"),
            ("DETAIL DUP", "HOST-DETAIL-DUP"),
            ("", "HOST-SEM-SERIE"),
        ])
        before = self.details.read_bytes()

        self.assertTrue(main.load_database())
        self.assertEqual(main.equipamentos_cache["S-1"]["apelido"], "HOST-UNICO")
        self.assertEqual(main.equipamentos_cache["S-1"]["invall_alias_status"], "aplicado")
        self.assertEqual(main.equipamentos_cache["DUP"]["apelido"], "ALIAS-DUP")
        self.assertEqual(main.equipamentos_cache["DUP"]["invall_alias_status"], "serie_invall_duplicada")
        self.assertEqual(main.equipamentos_cache["NO-SOURCE"]["apelido"], "ALIAS-KEEP")
        self.assertEqual(main.equipamentos_cache["NO-SOURCE"]["invall_alias_status"], "serie_nao_encontrada")
        self.assertEqual(main.equipamentos_cache["DETAIL-DUP"]["invall_alias_status"], "serie_detalhes_duplicada")
        self.assertEqual(main.equipamentos_cache["DETAIL DUP"]["invall_alias_status"], "serie_detalhes_duplicada")
        self.assertEqual(main.equipamentos_cache["DETAIL-DUP"]["apelido"], "ALIAS-DETAIL-DUP-A")
        self.assertEqual(main.get_equipamento("s-1")["apelido"], "HOST-UNICO")
        quick_results = main.buscar_seriais("HOST-UNICO")
        self.assertEqual(quick_results[0]["apelido"], "HOST-UNICO")
        review = main.endpoint_alias_review_items(
            main.read_endpoint_alias_index(), {"S1", "DUP", "NOSOURCE", "DETAILDUP"}, {"DETAILDUP"}
        )
        self.assertTrue(any(item["motivo"].startswith("Série duplicada") for item in review))
        self.assertTrue(any(item["serie"] == "" for item in review))
        self.assertTrue(any(item["serie"] == "NOSOURCE" and "não encontrada no endpoint" in item["motivo"] for item in review))
        self.assertEqual(self.details.read_bytes(), before)

    def test_export_changes_only_unique_changed_apelido_cells_and_keeps_other_content(self):
        make_details(self.details, [
            ("S-1", "LOCAL-ATUAL", "Modelo atual", "Maringá", "Notebook", "87000000"),
            ("S-2", "LOCAL-DOIS", "Modelo dois", "Curitiba", "Notebook", "80000000"),
        ])
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Parque"
        worksheet.append(["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP"])
        worksheet.append(["S-1", "ANTIGO", "Modelo antigo", "Londrina", "Notebook", "87000000"])
        worksheet.append(["S-2", "LOCAL-DOIS", "Modelo dois", "Curitiba", "Notebook", "80000000"])
        worksheet["B2"].fill = PatternFill("solid", fgColor="17365D")
        worksheet.freeze_panes = "A2"
        workbook.save(self.model)
        workbook.close()
        make_endpoint(self.endpoint, [("s-1", "HOST-NEW")])
        source_before = self.model.read_bytes()
        original_workbook = openpyxl.load_workbook(self.model, data_only=False)
        original_snapshot = main.workbook_snapshot(original_workbook)
        original_workbook.close()
        self.comparison_history.write_text(
            json.dumps([
                {
                    "id": "older-update",
                    "horario": "2026-10-01T14:00:00-03:00",
                    "comparacao": {"apelidos_atualizados": [{}, {}]},
                },
                {
                    "id": "restoration-with-type",
                    "tipo": "restauracao",
                    "horario": "2026-10-01T18:00:00-03:00",
                    "comparacao": {"tipo_operacao": "restauracao"},
                },
                {
                    "id": "restoration-legacy",
                    "horario": "2026-10-01T17:30:00-03:00",
                    "comparacao": {"tipo_operacao": "restauracao"},
                },
                {
                    "id": "latest-update",
                    "horario": "2026-10-01T16:30:00-03:00",
                    "comparacao": {"apelidos_atualizados": [{}, {}, {}]},
                },
            ]),
            encoding="utf-8",
        )

        response = main.exportar_equipamentos()
        output_bytes = asyncio.run(_collect_stream(response))
        output_path = self.root / "export-ficticio.xlsx"
        output_path.write_bytes(output_bytes)
        self.assertEqual(response.headers["x-total-alterados"], "1")
        self.assertEqual(response.headers["x-total-idcliente-alterados"], "0")
        self.assertEqual(response.headers["x-total-alteracoes-exportadas"], "1")
        self.assertEqual(response.headers["x-historico-apelidos-atualizados"], "3")
        self.assertEqual(response.headers["x-historico-horario"], "2026-10-01T16:30:00-03:00")
        self.assertIn(
            "x-historico-apelidos-atualizados",
            response.headers["access-control-expose-headers"].lower(),
        )
        self.assertEqual(self.model.read_bytes(), source_before)
        with zipfile.ZipFile(self.model) as original_package, zipfile.ZipFile(output_path) as exported_package:
            self.assertEqual(original_package.namelist(), exported_package.namelist())
            changed_parts = {
                name for name in original_package.namelist()
                if original_package.read(name) != exported_package.read(name)
            }
        self.assertEqual(changed_parts, {"xl/worksheets/sheet1.xml"})

        updated_workbook = openpyxl.load_workbook(output_path, data_only=False)
        updated_snapshot = main.workbook_snapshot(updated_workbook)
        expected_snapshot = copy.deepcopy(original_snapshot)
        old_alias_cell = expected_snapshot["worksheets"]["Parque"]["cells"]["B2"]
        expected_snapshot["worksheets"]["Parque"]["cells"]["B2"] = ("HOST-NEW", *old_alias_cell[1:])
        self.assertEqual(updated_snapshot, expected_snapshot)
        worksheet = updated_workbook["Parque"]
        self.assertEqual(worksheet["B2"].value, "HOST-NEW")
        self.assertEqual(worksheet["B2"].fill.fgColor.rgb, "0017365D")
        self.assertEqual(worksheet["C2"].value, "Modelo antigo")
        self.assertEqual(worksheet["D2"].value, "Londrina")
        self.assertEqual(worksheet["B3"].value, "LOCAL-DOIS")
        updated_workbook.close()

        make_details(self.details, [
            ("S-1", "LOCAL-ATUAL", "Modelo atual", "Maringá", "Notebook", "87000000"),
            ("S-3", "NOVO", "Modelo novo", "Curitiba", "Notebook", "80000000"),
        ])
        with self.assertRaises(main.HTTPException) as raised:
            main.exportar_equipamentos()
        self.assertEqual(raised.exception.status_code, 409)
        self.assertEqual(self.model.read_bytes(), source_before)

    def test_export_normalizes_blue_alias_and_idcliente_headers_without_changing_model(self):
        make_details(self.details, [
            ("S-1", "HOST-LOCAL", "Modelo A", "Maringá", "Notebook", "87000000", "ID-ATUAL"),
        ], headers=["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP", "IdCliente"])
        make_endpoint(self.endpoint, [("S-1", "HOST-ENDPOINT")])
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Parque"
        worksheet.append(["NumeroSerie", "Apelido", "IdCliente", "Modelo"])
        worksheet.append(["S-1", "HOST-ANTIGO", "ID-ANTIGO", "Modelo A"])
        gray = PatternFill("solid", fgColor="808080")
        blue = PatternFill("solid", fgColor="8EA9DB")
        worksheet["A1"].fill = gray
        worksheet["D1"].fill = gray
        worksheet["B1"].fill = blue
        worksheet["C1"].fill = blue
        workbook.save(self.model)
        workbook.close()
        original_bytes = self.model.read_bytes()

        response = main.exportar_equipamentos()
        output = openpyxl.load_workbook(io.BytesIO(asyncio.run(_collect_stream(response))))
        exported = output["Parque"]
        self.assertEqual(exported["B1"].style_id, exported["A1"].style_id)
        self.assertEqual(exported["C1"].style_id, exported["A1"].style_id)
        self.assertEqual(exported["D1"].style_id, exported["A1"].style_id)
        self.assertEqual(exported["B2"].value, "HOST-ENDPOINT")
        self.assertEqual(exported["C2"].value, "ID-ATUAL")
        output.close()
        self.assertEqual(self.model.read_bytes(), original_bytes)

    def test_export_availability_and_output_track_alias_and_idcliente_changes(self):
        cases = [
            ("nenhuma", "HOST-SAME", "HOST-SAME", "ID-SAME", "ID-SAME", 0, 0),
            ("somente-apelido", "HOST-NEW", "HOST-OLD", "ID-SAME", "ID-SAME", 1, 0),
            ("somente-idcliente", "HOST-SAME", "HOST-SAME", "ID-NEW", "ID-OLD", 0, 1),
            ("ambas", "HOST-NEW", "HOST-OLD", "ID-NEW", "ID-OLD", 1, 1),
        ]
        detail_headers = [
            "NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP", "IdCliente", "Observacao"
        ]
        model_headers = ["NumeroSerie", "Apelido", "IdCliente", "Observacao"]

        for name, current_alias, model_alias, current_id, model_id, expected_aliases, expected_ids in cases:
            with self.subTest(case=name):
                make_details(self.details, [
                    ("S-1", current_alias, "Modelo A", "Maringá", "Notebook", "87000000", current_id, "KEEP-UNRELATED")
                ], headers=detail_headers)
                make_endpoint(self.endpoint, [("S-1", current_alias)])
                workbook = openpyxl.Workbook()
                worksheet = workbook.active
                worksheet.title = "Parque"
                worksheet.append(model_headers)
                worksheet.append(["S-1", model_alias, model_id, "KEEP-UNRELATED"])
                workbook.save(self.model)
                workbook.close()
                source_before = self.model.read_bytes()

                expected_counts = {
                    "apelidos": expected_aliases,
                    "id_clientes": expected_ids,
                    "total": expected_aliases + expected_ids,
                }
                self.assertEqual(main.count_export_changes(), expected_counts)
                self.assertEqual(main.get_stats()["total_alteracoes_exportaveis"], expected_counts["total"])

                response = main.exportar_equipamentos()
                output_bytes = asyncio.run(_collect_stream(response))
                self.assertEqual(int(response.headers["x-total-alterados"]), expected_aliases)
                self.assertEqual(int(response.headers["x-total-idcliente-alterados"]), expected_ids)
                self.assertEqual(
                    int(response.headers["x-total-alteracoes-exportadas"]),
                    expected_aliases + expected_ids,
                )
                self.assertIn(
                    "x-total-idcliente-alterados",
                    response.headers["access-control-expose-headers"].lower(),
                )

                output_workbook = openpyxl.load_workbook(io.BytesIO(output_bytes), data_only=False)
                output_sheet = output_workbook["Parque"]
                self.assertEqual(output_sheet["B2"].value, current_alias)
                self.assertEqual(output_sheet["C2"].value, current_id)
                self.assertEqual(output_sheet["D2"].value, "KEEP-UNRELATED")
                output_workbook.close()
                self.assertEqual(self.model.read_bytes(), source_before)

    def test_partial_export_reports_review_reasons_and_sample_serials(self):
        details_headers = ["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP", "IdCliente"]
        make_details(self.details, [
            ("S1", "LOCAL-1", "Modelo 1", "Maringá", "Notebook", "87000000", "ID-1"),
            ("S2", "LOCAL-2", "Modelo 2", "Curitiba", "Notebook", "80000000", "ID-2"),
        ], headers=details_headers)
        make_endpoint(self.endpoint, [("S1", "HOST-1")])
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Parque"
        worksheet.append(["NumeroSerie", "Apelido", "IdCliente"])
        worksheet.append(["S1", "OLD-1", "OLD-ID-1"])
        worksheet.append(["S1", "OLD-1-DUP", "OLD-ID-DUP"])
        worksheet.append(["S2", "=CONCAT(\"HOST\",\"2\")", "OLD-ID-2"])
        workbook.save(self.model)
        workbook.close()

        response = main.exportar_equipamentos()
        asyncio.run(_collect_stream(response))

        self.assertEqual(response.headers["x-total-alterados"], "0")
        self.assertEqual(response.headers["x-total-idcliente-alterados"], "1")
        self.assertEqual(response.headers["x-total-conferencia"], "3")
        summary = json.loads(unquote(response.headers["x-conferencia-resumo"]))
        self.assertEqual(summary, [
            {
                "motivo": "Registros do Endpoint que exigem conferência",
                "total": 1,
                "series": ["S2"],
            },
            {
                "motivo": "Séries duplicadas no modelo de exportação",
                "total": 1,
                "series": ["S1"],
            },
            {
                "motivo": "Linhas protegidas por fórmulas",
                "total": 1,
                "series": ["S2"],
            },
        ])

    def test_serial_and_alias_header_variants_match_load_count_and_export(self):
        model_headers = ["NumeroSerie", "Apelido", "IdCliente", "Observacao"]

        for index, (serial_header, alias_header) in enumerate((
            ("NÚMERO-SERIE", "Nome do Computador"),
            ("Número de Série", "ComputerName"),
        )):
            with self.subTest(serial_header=serial_header, alias_header=alias_header):
                headers = [
                    serial_header, alias_header, "Modelo", "Cidade", "Tipo", "CEP",
                    "IdCliente", "Observacao",
                ]
                complete_headers = complete_details_headers(headers)
                make_details(
                    self.details,
                    [[
                        "SERIE-FICTICIA", "ALIAS-ATUAL", "Modelo A", "Maringá", "Notebook",
                        "87000000", "CLIENTE-ATUAL", "MANTER-DETALHES", "00.000.000/0000-00",
                    ]],
                    headers=complete_headers,
                )
                make_endpoint(self.endpoint, [("SERIE-FICTICIA", "HOST-EXPORTADO")])

                model_workbook = openpyxl.Workbook()
                model_sheet = model_workbook.active
                model_sheet.title = "Parque"
                model_sheet.append(model_headers)
                model_sheet.append(["SERIE-FICTICIA", "ALIAS-ANTERIOR", "CLIENTE-ANTERIOR", "MANTER-MODELO"])
                model_sheet["E2"] = "OUTRA-CELULA"
                model_workbook.create_sheet("Metadados")["A1"] = "ABA-PRESERVADA"
                model_workbook.save(self.model)
                model_workbook.close()
                model_before = self.model.read_bytes()

                uploaded = main.validate_uploaded_workbook(self.details.read_bytes(), "detalhes")
                self.assertEqual(uploaded["populated_rows"], 1)
                self.assertEqual(main.find_details_sheet(self.details), "EquipamentosDetalhes")
                self.assertTrue(main.load_database())
                self.assertEqual(main.get_equipamento("SERIE-FICTICIA")["serial"], "SERIE-FICTICIA")

                self.assertEqual(
                    main.count_export_changes(),
                    {"apelidos": 1, "id_clientes": 1, "total": 2},
                )

                response = main.exportar_equipamentos()
                output_bytes = asyncio.run(_collect_stream(response))
                self.assertEqual(response.headers["x-total-alterados"], "1")
                self.assertEqual(response.headers["x-total-idcliente-alterados"], "1")
                output_path = self.root / f"export-{index}.xlsx"
                output_path.write_bytes(output_bytes)
                output_workbook = openpyxl.load_workbook(output_path, data_only=False)
                output_sheet = output_workbook["Parque"]
                self.assertEqual(output_sheet["B2"].value, "HOST-EXPORTADO")
                self.assertEqual(output_sheet["C2"].value, "CLIENTE-ATUAL")
                self.assertEqual(output_sheet["D2"].value, "MANTER-MODELO")
                self.assertEqual(output_sheet["E2"].value, "OUTRA-CELULA")
                self.assertEqual(output_workbook["Metadados"]["A1"].value, "ABA-PRESERVADA")
                output_workbook.close()
                self.assertEqual(self.model.read_bytes(), model_before)

    def test_official_transfer_preserves_threaded_comments_and_changes_only_form_cells(self):
        template = make_transfer_template(self.input_dir / "modelo-transferencia-ficticio.xlsx")
        with self.assertRaises(main.WorkbookSafetyError):
            main.load_workbook_safely(template)
        payload = main.TransferenciaPayload(
            serial="SERIE-FICTICIA",
            descricao_equipamento="EQUIPAMENTO FICTICIO",
            contrato=16467,
            retirada_cnpj="00.000.000/0000-00",
            retirada_endereco="ORIGEM FICTICIA",
            retirada_contato="REMETENTE FICTICIO",
            retirada_email="remetente@example.invalid",
            retirada_telefone="41900000000",
            retirada_obs="OBSERVACAO FICTICIA",
            entrega_cnpj="00.000.000/0000-00",
            entrega_endereco="DESTINO FICTICIO",
            entrega_contato="DESTINATARIO FICTICIO",
            entrega_email="destinatario@example.invalid",
            entrega_telefone="41911112222",
            entrega_obs="OBSERVACAO FICTICIA",
            entrega_almoco="12:00:00",
            entrega_pausa_duracao="1h",
        )

        with patch.object(main, "TEMPLATE_FILE", template):
            response = main.gerar_planilha(payload)

        output_bytes = asyncio.run(_collect_stream(response))
        output_path = self.root / "transferencia-ficticia.xlsx"
        output_path.write_bytes(output_bytes)
        with zipfile.ZipFile(template) as original_package, zipfile.ZipFile(output_path) as generated_package:
            self.assertEqual(original_package.namelist(), generated_package.namelist())
            changed_parts = {
                name for name in original_package.namelist()
                if original_package.read(name) != generated_package.read(name)
            }
            for part in ("xl/threadedComments/threadedComment1.xml", "xl/persons/person.xml"):
                self.assertEqual(original_package.read(part), generated_package.read(part))
        self.assertEqual(changed_parts, {"xl/worksheets/sheet1.xml"})

        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Unknown extension is not supported and will be removed")
            generated_workbook = openpyxl.load_workbook(output_path, data_only=False)
        self.assertEqual(generated_workbook["Transferência"]["B8"].value, "SERIE-FICTICIA")
        self.assertEqual(generated_workbook["Transferência"]["B37"].value, "OBSERVACAO FICTICIA")
        self.assertEqual(generated_workbook["Transferência"]["J21"].value, "")
        self.assertEqual(generated_workbook["Transferência"]["J31"].value, "1h")
        generated_workbook.close()

    def test_delivery_pause_duration_is_required_except_when_office_stays_open(self):
        template = make_transfer_template(self.input_dir / "modelo-pausa-ficticio.xlsx")
        missing_pause = make_transfer_payload(entrega_pausa_duracao="")
        with patch.object(main, "TEMPLATE_FILE", template):
            with self.assertRaises(main.HTTPException) as raised:
                main.gerar_planilha(missing_pause)
        self.assertEqual(raised.exception.status_code, 422)
        self.assertIn("Duração da pausa da unidade de entrega", raised.exception.detail)

        office_stays_open = make_transfer_payload(
            entrega_almoco="Escritório não fecha",
            entrega_pausa_duracao="",
        )
        with patch.object(main, "TEMPLATE_FILE", template):
            response = main.gerar_planilha(office_stays_open)

        output_path = self.root / "pausa-opcional-ficticia.xlsx"
        output_path.write_bytes(asyncio.run(_collect_stream(response)))
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Unknown extension is not supported and will be removed")
            generated_workbook = openpyxl.load_workbook(output_path, data_only=False)
        self.assertEqual(generated_workbook["Transferência"]["F31"].value, "Escritório não fecha")
        self.assertEqual(generated_workbook["Transferência"]["J31"].value, "")
        generated_workbook.close()

    def test_transfer_rejects_empty_essential_fields_before_opening_the_template(self):
        template = make_transfer_template(self.input_dir / "modelo-validacao-ficticio.xlsx")
        invalid_fields = (
            ("serial", " ", "Número de série do equipamento"),
            ("serial", "---", "Número de série do equipamento precisa conter letras ou números"),
            ("retirada_cnpj", "  ", "CNPJ"),
            ("retirada_endereco", "", "Endereço"),
            ("entrega_cnpj", " ", "CNPJ"),
            ("entrega_endereco", "  ", "Endereço"),
        )

        with patch.object(main, "TEMPLATE_FILE", template):
            for field, empty_value, label in invalid_fields:
                with self.subTest(field=field):
                    with self.assertRaises(main.HTTPException) as raised:
                        main.gerar_planilha(make_transfer_payload(**{field: empty_value}))
                    self.assertEqual(raised.exception.status_code, 422)
                    self.assertIn(label, raised.exception.detail)

    def test_transfer_payload_restricts_contract_and_validates_contacts(self):
        with self.assertRaises(ValueError):
            make_transfer_payload(contrato=777)
        for invalid in (
            {"entrega_email": "sem-email"},
            {"retirada_telefone": "123"},
            {"entrega_cnpj": "123"},
            {"descricao_equipamento": "x" * 32768},
        ):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                make_transfer_payload(**invalid)

    def test_transfer_phone_formatting_handles_fixed_mobile_and_country_prefixes(self):
        for original, expected in PHONE_FORMAT_CASES:
            with self.subTest(original=original):
                self.assertEqual(main.formatar_telefone_brasileiro(original), expected)

    def test_saving_a_new_equipment_row_grows_dimensions_and_ignores_legacy_location_fields(self):
        path = self.input_dir / "equipamentos-novos-ficticio.xlsx"
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "EquipamentosDetalhes"
        worksheet.append(["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP"])
        workbook.save(path)
        workbook.close()
        add_structured_table(path, reference="A1:F1")
        workbook = openpyxl.load_workbook(path)
        workbook["EquipamentosDetalhes"].auto_filter.ref = "A1:F1"
        workbook.save(path)
        workbook.close()
        payload = main.EquipamentoUpdatePayload(
            serial="SERIE-NOVA-FICTICIA",
            apelido="APELIDO FICTICIO",
            modelo="MODELO FICTICIO",
            cidade="CURITIBA",
            tipo="Notebook",
            cep="80220000",
        )

        self.assertTrue(main.update_equipment_in_workbook_file(path, payload))

        updated_workbook = openpyxl.load_workbook(path, data_only=False)
        updated_sheet = updated_workbook["EquipamentosDetalhes"]
        self.assertEqual(updated_sheet.max_row, 2)
        self.assertEqual(updated_sheet.calculate_dimension(), "A1:F2")
        self.assertEqual(updated_sheet["A2"].value, "SERIE-NOVA-FICTICIA")
        self.assertEqual(updated_sheet["B2"].value, "APELIDO FICTICIO")
        self.assertEqual(updated_sheet.tables["EquipmentInventory"].ref, "A1:F2")
        self.assertEqual(updated_sheet.auto_filter.ref, "A1:F2")
        self.assertIsNone(updated_sheet["D2"].value)
        self.assertIsNone(updated_sheet["F2"].value)
        updated_workbook.close()

    def test_upload_preview_reports_only_aggregate_deltas_and_preserves_active_files(self):
        make_endpoint(self.endpoint, [
            ("SERIE-ANTIGA", "HOST-ANTIGO"),
            ("SERIE-MANTER", "HOST-MANTER"),
            ("SERIE-DUP", "HOST-DUP-A"),
            ("SERIE-DUP", "HOST-DUP-B"),
        ])
        self.details.write_bytes(_details_bytes([
            ("SERIE-ANTIGA", "ALIAS-ANTIGO", "Modelo A", "Maringá", "Notebook", "87000000"),
            ("SERIE-MANTER", "ALIAS-MANTER", "Modelo B", "Londrina", "Notebook", "86000000"),
            ("SERIE-REMOVER", "ALIAS-REMOVER", "Modelo C", "Curitiba", "Notebook", "80000000"),
        ]))
        self.comparison_history.write_text("[]", encoding="utf-8")
        endpoint_before = self.endpoint.read_bytes()
        details_before = self.details.read_bytes()
        history_before = self.comparison_history.read_bytes()
        incoming_endpoint = _endpoint_bytes([
            ("SERIE-NOVA", "HOST-NOVO"),
            ("SERIE-MANTER", "HOST-MANTER"),
            ("SERIE-DUP", "HOST-DUP-C"),
            ("SERIE-DUP", "HOST-DUP-D"),
        ])
        incoming_details = _details_bytes([
            ("SERIE-NOVA", "ALIAS-NOVO", "Modelo N", "Londrina", "Notebook", "86000000"),
            ("SERIE-MANTER", "ALIAS-MANTER", "Modelo B", "Londrina", "Notebook", "86000000"),
            ("SERIE-MANTER", "ALIAS-MANTER-DUP", "Modelo B", "Londrina", "Notebook", "86000000"),
        ])

        preview = asyncio.run(main.previsualizar_planilhas(
            UploadFile(filename="endpoint-sintetico.xlsx", file=io.BytesIO(incoming_endpoint)),
            UploadFile(filename="detalhes-sinteticos.xlsx", file=io.BytesIO(incoming_details)),
        ))

        self.assertEqual(preview["alteracoes"]["endpoint"], {
            "linhas_atuais": 4,
            "linhas_novas": 4,
            "series_novas": 1,
            "series_removidas": 1,
            "series_duplicadas_no_arquivo_novo": 1,
        })
        self.assertEqual(preview["alteracoes"]["detalhes"], {
            "linhas_atuais": 3,
            "linhas_novas": 3,
            "series_novas": 1,
            "series_removidas": 2,
            "series_duplicadas_no_arquivo_novo": 1,
        })
        self.assertNotIn("SERIE-", json.dumps(preview, ensure_ascii=False))
        self.assertEqual(self.endpoint.read_bytes(), endpoint_before)
        self.assertEqual(self.details.read_bytes(), details_before)
        self.assertEqual(self.comparison_history.read_bytes(), history_before)
        self.assertFalse(self.backups.exists())

        staged = {
            "endpoint": {"bytes": incoming_endpoint},
            "detalhes": {"bytes": incoming_details},
        }
        current = {"endpoint": endpoint_before, "detalhes": details_before}
        self.assertTrue(main.verify_inventory_preview_token(preview["preview_token"], staged, current))
        self.assertFalse(main.verify_inventory_preview_token(
            preview["preview_token"],
            {**staged, "endpoint": {"bytes": incoming_endpoint + b"mudou"}},
            current,
        ))
        self.assertFalse(main.verify_inventory_preview_token(
            preview["preview_token"], staged, {**current, "endpoint": current["endpoint"] + b"mudou"}
        ))

    def test_upload_confirmation_rejects_a_stale_preview_without_replacing_files(self):
        make_endpoint(self.endpoint, [("SERIE-ATUAL", "HOST-ATUAL")])
        self.details.write_bytes(_details_bytes([
            ("SERIE-ATUAL", "ALIAS-ATUAL", "Modelo A", "Maringá", "Notebook", "87000000"),
        ]))
        incoming_endpoint = _endpoint_bytes([("SERIE-NOVA", "HOST-NOVO")])
        incoming_details = _details_bytes([
            ("SERIE-NOVA", "ALIAS-NOVO", "Modelo N", "Londrina", "Notebook", "86000000"),
        ])
        preview = asyncio.run(main.previsualizar_planilhas(
            UploadFile(filename="endpoint-sintetico.xlsx", file=io.BytesIO(incoming_endpoint)),
            UploadFile(filename="detalhes-sinteticos.xlsx", file=io.BytesIO(incoming_details)),
        ))
        concurrent_endpoint = _endpoint_bytes([("SERIE-CONCORRENTE", "HOST-CONCORRENTE")])
        self.endpoint.write_bytes(concurrent_endpoint)
        details_before_confirm = self.details.read_bytes()

        with self.assertRaises(main.HTTPException) as raised:
            asyncio.run(main.atualizar_planilhas(
                UploadFile(filename="endpoint-sintetico.xlsx", file=io.BytesIO(incoming_endpoint)),
                UploadFile(filename="detalhes-sinteticos.xlsx", file=io.BytesIO(incoming_details)),
                preview_token=preview["preview_token"],
            ))

        self.assertEqual(raised.exception.status_code, 409)
        self.assertEqual(self.endpoint.read_bytes(), concurrent_endpoint)
        self.assertEqual(self.details.read_bytes(), details_before_confirm)
        self.assertFalse(self.backups.exists())

    def test_preview_confirmation_retry_recovers_receipt_after_preview_expires(self):
        endpoint = _endpoint_bytes([("SN-RETRY", "HOST-NEW")])
        details = _details_bytes([
            ("SN-RETRY", "HOST-NEW", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        preview = asyncio.run(main.previsualizar_planilhas(
            UploadFile(filename="endpoint.xlsx", file=io.BytesIO(endpoint)),
            UploadFile(filename="simpress.xlsx", file=io.BytesIO(details)),
        ))

        def submit(services):
            return asyncio.run(inventory_update.update_inventory_files(
                UploadFile(filename="endpoint.xlsx", file=io.BytesIO(endpoint)),
                UploadFile(filename="simpress.xlsx", file=io.BytesIO(details)),
                services, idempotency_key="preview-response-lost",
                preview_token=preview["preview_token"],
            ))

        first = submit(main._inventory_update_services)
        expired_services = replace(main._inventory_update_services, verify_inventory_preview_token=lambda *args: False)
        self.assertEqual(submit(expired_services), first)
        self.assertEqual(len(json.loads(self.comparison_history.read_text(encoding="utf-8"))), 1)
        self.assertEqual(len(list(self.backups.iterdir())), 1)

    def test_upload_appends_invall_alias_change_to_comparison_history_only(self):
        make_details(self.details, [("SN-1", "ALIAS-DETALHES", "Modelo A", "Maringá", "Notebook", "87000000")])
        make_endpoint(self.endpoint, [("SN-1", "PRNOTE-OLD")])
        make_model(self.model)
        self.comparison_history.write_text(json.dumps([{"id": "anterior", "comparacao": {}}]), encoding="utf-8")
        self.manual_history.write_text(json.dumps([{"id": "manual-antigo", "tipo": "apelido", "serial": "OUTRO"}]), encoding="utf-8")
        self.assertTrue(main.load_database())
        details_before_upload_alias = "ALIAS-DETALHES"
        comparison_before = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        manual_before = json.loads(self.manual_history.read_text(encoding="utf-8"))

        endpoint_bytes = _endpoint_bytes([("SN-1", "PRNOTE-NEW")])
        details_bytes = _details_bytes([
            ("SN-1", details_before_upload_alias, "Modelo A", "Maringá", "Notebook", "87000000"),
            ("", "ALIAS-SEM-SERIE", "Modelo B", "Londrina", "Notebook", "86000000"),
        ])
        result = asyncio.run(main.atualizar_planilhas(
            UploadFile(filename="InvAll_nova.xlsx", file=io.BytesIO(endpoint_bytes)),
            UploadFile(filename="Detalhes_novos.xlsx", file=io.BytesIO(details_bytes)),
        ))

        self.assertTrue(result["sucesso"])
        self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "PRNOTE-NEW")
        saved_details = openpyxl.load_workbook(self.details, data_only=True)
        self.assertEqual(saved_details["EquipamentosDetalhes"]["B2"].value, details_before_upload_alias)
        saved_details.close()
        comparison_after = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        manual_after = json.loads(self.manual_history.read_text(encoding="utf-8"))
        self.assertEqual(comparison_after[0], comparison_before[0])
        self.assertEqual(manual_after, manual_before)
        recorded = comparison_after[-1]["comparacao"]["apelidos_atualizados"]
        self.assertEqual(len(recorded), 1)
        self.assertEqual(recorded[0]["serie"], "SN-1")
        self.assertEqual(recorded[0]["apelido_anterior"], "PRNOTE-OLD")
        self.assertEqual(recorded[0]["apelido_novo"], "PRNOTE-NEW")
        self.assertEqual(recorded[0]["arquivos_origem"], {
            "endpoint": "InvAll_nova.xlsx",
            "detalhes": "Detalhes_novos.xlsx",
        })
        self.assertTrue(recorded[0]["horario"])
        self.assertTrue(any(
            item["serie"] == "" and "na Simpress" in item["motivo"]
            for item in comparison_after[-1]["comparacao"]["apelidos_sem_atualizacao"]
        ))

    def test_inventory_idempotency_replay_survives_service_reinitialization(self):
        self._prepare_inventory_upload_fixture()
        endpoint_bytes = _endpoint_bytes([("SN-IDEMPOTENT", "HOST-NEW")])
        details_bytes = _details_bytes([
            ("SN-IDEMPOTENT", "HOST-NEW", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])

        first = self._submit_inventory_update(endpoint_bytes, details_bytes, "upload-retry-01")
        services_after_restart = replace(main._inventory_update_services)
        repeated = self._submit_inventory_update(
            endpoint_bytes,
            details_bytes,
            "upload-retry-01",
            endpoint_name="renamed-endpoint.xlsx",
            details_name="renamed-details.xlsx",
            services=services_after_restart,
        )

        history = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        receipts = json.loads(
            (self.state_dir / "inventory_update_idempotency.json").read_text(encoding="utf-8")
        )
        backup_directories = [path for path in self.backups.iterdir() if path.is_dir()]
        self.assertEqual(repeated, first)
        self.assertEqual(len(history), 1)
        self.assertEqual(len(backup_directories), 1)
        self.assertEqual(len(receipts["records"]), 1)
        self.assertEqual(receipts["records"][0]["key"], "upload-retry-01")
        self.assertEqual(receipts["records"][0]["result"], first)

    def test_inventory_idempotency_hash_uses_role_and_committed_bytes_not_filename(self):
        staged = {
            "endpoint": {"name": "first.xlsx", "bytes": b"endpoint workbook"},
            "detalhes": {"name": "details.zip", "bytes": b"details workbook"},
        }
        renamed = {
            "endpoint": {"name": "renamed.xlsx", "bytes": b"endpoint workbook"},
            "detalhes": {"name": "renamed.zip", "bytes": b"details workbook"},
        }
        multipart_reordered = {
            "detalhes": {"name": "renamed.zip", "bytes": b"details workbook"},
            "endpoint": {"name": "renamed.xlsx", "bytes": b"endpoint workbook"},
        }
        role_swapped = {
            "endpoint": {"name": "renamed.xlsx", "bytes": b"details workbook"},
            "detalhes": {"name": "renamed.zip", "bytes": b"endpoint workbook"},
        }

        self.assertEqual(
            inventory_update._inventory_request_sha256(staged),
            inventory_update._inventory_request_sha256(renamed),
        )
        self.assertEqual(
            inventory_update._inventory_request_sha256(staged),
            inventory_update._inventory_request_sha256(multipart_reordered),
        )
        self.assertNotEqual(
            inventory_update._inventory_request_sha256(staged),
            inventory_update._inventory_request_sha256(role_swapped),
        )

    def test_inventory_idempotency_conflict_preserves_all_committed_state(self):
        self._prepare_inventory_upload_fixture()
        details_bytes = _details_bytes([
            ("SN-IDEMPOTENT", "HOST-NEW", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        first = self._submit_inventory_update(
            _endpoint_bytes([("SN-IDEMPOTENT", "HOST-NEW")]),
            details_bytes,
            "upload-conflict-01",
        )
        before = {
            "endpoint": self.endpoint.read_bytes(),
            "details": self.details.read_bytes(),
            "history": self.comparison_history.read_bytes(),
            "receipts": (self.state_dir / "inventory_update_idempotency.json").read_bytes(),
            "backups": sorted(path.name for path in self.backups.iterdir() if path.is_dir()),
        }

        with self.assertRaises(main.HTTPException) as raised:
            self._submit_inventory_update(
                _endpoint_bytes([("SN-IDEMPOTENT", "HOST-DIFFERENT")]),
                details_bytes,
                "upload-conflict-01",
            )

        self.assertEqual(raised.exception.status_code, 409)
        self.assertEqual(self.endpoint.read_bytes(), before["endpoint"])
        self.assertEqual(self.details.read_bytes(), before["details"])
        self.assertEqual(self.comparison_history.read_bytes(), before["history"])
        self.assertEqual(
            (self.state_dir / "inventory_update_idempotency.json").read_bytes(),
            before["receipts"],
        )
        self.assertEqual(
            sorted(path.name for path in self.backups.iterdir() if path.is_dir()),
            before["backups"],
        )
        self.assertEqual(first["historico_salvo"], True)

    def test_inventory_idempotency_replays_after_commit_when_response_is_lost(self):
        self._prepare_inventory_upload_fixture()
        endpoint_bytes = _endpoint_bytes([("SN-IDEMPOTENT", "HOST-NEW")])
        details_bytes = _details_bytes([
            ("SN-IDEMPOTENT", "HOST-NEW", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])

        def lose_confirmation(_prepared_state):
            raise RuntimeError("falha injetada depois da confirmação")

        services_with_lost_response = replace(
            main._inventory_update_services,
            publish_database_state=lose_confirmation,
        )
        with self.assertRaises(main.HTTPException) as raised:
            self._submit_inventory_update(
                endpoint_bytes,
                details_bytes,
                "upload-response-lost-01",
                services=services_with_lost_response,
            )

        self.assertEqual(raised.exception.status_code, 500)
        self.assertIn("mesma Idempotency-Key", raised.exception.detail)
        self.assertNotIn("restaurados", raised.exception.detail)
        history_after_commit = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        self.assertEqual(len(history_after_commit), 1)
        self.assertTrue((self.state_dir / "inventory_update_idempotency.json").is_file())

        repeated = self._submit_inventory_update(
            endpoint_bytes,
            details_bytes,
            "upload-response-lost-01",
            services=replace(main._inventory_update_services),
        )

        self.assertEqual(repeated["comparacao"]["evento_id"], history_after_commit[0]["id"])
        self.assertEqual(len(json.loads(self.comparison_history.read_text(encoding="utf-8"))), 1)
        self.assertEqual(len([path for path in self.backups.iterdir() if path.is_dir()]), 1)
        self.assertEqual(main.equipamentos_cache["SN-IDEMPOTENT"]["apelido"], "HOST-NEW")

    def test_inventory_idempotency_precommit_failure_can_be_retried(self):
        self._prepare_inventory_upload_fixture()
        endpoint_bytes = _endpoint_bytes([("SN-IDEMPOTENT", "HOST-NEW")])
        details_bytes = _details_bytes([
            ("SN-IDEMPOTENT", "HOST-NEW", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        original_endpoint = self.endpoint.read_bytes()
        original_details = self.details.read_bytes()
        original_commit = main.AtomicFileTransaction.commit

        def fail_before_commit(transaction):
            raise OSError("falha injetada antes da confirmação")

        with patch.object(main.AtomicFileTransaction, "commit", fail_before_commit):
            with self.assertRaises(main.HTTPException) as raised:
                self._submit_inventory_update(endpoint_bytes, details_bytes, "upload-retry-after-failure")

        self.assertEqual(raised.exception.status_code, 500)
        self.assertEqual(self.endpoint.read_bytes(), original_endpoint)
        self.assertEqual(self.details.read_bytes(), original_details)
        self.assertEqual(json.loads(self.comparison_history.read_text(encoding="utf-8")), [])
        self.assertFalse((self.state_dir / "inventory_update_idempotency.json").exists())
        failed_backup_directories = [path for path in self.backups.iterdir() if path.is_dir()]
        self.assertEqual(failed_backup_directories, [])

        with patch.object(main.AtomicFileTransaction, "commit", original_commit):
            result = self._submit_inventory_update(
                endpoint_bytes,
                details_bytes,
                "upload-retry-after-failure",
            )

        self.assertTrue(result["sucesso"])
        self.assertEqual(len(json.loads(self.comparison_history.read_text(encoding="utf-8"))), 1)
        receipts = json.loads(
            (self.state_dir / "inventory_update_idempotency.json").read_text(encoding="utf-8")
        )
        self.assertEqual([record["key"] for record in receipts["records"]], ["upload-retry-after-failure"])
        self.assertEqual(len([path for path in self.backups.iterdir() if path.is_dir()]), 1)

    def test_concurrent_inventory_requests_with_same_key_commit_once(self):
        self._prepare_inventory_upload_fixture()
        endpoint_bytes = _endpoint_bytes([("SN-IDEMPOTENT", "HOST-NEW")])
        details_bytes = _details_bytes([
            ("SN-IDEMPOTENT", "HOST-NEW", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        compare_barrier = threading.Barrier(2)
        original_compare = main.compare_uploaded_workbooks

        def synchronize_comparisons(*args, **kwargs):
            result = original_compare(*args, **kwargs)
            compare_barrier.wait(timeout=10)
            return result

        with (
            patch.object(main, "WORKBOOK_PARSE_ADMISSION", None),
            patch.object(main, "compare_uploaded_workbooks", synchronize_comparisons),
            ThreadPoolExecutor(max_workers=2) as pool,
        ):
            requests = [
                pool.submit(
                    self._submit_inventory_update,
                    endpoint_bytes,
                    details_bytes,
                    "upload-concurrent-01",
                    endpoint_name=f"endpoint-{index}.xlsx",
                    details_name=f"details-{index}.xlsx",
                )
                for index in range(2)
            ]
            results = [request.result(timeout=30) for request in requests]

        history = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        backup_directories = [path for path in self.backups.iterdir() if path.is_dir()]
        self.assertEqual(results[0], results[1])
        self.assertEqual(len(history), 1)
        self.assertEqual(len(backup_directories), 1)

    def test_inventory_idempotency_receipt_retention_is_bounded_and_fifo(self):
        store = {"version": 1, "records": []}
        for index in range(inventory_update.MAX_INVENTORY_IDEMPOTENCY_RECEIPTS + 1):
            store = inventory_update._append_inventory_idempotency_receipt(
                store,
                f"key-{index}",
                "a" * 64,
                {"sucesso": True},
            )

        keys = [record["key"] for record in store["records"]]
        self.assertEqual(len(keys), inventory_update.MAX_INVENTORY_IDEMPOTENCY_RECEIPTS)
        self.assertNotIn("key-0", keys)
        self.assertEqual(keys[0], "key-1")
        self.assertEqual(keys[-1], f"key-{inventory_update.MAX_INVENTORY_IDEMPOTENCY_RECEIPTS}")

    def test_inventory_update_migrates_a_recognized_alternate_details_filename(self):
        alternate = self.input_dir / "EquipamentosDetalhes-exportado.xlsx"
        make_details(alternate, [
            ("SN-ALT", "HOST-ANTIGO", "Modelo antigo", "Maringá", "Notebook", "87000000"),
        ])
        self.comparison_history.write_text("[]", encoding="utf-8")
        self.manual_history.write_text("[]", encoding="utf-8")
        comparison_before = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        manual_before = json.loads(self.manual_history.read_text(encoding="utf-8"))
        alternate_before = alternate.read_bytes()
        make_endpoint(self.endpoint, [("SN-ALT", "HOST-ANTIGO")])
        make_model(self.model)
        self.assertTrue(main.load_database())
        self.assertEqual(main.get_primary_equipment_file(), alternate)

        result = asyncio.run(main.atualizar_planilhas(
            UploadFile(filename="InvAll_nova.xlsx", file=io.BytesIO(_endpoint_bytes([("SN-ALT", "HOST-NOVO")]))),
            UploadFile(filename="Detalhes_novos.xlsx", file=io.BytesIO(_details_bytes([
                ("SN-ALT", "HOST-NOVO", "Modelo novo", "Londrina", "Notebook", "86000000"),
            ]))),
        ))

        self.assertTrue(result["sucesso"])
        self.assertEqual(alternate.read_bytes(), alternate_before)
        self.assertTrue(self.details.is_file())
        self.assertEqual(main.get_primary_equipment_file(), self.details)
        self.assertEqual(main.get_equipamento("SN-ALT")["modelo"], "Modelo novo")
        history = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        backup_dir = self.input_dir / history[-1]["comparacao"]["backup_diretorio"]
        restored_name = backup_dir / "detalhes" / self.details.name
        self.assertEqual(restored_name.read_bytes(), alternate_before)

        repeated = asyncio.run(main.atualizar_planilhas(
            UploadFile(filename="InvAll_repetida.xlsx", file=io.BytesIO(_endpoint_bytes([("SN-ALT", "HOST-NOVO")]))),
            UploadFile(filename="Detalhes_repetidos.xlsx", file=io.BytesIO(_details_bytes([
                ("SN-ALT", "HOST-NOVO", "Modelo novo", "Londrina", "Notebook", "86000000"),
            ]))),
        ))
        self.assertTrue(repeated["sucesso"])
        comparison_after_repeat = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        self.assertEqual(len(comparison_after_repeat[-1]["comparacao"]["apelidos_atualizados"]), 0)
        self.assertEqual(len(comparison_after_repeat), len(comparison_before) + 2)
        self.assertEqual(json.loads(self.manual_history.read_text(encoding="utf-8")), manual_before)

    def test_inventory_update_copies_annotation_only_to_an_exact_matching_row(self):
        make_details(self.details, [
            ("SN-ANNOTATION", "HOST-ANNOTATION", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-ANNOTATION", "HOST-ANNOTATION")])
        make_model(self.model)
        self.assertTrue(main.load_database())

        previous_row = {
            "serie": "SN-ANNOTATION",
            "computador": "HOST-ANNOTATION",
            "dias_sem_verificacao": 14,
            "ultima_verificacao": "2026-10-07T10:00:00-03:00",
        }
        previous_event = {
            "id": "evento-anotado-anterior",
            "horario": "2026-10-07T10:00:00-03:00",
            "comparacao": {"sem_acesso_7_dias": [previous_row]},
        }
        self.comparison_history.write_text(json.dumps([previous_event]), encoding="utf-8")
        previous_item_id = main.comparison_history_for_display(previous_event)["comparacao"][
            "sem_acesso_7_dias"
        ][0]["anotacao_id"]
        main.salvar_anotacao_comparacao_historica(
            previous_event["id"],
            previous_item_id,
            SimpleNamespace(anotacao="Acompanhar com a unidade."),
        )

        from domains.inventory_update import _commit_inventory_files

        def commit_test_upload(endpoint_name, details_name, automatic_import_id=None):
            endpoint_bytes = _endpoint_bytes([("SN-ANNOTATION", "HOST-ANNOTATION")])
            details_bytes = _details_bytes([
                ("SN-ANNOTATION", "HOST-ANNOTATION", "Modelo A", "Maringá", "Notebook", "87000000"),
            ])
            return _commit_inventory_files({
                "endpoint": {
                    "name": endpoint_name,
                    "bytes": endpoint_bytes,
                    "table": main.validate_uploaded_workbook(endpoint_bytes, "endpoint"),
                },
                "detalhes": {
                    "name": details_name,
                    "bytes": details_bytes,
                    "table": main.validate_uploaded_workbook(details_bytes, "detalhes"),
                },
            }, main._inventory_update_services, automatic_import_id)

        changed_row = {**previous_row, "dias_sem_verificacao": 15}
        with patch.object(
            main,
            "compare_uploaded_workbooks",
            return_value={"sem_acesso_7_dias": [copy.deepcopy(previous_row), changed_row]},
        ):
            result = commit_test_upload(
                "InvAll_nova.xlsx",
                "Detalhes_novos.xlsx",
                automatic_import_id="lote-automacao-teste",
            )

        self.assertTrue(result["sucesso"])
        history = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        new_event = history[-1]
        self.assertEqual(new_event["importacao_automatica_id"], "lote-automacao-teste")
        new_rows = new_event["comparacao"]["sem_acesso_7_dias"]
        self.assertEqual(new_rows[0]["anotacao"], "Acompanhar com a unidade.")
        self.assertNotEqual(new_rows[0]["anotacao_id"], previous_item_id)
        self.assertNotIn("anotacao", new_rows[1])
        shared_new_event_records = [
            record
            for record in main.load_history_annotations(self.shared_annotations)
            if record["evento_id"] == new_event["id"]
        ]
        self.assertEqual(shared_new_event_records, [])
        pending_new_event_records = [
            record
            for record in main.load_history_annotations(self.pending_annotations)
            if record["evento_id"] == new_event["id"]
        ]
        self.assertEqual(len(pending_new_event_records), 1)
        self.assertEqual(pending_new_event_records[0]["anotacao"], "Acompanhar com a unidade.")
        self.assertEqual(result["anotacoes_herdadas_enfileiradas"], 1)
        self.assertEqual(result["sincronizacao_anotacoes"], "pendente")

        with patch.object(main, "save_history_annotation", side_effect=OSError("pasta compartilhada indisponível")):
            failed_sync = main.flush_pending_history_annotations()
        self.assertEqual(failed_sync, {"pendentes": 1, "sincronizadas": 0})
        self.assertEqual(len(main.load_history_annotations(self.pending_annotations)), 1)
        synced = main.flush_pending_history_annotations()
        self.assertEqual(synced, {"pendentes": 0, "sincronizadas": 1})
        repeated_sync = main.flush_pending_history_annotations()
        self.assertEqual(repeated_sync, {"pendentes": 0, "sincronizadas": 0})
        shared_new_event_records = [
            record
            for record in main.load_history_annotations(self.shared_annotations)
            if record["evento_id"] == new_event["id"]
        ]
        self.assertEqual(len(shared_new_event_records), 1)
        self.assertEqual(shared_new_event_records[0]["anotacao"], "Acompanhar com a unidade.")

        # A clone without local comparison history can carry the note using
        # the exact row fingerprint saved with the shared annotation.
        self.comparison_history.write_text("[]", encoding="utf-8")
        with patch.object(
            main,
            "compare_uploaded_workbooks",
            return_value={"sem_acesso_7_dias": [copy.deepcopy(previous_row)]},
        ):
            clone_update = commit_test_upload("InvAll_clone.xlsx", "Detalhes_clone.xlsx")
        self.assertTrue(clone_update["sucesso"])
        clone_history = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        self.assertEqual(
            clone_history[-1]["comparacao"]["sem_acesso_7_dias"][0]["anotacao"],
            "Acompanhar com a unidade.",
        )

    def test_inherited_annotation_outbox_commits_and_recovers_with_inventory_history(self):
        from domains.inventory_update import _commit_inventory_files

        make_details(self.details, [
            ("SN-OUTBOX", "HOST-OUTBOX", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-OUTBOX", "HOST-OUTBOX")])
        make_model(self.model)
        self.assertTrue(main.load_database())

        category = "sem_acesso_7_dias"
        prior_row = {
            "serie": "SN-OUTBOX", "computador": "HOST-OUTBOX",
            "dias_sem_verificacao": 14, "ultima_verificacao": "2026-10-07T10:00:00-03:00",
        }
        prior_event = {"id": "evento-outbox-anterior", "horario": "2026-10-07T10:00:00-03:00",
                       "comparacao": {category: [prior_row]}}
        self.comparison_history.write_text(json.dumps([prior_event]), encoding="utf-8")
        prior_item_id = main.comparison_history_for_display(prior_event)["comparacao"][category][0]["anotacao_id"]
        main.salvar_anotacao_comparacao_historica(
            prior_event["id"], prior_item_id, SimpleNamespace(anotacao="Revisar com a unidade."),
        )

        endpoint_bytes = _endpoint_bytes([("SN-OUTBOX", "HOST-OUTBOX")])
        details_bytes = _details_bytes([
            ("SN-OUTBOX", "HOST-OUTBOX", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        staged = {
            "endpoint": {"name": "InvAll_outbox.xlsx", "bytes": endpoint_bytes,
                         "table": main.validate_uploaded_workbook(endpoint_bytes, "endpoint")},
            "detalhes": {"name": "Detalhes_outbox.xlsx", "bytes": details_bytes,
                         "table": main.validate_uploaded_workbook(details_bytes, "detalhes")},
        }
        comparison = {category: [copy.deepcopy(prior_row)]}

        def update():
            return _commit_inventory_files(staged, main._inventory_update_services)

        before_files = (self.endpoint.read_bytes(), self.details.read_bytes())
        history_before = self.comparison_history.read_bytes()
        original_apply = main.AtomicFileTransaction.apply

        def interrupt_before_confirmation(transaction):
            original_apply(transaction)
            raise RuntimeError("interrupção sintética antes da confirmação")

        with patch.object(main, "compare_uploaded_workbooks", return_value=copy.deepcopy(comparison)), \
             patch.object(main.AtomicFileTransaction, "commit", interrupt_before_confirmation), \
             self.assertRaisesRegex(RuntimeError, "antes da confirmação"):
            update()
        self.assertEqual((self.endpoint.read_bytes(), self.details.read_bytes()), before_files)
        self.assertEqual(self.comparison_history.read_bytes(), history_before)
        self.assertEqual(main.load_history_annotations(self.pending_annotations), [])

        with patch.object(main, "compare_uploaded_workbooks", return_value=copy.deepcopy(comparison)), \
             patch.object(main, "publish_database_state", side_effect=RuntimeError("interrupção após confirmação")), \
             self.assertRaisesRegex(RuntimeError, "após confirmação"):
            update()
        committed_history = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        new_event = committed_history[-1]
        self.assertEqual(new_event["comparacao"][category][0]["anotacao"], "Revisar com a unidade.")
        queued = main.load_history_annotations(self.pending_annotations)
        self.assertEqual(len(queued), 1)
        self.assertEqual(queued[0]["evento_id"], new_event["id"])

        # Simulate process startup recovery: the shared publisher is offline,
        # and the pending item is found and retried without opening a history route.
        main.recover_pending_transactions(self.root)
        with patch.object(main, "save_history_annotation", side_effect=OSError("rede indisponível")):
            offline = main.flush_pending_history_annotations()
        self.assertEqual(offline, {"pendentes": 1, "sincronizadas": 0})
        self.assertEqual(len(main.load_history_annotations(self.pending_annotations)), 1)
        online = main.flush_pending_history_annotations()
        self.assertEqual(online, {"pendentes": 0, "sincronizadas": 1})
        self.assertEqual(main.load_history_annotations(self.pending_annotations), [])
        resolved = resolve_history_annotations(main.load_history_annotations(self.shared_annotations))
        self.assertEqual(resolved[(new_event["id"], queued[0]["item_id"])]["anotacao"], "Revisar com a unidade.")

    def test_annotation_uses_stable_review_id_after_display_filter_and_stays_event_scoped(self):
        filtered_out = {
            "serie": "SERIE-LOCAL-1",
            "computador": "PRNOTE-LOCAL-1",
            "motivo": "Série não encontrada na InvAll; apelido atual preservado.",
        }
        visible_item = {
            "serie": "SERIE-DUP-1",
            "computador": "PRNOTE-DUP-1, PRNOTE-DUP-2",
            "motivo": "Série duplicada na InvAll; nenhum hostname foi escolhido.",
        }
        history = [
            {"id": "evento-anotacao-1", "comparacao": {"apelidos_sem_atualizacao": [filtered_out.copy(), visible_item.copy()]}},
            {"id": "evento-anotacao-2", "comparacao": {"apelidos_sem_atualizacao": [filtered_out.copy(), visible_item.copy()]}},
        ]
        self.comparison_history.write_text(json.dumps(history), encoding="utf-8")

        displayed = main.comparison_history_for_display(history[0])
        displayed_item = displayed["comparacao"]["apelidos_sem_atualizacao"][0]
        self.assertEqual(displayed_item["serie"], "SERIE-DUP-1")
        self.assertNotEqual(
            displayed_item["anotacao_id"],
            main.comparison_history_for_display(history[1])["comparacao"]["apelidos_sem_atualizacao"][0]["anotacao_id"],
        )

        result = main.salvar_anotacao_comparacao_historica(
            "evento-anotacao-1",
            displayed_item["anotacao_id"],
            SimpleNamespace(anotacao="Conferir com suporte."),
        )

        stored_history = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        first_items = stored_history[0]["comparacao"]["apelidos_sem_atualizacao"]
        second_items = stored_history[1]["comparacao"]["apelidos_sem_atualizacao"]
        self.assertTrue(result["sucesso"])
        self.assertNotIn("anotacao", first_items[0])
        self.assertEqual(first_items[1]["anotacao"], "Conferir com suporte.")
        self.assertEqual(first_items[1]["anotacao_id"], displayed_item["anotacao_id"])
        self.assertNotIn("anotacao", second_items[1])

    def test_history_annotations_sync_to_a_fresh_clone_and_can_be_edited_there(self):
        row = {
            "serie": "SERIE-SYNC-1",
            "computador": "PRNOTE-SYNC-1",
            "motivo": "Série duplicada na InvAll; nenhum hostname foi escolhido.",
        }
        event = {
            "id": "evento-sync-anotacoes",
            "horario": "2026-10-06T10:00:00-03:00",
            "comparacao": {"apelidos_sem_atualizacao": [row]},
        }
        self.comparison_history.write_text(json.dumps([event]), encoding="utf-8")
        item_id = main.comparison_history_for_display(event)["comparacao"]["apelidos_sem_atualizacao"][0]["anotacao_id"]

        main.salvar_anotacao_comparacao_historica(
            event["id"], item_id, SimpleNamespace(anotacao="Conferir no próximo inventário.")
        )
        shared_files = list(self.shared_annotations.glob("*.json"))
        self.assertEqual(len(shared_files), 1)
        shared_record = json.loads(shared_files[0].read_text(encoding="utf-8"))
        self.assertEqual(shared_record["anotacao"], "Conferir no próximo inventário.")
        self.assertEqual(set(shared_record), {
            "versao", "evento_id", "horario", "item_id", "serie", "computador",
            "motivo", "categoria", "anotacao", "atualizado_em", "linha_hash", "revisao_id", "base_revisoes",
        })

        fresh_clone_history = self.root / "clone_novo" / "historico_comparacoes.json"
        fresh_clone_history.parent.mkdir()
        fresh_clone_history.write_text("[]", encoding="utf-8")
        with patch.object(main, "COMPARISON_HISTORY_FILE", fresh_clone_history):
            listing = main.get_historico_comparacoes()
            self.assertEqual(listing["total"], 1)
            self.assertTrue(listing["eventos"][0]["somente_anotacoes_sincronizadas"])
            detail = main.get_comparacao_historica(event["id"])["evento"]
            displayed_row = detail["comparacao"]["apelidos_sem_atualizacao"][0]
            self.assertEqual(displayed_row["anotacao"], "Conferir no próximo inventário.")
            self.assertEqual(displayed_row["anotacao_id"], item_id)

            main.salvar_anotacao_comparacao_historica(
                event["id"], item_id, SimpleNamespace(anotacao="Atualizado no notebook novo.", base_revisoes=displayed_row["anotacao_revisoes"])
            )

        refreshed_record = resolve_history_annotations(main.load_history_annotations(self.shared_annotations))[(event["id"], item_id)]
        self.assertEqual(json.loads(shared_files[0].read_text(encoding="utf-8"))["anotacao"], "Conferir no próximo inventário.")
        self.assertEqual(refreshed_record["anotacao"], "Atualizado no notebook novo.")

        with patch.object(main, "COMPARISON_HISTORY_FILE", fresh_clone_history):
            main.salvar_anotacao_comparacao_historica(
                event["id"], item_id, SimpleNamespace(anotacao="", base_revisoes=refreshed_record["anotacao_revisoes"])
            )
            self.assertFalse(any(
                item.get("somente_anotacoes_sincronizadas")
                for item in main.get_historico_comparacoes()["eventos"]
            ))

    def test_annotations_save_for_all_requested_sections_and_appear_in_excel_export(self):
        event = {
            "id": "evento-todas-anotacoes",
            "horario": "2026-10-07T10:00:00-03:00",
            "resumo": {},
            "comparacao": {
                "sem_acesso_7_dias": [{"serie": "S-ACESSO", "computador": "PRNOTE-ACESSO", "dias_sem_verificacao": 14}],
                "ausentes_no_endpoint": [{"serie": "BRM-ENDPOINT", "computador": "PRNOTE-ENDPOINT", "local_previsto": "Maringá"}],
                "ausentes_em_detalhes": [{"serie": "S-SIMPRESS", "computador": "PRMAC-SIMPRESS", "local_previsto": "Londrina"}],
                "divergencias_local": [{"serie": "S-LOCAL", "computador": "PRNOTE-LOCAL", "local_previsto": "Maringá", "local_detalhes": "Londrina"}],
                "correspondencias_ambiguas": [{"tipo": "Série duplicada", "chave": "S-DUP", "fontes": [], "mensagem": "Revisar duplicidade"}],
                "registros_ambiguos": [{"serie": "S-AMB", "computador": "PRNOTE-AMB", "fontes": {"Endpoint": "ambíguo"}, "motivo": "Hostname duplicado"}],
                "series_sem_mapeamento": [{"serie": "ZZ-SEM-LOCAL", "quantidade": 1, "exemplos": ["PRNOTE-SEM-LOCAL"]}],
            },
        }
        self.comparison_history.write_text(json.dumps([event]), encoding="utf-8")
        sections_and_notes = {
            "sem_acesso_7_dias": "Verificar com suporte.",
            "ausentes_no_endpoint": "Confirmar cadastro no Endpoint.",
            "ausentes_em_detalhes": "Conferir com a Simpress.",
            "divergencias_local": "Validar endereço da unidade.",
            "correspondencias_ambiguas": "Conferir as séries duplicadas.",
            "registros_ambiguos": "Validar hostname com a equipe local.",
            "series_sem_mapeamento": "Cadastrar regra para o prefixo ZZ.",
        }

        displayed = main.get_comparacao_historica(event["id"])["evento"]["comparacao"]
        item_ids = {
            category: displayed[category][0]["anotacao_id"]
            for category in sections_and_notes
        }
        self.assertEqual(len(set(item_ids.values())), len(item_ids))
        for category, note in sections_and_notes.items():
            result = main.salvar_anotacao_comparacao_historica(
                event["id"], item_ids[category], SimpleNamespace(anotacao=note)
            )
            self.assertTrue(result["sucesso"])

        stored = json.loads(self.comparison_history.read_text(encoding="utf-8"))[0]["comparacao"]
        for category, note in sections_and_notes.items():
            self.assertEqual(stored[category][0]["anotacao"], note)
            self.assertEqual(stored[category][0]["anotacao_id"], item_ids[category])
        shared_records = [json.loads(path.read_text(encoding="utf-8")) for path in self.shared_annotations.glob("*.json")]
        self.assertEqual({record["categoria"] for record in shared_records}, set(sections_and_notes))

        class CapturedResponse:
            def __init__(self, content, media_type, headers):
                self.body = content.read()
                self.headers = headers

        with patch("routers.inventory_history.StreamingResponse", CapturedResponse):
            response = main.exportar_historico_comparacoes_excel(data="2026-10-07")
        workbook = openpyxl.load_workbook(io.BytesIO(response.body), data_only=True)
        try:
            for sheet, category in (
                ("Sem acesso", "sem_acesso_7_dias"),
                ("Ausências", "ausentes_no_endpoint"),
                ("Ausências", "ausentes_em_detalhes"),
                ("Localidades", "divergencias_local"),
                ("Localidades", "series_sem_mapeamento"),
                ("Ambiguidades", "correspondencias_ambiguas"),
                ("Ambiguidades", "registros_ambiguos"),
            ):
                cell_values = {
                    cell.value
                    for row in workbook[sheet].iter_rows()
                    for cell in row
                }
                self.assertIn(sections_and_notes[category], cell_values)
        finally:
            workbook.close()

    def test_non_review_annotation_is_available_on_a_clone_without_local_history(self):
        row = {
            "serie": "S-LOCAL-SYNC",
            "computador": "PRNOTE-LOCAL-SYNC",
            "local_previsto": "Maringá",
            "local_detalhes": "Londrina",
        }
        event = {
            "id": "evento-anotacao-local-sync",
            "horario": "2026-10-07T11:00:00-03:00",
            "comparacao": {"divergencias_local": [row]},
        }
        self.comparison_history.write_text(json.dumps([event]), encoding="utf-8")
        item_id = main.comparison_history_for_display(event)["comparacao"]["divergencias_local"][0]["anotacao_id"]
        main.salvar_anotacao_comparacao_historica(
            event["id"], item_id, SimpleNamespace(anotacao="Conferir este endereço.")
        )

        fresh_clone_history = self.root / "clone_sem_historico" / "historico_comparacoes.json"
        fresh_clone_history.parent.mkdir()
        fresh_clone_history.write_text("[]", encoding="utf-8")
        with patch.object(main, "COMPARISON_HISTORY_FILE", fresh_clone_history):
            detail = main.get_comparacao_historica(event["id"])["evento"]
            self.assertTrue(detail["somente_anotacoes_sincronizadas"])
            self.assertEqual(
                detail["comparacao"]["anotacoes_sincronizadas"][0]["categoria"],
                "divergencias_local",
            )
            synchronized_row = detail["comparacao"]["divergencias_local"][0]
            self.assertEqual(synchronized_row["anotacao"], "Conferir este endereço.")
            self.assertEqual(synchronized_row["anotacao_id"], item_id)
            main.salvar_anotacao_comparacao_historica(
                event["id"], item_id, SimpleNamespace(anotacao="Endereço validado no clone.", base_revisoes=synchronized_row["anotacao_revisoes"])
            )

        shared_record = resolve_history_annotations(main.load_history_annotations(self.shared_annotations))[(event["id"], item_id)]
        self.assertEqual(shared_record["categoria"], "divergencias_local")
        self.assertEqual(shared_record["anotacao"], "Endereço validado no clone.")

    def test_legacy_annotation_migrates_outside_the_history_read(self):
        event = {
            "id": "evento-legado-com-anotacao",
            "horario": "2026-10-05T10:00:00-03:00",
            "comparacao": {
                "apelidos_sem_atualizacao": [{
                    "serie": "BRJ-LEGADA-1",
                    "computador": "PRNOTE-LEGADO-1",
                    "motivo": "Série não encontrada na InvAll; apelido atual preservado.",
                    "anotacao": "Anotação salva pela versão anterior.",
                }],
            },
        }
        self.comparison_history.write_text(json.dumps([event]), encoding="utf-8")

        detail = main.get_comparacao_historica(event["id"])["evento"]
        displayed_row = detail["comparacao"]["apelidos_sem_atualizacao"][0]
        self.assertEqual(displayed_row["anotacao"], "Anotação salva pela versão anterior.")
        self.assertEqual(len(list(self.shared_annotations.glob("*.json"))), 0)
        self.assertEqual(len(list(self.pending_annotations.glob("*.json"))), 0)

        migration = main.migrate_legacy_history_annotations()
        self.assertEqual(migration["legadas_enfileiradas"], 1)
        self.assertEqual(migration["sincronizadas"], 1)
        self.assertEqual(len(list(self.shared_annotations.glob("*.json"))), 1)

    def test_pending_local_annotation_overrides_stale_shared_copy_until_retry(self):
        row = {
            "serie": "SERIE-DUP-PENDING",
            "computador": "PRNOTE-PENDING-A, PRNOTE-PENDING-B",
            "motivo": "Série duplicada na InvAll; nenhum hostname foi escolhido.",
        }
        event = {
            "id": "evento-pending-annotation",
            "horario": "2026-10-07T12:00:00-03:00",
            "comparacao": {"apelidos_sem_atualizacao": [row]},
        }
        self.comparison_history.write_text(json.dumps([event]), encoding="utf-8")
        item_id = main.comparison_history_for_display(event)["comparacao"]["apelidos_sem_atualizacao"][0]["anotacao_id"]
        main.save_history_annotation(
            self.shared_annotations,
            event_id=event["id"],
            event_time=event["horario"],
            item_id=item_id,
            serie=row["serie"],
            computador=row["computador"],
            motivo=row["motivo"],
            anotacao="Anotação compartilhada antiga.",
            categoria="apelidos_sem_atualizacao",
            linha_hash=main.history_annotation_row_fingerprint("apelidos_sem_atualizacao", row),
        )

        original_save = main.save_history_annotation

        def fail_shared_save(directory, **values):
            if Path(directory) == self.shared_annotations:
                raise OSError("OneDrive indisponível no teste")
            return original_save(directory, **values)

        with patch.object(main, "save_history_annotation", side_effect=fail_shared_save):
            saved = main.salvar_anotacao_comparacao_historica(
                event["id"], item_id, SimpleNamespace(anotacao="Anotação local mais recente.", base_revisoes=main.get_comparacao_historica(event["id"])["evento"]["comparacao"]["apelidos_sem_atualizacao"][0]["anotacao_revisoes"])
            )

        self.assertEqual(saved["sincronizacao"], "pendente")
        pending = main.load_history_annotations(self.pending_annotations)
        self.assertEqual(len(pending), 1)
        self.assertEqual(pending[0]["anotacao"], "Anotação local mais recente.")
        detail = main.get_comparacao_historica(event["id"])["evento"]
        displayed_row = detail["comparacao"]["apelidos_sem_atualizacao"][0]
        self.assertEqual(displayed_row["anotacao"], "Anotação local mais recente.")

        flushed = main.flush_pending_history_annotations()
        self.assertEqual(flushed, {"pendentes": 0, "sincronizadas": 1})
        self.assertEqual(list(self.pending_annotations.glob("*.json")), [])
        shared = resolve_history_annotations(main.load_history_annotations(self.shared_annotations))
        self.assertEqual(shared[(event["id"], item_id)]["anotacao"], "Anotação local mais recente.")

    def test_history_list_uses_saved_summary_without_transforming_rows_or_migrating_notes(self):
        event = {
            "id": "evento-resumo-rapido",
            "horario": "2026-10-07T13:00:00-03:00",
            "resumo": {"total_endpoint": 900, "ausentes_no_endpoint": 4},
            "comparacao": {
                "apelidos_sem_atualizacao": [{
                    "serie": "SN-LEGACY-SUMMARY",
                    "computador": "PRNOTE-LEGACY-SUMMARY",
                    "anotacao": "Anotação legada permanece na consulta detalhada.",
                }],
            },
        }
        self.comparison_history.write_text(json.dumps([event]), encoding="utf-8")

        with patch.object(
            main,
            "comparison_history_for_display",
            side_effect=AssertionError("a listagem não deve transformar todas as linhas"),
        ):
            listed = main.get_historico_comparacoes()

        self.assertEqual(listed["total"], 1)
        self.assertEqual(listed["eventos"][0]["resumo"], event["resumo"])
        self.assertEqual(list(self.shared_annotations.glob("*.json")), [])
        self.assertEqual(list(self.pending_annotations.glob("*.json")), [])

    def test_upload_validation_and_transaction_do_not_block_the_event_loop(self):
        make_details(self.details, [
            ("SN-1", "ALIAS-OLD", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-1", "PRNOTE-OLD")])
        make_model(self.model)
        self.comparison_history.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())

        validation_started = threading.Event()
        allow_validation = threading.Event()
        transaction_started = threading.Event()
        allow_transaction = threading.Event()
        retention_started = threading.Event()
        allow_retention = threading.Event()
        original_validator = main.validate_uploaded_workbook
        original_loader = main.load_database
        worker_threads = {}

        def pause_validation(file_bytes, role):
            if role == "endpoint":
                worker_threads["validation"] = threading.get_ident()
                validation_started.set()
                if not allow_validation.wait(3):
                    raise TimeoutError("o teste não liberou a validação pausada")
            return original_validator(file_bytes, role)

        def pause_transaction_reload(*args, **kwargs):
            worker_threads["transaction"] = threading.get_ident()
            transaction_started.set()
            if not allow_transaction.wait(3):
                raise TimeoutError("o teste não liberou a transação pausada")
            return original_loader(*args, **kwargs)

        def pause_retention():
            worker_threads["retention"] = threading.get_ident()
            retention_started.set()
            if not allow_retention.wait(2):
                raise TimeoutError("o teste não liberou a retenção pausada")

        upload_endpoint = UploadFile(
            filename="InvAll_nova.xlsx",
            file=io.BytesIO(_endpoint_bytes([("SN-1", "PRNOTE-NEW")])),
        )
        upload_details = UploadFile(
            filename="Detalhes_novos.xlsx",
            file=io.BytesIO(_details_bytes([
                ("SN-1", "ALIAS-NEW", "Modelo A", "Maringá", "Notebook", "87000000"),
            ])),
        )

        async def run_upload_with_heartbeats():
            loop = asyncio.get_running_loop()
            loop_thread = threading.get_ident()
            task = asyncio.create_task(main.atualizar_planilhas(upload_endpoint, upload_details))

            async def wait_for_event(event, description):
                deadline = loop.time() + 2
                while not event.is_set() and loop.time() < deadline:
                    await asyncio.sleep(0.005)
                self.assertTrue(event.is_set(), f"não iniciou {description}")

            try:
                await wait_for_event(validation_started, "a validação")
                await asyncio.wait_for(asyncio.sleep(0.025), timeout=0.25)
                self.assertFalse(task.done(), "a validação pausada bloqueou ou encerrou a atualização")
                allow_validation.set()

                await wait_for_event(transaction_started, "a recarga transacional")
                await asyncio.wait_for(asyncio.sleep(0.025), timeout=0.25)
                self.assertFalse(task.done(), "a recarga pausada bloqueou ou encerrou a atualização")
                allow_transaction.set()

                await wait_for_event(retention_started, "a retenção")
                await asyncio.wait_for(asyncio.sleep(0.025), timeout=0.25)
                self.assertFalse(task.done(), "a retenção pausada bloqueou ou encerrou a atualização")
            finally:
                allow_validation.set()
                allow_transaction.set()
                allow_retention.set()

            result = await task
            return loop_thread, result

        with (
            patch.object(main, "validate_uploaded_workbook", pause_validation),
            patch.object(main, "load_database", pause_transaction_reload),
            patch.object(main, "prune_expired_histories", pause_retention),
        ):
            loop_thread, result = asyncio.run(run_upload_with_heartbeats())

        self.assertTrue(result["sucesso"])
        self.assertNotEqual(worker_threads["validation"], loop_thread)
        self.assertNotEqual(worker_threads["transaction"], loop_thread)
        self.assertNotEqual(worker_threads["retention"], loop_thread)
        self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "PRNOTE-NEW")

    def test_interrupted_upload_finishes_atomic_transaction_before_closing_request(self):
        make_details(self.details, [
            ("SN-1", "ALIAS-OLD", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-1", "PRNOTE-OLD")])
        make_model(self.model)
        self.comparison_history.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())

        transaction_started = threading.Event()
        allow_transaction = threading.Event()
        original_loader = main.load_database

        def pause_transaction_reload(*args, **kwargs):
            transaction_started.set()
            if not allow_transaction.wait(3):
                raise TimeoutError("o teste não liberou a transação pausada")
            return original_loader(*args, **kwargs)

        endpoint_bytes = _endpoint_bytes([("SN-1", "PRNOTE-NEW")])
        details_bytes = _details_bytes([
            ("SN-1", "ALIAS-NEW", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        upload_endpoint = UploadFile(filename="InvAll_cancelada.xlsx", file=io.BytesIO(endpoint_bytes))
        upload_details = UploadFile(filename="Detalhes_cancelados.xlsx", file=io.BytesIO(details_bytes))

        async def cancel_and_release_upload():
            task = asyncio.create_task(main.atualizar_planilhas(upload_endpoint, upload_details))
            loop = asyncio.get_running_loop()
            deadline = loop.time() + 2
            try:
                while not transaction_started.is_set() and loop.time() < deadline:
                    await asyncio.sleep(0.005)
                self.assertTrue(transaction_started.is_set(), "a atualização não entrou na transação")

                task.cancel()
                await asyncio.sleep(0)
                task.cancel()
                await asyncio.sleep(0)
                task.cancel()
                await asyncio.sleep(0)
                self.assertFalse(task.done(), "o cancelamento liberou recursos antes da transação")
            finally:
                allow_transaction.set()

            with self.assertRaises(asyncio.CancelledError):
                await task

        with patch.object(main, "load_database", pause_transaction_reload):
            asyncio.run(cancel_and_release_upload())

        self.assertTrue(upload_endpoint.file.closed)
        self.assertTrue(upload_details.file.closed)
        self.assertEqual(self.endpoint.read_bytes(), endpoint_bytes)
        self.assertEqual(self.details.read_bytes(), details_bytes)
        history = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        self.assertEqual(len(history), 1)
        self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "PRNOTE-NEW")

    def test_repeated_upload_cancellation_keeps_parser_admission_until_worker_finishes(self):
        admission = threading.BoundedSemaphore(1)
        parser_started = threading.Event()
        release_parser = threading.Event()
        parser_finished = threading.Event()

        def blocked_parser(_data, _role):
            parser_started.set()
            try:
                if not release_parser.wait(3):
                    raise TimeoutError("o teste não liberou o parser")
                return {"records": []}
            finally:
                parser_finished.set()

        services = SimpleNamespace(
            max_upload_bytes=lambda: 1024,
            workbook_parse_admission=lambda: admission,
            validate_uploaded_workbook=blocked_parser,
        )
        upload = UploadFile(filename="endpoint-sintetico.xlsx", file=io.BytesIO(b"fixture"))

        async def run():
            task = asyncio.create_task(inventory_update.read_uploaded_xlsx(upload, "endpoint", services))
            try:
                self.assertTrue(await asyncio.wait_for(asyncio.to_thread(parser_started.wait, 3), 4))
                for _ in range(3):
                    task.cancel()
                    await asyncio.sleep(0)
                    self.assertFalse(task.done())
                    self.assertFalse(admission.acquire(blocking=False))
                release_parser.set()
                with self.assertRaises(asyncio.CancelledError):
                    await asyncio.wait_for(task, 4)
            finally:
                release_parser.set()
                if not task.done():
                    task.cancel()
                    await asyncio.gather(task, return_exceptions=True)

        asyncio.run(run())
        self.assertTrue(parser_finished.is_set())
        self.assertTrue(admission.acquire(blocking=False))
        admission.release()

    def test_preview_keeps_locked_snapshot_off_event_loop_until_repeated_cancellation(self):
        lock_started = threading.Event()
        release_lock = threading.Event()
        lock_finished = threading.Event()

        def blocked_lock_paths(_paths):
            from contextlib import contextmanager

            @contextmanager
            def hold_lock():
                lock_started.set()
                try:
                    if not release_lock.wait(3):
                        raise TimeoutError("o teste não liberou o lock da prévia")
                    yield
                finally:
                    lock_finished.set()

            return hold_lock()

        services = replace(
            main._inventory_update_services,
            get_primary_equipment_file=lambda: None,
        )
        endpoint_upload = UploadFile(
            filename="endpoint-sintetico.xlsx",
            file=io.BytesIO(_endpoint_bytes([("SN-PREVIEW-ASYNC", "HOST-1")])),
        )
        details_upload = UploadFile(
            filename="detalhes-sinteticos.xlsx",
            file=io.BytesIO(_details_bytes([
                ("SN-PREVIEW-ASYNC", "HOST-1", "Modelo A", "Maringá", "Notebook", "87000000"),
            ])),
        )

        async def run():
            task = asyncio.create_task(inventory_update.preview_inventory_files(
                endpoint_upload, details_upload, services,
            ))
            try:
                self.assertTrue(await asyncio.wait_for(asyncio.to_thread(lock_started.wait, 3), 4))
                loop_tick = asyncio.Event()
                asyncio.get_running_loop().call_soon(loop_tick.set)
                await asyncio.wait_for(loop_tick.wait(), 1)
                self.assertFalse(task.done(), "o lock da prévia não pode bloquear a event loop")

                task.cancel()
                await asyncio.sleep(0)
                task.cancel()
                await asyncio.sleep(0)
                self.assertFalse(task.done(), "a prévia cancelada deve aguardar a leitura protegida")
                release_lock.set()
                with self.assertRaises(asyncio.CancelledError):
                    await asyncio.wait_for(task, 4)
            finally:
                release_lock.set()
                if not task.done():
                    task.cancel()
                    await asyncio.gather(task, return_exceptions=True)

        with patch.object(inventory_update, "lock_paths", side_effect=blocked_lock_paths):
            asyncio.run(run())

        self.assertTrue(lock_finished.is_set())
        self.assertTrue(endpoint_upload.file.closed)
        self.assertTrue(details_upload.file.closed)

    def test_upload_replaces_inventory_snapshot_and_keeps_duplicate_source_rows(self):
        make_details(self.details, [
            ("SN-REMOVER", "ALIAS-ANTIGO", "Modelo antigo", "Maringá", "Notebook", "87000000"),
            ("SN-MANTER", "ALIAS-ANTIGO", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [
            ("SN-REMOVER", "HOST-ANTIGO"),
            ("SN-MANTER", "HOST-ANTIGO"),
        ])
        make_model(self.model)
        self.comparison_history.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())

        endpoint_bytes = _endpoint_bytes([
            ("SN-MANTER", "HOST-NOVO"),
            ("SN-DUP", "HOST-DUPLICADO"),
        ])
        details_bytes = _details_bytes([
            ("SN-MANTER", "ALIAS-DA-FONTE", "Modelo A", "Maringá", "Notebook", "87000000"),
            ("SN-DUP", "ALIAS-DUP-A", "Modelo B", "Londrina", "Notebook", "86000000"),
            ("SN-DUP", "ALIAS-DUP-B", "Modelo C", "Londrina", "Notebook", "86000000"),
        ])
        result = asyncio.run(main.atualizar_planilhas(
            UploadFile(filename="InvAll_retrato.xlsx", file=io.BytesIO(endpoint_bytes)),
            UploadFile(filename="Detalhes_retrato.xlsx", file=io.BytesIO(details_bytes)),
        ))

        self.assertTrue(result["sucesso"])
        saved_details = openpyxl.load_workbook(self.details, data_only=False)
        details_sheet = saved_details["EquipamentosDetalhes"]
        details_rows = [
            (row[0], row[1])
            for row in details_sheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in row)
        ]
        self.assertEqual(details_rows, [
            ("SN-MANTER", "ALIAS-DA-FONTE"),
            ("SN-DUP", "ALIAS-DUP-A"),
            ("SN-DUP", "ALIAS-DUP-B"),
        ])
        self.assertEqual(saved_details["Metadados"]["B1"].value, "=1+1")
        saved_details.close()

        saved_endpoint = openpyxl.load_workbook(self.endpoint, data_only=True)
        endpoint_rows = list(saved_endpoint["InvAllComputersSummary"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(endpoint_rows, [("SN-MANTER", "HOST-NOVO"), ("SN-DUP", "HOST-DUPLICADO")])
        saved_endpoint.close()

        self.assertNotIn("SN-REMOVER", main.equipamentos_cache)
        self.assertEqual(main.equipamentos_cache["SN-MANTER"]["apelido"], "HOST-NOVO")
        self.assertEqual(main.equipamentos_cache["SN-DUP"]["invall_alias_status"], "serie_detalhes_duplicada")
        self.assertNotEqual(main.equipamentos_cache["SN-DUP"]["apelido"], "HOST-DUPLICADO")

    def test_upload_history_failure_rolls_back_workbooks_history_and_memory(self):
        make_details(self.details, [("SN-1", "ALIAS-DETALHES", "Modelo A", "Maringá", "Notebook", "87000000")])
        make_endpoint(self.endpoint, [("SN-1", "PRNOTE-OLD")])
        make_model(self.model)
        self.comparison_history.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())
        before_bytes = {path: path.read_bytes() for path in (self.details, self.endpoint)}
        before_cache = copy.deepcopy(main.equipamentos_cache)
        endpoint_bytes = _endpoint_bytes([("SN-1", "PRNOTE-NEW")])
        details_bytes = _details_bytes([("SN-1", "ALIAS-DETALHES", "Modelo A", "Maringá", "Notebook", "87000000")])
        original_update = main.AtomicFileTransaction.update_json

        def fail_comparison_history(transaction, path, updater, default):
            if Path(path) == self.comparison_history:
                raise OSError("falha fictícia no histórico")
            return original_update(transaction, path, updater, default)

        with patch.object(main.AtomicFileTransaction, "update_json", fail_comparison_history):
            with self.assertRaises(main.HTTPException) as raised:
                asyncio.run(main.atualizar_planilhas(
                    UploadFile(filename="InvAll_falha.xlsx", file=io.BytesIO(endpoint_bytes)),
                    UploadFile(filename="Detalhes_falha.xlsx", file=io.BytesIO(details_bytes)),
                ))
        self.assertEqual(raised.exception.status_code, 500)
        self.assertIn("restaurados", raised.exception.detail)
        self.assertEqual({path: path.read_bytes() for path in before_bytes}, before_bytes)
        self.assertEqual(self.comparison_history.read_text(encoding="utf-8"), "[]")
        self.assertEqual(main.equipamentos_cache, before_cache)

    def test_upload_failure_preserves_alias_committed_while_upload_waited_for_lock(self):
        make_details(self.details, [("SN-1", "ALIAS-DETALHES", "Modelo A", "Maringá", "Notebook", "87000000")])
        make_endpoint(self.endpoint, [("SN-1", "PRNOTE-OLD")])
        self.comparison_history.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())

        original_transaction = main.AtomicFileTransaction
        upload_waiting = threading.Event()
        allow_upload = threading.Event()

        class PausingTransaction(original_transaction):
            def __enter__(self):
                required = {self.details_path, self.endpoint_path}
                if required.issubset(set(self.targets)):
                    upload_waiting.set()
                    if not allow_upload.wait(10):
                        raise TimeoutError("o teste não liberou a transação pausada")
                return super().__enter__()

        PausingTransaction.details_path = self.details.resolve()
        PausingTransaction.endpoint_path = self.endpoint.resolve()
        original_update = original_transaction.update_json

        def fail_comparison_history(transaction, path, updater, default):
            if Path(path).resolve() == self.comparison_history.resolve():
                raise OSError("falha fictícia no histórico")
            return original_update(transaction, path, updater, default)

        upload_outcome = {}

        def run_upload():
            try:
                upload_outcome["result"] = asyncio.run(main.atualizar_planilhas(
                    UploadFile(filename="InvAll_nova.xlsx", file=io.BytesIO(_endpoint_bytes([("SN-1", "PRNOTE-NEW")]))),
                    UploadFile(filename="Detalhes_novos.xlsx", file=io.BytesIO(_details_bytes([
                        ("SN-1", "ALIAS-DETALHES", "Modelo A", "Maringá", "Notebook", "87000000"),
                    ]))),
                ))
            except BaseException as error:
                upload_outcome["error"] = error

        with (
            patch.object(main, "AtomicFileTransaction", PausingTransaction),
            patch.object(original_transaction, "update_json", fail_comparison_history),
        ):
            worker = threading.Thread(target=run_upload)
            worker.start()
            try:
                self.assertTrue(upload_waiting.wait(10), "o upload não chegou ao bloqueio esperado")
                main.aplicar_alteracao_apelido("SN-1", "MANUAL-COMMITTED")
                committed_workbooks = {path: path.read_bytes() for path in (self.details, self.endpoint)}
                committed_overrides = self.overrides.read_text(encoding="utf-8")
                self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "MANUAL-COMMITTED")
            finally:
                allow_upload.set()
                worker.join(10)

        self.assertFalse(worker.is_alive(), "a transação de upload não terminou")
        self.assertIsInstance(upload_outcome.get("error"), main.HTTPException)
        self.assertEqual(upload_outcome["error"].status_code, 500)
        self.assertEqual({path: path.read_bytes() for path in committed_workbooks}, committed_workbooks)
        self.assertEqual(self.overrides.read_text(encoding="utf-8"), committed_overrides)
        self.assertEqual(self.comparison_history.read_text(encoding="utf-8"), "[]")
        self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "MANUAL-COMMITTED")

    def test_stale_restore_preview_does_not_revert_concurrent_alias_commit_in_cache(self):
        make_details(self.details, [("SN-1", "ALIAS-ATUAL", "Modelo A", "Maringá", "Notebook", "87000000")])
        make_endpoint(self.endpoint, [("SN-1", "PRNOTE-ATUAL")])
        backup_dir = self.backups / "restore-source"
        (backup_dir / "endpoint").mkdir(parents=True)
        (backup_dir / "detalhes").mkdir()
        make_endpoint(backup_dir / "endpoint" / self.endpoint.name, [("SN-1", "PRNOTE-ANTIGO")])
        backup_headers = complete_details_headers(
            ["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP"]
        )
        backup_row = ["SN-1", "ALIAS-ANTIGO", "Modelo A", "Maringá", "Notebook", "87000000"]
        backup_row.extend([None] * (len(backup_headers) - len(backup_row)))
        cnpj_header_keys = {
            main.normalize_workbook_header(header)
            for header in main.DETAILS_INSTALLATION_CNPJ_HEADER_ALIASES
        }
        cnpj_column = next(
            index for index, header in enumerate(backup_headers)
            if main.normalize_workbook_header(header) in cnpj_header_keys
        )
        backup_row[cnpj_column] = "12345678000195"
        make_details(backup_dir / "detalhes" / self.details.name, [
            backup_row,
        ], headers=backup_headers)
        self.comparison_history.write_text(json.dumps([{
            "id": "restore-source-event",
            "horario": "2026-10-01T12:00:00-03:00",
            "comparacao": {"backup_diretorio": str(backup_dir.relative_to(self.input_dir))},
        }]), encoding="utf-8")
        self.assertTrue(main.load_database())
        preview = main.visualizar_restauracao_comparacao("restore-source-event")

        original_transaction = main.AtomicFileTransaction
        restore_waiting = threading.Event()
        allow_restore = threading.Event()

        class PausingTransaction(original_transaction):
            def __enter__(self):
                required = {self.details_path, self.endpoint_path}
                if required.issubset(set(self.targets)):
                    restore_waiting.set()
                    if not allow_restore.wait(10):
                        raise TimeoutError("o teste não liberou a transação pausada")
                return super().__enter__()

        PausingTransaction.details_path = self.details.resolve()
        PausingTransaction.endpoint_path = self.endpoint.resolve()
        restore_outcome = {}

        def run_restore():
            try:
                restore_outcome["result"] = main.restaurar_comparacao_historica(
                    "restore-source-event",
                    main.ComparisonRestorePayload(preview_token=preview["preview_token"]),
                )
            except BaseException as error:
                restore_outcome["error"] = error

        with patch.object(main, "AtomicFileTransaction", PausingTransaction):
            worker = threading.Thread(target=run_restore)
            worker.start()
            try:
                self.assertTrue(restore_waiting.wait(10), "a restauração não chegou ao bloqueio esperado")
                main.aplicar_alteracao_apelido("SN-1", "MANUAL-COMMITTED")
                self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "MANUAL-COMMITTED")
            finally:
                allow_restore.set()
                worker.join(10)

        self.assertFalse(worker.is_alive(), "a restauração não terminou")
        self.assertNotIn("error", restore_outcome)
        self.assertTrue(restore_outcome["result"]["sucesso"])
        self.assertEqual(main.load_modified_apelidos(), {"SN-1": "MANUAL-COMMITTED"})
        self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "MANUAL-COMMITTED")

    def test_restore_counts_manual_aliases_left_without_equipment_and_keeps_them(self):
        make_details(self.details, [
            ("SN-1", "ALIAS-ATUAL", "Modelo A", "Maringá", "Notebook", "87000000"),
            ("SN-NOVO", "ALIAS-NOVO", "Modelo B", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-1", "PRNOTE-ATUAL"), ("SN-NOVO", "PRNOTE-NOVO")])
        backup_dir = self.backups / "restore-orphans"
        (backup_dir / "endpoint").mkdir(parents=True)
        (backup_dir / "detalhes").mkdir()
        make_endpoint(backup_dir / "endpoint" / self.endpoint.name, [("SN-1", "PRNOTE-ANTIGO")])
        backup_headers = complete_details_headers(
            ["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP"]
        )
        backup_row = ["SN-1", "ALIAS-ANTIGO", "Modelo A", "Maringá", "Notebook", "87000000"]
        backup_row.extend([None] * (len(backup_headers) - len(backup_row)))
        cnpj_header_keys = {
            main.normalize_workbook_header(header)
            for header in main.DETAILS_INSTALLATION_CNPJ_HEADER_ALIASES
        }
        cnpj_column = next(
            index for index, header in enumerate(backup_headers)
            if main.normalize_workbook_header(header) in cnpj_header_keys
        )
        backup_row[cnpj_column] = "12345678000195"
        make_details(backup_dir / "detalhes" / self.details.name, [backup_row], headers=backup_headers)
        self.comparison_history.write_text(json.dumps([{
            "id": "restore-orphans-event",
            "horario": "2026-10-01T12:00:00-03:00",
            "comparacao": {"backup_diretorio": str(backup_dir.relative_to(self.input_dir))},
        }]), encoding="utf-8")
        self.assertTrue(main.load_database())
        main.aplicar_alteracao_apelido("SN-1", "MANUAL-1")
        main.aplicar_alteracao_apelido("SN-NOVO", "MANUAL-NOVO")
        preview = main.visualizar_restauracao_comparacao("restore-orphans-event")

        result = main.restaurar_comparacao_historica(
            "restore-orphans-event",
            main.ComparisonRestorePayload(preview_token=preview["preview_token"]),
        )

        self.assertEqual(result["evento"]["apelidos_manuais_sem_equipamento"], 1)
        saved = next(
            event for event in json.loads(self.comparison_history.read_text(encoding="utf-8"))
            if event.get("tipo") == "restauracao"
        )
        self.assertEqual(saved["comparacao"]["apelidos_manuais_sem_equipamento"], 1)
        self.assertEqual(main.load_modified_apelidos(), {"SN-1": "MANUAL-1", "SN-NOVO": "MANUAL-NOVO"})

    def test_restore_keeps_old_cache_until_file_and_history_transaction_commits(self):
        self.details.write_bytes(_details_bytes([
            ("SN-RESTORE-PREPARE", "ALIAS-ATUAL", "Modelo atual", "Maringá", "Notebook", "87000000"),
        ]))
        self.endpoint.write_bytes(_endpoint_bytes([("SN-RESTORE-PREPARE", "HOST-ATUAL")]))
        backup_dir = self.backups / "restore-prepared-state"
        (backup_dir / "endpoint").mkdir(parents=True)
        (backup_dir / "detalhes").mkdir()
        (backup_dir / "endpoint" / self.endpoint.name).write_bytes(
            _endpoint_bytes([("SN-RESTORE-PREPARE", "HOST-ANTIGO")])
        )
        (backup_dir / "detalhes" / self.details.name).write_bytes(_details_bytes([
            ("SN-RESTORE-PREPARE", "ALIAS-ANTIGO", "Modelo antigo", "Londrina", "Notebook", "86000000"),
        ]))
        self.comparison_history.write_text(json.dumps([{
            "id": "restore-prepared-event",
            "horario": "2026-10-01T12:00:00-03:00",
            "comparacao": {"backup_diretorio": str(backup_dir.relative_to(self.input_dir))},
        }]), encoding="utf-8")
        self.assertTrue(main.load_database())
        preview = main.visualizar_restauracao_comparacao("restore-prepared-event")

        state_prepared = threading.Event()
        allow_commit = threading.Event()
        original_loader = main.load_database

        def pause_after_prepare(*args, **kwargs):
            state = original_loader(*args, **kwargs)
            if kwargs.get("publish") is False:
                state_prepared.set()
                if not allow_commit.wait(10):
                    raise TimeoutError("o teste não liberou a confirmação da restauração")
            return state

        outcome = {}

        def restore():
            try:
                outcome["result"] = main.restaurar_comparacao_historica(
                    "restore-prepared-event",
                    main.ComparisonRestorePayload(preview_token=preview["preview_token"]),
                )
            except BaseException as error:
                outcome["error"] = error

        with patch.object(main, "load_database", side_effect=pause_after_prepare):
            worker = threading.Thread(target=restore)
            worker.start()
            try:
                self.assertTrue(state_prepared.wait(10), "o estado restaurado não foi preparado")
                self.assertEqual(main.get_equipamento("SN-RESTORE-PREPARE")["modelo"], "Modelo atual")
            finally:
                allow_commit.set()
                worker.join(10)

        self.assertFalse(worker.is_alive(), "a restauração não terminou")
        self.assertNotIn("error", outcome)
        self.assertTrue(outcome["result"]["sucesso"])
        self.assertEqual(main.get_equipamento("SN-RESTORE-PREPARE")["modelo"], "Modelo antigo")

    def test_manual_alias_override_changes_export_only(self):
        detail_headers = ["Hostname", "Tipo", "Modelo", "Cidade", "CEP", "Serial"]
        make_details(
            self.details,
            [("HOST-SOURCE", "Notebook", "Modelo A", "Maringá", "87000000", "S-1")],
            headers=detail_headers,
        )
        make_endpoint(self.endpoint, [("S-1", "HOST-ENDPOINT")])

        export_model = self.input_dir / "modelo-exportacao.xlsx"
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Parque"
        worksheet.append(["NumeroSerie", "Apelido", "Observacao"])
        worksheet.append(["S-1", "HOST-MODELO", "MANTER"])
        workbook.save(export_model)
        workbook.close()

        original_details = self.details.read_bytes()
        original_export_model = export_model.read_bytes()
        self.assertTrue(main.load_database())

        class CapturedResponse:
            def __init__(self, content, media_type, headers):
                self.body = content.read()
                self.headers = headers

        with (
            patch.object(main, "EXPORT_MODEL_FILE", export_model),
            patch.object(main, "StreamingResponse", CapturedResponse),
        ):
            result = main.atualizar_apelido(main.ApelidoUpdatePayload(
                serial="S-1", apelido="HOST-MANUAL"
            ))

            self.assertTrue(result["salvo_no_disco"])
            self.assertEqual(result["equipamento"]["apelido"], "HOST-MANUAL (M)")
            self.assertEqual(main.load_modified_apelidos(), {"S-1": "HOST-MANUAL (M)"})
            self.assertEqual(self.details.read_bytes(), original_details)

            response = main.exportar_equipamentos()
            exported = openpyxl.load_workbook(io.BytesIO(response.body), data_only=False)
            self.assertEqual(exported["Parque"]["B2"].value, "HOST-MANUAL (M)")
            self.assertEqual(exported["Parque"]["C2"].value, "MANTER")
            exported.close()

        self.assertEqual(self.details.read_bytes(), original_details)
        self.assertEqual(export_model.read_bytes(), original_export_model)

    def test_validated_workbook_update_accepts_excel_column_width_float_round_trip(self):
        make_details(
            self.details,
            [("BRJ5512FGD", "HOST-ORIGINAL", "Modelo A", "Maringá", "Notebook", "87000000")],
        )
        workbook = openpyxl.load_workbook(self.details)
        worksheet = workbook["EquipamentosDetalhes"]
        worksheet.title = "Parque"
        for column, width in {
            "AE": 31.842256546020508,
            "AL": 30.294265747070312,
            "AO": 30.802078247070312,
            "AU": 33.097835540771484,
            "R": 30.873506546020508,
            "W": 30.635784149169922,
        }.items():
            worksheet.column_dimensions[column].width = width
        workbook.save(self.details)
        workbook.close()

        def update_alias(workbook):
            workbook["Parque"]["B2"] = "HOST-NOVO"
            return {("Parque", "B2"): "HOST-NOVO"}

        with main.AtomicFileTransaction(self.root, [self.details]) as transaction:
            staged = main.stage_workbook_update(
                transaction,
                self.details,
                update_alias,
            )
            self.assertTrue(staged)
            transaction.commit()

        saved = openpyxl.load_workbook(self.details, data_only=True)
        self.assertEqual(saved["Parque"]["B2"].value, "HOST-NOVO")
        saved.close()

    def test_manual_alias_rejects_changed_active_source_without_mutating_state(self):
        make_details(
            self.details,
            [("SN-1", "HOST-ORIGINAL", "Modelo A", "Maringá", "Notebook", "87000000")],
        )
        make_endpoint(self.endpoint, [("SN-1", "HOST-ENDPOINT")])
        self.overrides.write_text("{}", encoding="utf-8")
        self.manual_history.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())
        cached_before = copy.deepcopy(main.equipamentos_cache)
        original_source = Path(cached_before["SN-1"]["_source_file"]).resolve()

        older_mtime = time.time_ns() - 10_000_000_000
        os.utime(self.details, ns=(older_mtime, older_mtime))
        replacement_source = self.input_dir / "EquipamentosDetalhes_replacement.xlsx"
        make_details(
            replacement_source,
            [("SN-1", "HOST-REPLACEMENT", "Modelo A", "Maringá", "Notebook", "87000000")],
        )
        replacement_before = replacement_source.read_bytes()
        original_before = self.details.read_bytes()
        overrides_before = self.overrides.read_bytes()
        history_before = self.manual_history.read_bytes()

        self.assertEqual(main.get_primary_equipment_file().resolve(), replacement_source.resolve())
        self.assertEqual(original_source, self.details.resolve())
        with self.assertRaises(main.HTTPException) as raised:
            main.atualizar_apelido(main.ApelidoUpdatePayload(serial="SN-1", apelido="HOST-MANUAL"))

        self.assertEqual(raised.exception.status_code, 409)
        self.assertEqual(replacement_source.read_bytes(), replacement_before)
        self.assertEqual(self.details.read_bytes(), original_before)
        self.assertEqual(self.overrides.read_bytes(), overrides_before)
        self.assertEqual(self.manual_history.read_bytes(), history_before)
        self.assertEqual(main.equipamentos_cache, cached_before)

    def test_manual_alias_can_be_removed_after_inventory_refresh_with_hostname_table(self):
        details_headers = ["NumeroSerie", "Hostname", "Modelo", "Cidade", "Tipo", "CEP"]
        make_details(
            self.details,
            [("SN-1", "HOST-DETAILS-OLD", "Modelo A", "Maringá", "Notebook", "87000000")],
            headers=details_headers,
        )
        add_structured_table(self.details)
        make_endpoint(self.endpoint, [("SN-1", "HOST-INVALL-OLD")])
        make_model(self.model)
        self.comparison_history.write_text("[]", encoding="utf-8")
        self.assertTrue(main.load_database())

        manual_result = main.atualizar_apelido(main.ApelidoUpdatePayload(
            serial="SN-1", apelido="HOST-MANUAL"
        ))
        self.assertTrue(manual_result["salvo_no_disco"])
        workbook = openpyxl.load_workbook(self.details, data_only=True)
        self.assertEqual(workbook["EquipamentosDetalhes"]["B2"].value, "HOST-DETAILS-OLD")
        self.assertEqual(workbook["EquipamentosDetalhes"]["E2"].value, "Notebook")
        self.assertEqual(workbook["EquipamentosDetalhes"].tables["EquipmentInventory"].ref, "A1:F2")
        workbook.close()

        incoming_details = self.root / "detalhes-atualizados-ficticios.xlsx"
        make_details(
            incoming_details,
            [("SN-1", "HOST-DETAILS-NEW", "Modelo A", "Maringá", "Notebook", "87000000")],
            headers=details_headers,
        )
        add_structured_table(incoming_details)
        shutil.copyfile(incoming_details, self.details)
        make_endpoint(self.endpoint, [("SN-1", "HOST-INVALL-NEW")])
        self.assertTrue(main.load_database())
        event_id = "comparacao-manual-ficticia"
        divergence = {
            "serie": "SN-1",
            "apelido_manual": "HOST-MANUAL (M)",
            "apelido_equipamentos_detalhes": "HOST-DETAILS-NEW",
            "apelido_invall": "HOST-INVALL-NEW",
            "invall_encontrado": True,
            "resolucao": "pendente",
        }
        self.comparison_history.write_text(json.dumps([{
            "id": event_id,
            "comparacao": {"divergencias_apelidos_manuais": [divergence]},
        }]), encoding="utf-8")
        self.assertEqual(divergence["apelido_manual"], "HOST-MANUAL (M)")
        self.assertEqual(divergence["apelido_equipamentos_detalhes"], "HOST-DETAILS-NEW")
        self.assertEqual(divergence["apelido_invall"], "HOST-INVALL-NEW")

        resolution = main.resolver_divergencia_apelido_manual(
            event_id,
            "SN-1",
            main.ManualAliasConflictResolutionPayload(acao="excluir"),
        )

        self.assertEqual(resolution["resolucao"], "excluir")
        self.assertEqual(resolution["equipamento"]["apelido"], "HOST-INVALL-NEW")
        self.assertEqual(main.load_modified_apelidos(), {})
        saved_workbook = openpyxl.load_workbook(self.details, data_only=True)
        saved_sheet = saved_workbook["EquipamentosDetalhes"]
        self.assertEqual(saved_sheet["B2"].value, "HOST-DETAILS-NEW")
        self.assertEqual(saved_sheet["E2"].value, "Notebook")
        self.assertEqual(saved_sheet.tables["EquipmentInventory"].ref, "A1:F2")
        saved_workbook.close()

        history = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        stored_event = next(item for item in history if item["id"] == event_id)
        stored_divergence = stored_event["comparacao"]["divergencias_apelidos_manuais"][0]
        self.assertEqual(stored_divergence["resolucao"], "excluir")

    def test_manual_undo_reverses_latest_event_and_preserves_manual_override_priority(self):
        make_details(self.details, [("SN-1", "ALIAS-ORIGINAL", "Modelo A", "Maringá", "Notebook", "87000000")])
        make_endpoint(self.endpoint, [("SN-1", "HOST-INVALL")])
        make_model(self.model)
        self.comparison_history.write_text(json.dumps([{"id": "planilha-antiga"}]), encoding="utf-8")
        self.manual_history.write_text(json.dumps([{
            "id": "manual-anterior",
            "tipo": "apelido",
            "serial": "SN-1",
            "apelido_anterior": "VALOR-ANTES",
            "apelido_novo": "ALIAS-ORIGINAL",
        }]), encoding="utf-8")
        self.assertTrue(main.load_database())

        saved = main.atualizar_apelido(main.ApelidoUpdatePayload(serial="SN-1", apelido="ALIAS-MANUAL"))
        self.assertEqual(saved["equipamento"]["apelido"], "ALIAS-MANUAL (M)")
        history_after_save = json.loads(self.manual_history.read_text(encoding="utf-8"))
        latest_event = history_after_save[-1]
        self.assertEqual(latest_event["apelido_anterior"], "ALIAS-ORIGINAL")
        self.assertEqual(latest_event["apelido_novo"], "ALIAS-MANUAL (M)")

        undone = main.desfazer_alteracao_apelido(main.ApelidoUndoPayload(evento_id=latest_event["id"]))
        # Sem override antes da edição, o desfazer remove o override e o hostname do Endpoint volta a valer.
        self.assertEqual(undone["equipamento"]["apelido"], "HOST-INVALL")
        self.assertEqual(main.load_modified_apelidos(), {})
        self.assertEqual(len(json.loads(self.manual_history.read_text(encoding="utf-8"))), 3)
        self.assertEqual(json.loads(self.comparison_history.read_text(encoding="utf-8")), [{"id": "planilha-antiga"}])
        saved_details = openpyxl.load_workbook(self.details, data_only=True)
        self.assertEqual(saved_details["EquipamentosDetalhes"]["B2"].value, "ALIAS-ORIGINAL")
        saved_details.close()

        # A stale in-memory cache must not authorize undo if the persisted override changed elsewhere.
        latest_undo = json.loads(self.manual_history.read_text(encoding="utf-8"))[-1]
        self.overrides.write_text(json.dumps({"SN-1": "VALOR-EXTERNO"}), encoding="utf-8")
        details_before_stale_undo = self.details.read_bytes()
        with self.assertRaises(main.HTTPException) as raised:
            main.desfazer_alteracao_apelido(main.ApelidoUndoPayload(evento_id=latest_undo["id"]))
        self.assertEqual(raised.exception.status_code, 409)
        self.assertEqual(self.details.read_bytes(), details_before_stale_undo)
        self.assertEqual(main.load_modified_apelidos(), {"SN-1": "VALOR-EXTERNO"})

    def _prepare_alias_undo_fixture(self, manual_history=None):
        make_details(self.details, [("SN-1", "ALIAS-ORIGINAL", "Modelo A", "Maringá", "Notebook", "87000000")])
        make_endpoint(self.endpoint, [("SN-1", "HOST-INVALL")])
        make_model(self.model)
        self.comparison_history.write_text("[]", encoding="utf-8")
        if manual_history is not None:
            self.manual_history.write_text(json.dumps(manual_history), encoding="utf-8")

    def test_undo_without_previous_override_removes_override_and_restores_endpoint_hostname(self):
        self._prepare_alias_undo_fixture()
        self.assertTrue(main.load_database())
        self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "HOST-INVALL")

        main.atualizar_apelido(main.ApelidoUpdatePayload(serial="SN-1", apelido="ALIAS-MANUAL"))
        edit_event = json.loads(self.manual_history.read_text(encoding="utf-8"))[-1]
        self.assertIs(edit_event["tinha_override"], False)
        self.assertEqual(main.load_modified_apelidos(), {"SN-1": "ALIAS-MANUAL (M)"})

        undone = main.desfazer_alteracao_apelido(main.ApelidoUndoPayload(evento_id=edit_event["id"]))

        self.assertEqual(main.load_modified_apelidos(), {})
        self.assertEqual(undone["equipamento"]["apelido"], "HOST-INVALL")
        self.assertFalse(main.equipamentos_cache["SN-1"]["_apelido_manual_override"])
        undo_event = json.loads(self.manual_history.read_text(encoding="utf-8"))[-1]
        self.assertEqual(undo_event["acao"], "desfazer")
        self.assertEqual(undo_event["apelido_anterior"], "ALIAS-MANUAL (M)")
        self.assertEqual(undo_event["apelido_novo"], "ALIAS-ORIGINAL")
        self.assertIs(undo_event["tinha_override"], True)

        # Um novo hostname no Endpoint passa a valer, porque não sobrou override manual.
        make_endpoint(self.endpoint, [("SN-1", "HOST-INVALL-NOVO")])
        self.assertTrue(main.load_database())
        self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "HOST-INVALL-NOVO")

    def test_undo_with_previous_override_restores_that_override(self):
        self._prepare_alias_undo_fixture()
        self.overrides.write_text(json.dumps({"SN-1": "MANUAL-ANTIGO"}), encoding="utf-8")
        self.assertTrue(main.load_database())
        self.assertEqual(main.equipamentos_cache["SN-1"]["apelido"], "MANUAL-ANTIGO")

        main.atualizar_apelido(main.ApelidoUpdatePayload(serial="SN-1", apelido="MANUAL-NOVO"))
        edit_event = json.loads(self.manual_history.read_text(encoding="utf-8"))[-1]
        self.assertIs(edit_event["tinha_override"], True)

        undone = main.desfazer_alteracao_apelido(main.ApelidoUndoPayload(evento_id=edit_event["id"]))

        self.assertEqual(main.load_modified_apelidos(), {"SN-1": "MANUAL-ANTIGO"})
        self.assertEqual(undone["equipamento"]["apelido"], "MANUAL-ANTIGO")

    def test_undo_of_event_without_override_marker_keeps_legacy_behavior(self):
        legacy_event = {
            "id": "evento-legado",
            "tipo": "apelido",
            "serial": "SN-1",
            "apelido_anterior": "ALIAS-ORIGINAL",
            "apelido_novo": "ALIAS-MANUAL",
            "horario": "2026-10-01T10:00:00-03:00",
        }
        self._prepare_alias_undo_fixture(manual_history=[legacy_event])
        self.overrides.write_text(json.dumps({"SN-1": "ALIAS-MANUAL"}), encoding="utf-8")
        self.assertTrue(main.load_database())

        undone = main.desfazer_alteracao_apelido(main.ApelidoUndoPayload(evento_id="evento-legado"))

        self.assertEqual(main.load_modified_apelidos(), {"SN-1": "ALIAS-ORIGINAL"})
        self.assertEqual(undone["equipamento"]["apelido"], "ALIAS-ORIGINAL")

    def test_deleting_latest_alias_event_without_previous_override_removes_override(self):
        self._prepare_alias_undo_fixture()
        self.assertTrue(main.load_database())
        main.atualizar_apelido(main.ApelidoUpdatePayload(serial="SN-1", apelido="ALIAS-MANUAL"))
        edit_event = json.loads(self.manual_history.read_text(encoding="utf-8"))[-1]

        restored = main.excluir_evento_historico_apelido(edit_event["id"])

        self.assertEqual(main.load_modified_apelidos(), {})
        self.assertEqual(restored["equipamento"]["apelido"], "HOST-INVALL")
        remaining_ids = [item["id"] for item in json.loads(self.manual_history.read_text(encoding="utf-8"))]
        self.assertNotIn(edit_event["id"], remaining_ids)

    def test_address_history_restoration_remains_rejected_without_mutation(self):
        make_details(self.details, [
            ("SN-ENDERECO", "HOST-ENDERECO", "Modelo A", "Maringá", "Notebook", "87000000"),
        ])
        make_endpoint(self.endpoint, [("SN-ENDERECO", "HOST-ENDERECO")])
        address_event = {
            "id": "historico-endereco-ficticio",
            "tipo": "endereco",
            "serial": "SN-ENDERECO",
            "endereco_anterior": {"cidade": "Londrina"},
            "endereco_novo": {"cidade": "Maringá"},
        }
        self.manual_history.write_text(json.dumps([address_event]), encoding="utf-8")
        self.assertTrue(main.load_database())
        details_before = self.details.read_bytes()
        history_before = self.manual_history.read_bytes()
        equipment_before = copy.deepcopy(main.equipamentos_cache["SN-ENDERECO"])

        with self.assertRaises(main.HTTPException) as raised:
            main.excluir_evento_historico_apelido("historico-endereco-ficticio")

        self.assertEqual(raised.exception.status_code, 409)
        self.assertIn("atualizados pela planilha", raised.exception.detail)
        self.assertEqual(self.details.read_bytes(), details_before)
        self.assertEqual(self.manual_history.read_bytes(), history_before)
        self.assertEqual(main.equipamentos_cache["SN-ENDERECO"], equipment_before)

    def test_90_day_retention_preserves_recent_invalid_and_pinned_history(self):
        now = datetime(2026, 10, 4, 12, tzinfo=timezone.utc)
        old_timestamp = "2026-01-01T12:00:00+00:00"
        recent_timestamp = "2026-10-01T12:00:00+00:00"
        self.manual_history.write_text(json.dumps([
            {"id": "alias-old", "horario": old_timestamp},
            {"id": "alias-new", "horario": recent_timestamp},
            {"id": "alias-unknown-date", "horario": "invalid"},
        ]), encoding="utf-8")
        self.backups.mkdir()
        expired_backup = self.backups / "expired"
        recent_backup = self.backups / "recent"
        orphan_backup = self.backups / "orphan"
        for folder in (expired_backup, recent_backup, orphan_backup):
            folder.mkdir()
            (folder / "marker.txt").write_text(folder.name, encoding="utf-8")
        os.utime(orphan_backup, (946684800, 946684800))
        self.comparison_history.write_text(json.dumps([
            {"id": "comparison-old", "horario": old_timestamp, "comparacao": {"backup_diretorio": "_backups/expired"}},
            {"id": "comparison-new", "horario": recent_timestamp, "comparacao": {"backup_diretorio": "_backups/recent"}},
            {"id": main.PINNED_COMPARISON_EVENT_ID, "horario": main.PINNED_COMPARISON_EVENT_TIME, "comparacao": {}},
            {"id": "comparison-unknown-date", "horario": "invalid", "comparacao": {}},
        ]), encoding="utf-8")

        result = main.prune_expired_histories(now=now)

        aliases = json.loads(self.manual_history.read_text(encoding="utf-8"))
        comparisons = json.loads(self.comparison_history.read_text(encoding="utf-8"))
        self.assertEqual(result, {
            "apelidos_removidos": 1,
            "comparacoes_removidas": 1,
            "backups_removidos": 2,
        })
        self.assertEqual({item["id"] for item in aliases}, {"alias-new", "alias-unknown-date"})
        self.assertEqual(
            {item["id"] for item in comparisons},
            {"comparison-new", main.PINNED_COMPARISON_EVENT_ID, "comparison-unknown-date"},
        )
        self.assertFalse(expired_backup.exists())
        self.assertTrue(recent_backup.exists())
        self.assertFalse(orphan_backup.exists())

    def test_corrupt_local_json_is_reported_and_never_replaced_as_empty(self):
        self.overrides.write_text("{malformed", encoding="utf-8")
        original_overrides = self.overrides.read_bytes()
        with self.assertRaises(main.LocalJsonCorruptionError):
            main.load_modified_apelidos()
        # Production writers update aliases through a transaction, never a bare write.
        with self.assertRaises(main.FileTransactionError):
            with main.AtomicFileTransaction(main.TRANSACTION_ROOT, [self.overrides]) as transaction:
                transaction.update_json(
                    self.overrides,
                    lambda data: {**main._validate_json_object(data, "apelidos"), "SN-1": "HOST-1"},
                    {},
                )
                transaction.commit()
        self.assertEqual(self.overrides.read_bytes(), original_overrides)

        self.custom_units.write_text('{"wrong":"shape"}', encoding="utf-8")
        original_units = self.custom_units.read_bytes()
        with self.assertRaises(main.LocalJsonCorruptionError):
            main.load_custom_units()
        self.assertEqual(self.custom_units.read_bytes(), original_units)

        self.manual_history.write_text("[null]", encoding="utf-8")
        original_history = self.manual_history.read_bytes()
        with self.assertRaises(main.HTTPException) as raised:
            main.load_alias_history()
        self.assertEqual(raised.exception.status_code, 503)
        self.assertEqual(self.manual_history.read_bytes(), original_history)

    def test_search_reuses_pre_normalized_inventory_values(self):
        main.equipamentos_cache = {
            "SN-1": {
                "serial": "SN-1", "modelo": "Notebook One", "tipo": "Notebook",
                "apelido": "HOST-ONE", "retirada_cnpj": "12.345.678/0001-90",
            },
            "SN-2": {"serial": "SN-2", "modelo": "Monitor Two", "tipo": "Monitor", "apelido": "HOST-TWO"},
        }
        with main.STATE_LOCK:
            main._refresh_equipment_search_index_locked()
        with patch.object(main, "_normalizar_termo_busca", wraps=main._normalizar_termo_busca) as normalize:
            result = main._buscar_equipamentos_por_termo("HOST-ONE")

        self.assertEqual([item["serial"] for item in result], ["SN-1"])
        self.assertEqual(result[0]["retirada_cnpj"], "12.345.678/0001-90")
        self.assertEqual(normalize.call_count, 1)

    def test_type_stats_resolve_normalized_alias_headers_and_keep_empty_rows(self):

        rows = [
            ("SN-1", "HOST-1", "Modelo A", "Maringá", "Notebook", "87000000"),
            ("SN-2", "HOST-2", "Modelo B", "Londrina", None, "86000000"),
            (None, "HOST-3", "Modelo C", "Curitiba", "Monitor", "80000000"),
        ]
        expected = {"Notebook": 1, "": 1, "Monitor": 1}
        for type_header in ("Tipo", "TIPO", " Tipo ", "ClasseEquipamento"):
            with self.subTest(type_header=type_header):
                make_details(
                    self.details,
                    rows,
                    headers=["NumeroSerie", "Apelido", "Modelo", "Cidade", type_header, "CEP"],
                )
                self.assertTrue(main.load_database())
                self.assertEqual(main.stats_cache["tipos"], expected)

    def test_notebook_search_returns_installation_cnpj_for_every_synchronized_series(self):
        source_rows = [
            (
                "SN-1", "HOST-1", "Notebook Modelo 1", "Maringá", "Notebook",
                "87000000", "12.345.678/0001-90",
            ),
            (
                "SN-2", "HOST-2", "Notebook Modelo 2", "Londrina", "Notebook",
                "86000000", "98.765.432/0001-10",
            ),
            (
                "SN-3", "HOST-3", "Notebook Modelo 3", "Foz do Iguaçu", "Notebook",
                "85852010", "45.678.901/0001-23",
            ),
        ]
        make_details(
            self.details,
            source_rows,
            headers=[
                "NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP",
                "CNPJ Instalacao",
            ],
        )
        make_endpoint(self.endpoint, [(row[0], row[1]) for row in source_rows])
        self.assertTrue(main.load_database())

        results = main._buscar_equipamentos_por_termo("Notebook")

        self.assertEqual(
            {item["serial"]: item["retirada_cnpj"] for item in results},
            {row[0]: row[6] for row in source_rows},
        )

    def test_export_statistics_keep_unavailable_distinct_from_zero_changes(self):
        with (
            patch.object(main, "stats_cache", {"total_equipamentos": 12}),
            patch.object(main, "count_export_changes", return_value={"apelidos": 0, "id_clientes": 0, "total": 0}),
            patch.object(main, "get_export_status", return_value={"available": False, "reason": "modelo ausente"}),
        ):
            unavailable = main.get_stats()
        self.assertFalse(unavailable["exportacao_disponivel"])
        self.assertEqual(unavailable["exportacao_motivo"], "modelo ausente")
        self.assertEqual(unavailable["total_alteracoes_exportaveis"], 0)

        with (
            patch.object(main, "stats_cache", {"total_equipamentos": 12}),
            patch.object(main, "count_export_changes", return_value={"apelidos": 0, "id_clientes": 0, "total": 0}),
            patch.object(main, "get_export_status", return_value={"available": True, "reason": "OK"}),
        ):
            available = main.get_stats()
        self.assertTrue(available["exportacao_disponivel"])
        self.assertEqual(available["total_alteracoes_exportaveis"], 0)

async def _collect_stream(response):
    return b"".join([chunk async for chunk in response.body_iterator])


def _endpoint_bytes(rows):
    with tempfile.TemporaryDirectory(prefix=".upload-bytes-test-", dir=PROJECT_ROOT) as directory:
        path = Path(directory) / "endpoint-ficticio.xlsx"
        make_endpoint(path, rows)
        return path.read_bytes()


def _details_bytes(rows):
    with tempfile.TemporaryDirectory(prefix=".upload-bytes-test-", dir=PROJECT_ROOT) as directory:
        path = Path(directory) / "detalhes-ficticio.xlsx"
        headers = complete_details_headers(
            ["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP"]
        )
        installation_cnpj_headers = {
            main.normalize_workbook_header(header)
            for header in main.DETAILS_INSTALLATION_CNPJ_HEADER_ALIASES
        }
        cnpj_column = next(
            index for index, header in enumerate(headers)
            if main.normalize_workbook_header(header) in installation_cnpj_headers
        )
        complete_rows = []
        for row in rows:
            values = list(row)
            while len(values) <= cnpj_column:
                values.append(None)
            if not str(values[cnpj_column] or "").strip():
                values[cnpj_column] = "00.000.000/0000-00"
            complete_rows.append(values)
        make_details(
            path,
            complete_rows,
            headers=headers,
        )
        return path.read_bytes()


if __name__ == "__main__":
    unittest.main()
