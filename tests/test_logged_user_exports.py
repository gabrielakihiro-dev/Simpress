"""Logged user in the history export and the "Sem usuário" filter of the full search."""

from datetime import date
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from fastapi import HTTPException
import openpyxl

from domains.equipment_search import CATEGORIES, normalize_search_term
from domains.history_export import build_history_export_workbook, history_export_series
from routers.equipment_search import EquipmentSearchServices, create_equipment_search_router


def _event():
    return {
        "id": "evento-usuarios",
        "horario": "2026-10-20T09:00:00-03:00",
        "comparacao": {
            "apelidos_sem_atualizacao": [{"serie": "BRJ-ENDPOINT", "computador": "PRNOTE001", "motivo": "Conferir."}],
            "ausentes_no_endpoint": [{"serie": "BRJ-MANUAL", "computador": "PRNOTE002", "local_previsto": "Curitiba"}],
            "sem_acesso_7_dias": [{"serie": "BRJ-VAZIO", "computador": "PRNOTE003", "dias_sem_verificacao": 9}],
            "divergencias_local": [{"serie": "BRJ-ENDPOINT", "computador": "PRNOTE001", "local_previsto": "Curitiba"}],
            "series_sem_mapeamento": [{"serie": "PRX", "quantidade": 2, "exemplos": ["PRX1"]}],
        },
    }


USERS = {
    "BRJ-ENDPOINT": {"usuario": "ana.silva", "origem": "endpoint", "serial_cadastro": "BRJ-ENDPOINT"},
    "BRJ-MANUAL": {"usuario": "joao.souza", "origem": "manual", "serial_cadastro": "BRJ-MANUAL"},
    "BRJ-VAZIO": {"usuario": "", "origem": "", "serial_cadastro": "BRJ-VAZIO"},
}


def _sheet_rows(workbook, name):
    sheet = workbook[name]
    headers = [cell.value for cell in sheet[4]]
    return headers, [dict(zip(headers, (cell.value for cell in row))) for row in sheet.iter_rows(min_row=5)]


class HistoryExportLoggedUserTests(unittest.TestCase):
    def test_series_are_collected_once_in_first_seen_order(self):
        self.assertEqual(history_export_series([_event()]), ["BRJ-ENDPOINT", "BRJ-MANUAL", "BRJ-VAZIO"])

    def test_rows_show_the_user_and_its_source_before_the_annotation(self):
        output, _ = build_history_export_workbook([_event()], date(2026, 10, 20), USERS)
        workbook = openpyxl.load_workbook(output)
        try:
            for sheet in ("Apelidos", "Ausências", "Sem acesso", "Localidades"):
                headers, _ = _sheet_rows(workbook, sheet)
                self.assertEqual(headers[-3:], ["Usuário logado", "Origem do usuário", "Anotação"], sheet)
            _, aliases = _sheet_rows(workbook, "Apelidos")
            self.assertEqual((aliases[0]["Usuário logado"], aliases[0]["Origem do usuário"]), ("ana.silva", "Endpoint"))
            _, absences = _sheet_rows(workbook, "Ausências")
            self.assertEqual((absences[0]["Usuário logado"], absences[0]["Origem do usuário"]), ("joao.souza", "Manual"))
            _, no_access = _sheet_rows(workbook, "Sem acesso")
            self.assertIn(no_access[0]["Usuário logado"], (None, ""))
            self.assertIn(no_access[0]["Origem do usuário"], (None, ""))
            _, locations = _sheet_rows(workbook, "Localidades")
            self.assertEqual(locations[0]["Usuário logado"], "ana.silva")
            self.assertIn(locations[1]["Usuário logado"], (None, ""))
            self.assertEqual(locations[1]["Quantidade"], 2)
        finally:
            workbook.close()

    def test_export_without_users_keeps_blank_columns(self):
        output, _ = build_history_export_workbook([_event()], date(2026, 10, 20))
        workbook = openpyxl.load_workbook(output)
        try:
            _, aliases = _sheet_rows(workbook, "Apelidos")
            self.assertIn(aliases[0]["Usuário logado"], (None, ""))
            self.assertEqual(aliases[0]["Motivo"], "Conferir.")
        finally:
            workbook.close()


class HistoryExportRouteLoggedUserTests(unittest.TestCase):
    """``main`` wires the Home rule (``logged_users_for_series``) into the history export route."""

    def export(self, logged_users):
        import main

        class CapturedResponse:
            def __init__(self, content, media_type, headers):
                self.body = content.read()

        with tempfile.TemporaryDirectory() as directory:
            history = Path(directory) / "historico_comparacoes.json"
            history.write_text(json.dumps([_event()]), encoding="utf-8")
            with (
                patch.object(main, "COMPARISON_HISTORY_FILE", history),
                patch.object(main.user_changes, "logged_users_for_series", logged_users),
                patch("routers.inventory_history.StreamingResponse", CapturedResponse),
            ):
                response = main.exportar_historico_comparacoes_excel(data="2026-10-20")
        workbook = openpyxl.load_workbook(io.BytesIO(response.body))
        try:
            return _sheet_rows(workbook, "Ausências")[1][0]
        finally:
            workbook.close()

    def test_route_fills_the_current_user_of_each_series(self):
        calls = []

        def lookup(services, series):
            calls.append(list(series))
            return {"usuarios": USERS}

        row = self.export(lookup)
        self.assertEqual([sorted(series) for series in calls], [["BRJ-ENDPOINT", "BRJ-MANUAL", "BRJ-VAZIO"]])
        self.assertEqual((row["Usuário logado"], row["Origem do usuário"]), ("joao.souza", "Manual"))

    def test_route_still_exports_when_the_user_lookup_fails(self):
        def failing(services, series):
            raise RuntimeError("Endpoint ilegível")

        with self.assertLogs("routers.inventory_history", level="ERROR"):
            row = self.export(failing)
        self.assertIn(row["Usuário logado"], (None, ""))
        self.assertEqual(row["Série"], "BRJ-MANUAL")


ITEMS = [
    {"serial": "BRJ001", "modelo": "Notebook", "tipo": "Notebook", "categoria": "notebook", "retirada_cidade": "Curitiba", "usuarios_logados": "ana.silva", "usuarios_logados_origem": "endpoint"},
    {"serial": "BRJ002", "modelo": "Notebook", "tipo": "Notebook", "categoria": "notebook", "retirada_cidade": "Curitiba", "usuarios_logados": "", "usuarios_logados_origem": ""},
    {"serial": "BRJ003", "modelo": "Monitor", "tipo": "Monitor", "categoria": "monitor", "retirada_cidade": "Londrina", "usuarios_logados": "joao", "usuarios_logados_origem": "manual"},
    {"serial": "XYZ004", "modelo": "Monitor", "tipo": "Monitor", "categoria": "monitor", "retirada_cidade": "Londrina", "usuarios_logados": " ", "usuarios_logados_origem": ""},
    {"serial": "HS005", "modelo": "POLY-B8225", "tipo": "Headsets", "categoria": "fone", "retirada_cidade": "Curitiba", "usuarios_logados": "", "usuarios_logados_origem": ""},
    {"serial": "MC006", "modelo": "HP P24A G5", "tipo": "Micro Computador", "categoria": "", "retirada_cidade": "Curitiba", "usuarios_logados": "", "usuarios_logados_origem": ""},
    {"serial": "MAC007", "modelo": "MACBOOK PRO 14", "tipo": "Notebook", "categoria": "macbook", "retirada_cidade": "Curitiba", "usuarios_logados": "bia", "usuarios_logados_origem": "endpoint"},
    {"serial": "MAC008", "modelo": "MACBOOK PRO 14", "tipo": "Notebook", "categoria": "macbook", "retirada_cidade": "Curitiba", "usuarios_logados": "", "usuarios_logados_origem": ""},
    {"serial": "OUT009", "modelo": "Outro", "tipo": "Impressora", "categoria": "", "retirada_cidade": "Curitiba", "usuarios_logados": "", "usuarios_logados_origem": ""},
]


def _search_route():
    services = EquipmentSearchServices(
        search=lambda q: [item for item in ITEMS if normalize_search_term(q) in item["serial"]],
        normalize=normalize_search_term,
        categories=CATEGORIES,
        get_stats=dict,
        list_all=lambda: list(ITEMS),
    )
    _, exports = create_equipment_search_router(services)

    def call(q="", usuario=None, categoria=None, include_filters=True):
        return exports.buscar_seriais_expandida(
            q=q, offset=0, limit=50, categoria=categoria, tipo=None, cidade=None,
            usuario=usuario, include_filters=include_filters,
        )
    return call


class FullSearchUserFilterTests(unittest.TestCase):
    def test_listing_without_a_term_has_notebooks_without_user_and_other_equipment(self):
        page = _search_route()(usuario="sem")
        # No notebook with a user, no fones, no Micro Computador/monitor.
        self.assertEqual([item["serial"] for item in page["items"]], ["BRJ002", "MAC008", "OUT009"])
        self.assertEqual(page["total"], 3)
        self.assertEqual(page["filters"]["users"], {"sem": 3, "com": 0})
        self.assertEqual([entry["value"] for entry in page["filters"]["categories"]], ["notebook", "macbook"])
        self.assertEqual([entry["count"] for entry in page["filters"]["categories"]], [1, 1])

    def test_term_search_still_finds_fones_and_monitors(self):
        call = _search_route()
        self.assertEqual([item["serial"] for item in call("HS", "sem")["items"]], ["HS005"])
        self.assertEqual([item["serial"] for item in call("MC", None)["items"]], ["MC006"])

    def test_user_filter_combines_with_the_term_and_category(self):
        call = _search_route()
        self.assertEqual([item["serial"] for item in call("BRJ", "sem")["items"]], ["BRJ002"])
        self.assertEqual([item["serial"] for item in call("BRJ", "com")["items"]], ["BRJ001", "BRJ003"])
        self.assertEqual([item["serial"] for item in call("BRJ", "com", "monitor")["items"]], ["BRJ003"])
        self.assertEqual(call("BRJ", "invalido")["total"], 3)

    def test_a_term_is_still_required_except_for_the_without_user_listing(self):
        call = _search_route()
        for usuario in (None, "com", "qualquer"):
            with self.assertRaises(HTTPException) as raised:
                call("", usuario)
            self.assertEqual(raised.exception.status_code, 422)
        with self.assertRaises(HTTPException):
            call("B", None)


if __name__ == "__main__":
    unittest.main()
