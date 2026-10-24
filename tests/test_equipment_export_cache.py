import logging
from contextlib import nullcontext
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from domains import equipment_export


class EquipmentExportStatusCacheTests(unittest.TestCase):
    def test_temporary_read_error_is_retried_without_a_file_change(self):
        with tempfile.TemporaryDirectory(dir=tempfile.gettempdir()) as directory:
            root = Path(directory)
            model_file = root / "modelo.xlsx"
            details_file = root / "detalhes.xlsx"
            for path in (model_file, details_file):
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "EquipamentosDetalhes"
                sheet.append(["NumeroSerie", "Apelido"])
                sheet.append(["SERIE-1", ""])
                workbook.save(path)
                workbook.close()

            services = equipment_export.EquipmentExportServices(
                export_model_file=lambda: model_file,
                endpoint_upload_file=lambda: root / "endpoint.xlsx",
                modified_aliases_file=lambda: root / "aliases.json",
                transaction_root=lambda: root,
                read_lock_timeout=lambda: 1,
                get_parque_file=lambda: details_file,
                get_details_files=lambda: [details_file],
                find_details_sheet=lambda _path: "EquipamentosDetalhes",
                read_endpoint_alias_index=lambda: {"unique": {}},
                load_modified_apelidos=lambda: {},
                latest_comparison_update_entry=lambda: None,
                validate_xlsx=lambda _path: None,
                lock_paths=lambda *_args, **_kwargs: nullcontext(),
                logger=logging.getLogger(__name__),
            )
            load_workbook = openpyxl.load_workbook
            calls = 0

            def fail_once(path, *args, **kwargs):
                nonlocal calls
                calls += 1
                if calls == 1:
                    raise PermissionError("arquivo temporariamente ocupado")
                return load_workbook(path, *args, **kwargs)

            with patch.dict(equipment_export._EXPORT_COUNTS_CACHE, {}, clear=True), patch.object(
                equipment_export.openpyxl, "load_workbook", side_effect=fail_once
            ):
                first_counts, first_status = equipment_export.count_export_changes_with_status(
                    services, locked=True
                )
                second_counts, second_status = equipment_export.count_export_changes_with_status(
                    services, locked=True
                )

        self.assertEqual(first_counts["total"], 0)
        self.assertEqual(first_status["state"], "unavailable")
        self.assertEqual(second_counts["total"], 1)
        self.assertEqual(second_status["state"], "ready")


if __name__ == "__main__":
    unittest.main()
