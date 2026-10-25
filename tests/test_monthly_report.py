"""Monthly contract report: content rules (docs/operacao/relatorio-mensal.md), workbook and route.

Every serial, hostname and OS below is fictitious, except the Sebrae-owned serial that
AGENTS.md lists by name.
"""

import json
import tempfile
import unittest
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import openpyxl
from fastapi import FastAPI
from fastapi.responses import JSONResponse
from fastapi.testclient import TestClient

from core.errors import DomainError
from domains.inventory_comparison import comparison_history_for_display
from domains.monthly_report import (
    CLOSED,
    OPENED,
    PENDING,
    SECTION_ENDPOINT,
    SECTION_LOCATION,
    SECTION_SIMPRESS,
    InvalidReportMonthError,
    build_monthly_report,
    build_monthly_report_workbook,
    comparison_section,
    inventory_section,
    parse_report_month,
    summary_rows,
    tickets_section,
)
from routers.monthly_report import (
    MonthlyReportServices,
    create_monthly_report_router,
    generate_previous_month_report,
    previous_report_month,
)

TODAY = date(2026, 10, 17)
OCTOBER = parse_report_month("2026-10", TODAY)
SEPTEMBER = parse_report_month("2026-09", TODAY)


def equipment(serial, tipo, city, address):
    return {"serial": serial, "tipo": tipo, "retirada_cidade": city, "endereco_instalacao": address}


EQUIPMENT = [
    equipment("SN-FIC-001", "Notebook", "Curitiba", "R Caeté, 150 - SEDE - Prado Velho"),
    equipment("SN-FIC-002", "Notebook", "Curitiba", "R Caeté, 150 - SEDE - Prado Velho"),
    equipment("SN-FIC-003", "Monitor", "Londrina", "Av Ficticia, 10"),
    equipment("SN-FIC-004", "", "Maringá", "Rua Exemplo, 20"),
]
STATS = {
    "arquivo_origem": "EquipamentosDetalhes_atual.xlsx",
    "ultima_atualizacao": "17/10/2026 08:00",
    "linhas_sem_serie": 2,
    "series_duplicadas": {"SNFIC001": 2},
}


def update_event(event_id, moment, comparison):
    return {"id": event_id, "horario": moment, "tipo": "atualizacao", "comparacao": comparison, "resumo": {}}


def october_comparison():
    return {
        "ausentes_no_endpoint": [
            {"computador": "", "serie": "BRJFIC0001", "local_previsto": "Curitiba"},
        ],
        "ausentes_em_detalhes": [
            {"computador": "PRMACFIC01", "serie": "CFICMAC001", "local_previsto": "Curitiba"},
            # Sebrae-owned notebook: never counted as missing in the Simpress base.
            {"computador": "PRMACFIC02", "serie": "BRJ135P5XS", "local_previsto": "Curitiba"},
        ],
        "divergencias_local": [
            {"computador": "PRNOTEMARFIC1", "serie": "SN-FIC-010", "local_previsto": "Maringá", "local_detalhes": "Londrina"},
            # Installed in Curitiba although the base says São José dos Pinhais.
            {"computador": "PRNOTESJPFIC1", "serie": "SN-FIC-011", "local_previsto": "São José dos Pinhais", "local_detalhes": "Curitiba"},
            {"computador": "PRNOTERLFIC1", "serie": "SN-FIC-012", "local_previsto": "Rio Negro", "local_detalhes": "Curitiba"},
            {"computador": "PRNOTEARCFIC1", "serie": "SN-FIC-013", "local_previsto": "Araucária", "local_detalhes": "Curitiba"},
        ],
    }


HISTORY = [
    update_event("evt-setembro", "2026-09-10T10:00:00-03:00", {
        "ausentes_no_endpoint": [], "ausentes_em_detalhes": [], "divergencias_local": [],
    }),
    update_event("evt-outubro-1", "2026-10-02T09:00:00-03:00", {
        "ausentes_no_endpoint": [], "ausentes_em_detalhes": [], "divergencias_local": [],
    }),
    update_event("evt-outubro-2", "2026-10-15T09:00:00-03:00", october_comparison()),
    {"id": "evt-restauracao", "horario": "2026-10-20T09:00:00-03:00", "tipo": "restauracao", "comparacao": {}},
    update_event("evt-novembro", "2026-11-03T09:00:00-03:00", {
        "ausentes_no_endpoint": [{"computador": "", "serie": "BRJFIC0099"}], "ausentes_em_detalhes": [], "divergencias_local": [],
    }),
]


def oracle(os_number, created, concluded="", status="Gestão da Mesa de Operações", sub_status="Execução Pendente"):
    return {
        "solicitacao": f"SOL-{os_number}", "ordem_servico": os_number, "status": status, "sub_status": sub_status,
        "data_criacao": created, "data_conclusao": concluded, "ultima_atualizacao": created, "serie": "",
    }


def ticket(serial, oracle_items=(), simpress=None):
    return {"serie": serial, "relacao": "somente_oracle", "oracle": list(oracle_items), "simpress": simpress or {}}


def detailed_order(identifier, opened="", status="", delivered=""):
    stages = [
        {"chave": "data_pedido", "rotulo": "Data do pedido", "data": opened},
        {"chave": "entregue", "rotulo": "Pedido entregue", "data": delivered},
    ]
    return {
        "identificador": identifier,
        "status": [status] if status else [],
        "status_incompleto": not bool(status),
        "datas_pedido": [opened] if opened else [],
        "etapas_pedido": stages,
        "associacao_incompleta": False,
    }


SEPTEMBER_SNAPSHOT = {
    "id": "snap-1",
    "atualizado_em": "2026-09-25T10:00:00-03:00",
    "fontes": {},
    "resumo": {},
    "chamados": [ticket("SN-FIC-100", [oracle("9100", "20/09/2026 10:00")])],
}
OCTOBER_TICKETS = [
    ticket("SN-FIC-100", [oracle("9100", "20/09/2026 10:00", "05/10/2026 16:00", sub_status="Concluído")]),
    ticket("SN-FIC-101", [oracle("9101", "10/10/2026 08:30")]),
    # Closed without a conclusion date: neither pending nor concluded in the month.
    ticket("SN-FIC-102", [oracle("9102", "01/08/2026 09:00", status="Encerrado", sub_status="Cancelado")]),
    ticket("SN-FIC-103", simpress={
        "pedidos": ["P-FIC-1"], "datas_pedido": ["03/10/2026"], "status": ["Pedido em separação"],
    }),
    ticket("SN-FIC-104", simpress={
        "pedidos": ["P-FIC-2"], "datas_pedido": ["15/09/2026"], "status": ["Pedido entregue"],
        "etapas_pedido": [
            {"chave": "data_pedido", "rotulo": "Data do pedido", "data": "15/09/2026"},
            {"chave": "entregue", "rotulo": "Pedido entregue", "data": "02/10/2026"},
        ],
    }),
]
TICKET_REPORT = {
    "snapshot_id": "snap-2",
    "atualizado_em": "2026-10-12T18:00:00-03:00",
    "fontes": {},
    "resumo": {},
    "chamados": OCTOBER_TICKETS,
    "historico": [SEPTEMBER_SNAPSHOT],
}
def build(month=OCTOBER, display=comparison_history_for_display, **overrides):
    arguments = {
        "equipment": EQUIPMENT, "stats": STATS, "comparison_history": HISTORY,
        "comparison_display": display, "ticket_report": TICKET_REPORT,
        "generated_at": datetime(2026, 10, 17, 9, 30),
    }
    arguments.update(overrides)
    return build_monthly_report(month, **arguments)


class ReportMonthTests(unittest.TestCase):
    def test_accepts_current_and_past_months(self):
        self.assertEqual((OCTOBER.start, OCTOBER.end), (date(2026, 10, 1), TODAY))
        self.assertTrue(OCTOBER.is_partial)
        self.assertEqual(parse_report_month("2024-02", TODAY).end, date(2024, 2, 29))
        self.assertEqual(OCTOBER.label, "outubro de 2026")

    def test_rejects_invalid_or_future_months(self):
        for value in ("2026-11", "2027-01", "2026-13", "2026-00", "26-10", "2026/10", "", None, "1999-12"):
            with self.subTest(value=value), self.assertRaises(InvalidReportMonthError) as raised:
                parse_report_month(value, TODAY)
            self.assertEqual(raised.exception.status_code, 422)


class InventorySectionTests(unittest.TestCase):
    def test_counts_unique_serials_of_the_official_base_and_flags_gaps(self):
        section = inventory_section(EQUIPMENT, STATS)
        self.assertEqual(section["total"], 4)
        self.assertEqual(section["por_tipo"], [("Notebook", 2), ("Monitor", 1), ("Sem tipo", 1)])
        self.assertEqual(section["por_unidade"][0], ("Curitiba", "R Caeté, 150 - SEDE - Prado Velho", 2))
        self.assertEqual(section["unidades"], 3)
        self.assertEqual(section["linhas_sem_serie"], 2)
        self.assertEqual(section["series_duplicadas"], ["SNFIC001"])

    def test_breakdowns_add_up_to_the_total(self):
        section = inventory_section(EQUIPMENT, STATS)
        self.assertEqual(sum(count for _tipo, count in section["por_tipo"]), section["total"])
        self.assertEqual(sum(count for *_unit, count in section["por_unidade"]), section["total"])

    def test_missing_flags_default_to_zero(self):
        section = inventory_section([], {})
        self.assertEqual((section["total"], section["linhas_sem_serie"], section["series_duplicadas"]), (0, 0, []))


class ComparisonSectionTests(unittest.TestCase):
    def test_uses_latest_update_until_month_end_with_current_rules(self):
        section = comparison_section(HISTORY, OCTOBER, comparison_history_for_display)
        self.assertEqual(section["horario"].date(), date(2026, 10, 15))
        self.assertEqual(section["atualizacoes_no_mes"], 2)
        self.assertEqual(section["contagens"], {SECTION_ENDPOINT: 1, SECTION_SIMPRESS: 1, SECTION_LOCATION: 1})
        serials = {line["serie"] for line in section["linhas"]}
        self.assertNotIn("BRJ135P5XS", serials)
        self.assertFalse({"SN-FIC-011", "SN-FIC-012", "SN-FIC-013"} & serials)
        self.assertNotIn("BRJFIC0099", serials)

    def test_report_applies_sebrae_and_curitiba_rules_even_without_display_filters(self):
        section = comparison_section(HISTORY, OCTOBER, lambda event: event)
        self.assertEqual(section["contagens"][SECTION_SIMPRESS], 1)
        self.assertEqual(section["contagens"][SECTION_LOCATION], 1)

    def test_older_month_uses_its_own_last_update(self):
        section = comparison_section(HISTORY, SEPTEMBER, comparison_history_for_display)
        self.assertEqual(section["horario"].date(), date(2026, 9, 10))
        self.assertEqual(section["atualizacoes_no_mes"], 1)
        self.assertEqual(section["linhas"], [])

    def test_month_before_any_update_has_no_comparison(self):
        section = comparison_section(HISTORY, parse_report_month("2026-08", TODAY), comparison_history_for_display)
        self.assertIsNone(section["horario"])
        self.assertEqual(sum(section["contagens"].values()), 0)


class TicketSectionTests(unittest.TestCase):
    def test_october_counts_opened_concluded_and_pending(self):
        section = tickets_section(TICKET_REPORT, OCTOBER)
        self.assertEqual(section["oracle"], {"abertas": 1, "concluidas": 1, "pendentes": 1, "encerradas_sem_data": 1})
        self.assertEqual(section["simpress"], {"abertas": 1, "concluidas": 1, "pendentes": 1, "encerradas_sem_data": 0})
        self.assertEqual(section["importacoes"], 2)
        situations = {line["identificador"]: line["situacao"] for line in section["linhas"]}
        self.assertEqual(situations, {
            "9100": [CLOSED], "9101": [OPENED, PENDING],
            "P-FIC-1": [OPENED, PENDING], "P-FIC-2": [CLOSED],
        })

    def test_september_sees_tickets_still_open_at_month_end(self):
        section = tickets_section(TICKET_REPORT, SEPTEMBER)
        self.assertEqual(section["oracle"], {"abertas": 1, "concluidas": 0, "pendentes": 1, "encerradas_sem_data": 1})
        self.assertEqual(section["simpress"], {"abertas": 1, "concluidas": 0, "pendentes": 1, "encerradas_sem_data": 0})

    def test_os_repeated_across_imports_counts_once_with_newest_data(self):
        section = tickets_section(TICKET_REPORT, OCTOBER)
        rows = [line for line in section["linhas"] if line["identificador"] == "9100"]
        self.assertEqual(len(rows), 1)
        self.assertIn("Concluído", rows[0]["status"])

    def test_simpress_orders_are_counted_per_request_across_months_imports_and_equipment(self):
        september = {
            "id": "snap-setembro", "atualizado_em": "2026-09-30T18:00:00",
            "chamados": [{"serie": "SN-FIC-A", "simpress": {
                "pedidos": ["P-100"],
                "pedidos_detalhados": [detailed_order("P-100", "30/09/2026", "Em separação")],
            }}],
        }
        current = {
            "id": "snap-outubro", "atualizado_em": "2026-10-10T18:00:00",
            "chamados": [
                {"serie": "SN-FIC-A", "simpress": {"pedidos": ["P-200"], "pedidos_detalhados": [
                    detailed_order("P-100", "30/09/2026", "Pedido entregue", "02/10/2026"),
                    detailed_order("P-200", "05/10/2026", "Em separação"),
                ]}},
                {"serie": "SN-FIC-B", "simpress": {"pedidos": ["P-200"], "pedidos_detalhados": [
                    detailed_order("P-200", "05/10/2026", "Em separação"),
                ]}},
            ],
        }
        report = {
            "snapshot_id": current["id"], "atualizado_em": current["atualizado_em"],
            "chamados": current["chamados"], "historico": [september, current],
        }

        september_section = tickets_section(report, SEPTEMBER)
        self.assertEqual(september_section["simpress"], {
            "abertas": 1, "concluidas": 0, "pendentes": 1, "encerradas_sem_data": 0,
        })
        october_section = tickets_section(report, OCTOBER)
        self.assertEqual(october_section["simpress"], {
            "abertas": 1, "concluidas": 1, "pendentes": 1, "encerradas_sem_data": 0,
        })
        rows = [line for line in october_section["linhas"] if line["fonte"] == "Simpress"]
        self.assertEqual({line["identificador"] for line in rows}, {"P-100", "P-200"})
        self.assertEqual(sum(line["identificador"] == "P-200" for line in rows), 1)
        self.assertEqual(next(line["serie"] for line in rows if line["identificador"] == "P-200"), "SN-FIC-A, SN-FIC-B")
        p100 = next(line for line in rows if line["identificador"] == "P-100")
        self.assertEqual(p100["situacao"], [CLOSED])
        self.assertEqual(october_section["dados_insuficientes"], [])

    def test_missing_dates_do_not_infer_a_month_and_closed_without_delivery_is_flagged(self):
        report = {
            "snapshot_id": "snap-datas", "atualizado_em": "2026-10-12T18:00:00",
            "chamados": [
                {"serie": "SN-FIC-C", "simpress": {"pedidos": ["P-CANCEL"], "pedidos_detalhados": [
                    detailed_order("P-CANCEL", "20/09/2026", "Cancelado"),
                ]}},
                {"serie": "SN-FIC-D", "simpress": {"pedidos": ["P-SEM-DATA"], "pedidos_detalhados": [
                    detailed_order("P-SEM-DATA", "", "Cancelado"),
                ]}},
            ],
        }
        september_section = tickets_section(report, SEPTEMBER)
        self.assertEqual(september_section["simpress"]["abertas"], 1)
        self.assertEqual(september_section["simpress"]["encerradas_sem_data"], 1)
        october_section = tickets_section(report, OCTOBER)
        self.assertEqual(october_section["simpress"]["abertas"], 0)
        self.assertEqual(october_section["simpress"]["concluidas"], 0)
        self.assertEqual(october_section["simpress"]["pendentes"], 0)
        self.assertEqual(october_section["simpress"]["encerradas_sem_data"], 1)
        self.assertEqual([item["identificadores"] for item in october_section["dados_insuficientes"]], [["P-SEM-DATA"]])

    def test_legacy_multi_order_snapshot_is_reported_without_guessing_date_associations(self):
        legacy = {
            "snapshot_id": "legado-1", "atualizado_em": "2026-10-12T18:00:00",
            "chamados": [{"serie": "SN-FIC-LEGADO", "simpress": {
                "pedidos": ["P-LEG-1", "P-LEG-2"],
                "datas_pedido": ["02/10/2026", "10/10/2026"],
                "status": ["Em separação", "Pedido entregue"],
            }}],
        }
        section = tickets_section(legacy, OCTOBER)
        self.assertEqual(section["simpress"], {
            "abertas": 0, "concluidas": 0, "pendentes": 0, "encerradas_sem_data": 0,
        })
        self.assertEqual(len(section["dados_insuficientes"]), 1)
        self.assertEqual(section["dados_insuficientes"][0]["identificadores"], ["P-LEG-1", "P-LEG-2"])
        self.assertIn("legada", section["dados_insuficientes"][0]["motivo"])

    def test_empty_ticket_history_reports_zeros(self):
        section = tickets_section({}, OCTOBER)
        self.assertEqual(section["oracle"], {"abertas": 0, "concluidas": 0, "pendentes": 0, "encerradas_sem_data": 0})
        self.assertEqual((section["linhas"], section["importacoes"]), ([], 0))

    def test_latest_ambiguous_import_does_not_keep_counting_old_order_data(self):
        from copy import deepcopy

        report = {
            "snapshot_id": "latest-incomplete", "atualizado_em": "2026-10-12T18:00:00",
            "historico": [{
                "id": "older-valid", "atualizado_em": "2026-10-01T18:00:00",
                "chamados": [ticket("SN-FIC-A", simpress={"pedidos_detalhados": [
                    detailed_order("P-100", "01/10/2026", "Em separação"),
                ]})],
            }],
            "chamados": [ticket("SN-FIC-A", simpress={
                "pedidos": ["P-100", "P-200"], "datas_pedido": ["01/10/2026", "05/10/2026"],
                "status": ["Em separação", "Pedido entregue"],
            })],
        }
        original = deepcopy(report)
        section = tickets_section(report, OCTOBER)
        self.assertEqual(section["simpress"], {"abertas": 0, "concluidas": 0, "pendentes": 0, "encerradas_sem_data": 0})
        self.assertEqual(section["linhas"], [])
        self.assertEqual(section["dados_insuficientes"][0]["importacao"], "latest-incomplete")
        self.assertEqual(report, original)

    def test_empty_order_details_are_insufficient_instead_of_reusing_an_old_snapshot(self):
        report = {
            "snapshot_id": "empty-details", "atualizado_em": "2026-10-12T18:00:00",
            "historico": [{
                "id": "older-valid", "atualizado_em": "2026-10-01T18:00:00",
                "chamados": [ticket("SN-FIC-A", simpress={"pedidos_detalhados": [
                    detailed_order("P-100", "01/10/2026", "Em separação"),
                ]})],
            }],
            "chamados": [ticket("SN-FIC-A", simpress={"pedidos": ["P-100"], "pedidos_detalhados": []})],
        }
        section = tickets_section(report, OCTOBER)
        self.assertEqual(section["simpress"]["abertas"], 0)
        self.assertEqual(section["simpress"]["pendentes"], 0)
        self.assertEqual(section["dados_insuficientes"][0]["identificadores"], ["P-100"])

    def test_later_valid_import_can_count_an_order_again_after_insufficient_data(self):
        report = {
            "snapshot_id": "repaired", "atualizado_em": "2026-10-15T18:00:00",
            "historico": [{
                "id": "incomplete", "atualizado_em": "2026-10-12T18:00:00",
                "chamados": [ticket("SN-FIC-A", simpress={"pedidos": ["P-100"], "pedidos_detalhados": []})],
            }],
            "chamados": [ticket("SN-FIC-A", simpress={"pedidos_detalhados": [
                detailed_order("P-100", "01/10/2026", "Pedido entregue", "14/10/2026"),
            ]})],
        }
        section = tickets_section(report, OCTOBER)
        self.assertEqual(section["simpress"], {"abertas": 1, "concluidas": 1, "pendentes": 0, "encerradas_sem_data": 0})
        self.assertEqual([line["identificador"] for line in section["linhas"]], ["P-100"])


class WorkbookTests(unittest.TestCase):
    def workbook(self, report):
        output, filename = build_monthly_report_workbook(report)
        return openpyxl.load_workbook(BytesIO(output.getvalue())), filename

    def summary_values(self, workbook):
        return {row[0]: row[1] for row in workbook["Resumo"].iter_rows(values_only=True) if row and row[0]}

    def test_sheets_and_summary_values(self):
        workbook, filename = self.workbook(build())
        self.assertEqual(filename, "Relatorio-mensal-2026-10.xlsx")
        self.assertEqual(workbook.sheetnames, [
            "Resumo", "Inventário por tipo", "Inventário por unidade", "Comparação", "OS do mês", "OS sem dados",
        ])
        values = self.summary_values(workbook)
        self.assertEqual(values["Relatório mensal — outubro de 2026"], None)
        self.assertEqual(values["Equipamentos na base"], 4)
        self.assertEqual(values["Linhas sem número de série"], 2)
        self.assertEqual(values["Séries duplicadas"], 1)
        # The report shows the event in the server's local time zone, so CI (UTC) and Curitiba differ.
        local_time = datetime.fromisoformat("2026-10-15T09:00:00-03:00").astimezone()
        self.assertEqual(values["Data da comparação"], local_time.strftime("%d/%m/%Y %H:%M"))
        self.assertEqual(values[SECTION_ENDPOINT], 1)
        self.assertEqual(values[SECTION_SIMPRESS], 1)
        self.assertEqual(values[SECTION_LOCATION], 1)
        self.assertEqual(values["OS Oracle abertas no mês"], 1)
        self.assertEqual(values["OS Oracle pendentes no fim do mês"], 1)
        self.assertEqual(values["Pedidos Simpress entregues no mês"], 1)
        self.assertNotIn("Transferências geradas", values)
        self.assertTrue(any("16467" in str(cell.value) for cell in workbook["Resumo"]["A"]))

    def test_approved_source_labels_and_no_raw_source_names(self):
        workbook, _ = self.workbook(build())
        text = "\n".join(
            str(cell) for sheet in workbook for row in sheet.iter_rows(values_only=True) for cell in row if cell
        )
        self.assertIn("Sem correspondência — no endpoint", text)
        self.assertIn("Sem correspondência — na Simpress", text)
        self.assertNotIn("InvAll", text)
        self.assertNotIn("BRJ135P5XS", text)

    def test_detail_sheets_match_summary(self):
        workbook, _ = self.workbook(build())
        by_type = list(workbook["Inventário por tipo"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(sum(row[1] for row in by_type), 4)
        sections = [row[0] for row in workbook["Comparação"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sections, [SECTION_ENDPOINT, SECTION_SIMPRESS, SECTION_LOCATION])

    def test_spreadsheet_text_never_becomes_a_formula(self):
        hostile = [equipment("SN-FIC-900", "=HYPERLINK(\"x\")", "Curitiba", "+cmd")]
        workbook, _ = self.workbook(build(equipment=hostile))
        values = [cell for row in workbook["Inventário por tipo"].iter_rows(values_only=True) for cell in row]
        self.assertIn("'=HYPERLINK(\"x\")", values)

    def test_legacy_insufficient_rows_are_visible_in_workbook(self):
        legacy = {
            "snapshot_id": "legado-1", "atualizado_em": "2026-10-12T18:00:00",
            "chamados": [{"serie": "SN-FIC-LEGADO", "simpress": {
                "pedidos": ["P-LEG-1", "P-LEG-2"],
                "datas_pedido": ["02/10/2026", "10/10/2026"],
                "status": ["Em separação", "Pedido entregue"],
            }}],
        }
        workbook, _ = self.workbook(build(ticket_report=legacy))
        values = self.summary_values(workbook)
        self.assertEqual(values["Ocorrências sem dados por pedido"], 1)
        rows = list(workbook["OS sem dados"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(rows[0][1], "SN-FIC-LEGADO")
        self.assertEqual(rows[0][2], "P-LEG-1, P-LEG-2")

    def test_empty_month_still_generates_with_zeros(self):
        report = build(
            month=parse_report_month("2026-08", TODAY), comparison_history=[], ticket_report={},
        )
        rows = summary_rows(report)
        self.assertIn(("Última comparação", "Data da comparação", "Sem registros no mês",
                       "Atualização de inventário mais recente até o fim do mês."), rows)
        workbook, _ = self.workbook(report)
        self.assertEqual(workbook["OS do mês"]["A2"].value, "Sem registros no mês")


class MonthlyReportRouteTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.ticket_path = Path(self.tmp.name) / "historico_chamados.json"
        self.ticket_path.write_text(json.dumps(TICKET_REPORT), encoding="utf-8")
        self.equipment = list(EQUIPMENT)

    def tearDown(self):
        self.tmp.cleanup()

    def client(self):
        app = FastAPI()

        @app.exception_handler(DomainError)
        async def domain_error_handler(_request, error):
            return JSONResponse({"detail": error.detail}, status_code=error.status_code)

        app.include_router(create_monthly_report_router(MonthlyReportServices(
            inventory_snapshot=lambda: (self.equipment, dict(STATS)),
            load_comparison_history=lambda: HISTORY,
            comparison_history_for_display=comparison_history_for_display,
            ticket_report_path=lambda: self.ticket_path,
            today=lambda: TODAY,
            now=lambda: datetime(2026, 10, 17, 9, 30),
        )))
        return TestClient(app)

    def test_downloads_the_workbook_for_the_selected_month(self):
        response = self.client().get("/api/relatorio-mensal", params={"mes": "2026-10"})
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="Relatorio-mensal-2026-10.xlsx"', response.headers["content-disposition"])
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        self.assertEqual(workbook["Resumo"]["A1"].value, "Relatório mensal — outubro de 2026")

    def test_read_only_data_route_uses_the_same_monthly_report_builder(self):
        response = self.client().get("/api/relatorio-mensal/dados", params={"mes": "2026-10"})
        self.assertEqual(response.status_code, 200)
        report = response.json()
        self.assertEqual(report["mes"]["key"], "2026-10")
        self.assertEqual(report["inventario"]["total"], len(EQUIPMENT))
        self.assertEqual(report["comparacao"]["contagens"][SECTION_ENDPOINT], 1)
        self.assertEqual(report["os"]["oracle"]["concluidas"], 1)
        self.assertTrue(report["mes"]["parcial"])
        self.assertEqual(report["mes"]["data_corte"], TODAY.isoformat())
        self.assertIn("gerado_em", report)

    def test_invalid_or_future_month_is_rejected(self):
        client = self.client()
        for value in ("2026-11", "2026-13", "outubro"):
            with self.subTest(value=value):
                self.assertEqual(client.get("/api/relatorio-mensal", params={"mes": value}).status_code, 422)
        self.assertEqual(client.get("/api/relatorio-mensal").status_code, 422)

    def test_unloaded_base_answers_503(self):
        self.equipment = []
        response = self.client().get("/api/relatorio-mensal", params={"mes": "2026-10"})
        self.assertEqual(response.status_code, 503)
        self.assertIn("EquipamentosDetalhes", response.json()["detail"])

    def test_unreadable_ticket_history_answers_503(self):
        self.ticket_path.write_text("{", encoding="utf-8")
        self.assertEqual(self.client().get("/api/relatorio-mensal", params={"mes": "2026-10"}).status_code, 503)

    def test_missing_ticket_history_still_generates(self):
        self.ticket_path.unlink()
        response = self.client().get("/api/relatorio-mensal", params={"mes": "2026-10"})
        self.assertEqual(response.status_code, 200)


class AutomaticMonthlyReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.reports_dir = self.root / "data" / "relatorios_mensais"
        self.ticket_path = self.root / "historico_chamados.json"
        self.ticket_path.write_text(json.dumps(TICKET_REPORT), encoding="utf-8")
        self.equipment = list(EQUIPMENT)
        self.today = TODAY

    def tearDown(self):
        self.tmp.cleanup()

    def generate(self):
        services = MonthlyReportServices(
            inventory_snapshot=lambda: (self.equipment, dict(STATS)),
            load_comparison_history=lambda: HISTORY,
            comparison_history_for_display=comparison_history_for_display,
            ticket_report_path=lambda: self.ticket_path,
            today=lambda: self.today,
            now=lambda: datetime(2026, 10, 17, 9, 30),
        )
        return generate_previous_month_report(services, self.reports_dir, self.root)

    def test_previous_month_crosses_the_year(self):
        self.assertEqual(previous_report_month(date(2026, 10, 1)).key, "2026-09")
        self.assertEqual(previous_report_month(date(2027, 1, 31)).key, "2026-12")

    def test_saves_last_month_once_and_never_replaces_it(self):
        saved = self.generate()
        self.assertEqual(saved, self.reports_dir / "Relatorio-mensal-2026-09.xlsx")
        workbook = openpyxl.load_workbook(saved)
        self.assertEqual(workbook["Resumo"]["A1"].value, "Relatório mensal — setembro de 2026")
        content = saved.read_bytes()

        self.equipment = [EQUIPMENT[0]]
        self.assertIsNone(self.generate())
        self.assertEqual(saved.read_bytes(), content)
        self.assertEqual(sorted(path.name for path in self.reports_dir.glob("*.xlsx")), ["Relatorio-mensal-2026-09.xlsx"])
        self.assertFalse(any((self.root / ".transactions").iterdir()))

    def test_new_month_adds_its_own_file(self):
        self.generate()
        self.today = date(2026, 11, 1)
        self.assertEqual(self.generate().name, "Relatorio-mensal-2026-10.xlsx")
        self.assertEqual(len(list(self.reports_dir.glob("*.xlsx"))), 2)

    def test_waits_for_the_official_base_and_unreadable_tickets(self):
        self.equipment = []
        with self.assertLogs("routers.monthly_report", level="INFO"):
            self.assertIsNone(self.generate())
        self.assertFalse(self.reports_dir.exists())

        self.equipment = list(EQUIPMENT)
        self.ticket_path.write_text("{", encoding="utf-8")
        self.assertIsNone(self.generate())
        self.assertFalse(self.reports_dir.exists())


if __name__ == "__main__":
    unittest.main()
