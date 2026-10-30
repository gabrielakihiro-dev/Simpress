import json
import tempfile
import unittest
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from defusedxml import ElementTree
from fastapi import FastAPI
from fastapi.testclient import TestClient
import openpyxl

from domains.assistant_api import SimpressApi
from domains.assistant_openrouter import ModelReply, ModelToolCall, OpenRouterSettings
from domains.assistant_relatorios import FORMATOS, MAX_CELULA, PLANILHAS, TIPOS, GeradorRelatorios, RelatorioRecusado, _celula
from domains.assistant_tools import TOOLS
from routers.assistant import AssistantRouteServices, create_assistant_router

SEBRAE_SERIE = "BRJ149SCJG"
CPF = "529.982.247-25"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _os(abertas: int) -> Dict[str, Any]:
    return {"abertas": abertas, "concluidas": 1, "pendentes": 2, "encerradas_sem_data": 0}


class FakeTransport:
    """Answers the routes the reports read; records every call."""

    def __init__(self) -> None:
        self.calls: List[Tuple[str, str]] = []
        self.comparacao: Dict[str, Any] = {
            "divergencias_local": [
                {"serie": "AAA111", "computador": "PRNOTE01", "local_previsto": "Curitiba", "local_detalhes": "=HYPERLINK(\"http://x\")"},
                {"serie": "BBB222", "computador": "PRNOTE02", "local_previsto": "Maringá", "local_detalhes": f"Sala, CPF {CPF}"},
            ],
            "ausentes_no_endpoint": [{"serie": SEBRAE_SERIE, "computador": "PRNOTESB"}, {"serie": "CCC333", "computador": "PRNOTE03"}],
            "series_sem_mapeamento": [{"serie": "DDD444"}],
        }

    def __call__(self, method: str, path: str, params: Optional[Dict[str, Any]], body: Optional[Dict[str, Any]]) -> Tuple[int, Any]:
        self.calls.append((method, path))
        if path == "/api/relatorio-mensal/dados":
            return 200, {
                "mes": {"label": "Setembro de 2026", "key": params["mes"] if params else ""},
                "inventario": {"total": 10, "unidades": 3, "linhas_sem_serie": 0, "series_duplicadas": [], "arquivo_origem": "base.xlsx"},
                "comparacao": {
                    "horario": "2026-09-30T10:00:00", "atualizacoes_no_mes": 2,
                    "contagens": {"Sem correspondência — no endpoint": 1, "Sem correspondência — na Simpress": 2, "Localidade divergente": 3},
                },
                "os": {"oracle": _os(4), "simpress": _os(5), "dados_insuficientes": []},
            }
        if path == "/api/stats":
            return 200, {"total_equipamentos": 10, "total_unidades": 3, "base_importada": True, "arquivo_origem": "base.xlsx",
                         "ultima_atualizacao": "2026-10-01", "total_alteracoes_exportaveis": 0}
        if path == "/api/historico-comparacoes":
            return 200, {"eventos": [{"id": "e1", "horario": "2026-10-20T10:00:00", "resumo": {}}]}
        if path == "/api/historico-comparacoes/e1":
            return 200, {"evento": {"id": "e1", "comparacao": self.comparacao, "resumo": {}}}
        if path == "/api/sebrae/novos":
            return 200, {"importado": {"arquivo": "dash.xlsx", "total": 1, "abertos": [
                {"numero": "123", "titulo": "-Impressora parada", "valores": {"A": "Maria Souza", "B": "Suporte", "D": "EM ABERTO", "J": "Incidente"}},
                {"numero": "124", "titulo": "Teclado", "valores": {"B": "Suporte", "D": "", "J": "Requisição"}},
            ]}}
        if path == "/api/operacoes/pendencias":
            return 200, {"evento": {"id": "e1", "horario": "2026-10-20T10:00:00"}, "pendencias": [
                {"categoria": "divergencias_local", "secao": "Localidade divergente", "item": {"serie": "AAA111", "computador": "PRNOTE01", "motivo": "Local"}},
                {"categoria": "ausentes_no_endpoint", "secao": "Sem correspondência", "item": {"serie": SEBRAE_SERIE, "computador": "X"}},
            ]}
        return 404, {"detail": "não existe"}


class Relogio:
    def __init__(self) -> None:
        self.agora = 1000.0

    def __call__(self) -> float:
        return self.agora


def gerador(transporte: Optional[FakeTransport] = None, **extra: Any) -> GeradorRelatorios:
    return GeradorRelatorios(SimpressApi(transporte or FakeTransport()), agora=lambda: datetime(2026, 10, 29, 9, 30), **extra)


def textos_xlsx(dados: bytes) -> List[List[Any]]:
    celulas: List[List[Any]] = []
    for aba in openpyxl.load_workbook(BytesIO(dados)).worksheets:
        celulas += [list(linha) for linha in aba.iter_rows(values_only=True)]
    return celulas


def textos_docx(dados: bytes) -> List[str]:
    with zipfile.ZipFile(BytesIO(dados)) as arquivo:
        assert arquivo.testzip() is None
        raiz = ElementTree.fromstring(arquivo.read("word/document.xml"))
        ElementTree.fromstring(arquivo.read("[Content_Types].xml"))
        ElementTree.fromstring(arquivo.read("_rels/.rels"))
    return ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in raiz.iter(f"{W}p")]


def todo_o_texto(g: GeradorRelatorios, resultado: Dict[str, Any]) -> str:
    arquivo = g.obter(resultado["id"])
    assert arquivo is not None
    if arquivo.formato == "xlsx":
        return "\n".join(str(c) for linha in textos_xlsx(arquivo.dados) for c in linha)
    return "\n".join(textos_docx(arquivo.dados))


class GenerationTests(unittest.TestCase):
    def test_every_type_and_format_makes_a_valid_file(self) -> None:
        g = gerador()
        for tipo in TIPOS:
            for formato in FORMATOS:
                resultado = g.gerar(tipo, formato, "2026-09")
                with self.subTest(tipo=tipo, formato=formato):
                    arquivo = g.obter(resultado["id"])
                    assert arquivo is not None
                    self.assertEqual(resultado["tamanho"], len(arquivo.dados))
                    self.assertTrue(arquivo.nome.endswith("." + formato))
                    self.assertEqual(arquivo.dados[:2], b"PK")
                    self.assertTrue(todo_o_texto(g, resultado).strip())

    def test_xlsx_content_comes_from_the_routes(self) -> None:
        g = gerador()
        resumo = todo_o_texto(g, g.gerar("resumo_mensal", "xlsx", "2026-09"))
        self.assertIn("Equipamentos na base", resumo)
        self.assertIn("Sem correspondência — na Simpress", resumo)
        diverg = todo_o_texto(g, g.gerar("divergencias_localidade", "docx"))
        self.assertIn("AAA111", diverg)
        self.assertIn("Local na Simpress", diverg)
        self.assertIn("Comparação de 2026-10-20T10:00:00", diverg)

    def test_formula_injection_is_neutralised_in_both_formats(self) -> None:
        g = gerador()
        for formato in FORMATOS:
            texto = todo_o_texto(g, g.gerar("divergencias_localidade", formato))
            self.assertIn("'=HYPERLINK", texto)
            self.assertNotIn("\n=HYPERLINK", "\n" + texto)
        chamados = textos_xlsx(g.obter(g.gerar("chamados_sebrae_abertos", "xlsx")["id"]).dados)  # type: ignore[union-attr]
        self.assertIn("'-Impressora parada", [c for linha in chamados for c in linha])
        for inicio in "=+-@":
            self.assertEqual(_celula(f"{inicio}1+1"), f"'{inicio}1+1")
        self.assertEqual(_celula(-5), -5)
        self.assertEqual(len(_celula("x" * (MAX_CELULA + 50))), MAX_CELULA)

    def test_cpf_is_masked_and_control_characters_do_not_break_the_file(self) -> None:
        transporte = FakeTransport()
        transporte.comparacao["divergencias_local"][0]["computador"] = "PC\x00\x01 senha: segredo99"
        g = gerador(transporte)
        for formato in FORMATOS:
            texto = todo_o_texto(g, g.gerar("divergencias_localidade", formato))
            self.assertIn("***.***.***-**", texto)
            self.assertNotIn(CPF, texto)
            self.assertNotIn("segredo99", texto)

    def test_sebrae_series_and_responsible_names_stay_out(self) -> None:
        g = gerador()
        for tipo in ("ultima_comparacao", "pendencias_operacionais", "chamados_sebrae_abertos"):
            texto = todo_o_texto(g, g.gerar(tipo, "xlsx"))
            self.assertNotIn(SEBRAE_SERIE, texto, tipo)
            self.assertNotIn("Maria Souza", texto, tipo)
        self.assertIn("CCC333", todo_o_texto(g, g.gerar("ultima_comparacao", "xlsx")))

    def test_labels_follow_agents_md(self) -> None:
        transporte = FakeTransport()
        transporte.comparacao["apelidos_atualizados"] = [{"serie": "EEE555", "apelido_equipamentos_detalhes": "A", "apelido_invall": "B"}]
        g = gerador(transporte)
        texto = todo_o_texto(g, g.gerar("ultima_comparacao", "docx"))
        self.assertIn("Apelido na Simpress", texto)
        self.assertIn("Apelido no Endpoint", texto)
        self.assertNotIn("invall", texto.lower())
        self.assertNotIn("detalhes", texto.lower())

    def test_enum_and_month_are_validated(self) -> None:
        g = gerador()
        for tipo, formato, mes in [
            ("inventado", "xlsx", None), (None, "xlsx", None), ("resumo_mensal", "pdf", "2026-09"), ("resumo_mensal", None, "2026-09"),
            (["resumo_mensal"], "xlsx", None), ("resumo_mensal", "xlsx", None), ("resumo_mensal", "xlsx", "2026-13"),
            ("resumo_mensal", "xlsx", "../../x"), ("totais_inventario", "xlsx", "setembro"),
        ]:
            with self.subTest(tipo=tipo, formato=formato, mes=mes), self.assertRaises(RelatorioRecusado):
                g.gerar(tipo, formato, mes)

    def test_failed_route_gives_a_clear_error_and_stores_nothing(self) -> None:
        g = GeradorRelatorios(SimpressApi(lambda *a: (503, {"detail": "base fora"})))
        with self.assertRaises(RelatorioRecusado) as erro:
            g.gerar("totais_inventario", "xlsx")
        self.assertIn("503", str(erro.exception))

    def test_file_name_is_safe_ascii(self) -> None:
        g = gerador()
        for tipo in TIPOS:
            nome = g.obter(g.gerar(tipo, "docx", "2026-09")["id"]).nome  # type: ignore[union-attr]
            self.assertRegex(nome, r"^[A-Za-z0-9._-]+\.docx$")
        self.assertEqual(g.obter(g.gerar("resumo_mensal", "xlsx", "2026-09")["id"]).nome, "simpress-resumo-mensal-2026-09-20261029.xlsx")  # type: ignore[union-attr]


class StoreTests(unittest.TestCase):
    def test_ttl_expires_the_file(self) -> None:
        relogio = Relogio()
        g = gerador(ttl=60, relogio=relogio)
        id_ = g.gerar("totais_inventario", "xlsx")["id"]
        relogio.agora += 59
        self.assertIsNotNone(g.obter(id_))
        relogio.agora += 2
        self.assertIsNone(g.obter(id_))

    def test_count_limit_drops_the_oldest(self) -> None:
        g = gerador(max_relatorios=3)
        ids = [g.gerar("totais_inventario", "xlsx")["id"] for _ in range(5)]
        self.assertEqual([g.obter(i) is not None for i in ids], [False, False, True, True, True])

    def test_unknown_ids_and_clear(self) -> None:
        g = gerador()
        id_ = g.gerar("totais_inventario", "docx")["id"]
        self.assertIsNone(g.obter("0123456789abcdef"))
        g.limpar()
        self.assertIsNone(g.obter(id_))

    def test_ids_are_not_guessable(self) -> None:
        g = gerador()
        ids = {g.gerar("totais_inventario", "xlsx")["id"] for _ in range(4)}
        self.assertEqual(len(ids), 4)
        self.assertTrue(all(len(i) == 16 and int(i, 16) >= 0 for i in ids))


class ToolTests(unittest.TestCase):
    def test_schema_is_a_closed_enum_and_a_query_tool(self) -> None:
        tool = TOOLS["gerar_relatorio"]
        self.assertFalse(tool.changes_data)
        self.assertTrue(tool.usa_relatorios)
        self.assertEqual(tool.properties["tipo"]["enum"], list(TIPOS))
        self.assertEqual(tool.properties["formato"]["enum"], ["xlsx", "docx"])
        self.assertEqual(tool.required, ["tipo", "formato"])

    def test_tool_rejects_values_outside_the_enum(self) -> None:
        run = TOOLS["gerar_relatorio"].run
        for args in ({"tipo": "tudo", "formato": "xlsx"}, {"tipo": "totais_inventario", "formato": "pdf"}, {}):
            self.assertIn("erro", run(gerador(), args))

    def test_tool_answer_has_no_url_or_content(self) -> None:
        resultado = TOOLS["gerar_relatorio"].run(gerador(), {"tipo": "totais_inventario", "formato": "xlsx"})
        self.assertEqual(set(resultado), {"id", "nome_arquivo", "formato", "tamanho_bytes", "mensagem"})


class FakeModel:
    def __init__(self, *replies: ModelReply) -> None:
        self.replies = list(replies)
        self.sent: List[List[Dict[str, Any]]] = []

    def transcribe(self, wav: bytes) -> str:
        return ""

    def complete(self, messages: List[Dict[str, Any]], tools: List[Dict[str, Any]]) -> ModelReply:
        self.sent.append(json.loads(json.dumps(messages)))
        return self.replies.pop(0)


def pedido(**args: Any) -> ModelReply:
    return ModelReply(tool_calls=[ModelToolCall(id="id-1", name="gerar_relatorio", arguments=json.dumps(args))])


class ConversationAndRouteTests(unittest.TestCase):
    def make_client(self, model: FakeModel) -> TestClient:
        arquivo = Path(self._tmp.name) / "instrucoes.json"
        app = FastAPI()
        app.include_router(create_assistant_router(AssistantRouteServices(
            load_settings=lambda: OpenRouterSettings(api_key="chave-de-teste"), transport=FakeTransport(),
            client_factory=lambda settings: model,  # type: ignore[arg-type,return-value]
            instructions_file=arquivo,
        )))
        return TestClient(app)

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)

    def test_reply_carries_the_file_and_the_model_never_sees_the_id(self) -> None:
        model = FakeModel(pedido(tipo="divergencias_localidade", formato="docx"), ModelReply(content="Pronto, o arquivo está no botão abaixo."))
        client = self.make_client(model)
        resposta = client.post("/api/assistente/mensagem", json={"texto": "exporta as divergências em Word"})
        self.assertEqual(resposta.status_code, 200)
        corpo = resposta.json()
        self.assertEqual(corpo["ferramentas"], ["gerar_relatorio"])
        arquivo = corpo["arquivos"][0]
        self.assertEqual(set(arquivo), {"id", "nome", "formato", "tamanho"})
        self.assertNotIn(arquivo["id"], json.dumps(model.sent))  # the tool result sent back to the model has no id
        download = client.get(f"/api/assistente/relatorio/{arquivo['id']}")
        self.assertEqual(download.status_code, 200)
        self.assertEqual(download.content[:2], b"PK")
        self.assertEqual(len(download.content), arquivo["tamanho"])

    def test_download_headers(self) -> None:
        model = FakeModel(pedido(tipo="totais_inventario", formato="xlsx"), ModelReply(content="ok"))
        client = self.make_client(model)
        arquivo = client.post("/api/assistente/mensagem", json={"texto": "totais em Excel"}).json()["arquivos"][0]
        download = client.get(f"/api/assistente/relatorio/{arquivo['id']}")
        self.assertEqual(download.headers["content-type"], FORMATOS["xlsx"])
        self.assertEqual(download.headers["content-disposition"], f'attachment; filename="{arquivo["nome"]}"')
        self.assertEqual(download.headers["x-content-type-options"], "nosniff")
        self.assertEqual(download.headers["cache-control"], "no-store")
        self.assertNotIn("/", arquivo["nome"])

    def test_unknown_id_is_404_and_reply_without_report_has_no_files_key(self) -> None:
        model = FakeModel(ModelReply(content="oi"))
        client = self.make_client(model)
        self.assertEqual(client.get("/api/assistente/relatorio/0123456789abcdef").status_code, 404)
        self.assertEqual(client.get("/api/assistente/relatorio/..%2F..%2Fsegredo").status_code, 404)
        self.assertNotIn("arquivos", client.post("/api/assistente/mensagem", json={"texto": "oi"}).json())

    def test_failed_request_gives_no_button_and_new_conversation_drops_the_files(self) -> None:
        model = FakeModel(pedido(tipo="resumo_mensal", formato="xlsx"), ModelReply(content="Preciso do mês."),
                          pedido(tipo="totais_inventario", formato="docx"), ModelReply(content="Pronto."))
        client = self.make_client(model)
        primeira = client.post("/api/assistente/mensagem", json={"texto": "resumo mensal em Excel"}).json()
        self.assertNotIn("arquivos", primeira)  # the tool answered with an error: nothing was generated
        arquivo = client.post("/api/assistente/mensagem", json={"texto": "totais em Word"}).json()["arquivos"][0]
        self.assertEqual(client.post("/api/assistente/limpar").status_code, 200)
        self.assertEqual(client.get(f"/api/assistente/relatorio/{arquivo['id']}").status_code, 404)


if __name__ == "__main__":
    unittest.main()


class FakePlanilhas(FakeTransport):
    """Adds the four export routes; answers real .xlsx bytes like the app does."""

    def __call__(self, method: str, path: str, params: Optional[Dict[str, Any]], body: Optional[Dict[str, Any]]) -> Tuple[int, Any]:
        if path in ("/api/exportar-equipamentos", "/api/operacoes/chamados/exportar", "/api/usuarios-entra/exportar") or (
            path == "/api/historico-comparacoes/exportar/excel" and params == {"data": "2026-10-20"}
        ):
            self.calls.append((method, path))
            livro = openpyxl.Workbook()
            livro.active.append(["Série", "Obs"])
            livro.active.append(["AAA111", f"CPF {CPF}"])
            saida = BytesIO()
            livro.save(saida)
            return 200, saida.getvalue()
        if path == "/api/historico-comparacoes/exportar/excel":
            return 404, {"detail": "Não há registros de planilhas na data selecionada."}
        return super().__call__(method, path, params, body)


class SpreadsheetExportTests(unittest.TestCase):
    def test_each_sheet_is_kept_in_memory_and_model_gets_only_the_fixed_text(self) -> None:
        g = gerador(FakePlanilhas())
        for planilha in PLANILHAS:
            with self.subTest(planilha=planilha):
                resultado = TOOLS["exportar_planilha"].run(g, {"planilha": planilha, "data": "2026-10-20"})
                self.assertEqual(resultado["mensagem"], "Planilha Exportada com Sucesso!")
                self.assertEqual(set(resultado), {"id", "mensagem"})
                arquivo = g.obter(resultado["id"])
                assert arquivo is not None
                self.assertTrue(arquivo.nome.endswith(".xlsx"))
                self.assertEqual(textos_xlsx(arquivo.dados)[0], ["Série", "Obs"])

    def test_tool_is_a_query_and_refuses_bad_input(self) -> None:
        self.assertFalse(TOOLS["exportar_planilha"].changes_data)
        g = gerador(FakePlanilhas())
        for args in ({"planilha": "x"}, {"planilha": "historico_comparacoes"}, {"planilha": "historico_comparacoes", "data": "2026-13-01"}):
            with self.subTest(args=args):
                self.assertIn("erro", TOOLS["exportar_planilha"].run(g, args))

    def test_route_error_is_reported_without_file(self) -> None:
        resultado = TOOLS["exportar_planilha"].run(gerador(FakePlanilhas()), {"planilha": "historico_comparacoes", "data": "2026-10-21"})
        self.assertIn("Não há registros", resultado["erro"])
