import asyncio
import ast
import io
import json
import re
import shutil
import tempfile
import threading
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from core import config as app_config
import openpyxl
from core import openpyxl_compat
from fastapi import HTTPException
from starlette.routing import Match


with patch(
    "openpyxl.load_workbook",
    side_effect=RuntimeError("testes não acessam planilhas locais"),
):
    with redirect_stdout(io.StringIO()):
        import main


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _css_layer(selector: str) -> int:
    """z-index of a selector in static/app.css, resolved through the --z-* scale in tokens.css."""
    css = (app_config.STATIC_DIR / "app.css").read_text(encoding="utf-8")
    tokens = (app_config.STATIC_DIR / "styles" / "tokens.css").read_text(encoding="utf-8")
    scale = {name: int(value) for name, value in re.findall(r"(--z-[\w-]+):\s*(\d+);", tokens)}
    match = re.search(re.escape(selector) + r"\s*\{[^}]*?z-index:\s*var\((--z-[\w-]+)\);", css)
    if match is None:
        raise AssertionError(f"{selector} has no z-index from the --z-* scale")
    return scale[match.group(1)]


def _xlsx_with_worksheet_xml(*worksheet_xml: str) -> bytes:
    content = io.BytesIO()
    with zipfile.ZipFile(content, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(
            "xl/workbook.xml",
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
            + "".join(
                f'<sheet name="Sheet{index}" sheetId="{index}" r:id="rId{index}"/>'
                for index in range(1, len(worksheet_xml) + 1)
            )
            + "</sheets></workbook>",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + "".join(
                f'<Relationship Id="rId{index}" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
                f'Target="worksheets/sheet{index}.xml"/>'
                for index in range(1, len(worksheet_xml) + 1)
            )
            + "</Relationships>",
        )
        for index, xml in enumerate(worksheet_xml, start=1):
            archive.writestr(f"xl/worksheets/sheet{index}.xml", xml)
    return content.getvalue()


def _run_local_safety_middleware(
    host_values,
    origin=None,
    path="/api/stats",
    *,
    method="GET",
    body=b"",
    include_content_length=True,
    content_length=None,
):
    state = {"passed": False, "messages": []}

    async def app(scope, receive, send):
        state["passed"] = True
        await send({"type": "http.response.start", "status": 200, "headers": []})
        await send({"type": "http.response.body", "body": b"ok"})

    async def receive():
        return {"type": "http.request", "body": body, "more_body": False}

    async def send(message):
        state["messages"].append(message)

    headers = [(b"host", host.encode("ascii")) for host in host_values]
    if include_content_length:
        headers.append((b"content-length", str(len(body) if content_length is None else content_length).encode("ascii")))
    if origin is not None:
        headers.append((b"origin", origin.encode("ascii")))
    scope = {"type": "http", "path": path, "method": method, "headers": headers,
             "client": ("127.0.0.1", 12345)}
    middleware = main.LocalRequestSafetyMiddleware(app)
    asyncio.run(middleware(scope, receive, send))
    status = next(
        message["status"]
        for message in state["messages"]
        if message["type"] == "http.response.start"
    )
    return status, state["passed"]


def _frontend_script() -> str:
    """static/app.js and the feature modules it bootstraps, in index.html order."""
    html = (app_config.STATIC_DIR / "index.html").read_text(encoding="utf-8")
    files = re.findall(r'<script src="/static/((?:js/features/[^"]+|js/shared/history-format\.js|app\.js))"></script>', html)
    return "\n".join((app_config.STATIC_DIR / name).read_text(encoding="utf-8") for name in files)


class MaintenanceContractTests(unittest.TestCase):
    def test_both_inventory_upload_routes_share_the_request_size_limit(self):
        for path in main.INVENTORY_UPLOAD_PATHS:
            with self.subTest(path=path):
                status, passed = _run_local_safety_middleware(
                    ["127.0.0.1:8000"],
                    path=path,
                    content_length=main.MAX_UPLOAD_REQUEST_BYTES + 1,
                )
                self.assertEqual(status, 413)
                self.assertFalse(passed)

    def test_application_paths_remain_project_relative_and_reexported(self):
        self.assertEqual(app_config.BASE_DIR, PROJECT_ROOT.resolve())
        self.assertEqual(main.BASE_DIR, app_config.BASE_DIR)

        for name in (
            "STATIC_DIR",
            "DATA_DIR",
            "STATE_DIR",
            "INPUT_DIR",
            "TEMPLATE_FILE",
            "DETAILS_UPLOAD_FILE",
            "ENDPOINT_UPLOAD_FILE",
            "EXPORT_MODEL_FILE",
            "UPLOAD_BACKUPS_DIR",
            "CUSTOM_UNIDADES_FILE",
            "MODIFIED_APELIDOS_FILE",
            "ALIAS_HISTORY_FILE",
            "COMPARISON_HISTORY_FILE",
            "HISTORY_ANNOTATIONS_DIR",
            "HISTORY_ANNOTATIONS_SYNC_ENABLED",
            "LOCAL_HISTORY_ANNOTATION_OUTBOX_DIR",
            "TRANSACTION_ROOT",
        ):
            with self.subTest(path=name):
                self.assertEqual(getattr(main, name), getattr(app_config, name))

    def test_history_annotation_directory_prefers_override_then_authenticated_onedrive(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            override = root / "custom-shared-folder"
            onedrive = root / "OneDrive"
            onedrive.mkdir()

            resolved, sync_enabled = app_config.resolve_history_annotations_dir(
                root, {"SIMPRESS_ANNOTATIONS_DIR": str(override), "OneDriveCommercial": str(onedrive)}
            )
            self.assertEqual(resolved, override)
            self.assertTrue(sync_enabled)

            resolved, sync_enabled = app_config.resolve_history_annotations_dir(
                root, {"OneDriveCommercial": str(onedrive)}
            )
            self.assertEqual(resolved, onedrive / "Aplicativos" / "Simpress" / "AnotacoesHistorico")
            self.assertTrue(sync_enabled)

            resolved, sync_enabled = app_config.resolve_history_annotations_dir(
                root, {"OneDriveCommercial": str(root / "nao-existe")}
            )
            self.assertEqual(resolved, root / "anotacoes_historico")
            self.assertFalse(sync_enabled)

        vscode_settings = json.loads((PROJECT_ROOT / ".vscode" / "settings.json").read_text(encoding="utf-8"))
        self.assertEqual(
            vscode_settings["python.defaultInterpreterPath"],
            "${workspaceFolder}/.venv/Scripts/python.exe",
        )

    def test_openapi_keeps_existing_route_and_method_contracts(self):
        expected = {
            "/": {"get"},
            "/api/equipamento/{serial}": {"get"},
            "/api/equipamento/salvar": {"post"},
            "/api/equipamento/atualizar-apelido": {"post"},
            "/api/historico-alteracoes": {"get"},
            "/api/historico-apelidos": {"get"},
            "/api/historico-comparacoes": {"get"},
            "/api/historico-comparacoes/{evento_id}": {"get", "delete"},
            "/api/historico-comparacoes/{evento_id}/anotacoes/{item_id}": {"put"},
            "/api/operacoes/pendencias": {"get"},
            "/api/operacoes/evolucao": {"get"},
            "/api/operacoes/chamados": {"get"},
            "/api/operacoes/chamados/historico": {"get"},
            "/api/operacoes/chamados/historico/antigas": {"delete"},
            "/api/operacoes/chamados/snapshot/{snapshot_id}": {"get"},
            "/api/operacoes/chamados/evolucao": {"get"},
            "/api/operacoes/chamados/exportar": {"get"},
            "/api/operacoes/chamados/importar": {"post"},
            "/api/historico-alteracoes/{evento_id}": {"delete"},
            "/api/historico-apelidos/{evento_id}": {"delete"},
            "/api/equipamento/apelido/desfazer": {"post"},
            "/api/equipamento/apelidos-modificados": {"get"},
            "/api/exportar-equipamentos": {"get"},
            "/api/recarregar-base": {"post"},
            "/api/planilhas/previsualizar": {"post"},
            "/api/planilhas/atualizar": {"post"},
            "/api/unidade/salvar": {"post"},
            "/api/unidades": {"get"},
            "/api/buscar": {"get"},
            "/api/stats": {"get"},
            "/api/prontidao": {"get"},
            "/api/gerar-planilha": {"post"},
        }

        paths = main.app.openapi()["paths"]
        for path, methods in expected.items():
            with self.subTest(path=path):
                self.assertIn(path, paths)
                self.assertTrue(methods.issubset(paths[path]))
        upload = paths["/api/planilhas/atualizar"]["post"]
        body_ref = upload["requestBody"]["content"]["multipart/form-data"]["schema"]["$ref"]
        schema_name = body_ref.rsplit("/", 1)[-1]
        self.assertIn("preview_token", main.app.openapi()["components"]["schemas"][schema_name]["properties"])

    def test_fixed_modified_alias_route_is_dispatched_before_serial_parameter(self):
        scope = {
            "type": "http",
            "method": "GET",
            "path": "/api/equipamento/apelidos-modificados",
            "root_path": "",
            "headers": [],
        }
        first_match = None
        for route in main.app.router.routes:
            match, _ = route.matches(scope)
            if match is Match.FULL:
                first_match = route
                break

        self.assertIsNotNone(first_match)
        self.assertIs(first_match.endpoint, main.get_apelidos_modificados)
        with patch.object(main, "load_modified_apelidos", return_value={"SN-FICTICIO": "HOST-FICTICIO"}):
            self.assertEqual(first_match.endpoint(), {
                "total": 1,
                "apelidos": {"SN-FICTICIO": "HOST-FICTICIO"},
            })

    def test_toasts_render_above_modal_overlays(self):
        html = (app_config.STATIC_DIR / "index.html").read_text(encoding="utf-8")

        self.assertNotRegex(html, r'id="toast-container"[^>]*\bz-')
        self.assertGreater(_css_layer("#toast-container"), _css_layer(".modal-overlay"))

    def test_excel_exports_show_only_success_confirmation(self):
        script = _frontend_script()
        self.assertEqual(
            script.count('showToast("Planilha Exportada com Sucesso!", "success");'),
            2,
        )

    def test_tailwind_build_only_scans_fixed_local_assets(self):
        configuration = (PROJECT_ROOT / "config/tailwind.config.cjs").read_text(encoding="utf-8")
        self.assertIn('content: ["../static/index.html", "../static/app.js", "../static/js/**/*.js"]', configuration)
        self.assertNotIn("process.env", configuration)
        self.assertNotIn("userInput", configuration)

    def test_private_openpyxl_access_is_version_isolated(self):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "fixture"
        self.assertEqual([cell.value for cell in openpyxl_compat.stored_cells(worksheet)], ["fixture"])
        self.assertGreaterEqual(len(openpyxl_compat.style_catalog(workbook)), 1)
        self.assertGreaterEqual(len(openpyxl_compat.named_style_catalog(workbook)), 1)
        self.assertEqual(openpyxl_compat.chart_count(worksheet), 0)
        self.assertEqual(openpyxl_compat.image_count(worksheet), 0)
        workbook.close()

        integrity_source = (PROJECT_ROOT / "core/integrity.py").read_text(encoding="utf-8")
        for internal_name in ("_cells", "_charts", "_images", "_cell_styles", "_named_styles", "_cf_rules"):
            self.assertNotRegex(integrity_source, rf"(?<!\w){re.escape(internal_name)}\b")
        with patch.object(openpyxl_compat.openpyxl, "__version__", "99.0"):
            with self.assertRaises(RuntimeError):
                openpyxl_compat.stored_cells(worksheet)

    def test_importing_main_has_no_operational_startup_calls(self):
        module_ast = ast.parse((PROJECT_ROOT / "main.py").read_text(encoding="utf-8"))
        startup_calls = {"_recover_pending_transactions_at_startup", "load_database"}
        module_level_calls = {
            node.value.func.id
            for node in module_ast.body
            if isinstance(node, ast.Expr)
            and isinstance(node.value, ast.Call)
            and isinstance(node.value.func, ast.Name)
        }
        self.assertTrue((startup_calls | {"recover_pending_transactions"}).isdisjoint(module_level_calls))
        lifespan = next(node for node in module_ast.body if isinstance(node, ast.AsyncFunctionDef) and node.name == "lifespan")
        lifespan_calls = {
            node.func.id
            for node in ast.walk(lifespan)
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)
        }
        self.assertTrue(startup_calls.issubset(lifespan_calls))
        self.assertIn("_report_startup_readiness_best_effort", lifespan_calls)

    def test_history_router_exports_are_explicit_and_ci_covers_all_test_suites(self):
        main_source = (PROJECT_ROOT / "main.py").read_text(encoding="utf-8")
        self.assertNotIn("globals().update", main_source)
        workflow = (PROJECT_ROOT / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8")
        self.assertIn("node scripts/verify.cjs", workflow)
        verify_script = (PROJECT_ROOT / "scripts" / "verify.cjs").read_text(encoding="utf-8")
        self.assertIn("tests/run_tests.py", verify_script)
        self.assertIn('listFiles("tests/frontend", ".cjs", { prefix: "test_" })', verify_script)
        self.assertIn("scripts/check_secrets.py", verify_script)

    def test_transfer_contact_is_configured_and_workbook_metadata_is_anonymized(self):
        defaults = main.get_transfer_form_defaults()
        self.assertEqual(defaults["retirada_contato"], app_config.TRANSFER_ORIGIN_CONTACT)
        self.assertEqual(defaults["retirada_email"], app_config.TRANSFER_ORIGIN_EMAIL)
        self.assertEqual(defaults["retirada_telefone"], app_config.TRANSFER_ORIGIN_PHONE)
        html = (app_config.STATIC_DIR / "index.html").read_text(encoding="utf-8")
        script = _frontend_script()
        self.assertIn("/api/config/transferencia", script)
        self.assertNotIn("gansouza@simpress.com.br", html + script)

        template = app_config.TEMPLATE_FILE
        with zipfile.ZipFile(template) as archive:
            package_xml = "\n".join(
                archive.read(name).decode("utf-8", "replace")
                for name in archive.namelist()
                if name.endswith((".xml", ".vml"))
            ).lower()
        for personal_value in (
            "gabriel akihiro noda souza",
            "gansouza@simpress.com.br",
            "paula caroline evangelista",
            "pcevangelista@simpress.com.br",
            "rayra ibere vieira bicas",
        ):
            self.assertNotIn(personal_value, package_xml)
        self.assertIn("instruções do formulário", package_xml)

    def test_real_asgi_request_dispatches_the_fixed_modified_alias_route(self):
        messages = []
        sent_request = False

        async def receive():
            nonlocal sent_request
            if sent_request:
                return {"type": "http.disconnect"}
            sent_request = True
            return {"type": "http.request", "body": b"", "more_body": False}

        async def send(message):
            messages.append(message)

        scope = {
            "type": "http",
            "asgi": {"version": "3.0"},
            "http_version": "1.1",
            "method": "GET",
            "scheme": "http",
            "path": "/api/equipamento/apelidos-modificados",
            "raw_path": b"/api/equipamento/apelidos-modificados",
            "query_string": b"",
            "root_path": "",
            "headers": [(b"host", b"localhost:8000")],
            "server": ("localhost", 8000),
            "client": ("127.0.0.1", 12345),
        }
        with patch.object(main, "load_modified_apelidos", return_value={"SN-ASGI": "HOST-ASGI"}):
            asyncio.run(main.app(scope, receive, send))

        response_start = next(item for item in messages if item["type"] == "http.response.start")
        response_body = b"".join(item.get("body", b"") for item in messages if item["type"] == "http.response.body")
        self.assertEqual(response_start["status"], 200)
        self.assertEqual(json.loads(response_body), {"total": 1, "apelidos": {"SN-ASGI": "HOST-ASGI"}})

    def test_api_body_limit_and_inventory_admission_reject_overload(self):
        oversized = b"x" * (app_config.MAX_API_REQUEST_BYTES + 1)
        status, reached_app = _run_local_safety_middleware(
            ["localhost:8000"], path="/api/unidade/salvar", method="POST", body=oversized,
        )
        self.assertEqual(status, 413)
        self.assertFalse(reached_app)

        class BusyAdmission:
            def acquire(self, blocking=False):
                return False

            def release(self):
                raise AssertionError("A admissão recusada não deve ser liberada")

        with patch.object(main, "UPLOAD_ADMISSION", BusyAdmission()):
            status, reached_app = _run_local_safety_middleware(
                ["localhost:8000"], path="/api/recarregar-base", method="POST",
            )
        self.assertEqual(status, 503)
        self.assertFalse(reached_app)

    def test_gitignore_covers_environment_variants_and_private_key_files(self):
        ignore_rules = (PROJECT_ROOT / ".gitignore").read_text(encoding="utf-8")
        self.assertIn(".env.*", ignore_rules)
        self.assertIn("!.env.example", ignore_rules)
        for extension in ("*.pem", "*.key", "*.p12", "*.pfx", "*.keystore"):
            self.assertIn(extension, ignore_rules)

    def test_model_and_ticket_uploads_allow_multipart_beyond_small_json_limit(self):
        body = b"x" * (app_config.MAX_API_REQUEST_BYTES + 1)
        for path in ("/api/modelos/transferencia/importar", "/api/operacoes/chamados/importar"):
            with self.subTest(path=path):
                status, reached_app = _run_local_safety_middleware(
                    ["localhost:8000"], path=path, method="POST", body=body,
                )
                self.assertEqual(status, 200)
                self.assertTrue(reached_app)

    def test_model_and_ticket_uploads_share_bounded_admission(self):
        with patch.object(main, "UPLOAD_ADMISSION") as admission:
            admission.acquire.return_value = False
            for path in ("/api/modelos/exportacao/importar", "/api/operacoes/chamados/importar"):
                status, reached_app = _run_local_safety_middleware(
                    ["localhost:8000"], path=path, method="POST",
                )
                self.assertEqual(status, 503)
                self.assertFalse(reached_app)
            admission.release.assert_not_called()

    def test_readiness_endpoint_returns_the_existing_file_and_function_report(self):
        report = {
            "files": {"template_transferencia": {"available": True, "file": "modelo.xlsx", "reason": "OK"}},
            "functions": {"geracao_transferencia": {"available": True, "missing_files": []}},
        }
        with patch.object(main, "get_startup_readiness", return_value=report):
            self.assertEqual(main.get_readiness_status(), report)

    def test_readiness_panel_has_one_disclosure_and_renders_text_safely(self):
        html = (app_config.STATIC_DIR / "index.html").read_text(encoding="utf-8")
        script = _frontend_script()
        panel_start = script.index("async function carregarProntidao()")
        panel_end = script.index("function adicionarGrupoProntidao", panel_start)
        panel_script = script[panel_start:panel_end]

        self.assertIn('id="readiness-panel"', html)
        self.assertIn('id="readiness-summary"', html)
        self.assertIn('role="status" aria-live="polite"', html)
        self.assertNotIn('id="readiness-refresh"', html)
        self.assertNotIn("Ver detalhes de arquivos e funções", html)
        self.assertIn('fetch("/api/prontidao"', panel_script)
        self.assertIn("carregarProntidao();", script)
        self.assertIn("details.replaceChildren(content)", panel_script)
        self.assertIn("panel.setAttribute(\"aria-busy\", \"false\")", panel_script)
        self.assertIn("status.missing_files", script)
        self.assertIn("status.reason", script)
        self.assertNotIn("innerHTML", panel_script)

    def test_non_xlsx_upload_is_rejected_before_workbook_processing(self):
        with self.assertRaises(main.HTTPException) as raised:
            main.validate_uploaded_workbook(b"fictional invalid upload", "endpoint")

        self.assertEqual(raised.exception.status_code, 400)

    def test_local_request_middleware_rejects_untrusted_hosts_without_origin(self):
        for host in ("attacker.example:8000", "localhost.attacker.example:8000", "192.168.1.12:8000"):
            with self.subTest(host=host):
                status, reached_app = _run_local_safety_middleware([host])
                self.assertEqual(status, 403)
                self.assertFalse(reached_app)

    def test_local_request_middleware_accepts_loopback_hosts_and_originless_clients(self):
        for host in ("localhost:8000", "127.0.0.1:8000", "[::1]:8000"):
            with self.subTest(host=host):
                status, reached_app = _run_local_safety_middleware([host])
                self.assertEqual(status, 200)
                self.assertTrue(reached_app)

    def test_local_request_middleware_requires_api_origin_to_match_host(self):
        allowed = (
            ("localhost:8000", "http://localhost:8000"),
            ("127.0.0.1:8000", "http://127.0.0.1:8000"),
            ("[::1]:8000", "http://[::1]:8000"),
        )
        for host, origin in allowed:
            with self.subTest(host=host, origin=origin):
                status, reached_app = _run_local_safety_middleware([host], origin)
                self.assertEqual(status, 200)
                self.assertTrue(reached_app)

        for host, origin in (
            ("127.0.0.1:8000", "http://localhost:8000"),
            ("localhost:8000", "https://localhost:8000"),
            ("localhost:8000", "http://localhost.attacker.example:8000"),
        ):
            with self.subTest(host=host, origin=origin):
                status, reached_app = _run_local_safety_middleware([host], origin)
                self.assertEqual(status, 403)
                self.assertFalse(reached_app)

    def test_local_request_middleware_rejects_missing_or_duplicate_host_headers(self):
        missing_status, missing_reached_app = _run_local_safety_middleware([])
        duplicate_status, duplicate_reached_app = _run_local_safety_middleware(
            ["localhost:8000", "127.0.0.1:8000"]
        )

        self.assertEqual(missing_status, 400)
        self.assertFalse(missing_reached_app)
        self.assertEqual(duplicate_status, 400)
        self.assertFalse(duplicate_reached_app)

    def test_xlsx_extreme_declared_dimensions_are_rejected_before_openpyxl(self):
        package = _xlsx_with_worksheet_xml(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<dimension ref="A1:XFD1048576"/><sheetData/></worksheet>'
        )

        with patch("openpyxl.load_workbook", side_effect=AssertionError("openpyxl must not run")):
            with self.assertRaises(main.HTTPException) as raised:
                main.read_workbook_tables(package)

        self.assertEqual(raised.exception.status_code, 413)
        self.assertIn("limite", raised.exception.detail)

    def test_xlsx_auxiliary_xml_rejects_entity_declarations_before_openpyxl(self):
        original = _xlsx_with_worksheet_xml(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData/></worksheet>'
        )
        package = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(original), "r") as source, zipfile.ZipFile(package, "w") as target:
            for item in source.infolist():
                target.writestr(item, source.read(item.filename))
            target.writestr(
                "xl/styles.xml",
                '<!DOCTYPE styleSheet [<!ENTITY display "Fictional">]>'
                '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                '<fonts count="1"><font><name val="&display;"/></font></fonts></styleSheet>',
            )

        with patch("openpyxl.load_workbook", side_effect=AssertionError("openpyxl must not run")):
            with self.assertRaises(main.XlsxPackageError) as raised:
                main.validate_xlsx(package.getvalue())

        self.assertIn("declarações XML", str(raised.exception))

    def test_xlsx_dimension_check_uses_actual_cells_not_only_declared_range(self):
        package = _xlsx_with_worksheet_xml(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<dimension ref="A1:B2"/><sheetData><row r="250001">'
            '<c r="A250001" t="inlineStr"><is><t>fixture</t></is></c>'
            '</row></sheetData></worksheet>'
        )

        with self.assertRaises(main.HTTPException) as raised:
            main.validate_uploaded_workbook(package, "endpoint")

        self.assertEqual(raised.exception.status_code, 413)

    def test_xlsx_total_dimension_budget_applies_across_worksheets(self):
        sheet_xml = (
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<dimension ref="A1:Z200000"/><sheetData/></worksheet>'
        )
        package = _xlsx_with_worksheet_xml(sheet_xml, sheet_xml)

        with self.assertRaises(main.HTTPException) as raised:
            main.validate_xlsx_package(package)

        self.assertEqual(raised.exception.status_code, 413)

    def test_xlsx_wide_but_bounded_template_dimension_remains_accepted(self):
        package = _xlsx_with_worksheet_xml(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<dimension ref="A1:XFD13"/><sheetData/></worksheet>'
        )

        self.assertIsNone(main.validate_xlsx_package(package))

    def test_xlsx_without_dimension_uses_referenced_cells_and_accepts_small_sheet(self):
        package = _xlsx_with_worksheet_xml(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>header</t></is></c>'
            '</row><row r="2"><c r="A2" t="inlineStr"><is><t>value</t></is></c>'
            '</row></sheetData></worksheet>'
        )

        self.assertIsNone(main.validate_xlsx_package(package))

    def test_xlsx_implicit_cell_coordinate_cannot_exceed_column_limit(self):
        package = _xlsx_with_worksheet_xml(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData><row r="1"><c r="XFD1"><v>1</v></c><c><v>2</v></c></row></sheetData>'
            '</worksheet>'
        )

        with self.assertRaises(main.HTTPException) as raised:
            main.validate_xlsx_package(package)

        self.assertEqual(raised.exception.status_code, 413)
        self.assertIn("colunas", raised.exception.detail)

    def test_xlsx_implicit_cells_count_toward_rectangular_dimension_limit(self):
        implicit_cells = "<c><v>1</v></c>" * 52
        package = _xlsx_with_worksheet_xml(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData><row r="200000">' + implicit_cells + '</row></sheetData></worksheet>'
        )

        with self.assertRaises(main.HTTPException) as raised:
            main.validate_xlsx_package(package)

        self.assertEqual(raised.exception.status_code, 413)
        self.assertIn("área", raised.exception.detail)

    def test_xlsx_implicit_row_and_cell_coordinates_remain_accepted(self):
        package = _xlsx_with_worksheet_xml(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData><row r="1"><c r="B1"><v>1</v></c><c><v>2</v></c></row>'
            '<row><c><v>3</v></c></row></sheetData></worksheet>'
        )

        self.assertIsNone(main.validate_xlsx_package(package))

    def test_frontend_references_assets_served_from_the_existing_static_mount(self):
        index_file = app_config.STATIC_DIR / "index.html"
        html = index_file.read_text(encoding="utf-8")
        self.assertIn('<link rel="icon" type="image/svg+xml" href="/static/favicon.svg">', html)
        self.assertIn('<link rel="stylesheet" href="/static/vendor/fonts.css">', html)
        self.assertIn('<link rel="stylesheet" href="/static/vendor/tailwind-4.3.3.css">', html)
        self.assertIn('<script src="/static/vendor/lucide-1.48.0-subset.js"></script>', html)
        self.assertIn('<link rel="stylesheet" href="/static/app.css">', html)
        self.assertIn('<script src="/static/js/shared/transfer-utils.js"></script>', html)
        self.assertIn('<script src="/static/app.js"></script>', html)
        self.assertNotRegex(html, r"https://(?:fonts\.googleapis\.com|fonts\.gstatic\.com|cdn\.tailwindcss\.com|unpkg\.com)")
        self.assertTrue((app_config.STATIC_DIR / "app.css").is_file())
        self.assertTrue((app_config.STATIC_DIR / "app.js").is_file())
        self.assertTrue((app_config.STATIC_DIR / "js/shared/transfer-utils.js").is_file())
        self.assertTrue((app_config.STATIC_DIR / "vendor" / "fonts.css").is_file())
        self.assertTrue((app_config.STATIC_DIR / "vendor" / "tailwind-4.3.3.css").is_file())
        self.assertTrue((app_config.STATIC_DIR / "vendor" / "lucide-1.48.0-subset.js").is_file())
        favicon_file = app_config.STATIC_DIR / "favicon.svg"
        self.assertTrue(favicon_file.is_file())
        favicon_route = next(
            route for route in main.app.routes
            if getattr(route, "path", None) == "/favicon.ico"
        )
        favicon_response = favicon_route.endpoint()
        self.assertEqual(favicon_response.path, favicon_file)
        self.assertEqual(favicon_response.media_type, "image/svg+xml")

        mounts = [
            route for route in main.app.routes
            if getattr(route, "name", None) == "static"
        ]
        self.assertEqual(len(mounts), 1)
        self.assertEqual(mounts[0].path, "/static")

    def test_restore_invalidates_stale_equipment_and_refreshes_inventory_views(self):
        script = _frontend_script()
        invalidation_start = script.index("function invalidarDadosEquipamentoAposRestauracao()")
        invalidation_end = script.index("const inputSerialInicial", invalidation_start)
        invalidation = script[invalidation_start:invalidation_end]
        self.assertIn('serialInput.value = ""', invalidation)
        self.assertIn("invalidarBuscasEquipamento()", invalidation)
        self.assertIn("limparSelecaoEquipamento()", invalidation)
        self.assertIn("fecharAutocomplete()", invalidation)

        restore_start = script.index("async function restaurarHistoricoComparacao")
        restore_end = script.index("async function consultarHistoricoComparacao", restore_start)
        restore = script[restore_start:restore_end]
        for refresh in (
            "invalidarDadosEquipamentoAposRestauracao()",
            "carregarEstatisticas()",
            "carregarUnidades()",
            "carregarHistoricoComparacoes()",
            "carregarContagemHistoricoComparacoes()",
        ):
            with self.subTest(refresh=refresh):
                self.assertIn(refresh, restore)

    def test_dialog_overlays_stack_above_the_header_and_keep_keyboard_trapping(self):
        html = (app_config.STATIC_DIR / "index.html").read_text(encoding="utf-8")
        dialog_script = (app_config.STATIC_DIR / "js/ui/accessible-dialogs.js").read_text(encoding="utf-8")

        dialogs = re.findall(r'<div id="(modal-[^"]+)" class="([^"]+)"', html)
        self.assertGreater(len(dialogs), 0)
        for dialog_id, class_names in dialogs:
            with self.subTest(dialog=dialog_id):
                self.assertIn("modal-overlay", class_names)
                # .modal-overlay owns the layer; open order (accessible-dialogs.js) stacks dialogs.
                self.assertNotRegex(class_names, r"(?:^|\s)z-")

        self.assertGreater(_css_layer(".modal-overlay"), _css_layer(".site-header"))
        self.assertGreater(_css_layer(".site-header"), _css_layer(".quick-actions"))
        self.assertIn('dialog.setAttribute("aria-modal", "true")', dialog_script)
        self.assertIn('if (event.key === "Escape")', dialog_script)
        self.assertIn('if (event.key !== "Tab") return;', dialog_script)
        self.assertIn("event.shiftKey", dialog_script)
        self.assertIn("target.focus({ preventScroll: true })", dialog_script)

    def test_startup_readiness_maps_missing_files_to_unavailable_functions(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            details = root / "EquipamentosDetalhes_atual.xlsx"
            details.write_bytes(_xlsx_with_worksheet_xml(
                '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                '<dimension ref="A1"/><sheetData/></worksheet>'
            ))

            with (
                patch.object(main, "get_primary_equipment_file", return_value=details),
                patch.object(main, "get_details_files", return_value=[details]),
                patch.object(main, "INPUT_DIR", root),
                patch.object(main, "TRANSACTION_ROOT", root),
                patch.object(main, "TICKET_REPORT_FILE", root / "chamados.json"),
                patch.object(main, "load_modified_apelidos", return_value={}),
                patch.object(main, "load_custom_units", return_value=[]),
                patch.object(main, "ALIAS_HISTORY_FILE", root / "aliases.json"),
                patch.object(main, "COMPARISON_HISTORY_FILE", root / "comparison.json"),
                patch.object(main, "equipamentos_cache", {}),
                patch.object(main, "TEMPLATE_FILE", root / "template.xlsx"),
                patch.object(main, "DETAILS_UPLOAD_FILE", details),
                patch.object(main, "ENDPOINT_UPLOAD_FILE", root / "endpoint.xlsx"),
                patch.object(main, "EXPORT_MODEL_FILE", root / "export.xlsx"),
            ):
                report = main.get_startup_readiness()
                with patch.object(main, "get_transaction_recovery_block", return_value="journal divergente"):
                    blocked_report = main.get_startup_readiness()

            self.assertNotIn("gravacao_bloqueada", report)
            self.assertEqual(
                blocked_report["gravacao_bloqueada"],
                {"available": False, "state": "invalid", "reason": "journal divergente"},
            )
            self.assertEqual(blocked_report["functions"], report["functions"])
            self.assertFalse(report["files"]["inventario_equipamentos"]["available"])
            self.assertFalse(report["functions"]["geracao_transferencia"]["available"])
            self.assertEqual(
                report["functions"]["geracao_transferencia"]["missing_files"],
                ["EquipamentosDetalhes_atual.xlsx", "template.xlsx"],
            )
            self.assertFalse(report["functions"]["comparacao_atualizacao_inventarios"]["available"])
            self.assertEqual(
                report["functions"]["comparacao_atualizacao_inventarios"]["missing_files"],
                ["EquipamentosDetalhes_atual.xlsx", "endpoint.xlsx"],
            )
            self.assertFalse(report["functions"]["exportacao_equipamentos"]["available"])

    def test_startup_readiness_marks_corrupt_xlsx_package_unavailable(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "bad.xlsx"
            path.write_bytes(b"not an xlsx archive")

            status = main._check_startup_xlsx(path)

        self.assertFalse(status["available"])
        self.assertEqual(status["reason"], "arquivo inacessível ou XLSX inválido")

    def test_application_startup_reports_readiness_without_workbook_values(self):
        output = io.StringIO()

        with tempfile.TemporaryDirectory(prefix=".startup-test-", dir=PROJECT_ROOT) as temp_dir:
            root = Path(temp_dir)
            input_dir = root / "data" / "planilhas_entrada"
            state_dir = root / "data" / "state"
            input_dir.mkdir(parents=True)
            (state_dir / "apelidos_modificados").mkdir(parents=True)
            (state_dir / "historico_comparacoes").mkdir(parents=True)
            template = input_dir / "(MODELO) - TRANSFERENCIA SIMPRESS.xlsx"
            template.write_bytes(_xlsx_with_worksheet_xml(
                '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                '<dimension ref="A1"/><sheetData/></worksheet>'
            ))

            with (
                patch.object(main, "BASE_DIR", root),
                patch.object(main, "INPUT_DIR", input_dir),
                patch.object(main, "TEMPLATE_FILE", template),
                patch.object(main, "DETAILS_UPLOAD_FILE", input_dir / "EquipamentosDetalhes_atual.xlsx"),
                patch.object(main, "ENDPOINT_UPLOAD_FILE", input_dir / "InvAllComputersSummary_atual.xlsx"),
                patch.object(main, "EXPORT_MODEL_FILE", root / "modelos" / "EquipamentosDetalhes-exportacao.xlsx"),
                patch.object(main, "UPLOAD_BACKUPS_DIR", input_dir / "_backups"),
                patch.object(main, "CUSTOM_UNIDADES_FILE", state_dir / "unidades_custom.json"),
                patch.object(main, "MODIFIED_APELIDOS_FILE", state_dir / "apelidos_modificados" / "apelidos_modificados.json"),
                patch.object(main, "ALIAS_HISTORY_FILE", state_dir / "historico_alteracoes.json"),
                patch.object(main, "COMPARISON_HISTORY_FILE", state_dir / "historico_comparacoes" / "historico_comparacoes.json"),
                patch.object(main, "HISTORY_ANNOTATIONS_DIR", state_dir / "anotacoes_historico"),
                patch.object(main, "LOCAL_HISTORY_ANNOTATION_OUTBOX_DIR", state_dir / "anotacoes_historico_pendentes"),
                patch.object(main, "TRANSACTION_ROOT", root),
                patch.object(main, "TICKET_REPORT_FILE", state_dir / "chamados.json"),
                patch.object(main, "AUTOMATIC_INVENTORY_INBOX_DIR", state_dir / "inbox"),
                patch.object(main, "AUTOMATIC_TICKET_INBOX_DIR", state_dir / "ticket_inbox"),
                redirect_stdout(output),
            ):
                async def run_application_lifespan():
                    async with main.app.router.lifespan_context(main.app):
                        pass

                asyncio.run(run_application_lifespan())

        self.assertIn("[PRONTIDÃO] Verificação dos arquivos locais:", output.getvalue())
        self.assertIn(
            "[PRONTIDÃO] Geração de planilha de transferência: INDISPONÍVEL;",
            output.getvalue(),
        )

    def test_lifespan_prepares_disk_state_off_loop_before_starting_workers(self):
        events = []

        def record(name, result=None):
            def call(*_args, **_kwargs):
                events.append((name, threading.get_ident()))
                return result
            return call

        async def worker():
            events.append(("worker", threading.get_ident()))
            await asyncio.Event().wait()

        async def run_application_lifespan():
            loop_thread = threading.get_ident()
            async with main.app.router.lifespan_context(main.app):
                await asyncio.sleep(0)
            return loop_thread

        with (
            patch.object(main, "recover_pending_transactions", record("recover")),
            patch.object(main, "migrate_legacy_history_annotations", record("migrate", {})),
            patch.object(main, "prune_expired_histories", record("prune", {})),
            patch.object(main, "load_database", record("load", True)),
            patch.object(main, "_report_startup_readiness_best_effort", record("readiness")),
            patch.object(main, "_history_annotation_sync_worker", worker),
            patch.object(main, "_automatic_inventory_import_worker", worker),
            patch.object(main, "_automatic_ticket_import_worker", worker),
        ):
            loop_thread = asyncio.run(run_application_lifespan())

        names = [name for name, _ in events]
        self.assertEqual(names, ["recover", "migrate", "prune", "load", "readiness", "worker", "worker", "worker"])
        self.assertTrue(all(thread != loop_thread for _, thread in events[:5]))
        self.assertTrue(all(thread == loop_thread for _, thread in events[5:]))

    def test_diverged_recovery_journal_starts_read_only_until_resolved(self):
        from core import integrity

        self.addCleanup(setattr, main, "_TRANSACTION_RECOVERY_BLOCK", None)
        events = []

        def record(name, result=None):
            def call(*_args, **_kwargs):
                events.append(name)
                return result
            return call

        async def worker():
            events.append("worker")
            await asyncio.Event().wait()

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            target = root / "apelidos.json"
            target.write_bytes(b"before")
            interrupted = integrity.AtomicFileTransaction(root, [target])
            interrupted.__enter__()
            interrupted.stage_bytes(target, b"after")
            interrupted.apply()
            interrupted.close_after_interruption()
            # Matches neither recorded hash, so recovery must refuse to overwrite it.
            target.write_bytes(b"edited elsewhere")

            async def run_application_lifespan():
                async with main.app.router.lifespan_context(main.app):
                    await asyncio.sleep(0)
                    message = await asyncio.to_thread(main.get_transaction_recovery_block)
                    with self.assertRaises(HTTPException) as held:
                        await asyncio.to_thread(main._raise_if_transaction_recovery_blocked)
                return message, held.exception

            with (
                patch.object(main, "TRANSACTION_ROOT", root),
                patch.object(main, "recover_pending_transactions", integrity.recover_pending_transactions),
                patch.object(main, "migrate_legacy_history_annotations", record("migrate", {})),
                patch.object(main, "prune_expired_histories", record("prune", {})),
                patch.object(main, "load_database", record("load", True)),
                patch.object(main, "_report_startup_readiness_best_effort", record("readiness")),
                patch.object(main, "_history_annotation_sync_worker", worker),
                patch.object(main, "_automatic_inventory_import_worker", worker),
                patch.object(main, "_automatic_ticket_import_worker", worker),
                redirect_stdout(io.StringIO()) as output,
            ):
                message, held = asyncio.run(run_application_lifespan())

                self.assertEqual(target.read_bytes(), b"edited elsewhere")
                self.assertEqual(events, ["load", "readiness", "worker", "worker", "worker"])
                self.assertIn("não corresponde ao estado esperado", message)
                self.assertIn("[INDISPONÍVEL] Recuperação de gravação pendente:", output.getvalue())
                # 503 is the inboxes' transient status: batches stay pending instead of failing.
                self.assertEqual(held.status_code, 503)
                self.assertEqual(held.detail, message)

                # After manual review of the journal, writes resume without a restart.
                for journal in (root / ".transactions").iterdir():
                    shutil.rmtree(journal)
                self.assertIsNone(main.get_transaction_recovery_block())
                main._raise_if_transaction_recovery_blocked()

    def test_readiness_report_failure_is_logged_and_suppressed(self):
        with (
            patch.object(main, "report_startup_readiness", side_effect=OSError("readiness unavailable")),
            patch.object(main.logger, "exception") as log_exception,
        ):
            main._report_startup_readiness_best_effort()

        log_exception.assert_called_once()

    def test_windows_launcher_bootstraps_environment_without_gating_on_workbooks(self):
        launcher = (PROJECT_ROOT / "iniciar_servidor.bat").read_text(encoding="utf-8")

        self.assertIn("bootstrap_environment.py", launcher)
        self.assertNotIn("--postgres-pilot", launcher)
        self.assertIn('set "PROJECT_ROOT=%PROJECT_DIR%."', launcher)
        self.assertIn('--project-root "%PROJECT_ROOT%"', launcher)
        self.assertNotIn('if not exist "%TEMPLATE%" goto', launcher)
        self.assertNotIn(":missing_template", launcher)


if __name__ == "__main__":
    unittest.main()
