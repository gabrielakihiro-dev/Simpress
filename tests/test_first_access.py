"""First access and replacement tests use exclusively temporary synthetic data."""

import asyncio
import copy
import io
import json
import importlib.util
import struct
import unittest
import zipfile
from unittest.mock import patch

import openpyxl
from starlette.datastructures import UploadFile

from core.config import format_size_limit
from domains.local_setup import TRANSFER_CELLS, unpack_simpress_zip, validate_local_model
from core.xlsx_safety import XlsxPackageError, validate_xlsx
from tests import test_integrity_workflows as fixtures
from tests.test_integrity_workflows import complete_details_headers, main, make_details, make_endpoint


def workbook_bytes(rows, title="Parque"):
    wb = openpyxl.Workbook()
    wb.active.title = title
    for row in rows:
        wb.active.append(row)
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def transfer_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transferência"
    ws["B2"] = "TRANSFERÊNCIA DE EQUIPAMENTO"
    ws["B15"] = "DADOS DE RETIRADA"
    ws["B27"] = "DADOS DE ENTREGA"
    for address in TRANSFER_CELLS:
        ws[address] = "CAMPO FICTICIO"
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def archive_bytes(entries, compression=zipfile.ZIP_DEFLATED):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression) as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return stream.getvalue()


def corrupt_compressed_member(data, member_name):
    """Damage one synthetic member without changing its ZIP size or CRC."""
    damaged = bytearray(data)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        member = archive.getinfo(member_name)
        offset = member.header_offset
        filename_length, extra_length = struct.unpack_from("<HH", damaged, offset + 26)
        compressed_start = offset + 30 + filename_length + extra_length
    if member.compress_type == zipfile.ZIP_DEFLATED:
        # Reserved DEFLATE block type.
        damaged[compressed_start] = (damaged[compressed_start] & ~0x06) | 0x06
    elif member.compress_type == zipfile.ZIP_BZIP2:
        damaged[compressed_start] = 0
    elif member.compress_type == zipfile.ZIP_LZMA:
        # Invalid LZMA filter properties, after the four-byte ZIP-LZMA header.
        damaged[compressed_start + 4] = 0xFF
    else:
        raise ValueError("unsupported synthetic corruption method")
    return bytes(damaged)


class FirstAccessTests(unittest.TestCase):
    def setUp(self):
        fixtures.WorkflowIntegrityTests.setUp(self)
        self.template = self.input_dir / "template.xlsx"
        self.ticket_report = self.root / "chamados.json"
        for key, value in {
            "TEMPLATE_FILE": self.template,
            "TICKET_REPORT_FILE": self.ticket_report,
            "AUTOMATIC_INVENTORY_INBOX_DIR": self.root / "inbox",
            "AUTOMATIC_TICKET_INBOX_DIR": self.root / "ticket_inbox",
        }.items():
            self.stack.enter_context(patch.object(main, key, value))

    def tearDown(self):
        fixtures.WorkflowIntegrityTests.tearDown(self)

    def import_model(self, kind, data):
        return asyncio.run(main.importar_modelo_local(kind, UploadFile(filename="modelo.xlsx", file=io.BytesIO(data))))

    def import_pair(self, zipped=False, client=None):
        source = self.root / "origem.xlsx"
        endpoint = self.root / "endpoint.xlsx"
        make_details(source, [("S-1", "HOST-OLD", "Notebook A", "Curitiba", "Notebook", "80000000", "12.345.678/0001-90")],
                     complete_details_headers(["NumeroSerie", "Apelido", "Modelo", "Cidade", "Tipo", "CEP"]))
        make_endpoint(endpoint, [("S-1", "HOST-NEW")])
        data = source.read_bytes()
        if zipped:
            data = archive_bytes([("pasta/EquipamentosDetalhes.xlsx", data)])
        if client is not None:
            response = client.post("/api/planilhas/atualizar", files={
                "endpoint_file": ("endpoint.xlsx", endpoint.read_bytes()),
                "details_file": ("simpress.zip" if zipped else "simpress.xlsx", data),
            })
            self.assertEqual(response.status_code, 200, response.text)
            return response.json()
        return asyncio.run(main.atualizar_planilhas(
            endpoint_file=UploadFile(filename="endpoint.xlsx", file=io.BytesIO(endpoint.read_bytes())),
            details_file=UploadFile(filename="simpress.zip" if zipped else "simpress.xlsx", file=io.BytesIO(data)),
        ))

    def test_fresh_clone_starts_and_stops_all_workers_without_data(self):
        async def run():
            async with main.app.router.lifespan_context(main.app):
                await asyncio.sleep(0.05)
                report = main.get_startup_readiness()
                self.assertEqual(report["backend"], "files")
                self.assertTrue(all(not item["available"] for item in report["functions"].values()))
                self.assertEqual(report["functions"]["acompanhamento_chamados"]["state"], "awaiting_import")
                self.assertEqual(main._get_stats_snapshot()["total_equipamentos"], 0)
                self.assertFalse(main._get_stats_snapshot()["base_importada"])
        asyncio.run(run())

    def test_initial_pair_from_zip_loads_without_restart_and_preserves_original_name(self):
        self.template.write_bytes(transfer_bytes())
        result = self.import_pair(zipped=True)
        self.assertTrue(result["sucesso"])
        self.assertEqual(main._get_stats_snapshot()["total_equipamentos"], 1)
        self.assertTrue(main.unidades_cache)
        report = main.get_startup_readiness()
        for key in ("consulta_equipamentos", "geracao_transferencia", "comparacao_atualizacao_inventarios"):
            self.assertTrue(report["functions"][key]["available"], key)
        self.assertEqual(json.loads(self.comparison_history.read_text(encoding="utf-8"))[-1]["arquivos_enviados"]["detalhes"], "simpress.zip")

    def test_fresh_clone_previews_zip_before_confirming_initial_inventory(self):
        endpoint = fixtures._endpoint_bytes([("SN-PREVIEW", "HOST-NEW")])
        details = archive_bytes([("simpress.xlsx", fixtures._details_bytes([
            ("SN-PREVIEW", "HOST-OLD", "Modelo A", "Curitiba", "Notebook", "80000000"),
        ]))])
        preview = asyncio.run(main.previsualizar_planilhas(
            UploadFile(filename="endpoint.xlsx", file=io.BytesIO(endpoint)),
            UploadFile(filename="simpress.zip", file=io.BytesIO(details)),
        ))
        for role in ("endpoint", "detalhes"):
            self.assertEqual(preview["alteracoes"][role]["linhas_atuais"], 0)
            self.assertEqual(preview["alteracoes"][role]["series_novas"], 1)
        self.assertFalse(self.endpoint.exists())
        self.assertFalse(self.details.exists())
        result = asyncio.run(main.atualizar_planilhas(
            UploadFile(filename="endpoint.xlsx", file=io.BytesIO(endpoint)),
            UploadFile(filename="simpress.zip", file=io.BytesIO(details)),
            preview_token=preview["preview_token"],
        ))
        self.assertTrue(result["sucesso"])
        self.assertEqual(main._get_stats_snapshot()["total_equipamentos"], 1)

    def test_transfer_import_override_replacement_and_invalid_file_preserve_default(self):
        self.template.write_bytes(transfer_bytes())
        original = self.template.read_bytes()
        data = transfer_bytes()
        self.import_model("transferencia", data)
        target = main.get_local_transfer_template()
        self.assertNotEqual(target, self.template)
        self.import_model("transferencia", data)
        self.assertEqual(len(list((self.backups / "modelos").rglob("*.xlsx"))), 1)
        with self.assertRaises(main.HTTPException) as caught:
            self.import_model("transferencia", workbook_bytes([["qualquer planilha"]]))
        self.assertEqual(caught.exception.status_code, 400)
        self.assertEqual(target.read_bytes(), data)
        self.assertEqual(self.template.read_bytes(), original)

    def test_export_rejects_wrong_serial_set_and_rolls_back(self):
        self.import_pair()
        valid = workbook_bytes([["NumeroSerie", "Apelido"], ["S-1", "HOST-OLD"]])
        self.import_model("exportacao", valid)
        self.assertTrue(main.get_startup_readiness()["functions"]["exportacao_equipamentos"]["available"])
        with self.assertRaises(main.HTTPException) as caught:
            self.import_model("exportacao", workbook_bytes([["NumeroSerie", "Apelido"], ["OUTRA", "HOST"]]))
        self.assertEqual(caught.exception.status_code, 400)
        self.assertEqual(self.model.read_bytes(), valid)
        self.assertTrue(main.get_startup_readiness()["functions"]["exportacao_equipamentos"]["available"])

    def test_export_can_be_configured_before_inventory(self):
        self.import_model("exportacao", workbook_bytes([["NumeroSerie", "Apelido"], ["S-1", "HOST"]]))
        self.assertFalse(main.get_startup_readiness()["functions"]["exportacao_equipamentos"]["available"])
        self.import_pair()
        self.assertTrue(main.get_startup_readiness()["functions"]["exportacao_equipamentos"]["available"])

    def test_export_missing_writable_cell_is_rejected(self):
        with self.assertRaises(main.HTTPException) as caught:
            self.import_model("exportacao", workbook_bytes([["NumeroSerie", "Apelido"], ["S-1"]]))
        self.assertEqual(caught.exception.status_code, 400)
        self.assertFalse(self.model.exists())

    def test_zip_rejects_ambiguity_missing_xlsx_traversal_and_oversize(self):
        for entries, limit in [([("a.xlsx", b"a"), ("b.xlsx", b"b")], 100),
                               ([("a.txt", b"a")], 100), ([("../a.xlsx", b"a")], 100),
                               ([("a.xlsx", b"12345")], 4)]:
            with self.subTest(entries=entries), self.assertRaises(main.HTTPException):
                unpack_simpress_zip(archive_bytes(entries), limit)

    def test_size_limit_messages_follow_the_configured_limit(self):
        self.assertEqual(format_size_limit(60 * 1024 * 1024), "60 MB")
        self.assertEqual(format_size_limit(64 * 1024 * 1024), "64 MB")
        self.assertEqual(format_size_limit(2048), "2 KB")
        self.assertEqual(format_size_limit(4), "4 bytes")
        with self.assertRaises(main.HTTPException) as caught:
            unpack_simpress_zip(archive_bytes([("a.xlsx", b"12345")]), 4)
        self.assertIn("limite de 4 bytes", caught.exception.detail)

    def test_compression_corruption_maps_to_existing_xlsx_and_upload_errors(self):
        valid_xlsx = workbook_bytes([["NumeroSerie", "Apelido"], ["SINTETICO-1", "HOST-1"]])
        with zipfile.ZipFile(io.BytesIO(valid_xlsx)) as archive:
            parts = [(name, archive.read(name)) for name in archive.namelist()]
        for compression in (zipfile.ZIP_DEFLATED, zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA):
            repacked = archive_bytes(parts, compression)
            for member, expected in (
                ("xl/workbook.xml", "O arquivo contém metadados XLSX inválidos."),
                ("xl/_rels/workbook.xml.rels", "O arquivo contém metadados XLSX inválidos."),
                ("xl/worksheets/sheet1.xml", "Uma aba contém XML inválido ou incompleto."),
                ("docProps/core.xml", "A planilha contém uma parte XML inválida."),
            ):
                with self.subTest(compression=compression, member=member):
                    with self.assertRaises(XlsxPackageError) as workbook_error:
                        validate_xlsx(corrupt_compressed_member(repacked, member))
                    self.assertEqual(str(workbook_error.exception), expected)

            corrupt_outer_zip = corrupt_compressed_member(
                archive_bytes([("simpress.xlsx", valid_xlsx)], compression), "simpress.xlsx",
            )
            with self.subTest(compression=compression, member="simpress.xlsx"):
                with self.assertRaises(main.HTTPException) as zip_error:
                    unpack_simpress_zip(corrupt_outer_zip, 1_000_000)
                self.assertEqual(zip_error.exception.status_code, 400)
                self.assertEqual(zip_error.exception.detail, "Não foi possível ler o ZIP da Simpress.")

    def test_supported_compression_methods_preserve_text_identifiers(self):
        valid_xlsx = workbook_bytes([["NumeroSerie", "Apelido"], ["000123", "HOST-Á"]])
        with zipfile.ZipFile(io.BytesIO(valid_xlsx)) as archive:
            parts = [(name, archive.read(name)) for name in archive.namelist()]
        for compression in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED, zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA):
            with self.subTest(compression=compression):
                repacked = archive_bytes(parts, compression)
                validate_xlsx(repacked)
                workbook = openpyxl.load_workbook(io.BytesIO(repacked), read_only=True)
                try:
                    self.assertEqual(workbook.active["A2"].value, "000123")
                    self.assertEqual(workbook.active["B2"].value, "HOST-Á")
                finally:
                    workbook.close()
                unpacked, name = unpack_simpress_zip(
                    archive_bytes([("pasta/simpress.xlsx", repacked)], compression), len(repacked),
                )
                self.assertEqual(unpacked, repacked)
                self.assertEqual(name, "pasta/simpress.xlsx")

    @unittest.skipUnless(importlib.util.find_spec("httpx"), "optional HTTP test dependency not installed")
    def test_http_corrupt_compression_preserves_active_files_history_and_cache(self):
        from fastapi.testclient import TestClient

        self.import_pair()
        tracked = [self.endpoint, self.details, self.comparison_history, self.overrides]
        before = {path: path.read_bytes() for path in tracked if path.is_file()}
        cache_before = copy.deepcopy(main.equipamentos_cache)
        stats_before = main._get_stats_snapshot()
        valid_xlsx = self.details.read_bytes()
        with zipfile.ZipFile(io.BytesIO(valid_xlsx)) as archive:
            parts = [(name, archive.read(name)) for name in archive.namelist()]
        # No lifespan: all request state is supplied by temporary fixtures.
        client = TestClient(main.app, base_url="http://127.0.0.1:8000", client=("127.0.0.1", 1234))
        try:
            for compression in (zipfile.ZIP_DEFLATED, zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA):
                corrupt_zip = corrupt_compressed_member(
                    archive_bytes([("simpress.xlsx", valid_xlsx)], compression), "simpress.xlsx",
                )
                corrupt_xlsx = corrupt_compressed_member(
                    archive_bytes(parts, compression), "xl/worksheets/sheet1.xml",
                )
                for route in ("/api/planilhas/previsualizar", "/api/planilhas/atualizar"):
                    for name, payload in (("simpress.zip", corrupt_zip), ("simpress.xlsx", corrupt_xlsx)):
                        with self.subTest(compression=compression, route=route, name=name):
                            response = client.post(route, files={
                                "endpoint_file": ("endpoint.xlsx", before[self.endpoint]),
                                "details_file": (name, payload),
                            })
                            self.assertEqual(response.status_code, 400, response.text)
                            if name.endswith(".zip"):
                                self.assertEqual(response.json()["detail"], "Não foi possível ler o ZIP da Simpress.")
                            self.assertEqual({path: path.read_bytes() for path in before}, before)
                            self.assertEqual(main.equipamentos_cache, cache_before)
                            self.assertEqual(main._get_stats_snapshot(), stats_before)
        finally:
            client.close()

    def test_invalid_zip_cannot_replace_inventory_or_cache(self):
        self.import_pair()
        before = self.details.read_bytes()
        with self.assertRaises(main.HTTPException):
            asyncio.run(main.atualizar_planilhas(
                endpoint_file=UploadFile(filename="endpoint.xlsx", file=io.BytesIO(self.endpoint.read_bytes())),
                details_file=UploadFile(filename="simpress.zip", file=io.BytesIO(b"broken zip")),
            ))
        self.assertEqual(before, self.details.read_bytes())
        self.assertEqual(main._get_stats_snapshot()["total_equipamentos"], 1)

    def test_merged_transfer_field_is_rejected(self):
        wb = openpyxl.load_workbook(io.BytesIO(transfer_bytes()))
        wb.active.merge_cells("A5:B5")
        data = io.BytesIO()
        wb.save(data)
        wb.close()
        with self.assertRaises(ValueError):
            validate_local_model(io.BytesIO(data.getvalue()), "transferencia")

    def test_missing_transfer_cell_is_rejected_before_generation(self):
        wb = openpyxl.load_workbook(io.BytesIO(transfer_bytes()))
        del wb.active["B5"]
        data = io.BytesIO()
        wb.save(data)
        wb.close()
        with self.assertRaises(ValueError):
            validate_local_model(io.BytesIO(data.getvalue()), "transferencia")

    def test_readiness_cache_detects_external_invalid_replacement(self):
        self.template.write_bytes(transfer_bytes())
        self.assertTrue(main.get_startup_readiness()["files"]["template_transferencia"]["available"])
        self.template.write_bytes(workbook_bytes([["wrong layout"]]))
        status = main.get_startup_readiness()["files"]["template_transferencia"]
        self.assertFalse(status["available"])
        self.assertEqual(status["state"], "invalid")

    def test_ticket_valid_empty_import_is_ready(self):
        router = main.operations_router
        handler = next(route.endpoint for route in router.routes if route.path == "/api/operacoes/chamados/importar")
        handler(
            simpress_file=UploadFile(filename="simpress.xlsx", file=io.BytesIO(workbook_bytes([
                ["NumeroPedido", "NumeroSerie", "OrdemServico", "StatusAtual", "DataPedido"]]))),
            oracle_file=UploadFile(filename="oracle.xlsx", file=io.BytesIO(workbook_bytes([
                ["Solicitação", "ID OS Corporate", "Número de Série", "Data de Criação", "Status"]]))),
        )
        self.assertEqual(json.loads(self.ticket_report.read_text(encoding="utf-8"))["chamados"], [])
        self.assertTrue(main.get_startup_readiness()["functions"]["acompanhamento_chamados"]["available"])

    @unittest.skipUnless(importlib.util.find_spec("httpx"), "optional HTTP test dependency not installed")
    def test_http_first_access_upload_pair_validation_and_model_replacement(self):
        from fastapi.testclient import TestClient
        with TestClient(main.app, base_url="http://127.0.0.1:8000", client=("127.0.0.1", 1234)) as client:
            self.assertEqual(client.get("/").status_code, 200)
            self.assertEqual(client.get("/api/stats").json()["total_equipamentos"], 0)
            self.assertEqual(client.post("/api/planilhas/atualizar", files={
                "details_file": ("simpress.xlsx", b"empty"),
            }).status_code, 422)
            self.assertEqual(client.post("/api/modelos/desconhecido/importar", files={
                "file": ("modelo.xlsx", transfer_bytes()),
            }).status_code, 404)
            self.assertEqual(client.post("/api/modelos/transferencia/importar", headers={"Origin": "https://outside.example"},
                files={"file": ("modelo.xlsx", transfer_bytes())}).status_code, 403)
            response = client.post("/api/modelos/transferencia/importar", files={"file": ("modelo.xlsx", transfer_bytes())})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertFalse(client.get("/api/prontidao").json()["functions"]["geracao_transferencia"]["available"])
            self.import_pair(zipped=True, client=client)
            generated = client.post("/api/gerar-planilha", json=fixtures.make_transfer_payload(serial="S-1").model_dump())
            self.assertEqual(generated.status_code, 200, generated.text if generated.status_code != 200 else "")
            wb = openpyxl.load_workbook(io.BytesIO(generated.content))
            self.assertEqual(wb.active["B8"].value, "S-1")
            wb.close()
            response = client.post("/api/modelos/exportacao/importar", files={"file": ("modelo.xlsx", workbook_bytes([
                ["NumeroSerie", "Apelido"], ["S-1", "HOST-OLD"],
            ]))})
            self.assertEqual(response.status_code, 200, response.text)
            output = client.get("/api/exportar-equipamentos")
            self.assertEqual(output.status_code, 200, output.text if output.status_code != 200 else "")
            wb = openpyxl.load_workbook(io.BytesIO(output.content))
            self.assertEqual(wb.active["B2"].value, "HOST-NEW")
            wb.close()
            usage = client.get("/api/painel").json()["uso_30_dias"]["totais"]
            self.assertEqual((usage["transferencias"], usage["comparacoes"], usage["exportacoes"]), (1, 1, 1))
        metrics = json.loads((self.state_dir / "metricas_uso.json").read_text(encoding="utf-8"))
        self.assertNotIn("S-1", json.dumps(metrics))

    def test_failed_durable_replacement_rolls_back_and_releases_parser(self):
        self.import_model("transferencia", transfer_bytes())
        before = main.get_local_transfer_template().read_bytes()
        with patch("core.integrity._replace_with_retries", side_effect=main.FileBusyError("ocupado")):
            with self.assertRaises(main.HTTPException) as caught:
                self.import_model("transferencia", transfer_bytes())
        self.assertEqual(caught.exception.status_code, 409)
        self.assertEqual(main.get_local_transfer_template().read_bytes(), before)
        self.import_model("transferencia", transfer_bytes())

    def test_versioned_transfer_model_passes_layout_validation(self):
        source = fixtures.PROJECT_ROOT / "data" / "planilhas_entrada" / "(MODELO) - TRANSFERENCIA SIMPRESS.xlsx"
        validate_local_model(source, "transferencia")

    def test_binary_and_unsupported_models_are_rejected_before_replacement(self):
        self.import_model("transferencia", transfer_bytes())
        before = main.get_local_transfer_template().read_bytes()
        with zipfile.ZipFile(io.BytesIO(transfer_bytes())) as source:
            content = [(item.filename, source.read(item)) for item in source.infolist()]
        content.append(("xl/embeddings/unsafe.bin", b"synthetic"))
        for data in (b"not an XLSX", archive_bytes(content)):
            with self.subTest(data=data[:20]), self.assertRaises(main.HTTPException) as caught:
                self.import_model("transferencia", data)
            self.assertEqual(caught.exception.status_code, 400)
        self.assertEqual(main.get_local_transfer_template().read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
