"""Presentation labels preserve stored comparison records and source filenames."""

from copy import deepcopy
from datetime import date
import json
from pathlib import Path
import subprocess
import unittest

import openpyxl

from domains.history_export import build_history_export_workbook
from domains.inventory_comparison import (
    PINNED_COMPARISON_EVENT_ID,
    PINNED_COMPARISON_EVENT_TIME,
    _comparison_source_text_for_display,
    comparison_history_for_display,
    filter_update_review_items,
)


class HistoryDisplayLabelTests(unittest.TestCase):
    def test_prepositions_legacy_sources_and_filenames_without_changing_raw_records(self):
        cases = [
            (f"Dados {prep} {source}", f"Dados {prep} {expected}")
            for prep in ("de", "da", "do", "das", "dos", "na", "no", "pela", "pelo", "em")
            for source, name in (
                ("Simpress", "Simpress"), ("Endpoint", "Endpoint"),
            )
            for expected in ("endpoint" if name == "Endpoint" and prep == "no" else name,)
        ]
        cases += [
            ("Simpress", "na Simpress"), ("Endpoint", "no endpoint"),
            ("EquipamentosDetalhes", "na Simpress"), ("InvAll", "no endpoint"),
            ("Série duplicada no InvAll", "Série duplicada no endpoint"),
            ("Série não encontrada na InvAll", "Série não encontrada no endpoint"),
            ("Dados de EquipamentosDetalhes", "Dados na Simpress"),
            ("EquipamentosDetalhes_atual.xlsx", "Simpress_atual.xlsx"),
            ("InvAllComputersSummary_atual.xlsx", "EndpointComputersSummary_atual.xlsx"),
            ("EquipamentosDetalhes (1).xlsx", "Simpress (1).xlsx"),
            ("InvAllComputersSummary (2).xlsx", "EndpointComputersSummary (2).xlsx"),
            ("Simpress.xlsx", "Simpress.xlsx"), ("Endpoint.xlsx", "Endpoint.xlsx"),
            ("endpoint.xlsx", "endpoint.xlsx"), ("simpress.xlsx", "simpress.xlsx"),
            ("historico_EquipamentosDetalhes_atual.xlsx", "historico_Simpress_atual.xlsx"),
            ("Arquivo InvAll-20261017.xlsx", "Arquivo Endpoint-20261017.xlsx"),
            ("Conferir Simpress.", "Conferir na Simpress."),
            ("Conferir Endpoint.", "Conferir no endpoint."),
            ("PRNOTE-ENDPOINT e PRNOTE-SIMPRESS", "PRNOTE-ENDPOINT e PRNOTE-SIMPRESS"),
            ("Planilha Simpress EquipamentosDetalhes_atual.xlsx", "Planilha na Simpress Simpress_atual.xlsx"),
            ("Dados DA SIMPRESS e PELO ENDPOINT", "Dados DA Simpress e PELO Endpoint"),
        ]
        frontend = subprocess.run(
            ["node", "-e", "const fs=require('node:fs');const f=require('./static/js/shared/history-format.js');const values=JSON.parse(fs.readFileSync(0,'utf8'));process.stdout.write(JSON.stringify(values.map(f.formatarNomeFonteHistorico)));"],
            input=json.dumps([raw for raw, _ in cases]),
            cwd=Path(__file__).resolve().parents[1], capture_output=True,
            text=True, encoding="utf-8", check=True,
        )
        self.assertEqual(
            json.loads(frontend.stdout),
            [_comparison_source_text_for_display(raw) for raw, _ in cases],
        )
        for raw, expected in cases:
            with self.subTest(raw=raw):
                event = {
                    "id": "event_InvAll_EquipamentosDetalhes_20261017",
                    "horario": PINNED_COMPARISON_EVENT_TIME,
                    "arquivos_enviados": {"detalhes": raw},
                    "comparacao": {"apelidos_sem_atualizacao": [{"serie": "", "motivo": raw}]},
                }
                original = deepcopy(event)
                self.assertEqual(_comparison_source_text_for_display(raw), expected)
                self.assertEqual(_comparison_source_text_for_display(expected), expected)
                displayed = comparison_history_for_display(event)
                self.assertEqual(displayed["id"], event["id"])
                self.assertEqual(displayed["comparacao"]["apelidos_sem_atualizacao"][0]["motivo"], expected)
                output, _ = build_history_export_workbook([displayed], date(2026, 10, 17))
                workbook = openpyxl.load_workbook(output)
                self.assertEqual(workbook["Histórico"].cell(5, 4).value, event["id"])
                self.assertEqual(workbook["Histórico"].cell(5, 6).value, expected)
                workbook.close()
                self.assertEqual(event, original)

    def test_section_labels_match_actual_screen_formatters_for_current_and_legacy_events(self):
        # Evaluate the existing section functions, so parity is checked against
        # the frontend code rather than a Python copy of its replacements.
        script = r"""
const fs = require('node:fs');
const vm = require('node:vm');
const format = require('./static/js/shared/history-format.js');
const source = fs.readFileSync('static/js/ui/comparison-result-renderer.js', 'utf8');
const functions = source.slice(source.indexOf('  const formatarMotivoRevisao'), source.indexOf('  const atributosLinhaComparacao'));
const labels = JSON.parse(fs.readFileSync(0, 'utf8'));
const headers = source.match(/aliasReviewHeaders: (\[[^\n]+\])/)[1];
const result = vm.runInNewContext(functions + `;({
  review: labels.map(formatarMotivoRevisao),
  absence: labels.map(formatarFonteAusencia),
  ambiguity: labels.map(formatarFonteAmbigua),
  headers: ${headers},
  simpressAbsence: formatarFonteAusencia('EquipamentosDetalhes')
})`, { ...format, labels, permitirAnotacoes: false });
process.stdout.write(JSON.stringify(result));
"""
        labels = ["Endpoint", "InvAll", "no endpoint", "no InvAll", "NO ENDPOINT", "NO INVALL", "Série duplicada NO ENDPOINT"]
        result = subprocess.run(
            ["node", "-e", script], input=json.dumps(labels),
            cwd=Path(__file__).resolve().parents[1], capture_output=True,
            text=True, encoding="utf-8", check=True,
        )
        screen = json.loads(result.stdout)
        for index, label in enumerate(labels):
            with self.subTest(source=label):
                event = {
                    "id": "synthetic-event", "horario": "2026-10-17T10:00:00-03:00",
                    "comparacao": {
                        "apelidos_sem_atualizacao": [{"serie": "S1", "motivo": label}],
                        "ausentes_no_endpoint": [{"serie": "S1", "ausente_em": [label]}],
                        "ausentes_em_detalhes": [{"serie": "S2", "fonte_ausente": "EquipamentosDetalhes"}],
                        "correspondencias_ambiguas": [{
                            "chave": "S1", "fontes": [{"fonte": label}], "mensagem": label,
                        }],
                        "registros_ambiguos": [{"serie": "S1", "fontes": {label: "ambíguo"}, "motivo": label}],
                    },
                }
                original = deepcopy(event)
                output, _ = build_history_export_workbook([event], date(2026, 10, 17))
                workbook = openpyxl.load_workbook(output)
                self.assertEqual(workbook["Apelidos"].cell(4, 5).value, screen["headers"][1])
                self.assertEqual(workbook["Apelidos"].cell(5, 11).value, screen["review"][index])
                self.assertEqual(workbook["Ausências"].cell(5, 3).value, screen["absence"][index])
                self.assertEqual(workbook["Ausências"].cell(6, 3).value, screen["simpressAbsence"])
                for row in (5, 6):
                    self.assertEqual(workbook["Ambiguidades"].cell(row, 4).value, screen["ambiguity"][index])
                    self.assertEqual(workbook["Ambiguidades"].cell(row, 6).value, screen["ambiguity"][index])
                workbook.close()
                self.assertEqual(event, original)

    def test_legacy_preposition_keeps_review_classification_compatible(self):
        row = {"serie": "BRM123", "motivo": "Série não encontrada na InvAll"}
        original = deepcopy(row)
        self.assertEqual(filter_update_review_items([row]), [row])
        self.assertEqual(row, original)

    def test_both_sources_are_bare_in_ambiguities_and_absence_lists(self):
        event = {"comparacao": {
            "ausentes_no_endpoint": [{"ausente_em": ["InvAll", "EquipamentosDetalhes"]}],
            "correspondencias_ambiguas": [{"fontes": [
                {"fonte": "EquipamentosDetalhes", "quantidade": 2}, {"fonte": "InvAll", "serie": "S1"},
            ]}],
            "registros_ambiguos": [{"fontes": {"EquipamentosDetalhes": "ambíguo", "InvAll": "presente"}}],
        }}
        original = deepcopy(event)
        output, _ = build_history_export_workbook([event], date(2026, 10, 17))
        workbook = openpyxl.load_workbook(output)
        self.assertEqual(workbook["Ausências"].cell(5, 3).value, "Endpoint, Simpress")
        self.assertEqual(workbook["Ambiguidades"].cell(5, 4).value, "Simpress (2 registros); Endpoint — série S1")
        self.assertEqual(workbook["Ambiguidades"].cell(6, 4).value, "Simpress")
        workbook.close()
        self.assertEqual(event, original)


    def test_user_changes_export_previous_and_new_user_with_manual_marker(self):
        event = {"comparacao": {"mudancas_usuarios": [{
            "tipo": "manual_voltou", "serie": "S1", "computador": "PRNOTE01",
            "usuario_anterior": "ana.endpoint", "origem_anterior": "endpoint",
            "usuario_novo": "ana.manual", "origem_nova": "manual",
        }]}}
        output, _ = build_history_export_workbook([event], date(2026, 10, 17))
        workbook = openpyxl.load_workbook(output)
        values = [cell.value for cell in workbook["Usuários"][5]]
        workbook.close()
        self.assertIn("Manual voltou a valer", values)
        self.assertIn("ana.endpoint", values)
        self.assertIn("ana.manual (M)", values)

if __name__ == "__main__":
    unittest.main()
