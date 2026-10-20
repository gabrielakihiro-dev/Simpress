import asyncio
import io
import filecmp
import hashlib
import hmac
import json
import logging
import os
import re
import shutil
import threading
import time
from copy import deepcopy
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Any, Set, Tuple

import openpyxl
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from pydantic import BaseModel, Field, field_validator
from starlette.datastructures import UploadFile as StarletteUploadFile
from routers.frontend import RevalidatingStaticFiles, frontend_response
from core.integrity import (
    AtomicFileTransaction,
    FileBusyError,
    FileTransactionError,
    TransactionConflictError,
    WorkbookSafetyError,
    lock_paths,
    load_workbook_safely,
    recover_pending_transactions,
    stage_workbook_update,
    workbook_snapshot,
)
from core.errors import DomainError
from core.config import (
    ALIAS_HISTORY_FILE,
    AUTOMATIC_INVENTORY_IMPORT_POLL_SECONDS,
    AUTOMATIC_INVENTORY_INBOX_DIR,
    AUTOMATIC_TICKET_IMPORT_POLL_SECONDS,
    AUTOMATIC_TICKET_INBOX_DIR,
    BASE_DIR,
    COMPARISON_HISTORY_FILE,
    CUSTOM_UNIDADES_FILE,
    DATA_DIR,
    DETAILS_UPLOAD_FILE,
    ENDPOINT_UPLOAD_FILE,
    EXPORT_MODEL_FILE,
    HISTORY_ANNOTATIONS_DIR,
    HISTORY_ANNOTATIONS_SYNC_ENABLED,
    LOCAL_HISTORY_ANNOTATION_OUTBOX_DIR,
    INPUT_DIR,
    MAX_API_REQUEST_BYTES,
    MAX_UPLOAD_BYTES,
    MAX_UPLOAD_REQUEST_BYTES,
    MAX_CONCURRENT_UPLOAD_REQUESTS,
    MAX_CONCURRENT_WORKBOOK_PARSERS,
    HISTORY_RETENTION_DAYS,
    LOG_DIR,
    MANUAL_USERS_FILE,
    MODIFIED_APELIDOS_FILE,
    REMOVED_ENDPOINT_HOSTNAMES_FILE,
    MONTHLY_REPORT_CHECK_SECONDS,
    MONTHLY_REPORT_FIRST_CHECK_SECONDS,
    MONTHLY_REPORTS_DIR,
    STALE_BASE_ALERT_DAYS,
    STALE_TICKET_ALERT_DAYS,
    STATIC_DIR,
    STATE_DIR,
    TEMPLATE_FILE,
    TICKET_REPORT_FILE,
    TRANSACTION_ROOT,
    TRANSFER_HISTORY_FILE,
    UPLOAD_BACKUPS_DIR,
    SUPPORTED_CONTRACT,
    SUPPORTED_CONTRACT_LEGAL_NAME,
    TRANSFER_ORIGIN_CONTACT_FILE,
)
from routers.inventory_history import (
    ComparisonAnnotationPayload,
    ComparisonRestorePayload,
    InventoryHistoryServices,
    ManualAliasConflictResolutionPayload,
    create_inventory_history_router,
)
from routers.transfers import (
    TransferenciaPayload,
    create_transfer_router,
    formatar_telefone_brasileiro,
)
from routers.transfer_history import TransferHistoryServices, create_transfer_history_router
from routers.dashboard import DashboardServices, create_dashboard_router
from routers.operations import MAX_TICKET_WORKBOOK_BYTES, create_operations_router
from routers.local_setup import create_local_setup_router
from routers.monthly_report import (
    MonthlyReportServices,
    create_monthly_report_router,
    generate_previous_month_report,
)
from domains.local_setup import load_origin_contact, validate_local_model
from core.logging_setup import configure_logging
from core.xlsx_safety import (
    XlsxLimitExceeded,
    XlsxPackageError,
    set_cell_value_as_text,
    validate_xlsx,
)

from domains.inventory_comparison import (
    COMPARISON_SOURCE_LABELS,
    CURITIBA_HOSTNAME_PREFIXES,
    PINNED_COMPARISON_EVENT_ID,
    PINNED_COMPARISON_EVENT_TIME,
    SEBRAE_OWNED_NOTEBOOK_SERIALS,
    SERIE_LOCAL_MAP,
    _comparison_source_records,
    _record_value,
    _unique_upload_records,
    _upload_alias_value,
    _upload_serial_value,
    _uploaded_aliases_by_serial,
    comparison_annotation_item_id,
    compare_uploaded_workbooks,
    comparison_history_for_display,
    comparison_review_item_id,
    filter_curitiba_hostname_location_mismatches,
    filter_sebrae_owned_unmatched_items,
    filter_update_review_items,
    filter_update_unmatched_details,
    filter_update_unmatched_series,
    is_missing_cell_value,
    is_curitiba_hostname_exception,
    is_prnoteti_curitiba_location_match,
    is_sebrae_owned_notebook_serial,
    location_for_computer_name,
    normalize_reference_value,
)
from domains.inventory_inbox import (
    AutomaticInventoryInbox,
    automatic_inventory_import_worker as run_automatic_inventory_import_worker,
)
from domains.ticket_inbox import (
    AutomaticTicketInbox,
    automatic_ticket_import_worker as run_automatic_ticket_import_worker,
)
from domains.history_annotations import (
    build_history_annotation_record,
    history_annotation_row_fingerprint,
    history_annotation_record_path,
    load_history_annotations,
    save_history_annotation,
)
from domains import effective_alias
from domains import alias_changes, equipment_export, inventory_state, user_changes
from domains.alias_changes import (
    compare_equipment_aliases,
    snapshot_equipment_aliases,
    validate_equipment_serial as _validate_equipment_serial,
)
from domains.inventory_state import (
    STATE_LOCK,
    _coerce_unit_id,
    _unit_identity_key,
    format_cep,
    get_suggested_description,
    normalize_contract_value,
)
from domains import readiness, uploaded_workbooks
from domains.uploaded_workbooks import (
    DETAILS_ALIAS_HEADER_ALIASES,
    DETAILS_COMPLETE_SOURCE_COLUMN_COUNT,
    DETAILS_INSTALLATION_CNPJ_HEADER_ALIASES,
    DETAILS_REQUIRED_HEADERS,
    DETAILS_SERIAL_HEADER_ALIASES,
    WorkbookHeaderError,
    normalize_workbook_header,
    select_upload_table,
    validate_unambiguous_workbook_headers,
)
from domains.inventory_update import InventoryUpdateServices
from domains.usage_metrics import UsageMetricsMiddleware, UsageMetricsStore
from domains.inventory_update import preview_inventory_files
from domains.inventory_update import update_inventory_files
from domains.equipment_search import (
    CATEGORIES as _CATEGORIAS_EQUIPAMENTO,
    build_search_index,
    categorize_equipment,
    normalize_search_term,
    search_normalized_index as search_normalized_equipment_index,
)
from routers.equipment_search import EquipmentSearchServices, create_equipment_search_router
from routers.system import SystemRouteServices, create_system_router
from routers.units import UnidadePayload, UnitRouteServices, create_units_router
from routers.aliases import ApelidoUndoPayload, ApelidoUpdatePayload, UsuarioUpdatePayload, create_alias_router
from routers.endpoint_removals import create_endpoint_removals_router
from domains import endpoint_removals
from core.inventory_requests import (
    InventoryUpdateMiddleware,
    idempotency_key,
)
from security.local_requests import (
    LocalRequestSafetyMiddleware as _LocalRequestSafetyMiddleware,
    LocalRequestSafetyServices,
    RequestBodyLimits,
    RequestBodyTooLarge,
    allowed_local_authorities,
    parse_request_authority,
)

# Runtime synchronization and in-memory state remain in the application entry point.
# Lock hierarchy (docs/arquitetura/auditoria_backend/hierarquia_de_locks.md): the global
# ".atomic-transactions" lock, then file locks, then these Python locks as leaves.
# Hold a leaf lock only around memory or a single file read; never open a
# transaction, call lock_paths or load_database while holding one.
# Read-only routes polled by the page give up early (409 + Retry-After) instead of
# holding a worker thread for the full 15 s while an import owns the global lock.
# STATE_LOCK guards the inventory caches in domains/inventory_state.py.
READ_LOCK_TIMEOUT_SECONDS = 3.0
MODIFIED_APELIDOS_LOCK = threading.RLock()
ALIAS_HISTORY_LOCK = threading.RLock()
COMPARISON_HISTORY_LOCK = threading.RLock()
UPLOAD_ADMISSION = threading.BoundedSemaphore(MAX_CONCURRENT_UPLOAD_REQUESTS)
WORKBOOK_PARSE_ADMISSION = threading.BoundedSemaphore(MAX_CONCURRENT_WORKBOOK_PARSERS)
logger = logging.getLogger(__name__)


class LocalJsonCorruptionError(ValueError):
    """A persisted JSON file is present but cannot be trusted or updated safely."""


def _read_json_object(path: Path, label: str) -> Dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise LocalJsonCorruptionError(f"O arquivo local de {label} está ilegível; os dados foram preservados.") from error
    if not isinstance(value, dict) or any(not isinstance(key, str) or not isinstance(item, str) for key, item in value.items()):
        raise LocalJsonCorruptionError(f"O arquivo local de {label} precisa conter um objeto com chaves e valores textuais.")
    return value


def _validate_json_object(value: Any, label: str) -> Dict[str, Any]:
    if not isinstance(value, dict) or any(not isinstance(key, str) or not isinstance(item, str) for key, item in value.items()):
        raise LocalJsonCorruptionError(f"O arquivo local de {label} precisa conter um objeto com chaves e valores textuais.")
    return value


def _read_json_history(path: Path, label: str) -> List[Dict[str, Any]]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise LocalJsonCorruptionError(f"O histórico local de {label} está ilegível; os dados foram preservados.") from error
    if not isinstance(value, list) or any(not isinstance(item, dict) for item in value):
        raise LocalJsonCorruptionError(f"O histórico local de {label} precisa conter uma lista de registros.")
    return value
UPLOAD_PREVIEW_TOKEN_SECRET = os.urandom(32)
UPLOAD_PREVIEW_TOKEN_TTL_SECONDS = 10 * 60
INVENTORY_UPLOAD_PATHS = frozenset({
    "/api/planilhas/previsualizar",
    "/api/planilhas/atualizar",
})


@contextmanager
def inventory_state_transaction(targets):
    """Keep an in-memory inventory snapshot aligned with a file transaction."""
    with AtomicFileTransaction(TRANSACTION_ROOT, targets) as transaction:
        with STATE_LOCK:
            state_before = inventory_state.snapshot_state()
        try:
            yield transaction
        except BaseException:
            with STATE_LOCK:
                inventory_state.restore_state(state_before)
            raise


def _json_bytes(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, indent=2).encode("utf-8")


def load_modified_apelidos() -> Dict[str, str]:
    with MODIFIED_APELIDOS_LOCK:
        if not MODIFIED_APELIDOS_FILE.exists():
            return {}
        return _read_json_object(MODIFIED_APELIDOS_FILE, "apelidos")

def load_removed_endpoint_hostnames() -> List[Dict[str, Any]]:
    with MODIFIED_APELIDOS_LOCK:
        if not REMOVED_ENDPOINT_HOSTNAMES_FILE.exists():
            return []
        return _read_json_history(REMOVED_ENDPOINT_HOSTNAMES_FILE, "hostnames removidos")


def removed_endpoint_hostname_keys() -> frozenset:
    """Pares (série, hostname) descartados; um arquivo ilegível não bloqueia o inventário."""
    try:
        return endpoint_removals.removal_keys(load_removed_endpoint_hostnames())
    except LocalJsonCorruptionError as error:
        logger.warning("Hostnames removidos ignorados: %s", error)
        return frozenset()


def load_manual_users() -> Dict[str, str]:
    with MODIFIED_APELIDOS_LOCK:
        if not MANUAL_USERS_FILE.exists():
            return {}
        return _read_json_object(MANUAL_USERS_FILE, "usuários manuais")

def load_alias_history() -> List[Dict[str, Any]]:
    with ALIAS_HISTORY_LOCK:
        if not ALIAS_HISTORY_FILE.exists():
            return []
        try:
            return _read_json_history(ALIAS_HISTORY_FILE, "apelidos")
        except LocalJsonCorruptionError as error:
            raise HTTPException(status_code=503, detail=str(error)) from error

def load_comparison_history() -> List[Dict[str, Any]]:
    with COMPARISON_HISTORY_LOCK:
        if not COMPARISON_HISTORY_FILE.exists():
            return []
        try:
            return _read_json_history(COMPARISON_HISTORY_FILE, "comparações")
        except LocalJsonCorruptionError as error:
            raise HTTPException(status_code=503, detail=str(error)) from error


def pinned_comparison_history_entry(history: Optional[List[Dict[str, Any]]] = None) -> Optional[Dict[str, Any]]:
    entries = load_comparison_history() if history is None else history
    if not isinstance(entries, list):
        return None
    return next(
        (
            entry for entry in entries
            if isinstance(entry, dict)
            and entry.get("id") == PINNED_COMPARISON_EVENT_ID
            and entry.get("horario") == PINNED_COMPARISON_EVENT_TIME
        ),
        None,
    )


def latest_comparison_update_entry(history: Optional[List[Dict[str, Any]]] = None) -> Optional[Dict[str, Any]]:
    """Returns the newest comparison update while leaving the pinned baseline available."""
    entries = load_comparison_history() if history is None else history
    candidates = [
        entry for entry in entries
        if isinstance(entry, dict)
        and entry.get("id") != PINNED_COMPARISON_EVENT_ID
        and entry.get("tipo") != "restauracao"
        and not (
            isinstance(entry.get("comparacao"), dict)
            and entry["comparacao"].get("tipo_operacao") == "restauracao"
        )
    ] if isinstance(entries, list) else []
    if not candidates:
        return None

    def event_timestamp(entry: Dict[str, Any]) -> float:
        try:
            return datetime.fromisoformat(str(entry.get("horario") or "").replace("Z", "+00:00")).timestamp()
        except (TypeError, ValueError, OverflowError):
            return float("-inf")

    return max(candidates, key=event_timestamp)


def _equipment_revision(serial: Any, modelo: Any, tipo: Any, apelido: Any, id_cliente: Any) -> str:
    values = []
    for index, value in enumerate((serial, modelo, tipo, apelido, id_cliente)):
        if value is None:
            values.append("")
            continue
        text = str(value).strip()
        if index == 4 and re.fullmatch(r"\d+\.0+", text):
            text = text.split(".", 1)[0]
        values.append(text)
    payload = json.dumps(values, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


_ensure_workbook_readable = uploaded_workbooks.ensure_workbook_readable


def _read_details_sheet_records(file_path: Path, sheet_name: str):
    """Read inventory rows with openpyxl, preserving the cells used by the application."""
    return uploaded_workbooks.read_details_sheet_records(_workbook_reader_services, file_path, sheet_name)


def find_details_sheet(file_path: Path) -> Optional[str]:
    """Retorna uma aba EquipamentosDetalhes com os campos usados pelo site.

    Levanta FileBusyError quando o arquivo existe mas não pode ser lido; ``None``
    significa somente que o arquivo não tem uma aba de inventário válida.
    """
    return uploaded_workbooks.find_details_sheet(_workbook_reader_services, file_path)


def validate_xlsx_package(file_bytes: bytes) -> None:
    """Rejects non-XLSX, malformed, or excessively large ZIP packages before parsing."""
    uploaded_workbooks.validate_xlsx_package(file_bytes, validate_xlsx)


def read_workbook_tables(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Lê abas não vazias de um XLSX e retorna cabeçalhos normalizados e linhas."""
    return uploaded_workbooks.read_workbook_tables(_workbook_reader_services, file_bytes)


_workbook_reader_services = uploaded_workbooks.WorkbookReaderServices(
    validate_xlsx=lambda *args, **kwargs: validate_xlsx(*args, **kwargs),
    validate_xlsx_package=lambda file_bytes: validate_xlsx_package(file_bytes),
    ensure_workbook_readable=lambda file_path: _ensure_workbook_readable(file_path),
    logger=logger,
)


def get_details_files() -> List[Path]:
    """Lista arquivos EquipamentosDetalhes enviados ou mantidos na pasta de entrada."""
    if not INPUT_DIR.is_dir():
        return []
    return sorted(
        (
            path for path in INPUT_DIR.glob("EquipamentosDetalhes*.xlsx")
            if path.is_file() and not path.name.startswith("~$")
        ),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )


def get_primary_equipment_file() -> Optional[Path]:
    """Retorna o EquipamentosDetalhes válido mais recente, fonte única do cadastro.

    Um candidato ilegível interrompe a busca com FileBusyError em vez de deixar
    um arquivo mais antigo assumir a base ativa.
    """
    return next(
        (path for path in get_details_files() if find_details_sheet(path) is not None),
        None,
    )


def _check_startup_xlsx(path: Optional[Path]) -> Dict[str, Any]:
    """Check the local XLSX container without loading or logging workbook data."""
    return readiness.check_startup_xlsx(path, validate_xlsx)


_READY_WORKBOOK_CACHE = readiness.ReadyWorkbookCache()
_READINESS_VALIDATION_CACHE = _READY_WORKBOOK_CACHE.entries
_READINESS_VALIDATION_LOCK = _READY_WORKBOOK_CACHE.lock
_READINESS_REPORT_CACHE = readiness.ReadinessReportCache()


def get_local_transfer_template() -> Path:
    override = INPUT_DIR / "_modelos" / "transferencia.xlsx"
    return override if override.exists() else TEMPLATE_FILE


def _check_ready_workbook(path: Path, role: str) -> Dict[str, Any]:
    return readiness.check_ready_workbook(path, role, _ready_workbook_services, _READY_WORKBOOK_CACHE)


_check_startup_ticket_report = readiness.check_startup_ticket_report


def _readiness_loaded_state() -> Tuple[bool, bool]:
    with STATE_LOCK:
        return bool(inventory_state.equipamentos_cache), bool(inventory_state.unidades_cache)


def _validate_local_state_for_readiness() -> None:
    load_modified_apelidos()
    load_custom_units()
    if ALIAS_HISTORY_FILE.exists():
        load_alias_history()
    if COMPARISON_HISTORY_FILE.exists():
        load_comparison_history()


def _refresh_export_status_for_readiness() -> Dict[str, Any]:
    # Readiness reads without the transaction lock; files are replaced atomically.
    count_export_changes(locked=True)
    return get_export_status()


def _readiness_signature() -> Tuple[Any, ...]:
    """Everything the readiness report reads, so a cached report is never stale."""
    try:
        details_files: Any = tuple(get_details_files())
    except OSError as error:
        details_files = type(error).__name__
    watched = (
        TICKET_REPORT_FILE,
        DETAILS_UPLOAD_FILE,
        ENDPOINT_UPLOAD_FILE,
        EXPORT_MODEL_FILE,
        MODIFIED_APELIDOS_FILE,
        CUSTOM_UNIDADES_FILE,
        ALIAS_HISTORY_FILE,
        COMPARISON_HISTORY_FILE,
        INPUT_DIR / "_modelos" / "transferencia.xlsx",
        TEMPLATE_FILE,
    )
    details_signature = (
        readiness.file_signature(details_files) if isinstance(details_files, tuple) else details_files
    )
    return (
        readiness.file_signature(watched),
        details_signature,
        _readiness_loaded_state(),
        get_transaction_recovery_block(),
    )


def get_startup_readiness() -> Dict[str, Any]:
    """Readiness without the global transaction lock, cached while its inputs stay the same."""
    return _READINESS_REPORT_CACHE.get_or_compute(_readiness_signature(), _get_startup_readiness_unlocked)


def _get_startup_readiness_unlocked() -> Dict[str, Any]:
    """Map essential local files to the application functions that need them."""
    return readiness.build_readiness_report(_readiness_services)


_ready_workbook_services = readiness.ReadyWorkbookServices(
    validate_local_model=lambda *args, **kwargs: validate_local_model(*args, **kwargs),
    validate_uploaded_workbook=lambda *args, **kwargs: validate_uploaded_workbook(*args, **kwargs),
    max_upload_bytes=lambda: MAX_UPLOAD_BYTES,
    logger=logger,
)

_readiness_services = readiness.ReadinessServices(
    ticket_report_file=lambda: TICKET_REPORT_FILE,
    details_upload_file=lambda: DETAILS_UPLOAD_FILE,
    endpoint_upload_file=lambda: ENDPOINT_UPLOAD_FILE,
    export_model_file=lambda: EXPORT_MODEL_FILE,
    get_details_files=lambda: get_details_files(),
    get_primary_equipment_file=lambda: get_primary_equipment_file(),
    get_local_transfer_template=lambda: get_local_transfer_template(),
    check_startup_xlsx=lambda path: _check_startup_xlsx(path),
    check_ready_workbook=lambda path, role: _check_ready_workbook(path, role),
    check_startup_ticket_report=lambda path: _check_startup_ticket_report(path),
    loaded_state=lambda: _readiness_loaded_state(),
    validate_local_state=lambda: _validate_local_state_for_readiness(),
    local_state_errors=(LocalJsonCorruptionError, HTTPException),
    refresh_export_status=lambda: _refresh_export_status_for_readiness(),
    get_transaction_recovery_block=lambda: get_transaction_recovery_block(),
    logger=logger,
)


def report_startup_readiness() -> None:
    """Log a clear feature impact summary without exposing workbook contents."""
    for line in readiness.readiness_log_lines(get_startup_readiness()):
        logger.info("%s", line)


def _report_startup_readiness_best_effort() -> None:
    """Log feature readiness when possible without making it a startup requirement."""
    try:
        report_startup_readiness()
    except Exception:
        logger.exception(
            "Não foi possível calcular a prontidão local; o servidor continuará em execução."
        )


# Message of the startup recovery that refused to touch a diverged file; None when writes are allowed.
_TRANSACTION_RECOVERY_BLOCK: Optional[str] = None
_TRANSACTION_RECOVERY_BLOCK_LOCK = threading.Lock()


def _recover_pending_transactions_at_startup() -> bool:
    """Recover interrupted writes, or start read-only when the journal no longer matches the files."""
    global _TRANSACTION_RECOVERY_BLOCK
    try:
        recover_pending_transactions(TRANSACTION_ROOT)
    except TransactionConflictError as error:
        message = (
            f"{error} As gravações ficam bloqueadas até a conferência manual da pasta "
            f"'.transactions'; os dados exibidos podem refletir uma gravação incompleta."
        )
        with _TRANSACTION_RECOVERY_BLOCK_LOCK:
            _TRANSACTION_RECOVERY_BLOCK = message
        logger.error("Recuperação de gravação pendente indisponível: %s", message)
        return False
    with _TRANSACTION_RECOVERY_BLOCK_LOCK:
        _TRANSACTION_RECOVERY_BLOCK = None
    return True


def _has_pending_transaction_journal() -> bool:
    try:
        return any(
            (entry / "manifest.json").is_file()
            for entry in (TRANSACTION_ROOT / ".transactions").iterdir()
            if entry.is_dir()
        )
    except FileNotFoundError:
        return False
    except OSError:
        # An unreadable journal folder cannot prove that recovery was resolved.
        return True


def get_transaction_recovery_block() -> Optional[str]:
    """Return why writes are blocked, clearing the block once the journal was resolved manually."""
    global _TRANSACTION_RECOVERY_BLOCK
    with _TRANSACTION_RECOVERY_BLOCK_LOCK:
        message = _TRANSACTION_RECOVERY_BLOCK
        if message is None:
            return None
        # While blocked no transaction can open, so any remaining journal is the unresolved one.
        if not _has_pending_transaction_journal():
            _TRANSACTION_RECOVERY_BLOCK = None
            return None
        return message


def _raise_if_transaction_recovery_blocked() -> None:
    """Keep automatic batches pending (503 is transient for the inboxes) instead of failing them."""
    if (message := get_transaction_recovery_block()) is not None:
        raise HTTPException(status_code=503, detail=message)


def _as_domain_error(error: HTTPException) -> DomainError:
    """Hand an HTTP error to the automatic inboxes, which only know domain errors."""
    return DomainError(error.detail, status_code=error.status_code)


def inventory_update_preview_summary(
    current_tables: Dict[str, Dict[str, Any]],
    incoming_tables: Dict[str, Dict[str, Any]],
) -> Dict[str, Dict[str, int]]:
    """Return counts only; never include row values or serials in an upload preview."""
    summary = {}
    for role in ("endpoint", "detalhes"):
        def count_serials(table):
            counts: Dict[str, int] = {}
            for record in table.get("records", []):
                serial_key = normalize_reference_value(_upload_serial_value(record, role))
                if serial_key:
                    counts[serial_key] = counts.get(serial_key, 0) + 1
            return counts

        current_counts = count_serials(current_tables[role])
        incoming_counts = count_serials(incoming_tables[role])
        current_serials = set(current_counts)
        incoming_serials = set(incoming_counts)
        summary[role] = {
            "linhas_atuais": len(current_tables[role].get("records", [])),
            "linhas_novas": len(incoming_tables[role].get("records", [])),
            "series_novas": len(incoming_serials - current_serials),
            "series_removidas": len(current_serials - incoming_serials),
            "series_duplicadas_no_arquivo_novo": sum(
                1 for count in incoming_counts.values() if count > 1
            ),
        }
    return summary


def _inventory_preview_token_payload(issued_at: int, staged, current_bytes) -> bytes:
    payload = {
        "versao": 1,
        "emitido_em": issued_at,
        "uploads": {
            role: hashlib.sha256(staged[role]["bytes"]).hexdigest()
            for role in ("endpoint", "detalhes")
        },
        "bases_atuais": {
            role: hashlib.sha256(current_bytes[role]).hexdigest()
            for role in ("endpoint", "detalhes")
        },
    }
    return json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")


def create_inventory_preview_token(staged, current_bytes) -> str:
    issued_at = int(time.time())
    payload = _inventory_preview_token_payload(issued_at, staged, current_bytes)
    signature = hmac.new(UPLOAD_PREVIEW_TOKEN_SECRET, payload, hashlib.sha256).hexdigest()
    return f"{issued_at}.{signature}"


def verify_inventory_preview_token(token, staged, current_bytes) -> bool:
    try:
        issued_at_text, supplied_signature = str(token).split(".", 1)
        issued_at = int(issued_at_text)
        age = int(time.time()) - issued_at
        if age < 0 or age > UPLOAD_PREVIEW_TOKEN_TTL_SECONDS:
            return False
        payload = _inventory_preview_token_payload(issued_at, staged, current_bytes)
        expected_signature = hmac.new(
            UPLOAD_PREVIEW_TOKEN_SECRET, payload, hashlib.sha256
        ).hexdigest()
        return hmac.compare_digest(supplied_signature, expected_signature)
    except (AttributeError, TypeError, ValueError):
        return False


def _manual_alias_source_snapshots(
    overrides: Dict[str, str],
    endpoint_table: Dict[str, Any],
    details_table: Dict[str, Any],
) -> List[Dict[str, Any]]:
    manual_aliases = _manual_aliases_by_serial(overrides)
    endpoint_rows = _unique_upload_records(endpoint_table, "endpoint")
    details_rows = _unique_upload_records(details_table, "detalhes")
    divergences = []

    for serial_key, manual_alias in manual_aliases.items():
        endpoint_record = endpoint_rows.get(serial_key)
        details_record = details_rows.get(serial_key)
        if endpoint_record is None and details_record is None:
            continue

        # O marcador (M) conta na comparação: todo apelido manual presente nas planilhas
        # enviadas fica pendente no histórico até o usuário decidir manter ou excluir.
        invall_alias = _upload_alias_value(endpoint_record, "endpoint") if endpoint_record else ""
        details_alias = _upload_alias_value(details_record, "detalhes") if details_record else ""
        details_differs = (
            details_record is not None
            and normalize_reference_value(details_alias) != normalize_reference_value(manual_alias)
        )
        invall_differs = (
            endpoint_record is not None
            and normalize_reference_value(invall_alias) != normalize_reference_value(manual_alias)
        )
        source_serial = (
            _upload_serial_value(details_record, "detalhes") if details_record else ""
        ) or (
            _upload_serial_value(endpoint_record, "endpoint") if endpoint_record else ""
        ) or serial_key
        divergences.append({
            "serie": str(source_serial).strip(),
            "apelido_manual": manual_alias,
            "apelido_equipamentos_detalhes": details_alias,
            "apelido_invall": invall_alias,
            "detalhes_encontrado": details_record is not None,
            "invall_encontrado": endpoint_record is not None,
            "diverge_detalhes": details_differs,
            "diverge_invall": invall_differs,
            "resolucao": "pendente",
        })

    return sorted(divergences, key=lambda item: normalize_reference_value(item["serie"]))


def _manual_alias_divergences(
    overrides: Dict[str, str],
    endpoint_table: Dict[str, Any],
    details_table: Dict[str, Any],
) -> List[Dict[str, Any]]:
    return [
        item for item in _manual_alias_source_snapshots(overrides, endpoint_table, details_table)
        if item["diverge_detalhes"] or item["diverge_invall"]
    ]


def validate_uploaded_workbook(file_bytes: bytes, role: str) -> Dict[str, Any]:
    return select_upload_table(read_workbook_tables(file_bytes), role)


def load_export_model_aliases() -> Dict[str, str]:
    """Lê os apelidos do modelo fixo usado como base de exportação."""
    return equipment_export.load_export_model_aliases(_equipment_export_services)


def load_endpoint_aliases() -> Dict[str, str]:
    """Lê somente séries únicas com hostname preenchido na InvAll ativa."""
    return read_endpoint_alias_index()["unique"]


def read_endpoint_alias_index() -> Dict[str, Any]:
    """Retorna associações unívocas e linhas que precisam de conferência."""
    index = effective_alias.read_endpoint_alias_index(_effective_alias_services, removed_endpoint_hostname_keys())
    try:
        index["manual_users"] = user_changes.manual_users_by_serial(load_manual_users())
    except LocalJsonCorruptionError as error:
        # Um arquivo manual ilegível não pode impedir o carregamento do inventário.
        logger.warning("Usuários manuais ignorados: %s", error)
        index["manual_users"] = {}
    return index


endpoint_alias_review_items = effective_alias.endpoint_alias_review_items
_manual_alias_for_serial = effective_alias.manual_alias_for_serial
_manual_aliases_by_serial = effective_alias.manual_aliases_by_serial


def _apply_effective_alias(
    equipment: Dict[str, Any],
    endpoint_index: Optional[Dict[str, Any]] = None,
    *,
    details_duplicate: bool = False,
) -> None:
    if endpoint_index is None:
        with lock_paths([TRANSACTION_ROOT / ".atomic-transactions"], timeout=READ_LOCK_TIMEOUT_SECONDS):
            with lock_paths([ENDPOINT_UPLOAD_FILE], timeout=READ_LOCK_TIMEOUT_SECONDS):
                endpoint_index = read_endpoint_alias_index()
    effective_alias.apply_effective_alias(equipment, endpoint_index, details_duplicate=details_duplicate)


_effective_alias_services = effective_alias.EffectiveAliasServices(
    endpoint_upload_file=lambda: ENDPOINT_UPLOAD_FILE,
    validate_uploaded_workbook=lambda *args, **kwargs: validate_uploaded_workbook(*args, **kwargs),
    logger=logger,
)


# Late-bound so tests that patch these main attributes keep affecting the export.
_equipment_export_services = equipment_export.EquipmentExportServices(
    export_model_file=lambda: EXPORT_MODEL_FILE,
    endpoint_upload_file=lambda: ENDPOINT_UPLOAD_FILE,
    modified_aliases_file=lambda: MODIFIED_APELIDOS_FILE,
    transaction_root=lambda: TRANSACTION_ROOT,
    read_lock_timeout=lambda: READ_LOCK_TIMEOUT_SECONDS,
    get_parque_file=lambda: get_parque_file(),
    get_details_files=lambda: get_details_files(),
    find_details_sheet=lambda file_path: find_details_sheet(file_path),
    read_endpoint_alias_index=lambda: read_endpoint_alias_index(),
    load_modified_apelidos=lambda: load_modified_apelidos(),
    latest_comparison_update_entry=lambda: latest_comparison_update_entry(),
    validate_xlsx=lambda file_path: validate_xlsx(file_path),
    lock_paths=lambda *args, **kwargs: lock_paths(*args, **kwargs),
    logger=logger,
    external_errors=(HTTPException,),
    removed_hostnames_file=lambda: REMOVED_ENDPOINT_HOSTNAMES_FILE,
)


def get_export_status() -> Dict[str, Any]:
    return equipment_export.get_export_status()


def count_export_changes(
    locked: bool = False,
    endpoint_index: Optional[Dict[str, Any]] = None,
) -> Dict[str, int]:
    """Count writable alias and IdCliente cells that export would change."""
    return equipment_export.count_export_changes(_equipment_export_services, locked, endpoint_index)


def get_equipamentos_template_file() -> Optional[Path]:
    """Retorna a planilha EquipamentosDetalhes usada pelo site."""
    return get_primary_equipment_file()

def get_parque_file() -> Optional[Path]:
    """Retorna a fonte EquipamentosDetalhes carregada pelo site."""
    return get_equipamentos_template_file()

@asynccontextmanager
async def lifespan(app: FastAPI):
    def prepare_local_state() -> None:
        """Recover, migrate, prune and load disk state in one worker thread."""
        # A diverged journal keeps the server up in read-only mode; maintenance writes wait.
        if _recover_pending_transactions_at_startup():
            try:
                migration_result = migrate_legacy_history_annotations()
                if migration_result.get("legadas_enfileiradas"):
                    logger.info(
                        "Anotações legadas preparadas para sincronização: %s.",
                        migration_result,
                    )
            except Exception:
                logger.exception("Falha ao migrar anotações legadas do histórico de planilhas.")
            try:
                prune_result = prune_expired_histories()
                if any(prune_result.values()):
                    logger.info("Registros e backups removidos conforme política de %s dias: %s.", HISTORY_RETENTION_DAYS, prune_result)
            except Exception as error:
                logger.warning("A retenção automática não pôde ser concluída (%s); os arquivos foram preservados.", type(error).__name__)
        try:
            load_database()
        except LocalJsonCorruptionError as error:
            logger.error("Estado local preservado: %s", error)
        _report_startup_readiness_best_effort()

    # Workers start only after the base is loaded, as before; disk work stays off the loop.
    await asyncio.to_thread(prepare_local_state)
    background_tasks: List[asyncio.Task] = []
    try:
        background_tasks.extend((
            asyncio.create_task(_history_annotation_sync_worker()),
            asyncio.create_task(_automatic_inventory_import_worker()),
            asyncio.create_task(_automatic_ticket_import_worker()),
            asyncio.create_task(_external_change_reload_worker()),
            asyncio.create_task(_monthly_report_worker()),
        ))
        yield
    finally:
        for task in background_tasks:
            task.cancel()
        await asyncio.gather(
            *background_tasks,
            return_exceptions=True,
        )


def reload_database_if_changed_externally() -> bool:
    """Reload the cache when the active files differ from the published ones."""
    if get_transaction_recovery_block() is not None:
        return False
    with STATE_LOCK:
        loaded = inventory_state.loaded_files_signature
    if loaded is None:
        return False
    try:
        if inventory_state.inventory_files_signature(_inventory_state_services) == loaded:
            return False
    except OSError:
        return False
    logger.info("Planilhas ou configurações alteradas fora do site; recarregando a base.")
    try:
        return bool(load_database())
    except (FileBusyError, LocalJsonCorruptionError) as error:
        logger.warning("Recarga adiada (%s); nova tentativa no próximo ciclo.", type(error).__name__)
        return False


async def _external_change_reload_worker() -> None:
    while True:
        await asyncio.sleep(EXTERNAL_CHANGE_POLL_SECONDS)
        try:
            await asyncio.to_thread(reload_database_if_changed_externally)
        except Exception:
            logger.exception("Falha inesperada ao conferir alterações externas da base.")


def _flush_pending_history_annotations_unless_blocked() -> None:
    # The outbox keeps every note; retrying against a blocked journal would only repeat the error.
    if get_transaction_recovery_block() is None:
        flush_pending_history_annotations()


async def _history_annotation_sync_worker() -> None:
    while True:
        await asyncio.sleep(60)
        try:
            await asyncio.to_thread(_flush_pending_history_annotations_unless_blocked)
        except Exception:
            logger.exception("Falha inesperada ao tentar sincronizar anotações pendentes.")


async def _automatic_inventory_import_worker() -> None:
    inbox = AutomaticInventoryInbox(
        root=AUTOMATIC_INVENTORY_INBOX_DIR,
        max_upload_bytes=MAX_UPLOAD_BYTES,
        load_history=load_comparison_history,
        import_batch=_import_automatic_inventory_batch,
    )
    logger.info("Entrada automática de planilhas ativa em %s.", AUTOMATIC_INVENTORY_INBOX_DIR)
    await run_automatic_inventory_import_worker(
        inbox,
        poll_interval_seconds=AUTOMATIC_INVENTORY_IMPORT_POLL_SECONDS,
    )


async def _automatic_ticket_import_worker() -> None:
    endpoint = next((
        route.endpoint
        for route in operations_router.routes
        if getattr(route, "path", None) == "/api/operacoes/chamados/importar"
        and "POST" in getattr(route, "methods", set())
    ), None)
    if endpoint is None:
        raise RuntimeError("A rota de importação de chamados não está registrada.")

    def import_batch(batch_id: str, simpress_name: str, simpress_bytes: bytes, oracle_name: str, oracle_bytes: bytes):
        try:
            _raise_if_transaction_recovery_blocked()
            simpress_upload = StarletteUploadFile(filename=simpress_name, file=io.BytesIO(simpress_bytes))
            oracle_upload = StarletteUploadFile(filename=oracle_name, file=io.BytesIO(oracle_bytes))
            return endpoint(
                simpress_file=simpress_upload,
                oracle_file=oracle_upload,
                automatic_import_id=batch_id,
            )
        except HTTPException as error:
            raise _as_domain_error(error) from error

    inbox = AutomaticTicketInbox(
        root=AUTOMATIC_TICKET_INBOX_DIR,
        report_path=TICKET_REPORT_FILE,
        max_upload_bytes=MAX_TICKET_WORKBOOK_BYTES,
        import_batch=import_batch,
    )
    logger.info("Entrada automática de chamados ativa em %s.", AUTOMATIC_TICKET_INBOX_DIR)
    await run_automatic_ticket_import_worker(
        inbox,
        poll_interval_seconds=AUTOMATIC_TICKET_IMPORT_POLL_SECONDS,
    )


# Interactive docs expose the API contract to any local process; developers opt in.
API_DOCS_ENABLED = os.environ.get("SIMPRESS_ENABLE_DOCS", "").strip() == "1"

app = FastAPI(
    title="Simpress Transferência de Equipamentos - SEBRAE/PR",
    description="Sistema inteligente para emissão automatizada de planilhas oficiais de transferência de equipamentos.",
    version="1.3.0",
    lifespan=lifespan,
    docs_url="/docs" if API_DOCS_ENABLED else None,
    redoc_url="/redoc" if API_DOCS_ENABLED else None,
    openapi_url="/openapi.json" if API_DOCS_ENABLED else None,
)


@app.exception_handler(DomainError)
async def domain_error_handler(_request: Request, error: DomainError) -> JSONResponse:
    """Answer domain errors with the same body FastAPI produces for HTTPException."""
    return JSONResponse({"detail": error.detail}, status_code=error.status_code)

app.add_middleware(InventoryUpdateMiddleware)

# Daily counts only (no serials, hostnames or terms); see docs/operacao/metricas-uso.md.
usage_metrics = UsageMetricsStore(lambda: STATE_DIR / "metricas_uso.json", lambda: TRANSACTION_ROOT)
app.add_middleware(UsageMetricsMiddleware, store=usage_metrics)


SERVER_PORT = int(os.environ.get("SIMPRESS_PORT", "8000"))
IDEMPOTENT_UPLOAD_ADMISSION_WAIT_SECONDS = 120

_ALLOWED_LOCAL_AUTHORITIES = allowed_local_authorities(SERVER_PORT)
_parse_request_authority = parse_request_authority

_local_request_safety_services = LocalRequestSafetyServices(
    allowed_authorities=lambda: _ALLOWED_LOCAL_AUTHORITIES,
    inventory_upload_paths=lambda: INVENTORY_UPLOAD_PATHS,
    body_limits=lambda: RequestBodyLimits(
        upload_request_bytes=MAX_UPLOAD_REQUEST_BYTES,
        upload_bytes=MAX_UPLOAD_BYTES,
        ticket_workbook_bytes=MAX_TICKET_WORKBOOK_BYTES,
        api_request_bytes=MAX_API_REQUEST_BYTES,
    ),
    upload_admission=lambda: UPLOAD_ADMISSION,
    idempotent_upload_wait_seconds=lambda: IDEMPOTENT_UPLOAD_ADMISSION_WAIT_SECONDS,
    service_unavailable_errors=(LocalJsonCorruptionError,),
)


class LocalRequestSafetyMiddleware(_LocalRequestSafetyMiddleware):
    """Restrict local peers and authorities, cap uploads, and hide exceptions."""

    def __init__(self, app):
        super().__init__(app, _local_request_safety_services)


app.add_middleware(LocalRequestSafetyMiddleware)

def load_custom_units() -> List[Dict[str, Any]]:
    if not CUSTOM_UNIDADES_FILE.exists():
        return []
    try:
        value = json.loads(CUSTOM_UNIDADES_FILE.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise LocalJsonCorruptionError("O catálogo local de unidades está ilegível; os dados foram preservados.") from error
    if not isinstance(value, list) or any(not isinstance(item, dict) for item in value):
        raise LocalJsonCorruptionError("O catálogo local de unidades precisa conter uma lista de registros.")
    for item in value:
        if any(field in item and not isinstance(item[field], (str, int, float, bool, type(None))) for field in ("id", "cidade", "endereco", "uf", "cep", "cnpj")):
            raise LocalJsonCorruptionError("O catálogo local de unidades contém campos com tipos inválidos.")
    return value


def _history_event_datetime(event: Dict[str, Any]) -> Optional[datetime]:
    try:
        value = datetime.fromisoformat(str(event.get("horario") or "").replace("Z", "+00:00"))
    except (TypeError, ValueError, OverflowError):
        return None
    return value.astimezone()


def _backup_directory_from_history(event: Dict[str, Any]) -> Optional[Path]:
    comparison = event.get("comparacao")
    relative_value = comparison.get("backup_diretorio") if isinstance(comparison, dict) else None
    if not isinstance(relative_value, str) or not relative_value.strip():
        return None
    relative = Path(relative_value)
    if relative.is_absolute() or any(part in {"", ".", ".."} for part in relative.parts):
        return None
    candidate = (INPUT_DIR / relative).resolve(strict=False)
    backup_root = UPLOAD_BACKUPS_DIR.resolve(strict=False)
    if candidate.parent != backup_root:
        return None
    return candidate


def prune_expired_histories(now: Optional[datetime] = None) -> Dict[str, int]:
    """Apply the HISTORY_RETENTION_DAYS policy and remove only verified, unreferenced backups."""
    current_time = (now or datetime.now().astimezone()).astimezone()
    cutoff = current_time - timedelta(days=HISTORY_RETENTION_DAYS)
    results = {"apelidos_removidos": 0, "comparacoes_removidas": 0, "backups_removidos": 0}
    expired_backup_paths: Set[Path] = set()
    active_backup_paths: Set[Path] = set()
    targets = [path for path in (ALIAS_HISTORY_FILE, COMPARISON_HISTORY_FILE) if path.exists()]

    # The transaction's file locks serialize the histories; holding the Python
    # history locks here would invert the global → file → leaf order.
    if targets:
        with AtomicFileTransaction(TRANSACTION_ROOT, targets) as transaction:
            if ALIAS_HISTORY_FILE in targets:
                def keep_alias_events(value):
                    events = _validate_json_history(value, "apelidos")
                    kept = []
                    for event in events:
                        timestamp = _history_event_datetime(event)
                        if timestamp is not None and timestamp < cutoff:
                            results["apelidos_removidos"] += 1
                        else:
                            kept.append(event)
                    return kept
                transaction.update_json(ALIAS_HISTORY_FILE, keep_alias_events, [])

            if COMPARISON_HISTORY_FILE in targets:
                def keep_comparison_events(value):
                    events = _validate_json_history(value, "comparações")
                    kept = []
                    for event in events:
                        timestamp = _history_event_datetime(event)
                        is_pinned = event.get("id") == PINNED_COMPARISON_EVENT_ID
                        if not is_pinned and timestamp is not None and timestamp < cutoff:
                            backup_path = _backup_directory_from_history(event)
                            if backup_path is not None:
                                expired_backup_paths.add(backup_path)
                            results["comparacoes_removidas"] += 1
                            continue
                        kept.append(event)
                        backup_path = _backup_directory_from_history(event)
                        if backup_path is not None:
                            active_backup_paths.add(backup_path)
                    return kept
                transaction.update_json(COMPARISON_HISTORY_FILE, keep_comparison_events, [])
            transaction.commit()

    if UPLOAD_BACKUPS_DIR.is_symlink() or not UPLOAD_BACKUPS_DIR.is_dir():
        return results
    backup_root = UPLOAD_BACKUPS_DIR.resolve(strict=False)
    if backup_root.parent != INPUT_DIR.resolve(strict=False):
        return results
    expired_directories: List[Path] = []
    for path in backup_root.iterdir():
        if path.is_symlink() or not path.is_dir():
            continue
        resolved = path.resolve(strict=False)
        if resolved.parent != UPLOAD_BACKUPS_DIR.resolve(strict=False) or resolved in active_backup_paths:
            continue
        expired = resolved in expired_backup_paths
        if not expired:
            try:
                modified = datetime.fromtimestamp(path.stat().st_mtime).astimezone()
                expired = modified < cutoff
            except OSError:
                continue
        if expired:
            expired_directories.append(path)

    # Remove after enumeration; deleting while iterating a Windows directory can
    # cause the iterator to stop before other expired directories are visited.
    for path in expired_directories:
        try:
            shutil.rmtree(path)
            results["backups_removidos"] += 1
        except OSError as error:
            logger.warning("Não foi possível remover backup vencido '%s': %s.", path.name, type(error).__name__)
    return results


def _validate_json_history(value: Any, label: str) -> List[Dict[str, Any]]:
    if not isinstance(value, list) or any(not isinstance(item, dict) for item in value):
        raise LocalJsonCorruptionError(f"O histórico local de {label} precisa conter uma lista de registros.")
    return value


def load_database(transaction_locked: bool = False, publish: bool = True) -> Any:
    if transaction_locked:
        return inventory_state.load_database_unlocked(_inventory_state_services, publish=publish)
    with lock_paths([TRANSACTION_ROOT / ".atomic-transactions"]):
        detail_files = get_details_files()
        with lock_paths([
            *detail_files,
            ENDPOINT_UPLOAD_FILE,
            MODIFIED_APELIDOS_FILE,
            REMOVED_ENDPOINT_HOSTNAMES_FILE,
            CUSTOM_UNIDADES_FILE,
            EXPORT_MODEL_FILE,
        ]):
            return inventory_state.load_database_unlocked(_inventory_state_services, publish=publish)


def prepare_database_state(
    equipment_cache: Dict[str, Dict[str, Any]],
    unit_cache: List[Dict[str, Any]],
    stats: Dict[str, Any],
) -> Dict[str, Any]:
    """Build every derived cache before a disk transaction is committed."""
    return inventory_state.prepare_database_state(equipment_cache, unit_cache, stats, _inventory_state_services)


# Files behind the published cache; a change made by another process (the inbox
# scripts in scripts/) is noticed by _external_change_reload_worker.
EXTERNAL_CHANGE_POLL_SECONDS = 10


def publish_database_state(state: Dict[str, Any]) -> None:
    inventory_state.publish_database_state(state, _inventory_state_services)


# Late-bound so tests that patch these main attributes keep affecting the load.
_inventory_state_services = inventory_state.InventoryStateServices(
    get_parque_file=lambda: get_parque_file(),
    input_dir=lambda: INPUT_DIR,
    find_details_sheet=lambda file_path: find_details_sheet(file_path),
    read_details_sheet_records=lambda file_path, sheet_name: _read_details_sheet_records(file_path, sheet_name),
    load_custom_units=lambda: load_custom_units(),
    load_modified_apelidos=lambda: load_modified_apelidos(),
    read_endpoint_alias_index=lambda: read_endpoint_alias_index(),
    equipment_revision=lambda *args: _equipment_revision(*args),
    apply_effective_alias=lambda *args, **kwargs: _apply_effective_alias(*args, **kwargs),
    count_export_changes=lambda *args, **kwargs: count_export_changes(*args, **kwargs),
    get_export_status=lambda: get_export_status(),
    prepare_database_state=lambda *args: prepare_database_state(*args),
    publish_database_state=lambda state: publish_database_state(state),
    build_search_index=lambda *args: build_search_index(*args),
    normalize_search_term=lambda value: _normalizar_termo_busca(value),
    inventory_files=lambda: (
        *get_details_files(),
        ENDPOINT_UPLOAD_FILE,
        MODIFIED_APELIDOS_FILE,
        REMOVED_ENDPOINT_HOSTNAMES_FILE,
        CUSTOM_UNIDADES_FILE,
        EXPORT_MODEL_FILE,
    ),
    logger=logger,
)

# Models for API
class EquipamentoUpdatePayload(BaseModel):
    serial: str = Field(..., min_length=1, max_length=128, description="Número de Série do Equipamento")
    novo: bool = Field(False, description="Indica que o formulário está cadastrando um equipamento novo")
    modelo: Optional[str] = Field("", max_length=300, description="Modelo do Equipamento")
    tipo: Optional[str] = Field("", max_length=120, description="Tipo (Notebook, Micro Computador, Headsets...)")
    apelido: Optional[str] = Field("", max_length=160, description="Apelido / Hostname")
    idCliente: Optional[str] = Field("", max_length=128, description="ID do Cliente")
    revisao: Optional[str] = Field(None, min_length=64, max_length=64)
    # Mantidos para compatibilidade com clientes antigos; a rota ignora esses campos.
    endereco: Optional[str] = Field("", max_length=500, description="Endereço carregado da Simpress")
    cidade: Optional[str] = Field("", max_length=160, description="Cidade carregada da Simpress")
    uf: Optional[str] = Field("PR", max_length=2, description="UF carregada da Simpress")
    cep: Optional[str] = Field("", max_length=12, description="CEP carregado da Simpress")
    cnpj_instalacao: Optional[str] = Field("", max_length=24, description="CNPJ carregado da Simpress")

    @field_validator("serial")
    @classmethod
    def validate_serial(cls, value: str) -> str:
        return _validate_equipment_serial(value)


def _public_equipment(equipment: Dict[str, Any]) -> Dict[str, Any]:
    public = {key: value for key, value in equipment.items() if not key.startswith("_")}
    revision = equipment.get("_revisao")
    if revision:
        public["revisao"] = revision
    return public


def _edit_equipment_workbook(
    wb,
    payload: EquipamentoUpdatePayload,
    target_sheet_name: Optional[str] = None,
    record_must_exist: Optional[bool] = None,
    fields_to_update: Optional[Set[str]] = None,
    expected_revision: Optional[str] = None,
):
    if target_sheet_name is None:
        target_sheet_name = next(
            (name for name in wb.sheetnames if "equipamento" in name.lower() or "parque" in name.lower()),
            wb.sheetnames[0] if wb.sheetnames else None,
        )
    elif target_sheet_name not in wb.sheetnames:
        return None
    if target_sheet_name is None:
        return None
    ws = wb[target_sheet_name]
    header_row = [cell.value for cell in next(ws.iter_rows(max_row=1), ())]
    col_map = {
        normalize_workbook_header(value): index + 1
        for index, value in enumerate(header_row)
        if value is not None
    }
    serial_col = next(
        (col_map.get(alias) for alias in DETAILS_SERIAL_HEADER_ALIASES if col_map.get(alias)),
        None,
    )
    if not serial_col:
        return None

    target_serial = normalize_reference_value(payload.serial)
    matching_rows = [
        row for row in range(2, ws.max_row + 1)
        if normalize_reference_value(ws.cell(row=row, column=serial_col).value) == target_serial
        and target_serial
    ]
    if len(matching_rows) > 1:
        raise HTTPException(
            status_code=409,
            detail=f"A série {payload.serial.strip()} aparece mais de uma vez na aba '{ws.title}'; nenhum registro foi alterado.",
        )
    if matching_rows:
        contract_col = col_map.get("CODIGOCONTRATO") or col_map.get("CODIGODOCONTRATO")
        if contract_col:
            raw_contract = ws.cell(row=matching_rows[0], column=contract_col).value
            normalized_contract = normalize_contract_value(raw_contract)
            if normalized_contract and normalized_contract != str(SUPPORTED_CONTRACT):
                raise HTTPException(
                    status_code=409,
                    detail=(
                        f"O equipamento pertence ao contrato {normalized_contract}; "
                        f"este sistema só edita o contrato {SUPPORTED_CONTRACT}."
                    ),
                )
    if matching_rows and expected_revision:
        target_row_for_revision = matching_rows[0]

        def row_value(*aliases: str) -> Any:
            column = next(
                (col_map.get(normalize_workbook_header(alias)) for alias in aliases
                 if col_map.get(normalize_workbook_header(alias))),
                None,
            )
            return ws.cell(row=target_row_for_revision, column=column).value if column else ""

        current_revision = _equipment_revision(
            ws.cell(row=target_row_for_revision, column=serial_col).value,
            row_value("MODELO"),
            row_value("TIPO", "CLASSEEQUIPAMENTO"),
            row_value(*DETAILS_ALIAS_HEADER_ALIASES),
            row_value("IDCLIENTE", "ID CLIENTE"),
        )
        if not hmac.compare_digest(expected_revision, current_revision):
            raise HTTPException(
                status_code=409,
                detail="O equipamento foi alterado desde a leitura; recarregue os dados antes de salvar.",
            )
    if record_must_exist is True and not matching_rows:
        raise HTTPException(
            status_code=409,
            detail=f"A série {payload.serial.strip()} não está mais na planilha ativa; recarregue o inventário antes de editar.",
        )
    if record_must_exist is False and matching_rows:
        raise HTTPException(
            status_code=409,
            detail=f"A série {payload.serial.strip()} já existe na planilha ativa; recarregue o inventário antes de cadastrar.",
        )
    target_row = matching_rows[0] if matching_rows else ws.max_row + 1
    expected: Dict[tuple, Any] = {}
    requested_fields = set(fields_to_update if fields_to_update is not None else payload.model_fields_set)

    def set_value(column: Optional[int], value: Any) -> None:
        if column is None:
            return
        cell = ws.cell(row=target_row, column=column)
        persisted_value = "" if value is None else value
        set_cell_value_as_text(cell, persisted_value)
        expected[(ws.title, cell.coordinate)] = persisted_value

    if not matching_rows:
        set_value(serial_col, payload.serial.strip().upper())
        set_value(col_map.get("CODIGOCONTRATO"), SUPPORTED_CONTRACT)
        set_value(col_map.get("ITEMDECONTRATO"), target_row - 1)
        set_value(col_map.get("RAZAOSOCIAL"), SUPPORTED_CONTRACT_LEGAL_NAME)

    fields = (
        ("apelido", col_map.get("APELIDO")),
        ("modelo", col_map.get("MODELO")),
        ("tipo", col_map.get("TIPO") or col_map.get("CLASSEEQUIPAMENTO")),
        ("idCliente", col_map.get("IDCLIENTE")),
    )
    for field_name, column in fields:
        if matching_rows and field_name not in requested_fields:
            continue
        value = getattr(payload, field_name)
        set_value(column, value.strip() if value is not None else "")

    if not matching_rows:
        from openpyxl.utils.cell import get_column_letter, range_boundaries

        for table in ws.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            if min_row == 1 and max_row == target_row - 1 and min_col <= serial_col <= max_col:
                expanded_ref = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(max_col)}{target_row}"
                )
                table.ref = expanded_ref
                if table.autoFilter is not None:
                    table.autoFilter.ref = expanded_ref

        if ws.auto_filter.ref:
            min_col, min_row, max_col, max_row = range_boundaries(ws.auto_filter.ref)
            if min_row == 1 and max_row == target_row - 1 and min_col <= serial_col <= max_col:
                ws.auto_filter.ref = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(max_col)}{target_row}"
                )
    return expected


def update_equipment_in_workbook_file(
    file_path: Path,
    payload: EquipamentoUpdatePayload,
    transaction: Optional[AtomicFileTransaction] = None,
    target_sheet_name: Optional[str] = None,
    record_must_exist: Optional[bool] = None,
    fields_to_update: Optional[Set[str]] = None,
    expected_revision: Optional[str] = None,
) -> bool:
    """Prepare and validate a workbook update before its atomic replacement."""
    file_path = Path(file_path)
    if not file_path.exists():
        return False
    try:
        if transaction is not None:
            return stage_workbook_update(
                transaction,
                file_path,
                lambda wb: _edit_equipment_workbook(
                    wb, payload, target_sheet_name, record_must_exist, fields_to_update, expected_revision
                ),
                allow_dimension_growth=True,
                allow_table_growth=True,
            )
        with AtomicFileTransaction(TRANSACTION_ROOT, [file_path]) as own_transaction:
            saved = stage_workbook_update(
                own_transaction,
                file_path,
                lambda wb: _edit_equipment_workbook(
                    wb, payload, target_sheet_name, record_must_exist, fields_to_update, expected_revision
                ),
                allow_dimension_growth=True,
                allow_table_growth=True,
            )
            if saved:
                own_transaction.commit()
            return saved
    except FileBusyError as error:
        raise HTTPException(status_code=409, detail=str(error)) from error
    except WorkbookSafetyError as error:
        raise HTTPException(status_code=409, detail=f"Gravação cancelada: {error}") from error
    except FileTransactionError as error:
        raise HTTPException(status_code=500, detail=f"Não foi possível salvar '{file_path.name}': {error}") from error


def normalizar_endereco_instalacao(payload: EquipamentoUpdatePayload) -> Dict[str, str]:
    endereco = payload.endereco.strip() if payload.endereco else ""
    cidade = payload.cidade.strip() if payload.cidade else ""
    uf = (payload.uf.strip() if payload.uf else "PR").upper()
    cep = format_cep(payload.cep)
    cnpj = payload.cnpj_instalacao.strip() if payload.cnpj_instalacao else ""

    is_sede = (
        ("caeté" in endereco.lower() or "caete" in endereco.lower() or "sede" in endereco.lower() or "sede" in cidade.lower())
        and ("curitiba" in cidade.lower() or "sede" in cidade.lower())
    )
    if is_sede:
        endereco = "R Caeté, 150 - SEDE - Prado Velho"
        cidade = "Curitiba"
        uf = "PR"
        cep = "80220-300"

    return {
        "endereco": endereco,
        "cidade": cidade,
        "uf": uf,
        "cep": cep,
        "cnpj": cnpj,
    }


def snapshot_endereco_instalacao(equipamento: Dict[str, Any]) -> Dict[str, str]:
    return {
        "endereco": str(equipamento.get("endereco_instalacao") or "").strip(),
        "cidade": str(equipamento.get("retirada_cidade") or "").strip(),
        "uf": str(equipamento.get("retirada_uf") or "").strip(),
        "cep": str(equipamento.get("retirada_cep") or "").strip(),
        "cnpj": str(equipamento.get("retirada_cnpj") or "").strip(),
    }


def save_equipment_to_excel(
    payload: EquipamentoUpdatePayload,
    json_mutations: Optional[Dict[Path, Any]] = None,
    fields_to_update: Optional[Set[str]] = None,
):
    """Salva um equipamento somente na fonte e na aba associadas ao cache atual."""
    target_serial_upper = payload.serial.strip().upper()
    detail_files = [path for path in get_details_files() if path.is_file()]
    if not detail_files:
        raise HTTPException(status_code=404, detail="Nenhuma planilha da Simpress encontrada no servidor.")

    requested_fields = set(fields_to_update if fields_to_update is not None else payload.model_fields_set)
    json_mutations = dict(json_mutations or {})
    try:
        # Lock every candidate while resolving the current primary source. Only
        # that source is staged below; legacy copies remain untouched.
        transaction_targets = list(dict.fromkeys([
            *detail_files,
            ENDPOINT_UPLOAD_FILE,
            MODIFIED_APELIDOS_FILE,
            *json_mutations.keys(),
        ]))
        with AtomicFileTransaction(TRANSACTION_ROOT, transaction_targets) as transaction:
            with STATE_LOCK:
                next_equipment_cache = deepcopy(inventory_state.equipamentos_cache)
                next_unit_cache = deepcopy(inventory_state.unidades_cache)
                next_stats_cache = deepcopy(inventory_state.stats_cache)
            cached_equipment = deepcopy(next_equipment_cache.get(target_serial_upper))
            if cached_equipment is None and not payload.novo:
                raise HTTPException(
                    status_code=409,
                    detail="Este equipamento não está mais na base carregada; recarregue o inventário antes de editar.",
                )
            expected_revision = payload.revisao
            if cached_equipment is not None:
                current_revision = str(cached_equipment.get("_revisao") or "")
                if not expected_revision:
                    raise HTTPException(
                        status_code=409,
                        detail="A revisão deste equipamento não foi informada; recarregue-o antes de salvar.",
                    )
                if not hmac.compare_digest(expected_revision, current_revision):
                    raise HTTPException(
                        status_code=409,
                        detail="O equipamento foi alterado desde a leitura; recarregue os dados antes de salvar.",
                    )
            if cached_equipment is not None and payload.novo:
                raise HTTPException(
                    status_code=409,
                    detail="A série já pertence a um equipamento carregado; atualize o inventário antes de cadastrar outra vez.",
                )
            if cached_equipment is not None:
                if cached_equipment.get("invall_alias_status") == "serie_detalhes_duplicada":
                    raise HTTPException(
                        status_code=409,
                        detail=f"A série {target_serial_upper} está duplicada na base ativa; nenhum registro foi alterado.",
                    )
                # Preserve the latest local alias when this save waited behind
                # a concurrent alias edit, and discard a now-stale new-record override.
                payload.apelido = cached_equipment.get(
                    "_apelido_local", cached_equipment.get("_apelido_base", "")
                )
                json_mutations.pop(MODIFIED_APELIDOS_FILE, None)

            active_file = get_primary_equipment_file()
            if active_file is None or not active_file.is_file():
                raise HTTPException(status_code=404, detail="Nenhuma planilha válida da Simpress está ativa.")

            target_file = active_file.resolve()
            locked_files = {path.resolve() for path in detail_files}
            if target_file not in locked_files:
                raise HTTPException(
                    status_code=409,
                    detail="A lista de inventários mudou durante a edição; recarregue a base e tente novamente.",
                )

            cached_source = cached_equipment.get("_source_file") if cached_equipment else None
            if cached_source:
                if Path(cached_source).resolve() != target_file:
                    raise HTTPException(
                        status_code=409,
                        detail="A base ativa mudou depois que este equipamento foi carregado; recarregue-o antes de editar.",
                    )

            current_sheet = find_details_sheet(target_file)
            if current_sheet is None:
                raise HTTPException(status_code=409, detail="A base ativa não contém uma aba de equipamentos válida.")
            cached_sheet = cached_equipment.get("_source_sheet") if cached_equipment else None
            if cached_sheet and cached_sheet != current_sheet:
                raise HTTPException(
                    status_code=409,
                    detail="A aba de origem mudou depois que este equipamento foi carregado; recarregue-o antes de editar.",
                )
            target_sheet = cached_sheet or current_sheet

            saved = update_equipment_in_workbook_file(
                target_file,
                payload,
                transaction,
                target_sheet_name=target_sheet,
                record_must_exist=not payload.novo,
                fields_to_update=requested_fields,
                expected_revision=expected_revision,
            )
            if not saved:
                raise HTTPException(status_code=500, detail="Não foi possível atualizar o equipamento na planilha.")
            staged_json: Dict[Path, Any] = {}
            for path, updater in json_mutations.items():
                staged_json[path] = transaction.update_json(
                    path,
                    updater,
                    [] if path == ALIAS_HISTORY_FILE else {},
                )

            endereco_info = (
                snapshot_endereco_instalacao(cached_equipment)
                if cached_equipment
                else normalizar_endereco_instalacao(EquipamentoUpdatePayload(
                    serial=target_serial_upper,
                    novo=False,
                    modelo="",
                    tipo="",
                    apelido="",
                    idCliente="",
                    revisao=None,
                    endereco="",
                    cidade="",
                    uf="PR",
                    cep="",
                    cnpj_instalacao="",
                ))
            )
            endereco = endereco_info["endereco"]
            cidade = endereco_info["cidade"]
            uf = endereco_info["uf"]
            cep_formatado = endereco_info["cep"]
            cnpj_instalacao = endereco_info["cnpj"]
            retirada_endereco = f"{endereco} - {cidade}/{uf} - CEP {cep_formatado}" if endereco else f"{cidade}/{uf}"

            def updated_value(field_name: str, cache_key: str) -> str:
                if cached_equipment is not None and field_name not in requested_fields:
                    return str(cached_equipment.get(cache_key) or "").strip()
                value = getattr(payload, field_name)
                return value.strip() if value is not None else ""

            base_alias = updated_value("apelido", "_apelido_base")
            modelo = updated_value("modelo", "modelo")
            tipo = updated_value("tipo", "tipo")
            id_cliente = updated_value("idCliente", "idCliente")
            desc_sugerida = get_suggested_description(tipo, modelo)

            modified_aliases = staged_json.get(MODIFIED_APELIDOS_FILE)
            if not isinstance(modified_aliases, dict):
                modified_aliases = load_modified_apelidos()
            manual_aliases = _manual_aliases_by_serial(modified_aliases)
            normalized_serial = normalize_reference_value(target_serial_upper)
            local_alias = manual_aliases.get(normalized_serial, base_alias)
            updated_equipment = {
                "serial": target_serial_upper,
                "_source_file": str(target_file),
                "_source_sheet": target_sheet,
                "modelo": modelo,
                "tipo": tipo,
                "apelido": local_alias,
                "_apelido_base": base_alias,
                "_apelido_local": local_alias,
                "_apelido_manual_override": normalized_serial in manual_aliases,
                "_revisao": _equipment_revision(
                    target_serial_upper, modelo, tipo, base_alias, id_cliente
                ),
                "idCliente": id_cliente,
                "descricao_sugerida": desc_sugerida,
                "endereco_instalacao": endereco,
                "retirada_cnpj": cnpj_instalacao,
                "retirada_endereco": retirada_endereco,
                "retirada_cidade": cidade,
                "retirada_uf": uf,
                "retirada_cep": cep_formatado,
                "contrato": (
                    cached_equipment.get("contrato", str(SUPPORTED_CONTRACT))
                    if cached_equipment else str(SUPPORTED_CONTRACT)
                ),
                "razao_social": (
                    cached_equipment.get("razao_social", SUPPORTED_CONTRACT_LEGAL_NAME)
                    if cached_equipment else SUPPORTED_CONTRACT_LEGAL_NAME
                ),
            }

            endpoint_index = read_endpoint_alias_index()
            _apply_effective_alias(
                updated_equipment,
                endpoint_index,
                details_duplicate=bool(
                    cached_equipment
                    and cached_equipment.get("invall_alias_status") == "serie_detalhes_duplicada"
                ),
            )
            next_equipment_cache[target_serial_upper] = updated_equipment
            next_stats_cache["total_equipamentos"] = len(next_equipment_cache)
            next_stats_cache["ultima_atualizacao"] = datetime.now().strftime("%d/%m/%Y %H:%M")
            prepared_state = prepare_database_state(
                next_equipment_cache,
                next_unit_cache,
                next_stats_cache,
            )

            # All reads and derived data are prepared before the durable commit.
            # Publishing afterward only swaps already-built references.
            transaction.commit()
            publish_database_state(prepared_state)
            return updated_equipment, len(modified_aliases)
    except FileBusyError as error:
        raise HTTPException(status_code=409, detail=str(error)) from error
    except WorkbookSafetyError as error:
        raise HTTPException(status_code=409, detail=f"Gravação cancelada: {error}") from error
    except FileTransactionError as error:
        raise HTTPException(status_code=500, detail=f"Não foi possível concluir a gravação: {error}") from error


# Late-bound so tests that patch these main attributes keep affecting alias changes.
_alias_change_services = alias_changes.AliasChangeServices(
    modified_aliases_file=lambda: MODIFIED_APELIDOS_FILE,
    alias_history_file=lambda: ALIAS_HISTORY_FILE,
    transaction_root=lambda: TRANSACTION_ROOT,
    atomic_file_transaction=lambda root, targets: AtomicFileTransaction(root, targets),
    get_details_files=lambda: get_details_files(),
    get_primary_equipment_file=lambda: get_primary_equipment_file(),
    find_details_sheet=lambda file_path: find_details_sheet(file_path),
    load_alias_history=lambda: load_alias_history(),
    load_modified_apelidos=lambda: load_modified_apelidos(),
    validate_json_object=lambda value, label: _validate_json_object(value, label),
    validate_json_history=lambda value, label: _validate_json_history(value, label),
    read_endpoint_alias_index=lambda: read_endpoint_alias_index(),
    apply_effective_alias=lambda *args, **kwargs: _apply_effective_alias(*args, **kwargs),
    prepare_database_state=lambda *args: prepare_database_state(*args),
    publish_database_state=lambda state: publish_database_state(state),
    prune_expired_histories=lambda: prune_expired_histories(),
    public_equipment=lambda equipment: _public_equipment(equipment),
    logger=logger,
    external_errors=(HTTPException,),
    manual_users_file=lambda: MANUAL_USERS_FILE,
    load_manual_users=lambda: load_manual_users(),
)

# Registered directly on the app, before /api/equipamento/{serial}, so the fixed
# "apelidos-modificados" path is matched first and is not read as a serial.
_alias_router, _alias_exports = create_alias_router(_alias_change_services, app.router)
get_apelidos_modificados = _alias_exports.get_apelidos_modificados
atualizar_apelido = _alias_exports.atualizar_apelido
get_historico_apelidos = _alias_exports.get_historico_apelidos
excluir_evento_historico_apelido = _alias_exports.excluir_evento_historico_apelido
desfazer_alteracao_apelido = _alias_exports.desfazer_alteracao_apelido
atualizar_usuario = _alias_exports.atualizar_usuario
consultar_usuarios_logados = _alias_exports.consultar_usuarios_logados

_endpoint_removal_services = endpoint_removals.EndpointRemovalServices(
    removals_file=lambda: REMOVED_ENDPOINT_HOSTNAMES_FILE,
    load_removals=lambda: load_removed_endpoint_hostnames(),
    validate_json_history=lambda value, label: _validate_json_history(value, label),
    transaction_root=lambda: TRANSACTION_ROOT,
    atomic_file_transaction=lambda root, targets: AtomicFileTransaction(root, targets),
    read_endpoint_alias_index=lambda: read_endpoint_alias_index(),
    reload_database=lambda: load_database(),
    external_errors=(HTTPException,),
)
app.include_router(create_endpoint_removals_router(_endpoint_removal_services))


# Endpoints
@app.get("/api/equipamento/{serial}")
def get_equipamento(serial: str):
    """Busca um equipamento pelo número de série (sem distinção de maiúsculas/minúsculas)."""
    clean_serial = serial.strip().upper()
    with STATE_LOCK:
        equipment = deepcopy(inventory_state.equipamentos_cache.get(clean_serial))
    if equipment is not None:
        return _public_equipment(equipment)

    raise HTTPException(status_code=404, detail=f"Equipamento com serial '{serial}' não encontrado na base de dados.")


@app.post("/api/equipamento/salvar")
def salvar_equipamento(payload: EquipamentoUpdatePayload):
    """Atualiza ou cadastra um equipamento tanto na memória do site quanto diretamente na planilha Excel."""
    if not payload.serial or not payload.serial.strip():
        raise HTTPException(status_code=400, detail="Número de série é obrigatório.")
    if not normalize_reference_value(payload.serial):
        raise HTTPException(status_code=422, detail="O número de série precisa conter letras ou números.")
    
    target_serial = payload.serial.strip().upper()
    requested_fields = set(payload.model_fields_set)
    with STATE_LOCK:
        equipamento_existente = deepcopy(inventory_state.equipamentos_cache.get(target_serial))
    if payload.novo and equipamento_existente is not None:
        raise HTTPException(
            status_code=409,
            detail="A série já está cadastrada; atualize a base antes de criar um novo equipamento.",
        )
    if not payload.novo and equipamento_existente is None:
        raise HTTPException(
            status_code=409,
            detail="Este equipamento não está mais na base carregada; recarregue o inventário antes de editar.",
        )
    if equipamento_existente is not None and not payload.revisao:
        raise HTTPException(
            status_code=409,
            detail="A revisão deste equipamento não foi informada; recarregue-o antes de salvar.",
        )
    if equipamento_existente is not None:
        # Para equipamentos já cadastrados, o apelido só pode ser alterado pela rota dedicada.
        # Isso impede que o formulário de dados/endereço sobrescreva apelidos acidentalmente.
        # Keep the stored EquipamentosDetalhes alias, not its InvAll presentation value.
        payload.apelido = equipamento_existente.get("_apelido_local", equipamento_existente.get("_apelido_base", ""))

    json_mutations = {}
    if equipamento_existente is None and payload.apelido is not None and payload.apelido.strip():
        new_alias = payload.apelido.strip()
        def save_new_alias(data):
            aliases = dict(_validate_json_object(data, "apelidos"))
            for key in list(aliases):
                if normalize_reference_value(key) == normalize_reference_value(target_serial):
                    aliases.pop(key, None)
            aliases[target_serial] = new_alias
            return aliases
        json_mutations[MODIFIED_APELIDOS_FILE] = save_new_alias

    saved_equipment, modified_alias_count = save_equipment_to_excel(
        payload,
        json_mutations=json_mutations,
        fields_to_update=requested_fields,
    )
    return {
        "sucesso": True,
        "mensagem": f"Equipamento {target_serial} atualizado com sucesso no site e na planilha Excel!",
        "equipamento": _public_equipment(saved_equipment),
        "total_modificados": modified_alias_count,
    }


def aplicar_alteracao_apelido(
    target_serial: str,
    novo_apelido: str,
    action: str = "alteracao",
    reference_id: Optional[str] = None,
    registrar_historico: bool = True,
    expected_current_local: Optional[str] = None,
    expected_latest_event_id: Optional[str] = None,
    remove_history_event_id: Optional[str] = None,
    drop_override: bool = False,
) -> Dict[str, Any]:
    """Save a manual alias, or with ``drop_override`` remove it so the sources apply again."""
    return alias_changes.apply_alias_change(
        _alias_change_services,
        target_serial,
        novo_apelido,
        action=action,
        reference_id=reference_id,
        registrar_historico=registrar_historico,
        expected_current_local=expected_current_local,
        expected_latest_event_id=expected_latest_event_id,
        remove_history_event_id=remove_history_event_id,
        drop_override=drop_override,
    )


@app.get("/api/exportar-equipamentos")
def exportar_equipamentos():
    """Exporta a apresentação de referência com os registros da base ativa."""
    export = equipment_export.build_equipment_export(_equipment_export_services)
    return StreamingResponse(
        io.BytesIO(export.content),
        media_type=equipment_export.XLSX_MEDIA_TYPE,
        headers=export.headers,
    )


@app.post("/api/recarregar-base")
def recarregar_base():
    """Recarrega a planilha válida mais recente da pasta de entrada e informa arquivo e total."""
    if not load_database():
        raise HTTPException(
            status_code=400,
            detail=f"Nenhuma planilha válida pôde ser carregada em '{INPUT_DIR.name}'."
        )

    filename = inventory_state.stats_cache["arquivo_origem"]
    total_equipamentos = inventory_state.stats_cache["total_equipamentos"]
    return {
        "sucesso": True,
        "arquivo_origem": filename,
        "total_equipamentos": total_equipamentos,
        "mensagem": f"Planilha '{filename}' carregada com {total_equipamentos} equipamentos.",
        "stats": inventory_state.stats_cache,
        "total_unidades": len(inventory_state.unidades_cache)
    }


def _history_equipment_snapshot() -> Dict[str, Dict[str, Any]]:
    with STATE_LOCK:
        return deepcopy(inventory_state.equipamentos_cache)


_inventory_update_services = InventoryUpdateServices(
    max_upload_bytes=lambda: MAX_UPLOAD_BYTES,
    workbook_parse_admission=lambda: WORKBOOK_PARSE_ADMISSION,
    validate_uploaded_workbook=lambda *args, **kwargs: validate_uploaded_workbook(*args, **kwargs),
    input_directory=lambda: INPUT_DIR,
    endpoint_upload_file=lambda: ENDPOINT_UPLOAD_FILE,
    details_upload_file=lambda: DETAILS_UPLOAD_FILE,
    comparison_history_file=lambda: COMPARISON_HISTORY_FILE,
    inventory_idempotency_file=lambda: STATE_DIR / "inventory_update_idempotency.json",
    upload_backups_directory=lambda: UPLOAD_BACKUPS_DIR,
    history_annotations_directory=lambda: HISTORY_ANNOTATIONS_DIR,
    local_annotation_outbox_directory=lambda: LOCAL_HISTORY_ANNOTATION_OUTBOX_DIR,
    build_history_annotation_record=lambda *args, **kwargs: build_history_annotation_record(*args, **kwargs),
    history_annotation_record_path=history_annotation_record_path,
    get_primary_equipment_file=lambda: get_primary_equipment_file(),
    compare_uploaded_workbooks=lambda *args, **kwargs: compare_uploaded_workbooks(*args, **kwargs),
    inventory_state_transaction=lambda targets: inventory_state_transaction(targets),
    get_equipment_snapshot=_history_equipment_snapshot,
    snapshot_equipment_aliases=lambda *args, **kwargs: snapshot_equipment_aliases(*args, **kwargs),
    load_modified_apelidos=lambda: load_modified_apelidos(),
    manual_alias_source_snapshots=lambda *args, **kwargs: _manual_alias_source_snapshots(*args, **kwargs),
    normalize_reference_value=lambda *args, **kwargs: normalize_reference_value(*args, **kwargs),
    upload_serial_value=_upload_serial_value,
    upload_alias_value=_upload_alias_value,
    load_database=lambda *args, **kwargs: load_database(*args, **kwargs),
    read_endpoint_alias_index=lambda: read_endpoint_alias_index(),
    compare_equipment_aliases=lambda *args, **kwargs: compare_equipment_aliases(*args, **kwargs),
    endpoint_alias_review_items=lambda *args, **kwargs: endpoint_alias_review_items(*args, **kwargs),
    publish_database_state=lambda *args, **kwargs: publish_database_state(*args, **kwargs),
    comparison_history_for_display=lambda *args, **kwargs: comparison_history_for_display(*args, **kwargs),
    load_history_annotations=lambda *args, **kwargs: load_history_annotations(*args, **kwargs),
    comparison_annotation_item_id=lambda *args, **kwargs: comparison_annotation_item_id(*args, **kwargs),
    inventory_update_preview_summary=inventory_update_preview_summary,
    create_inventory_preview_token=create_inventory_preview_token,
    verify_inventory_preview_token=verify_inventory_preview_token,
    external_errors=(HTTPException,),
    alias_history_file=lambda: ALIAS_HISTORY_FILE,
    removed_endpoint_hostnames=lambda: removed_endpoint_hostname_keys(),
)


async def _import_automatic_inventory_batch(
    batch_id: str,
    endpoint_name: str,
    endpoint_bytes: bytes,
    simpress_name: str,
    simpress_bytes: bytes,
):
    try:
        _raise_if_transaction_recovery_blocked()
        endpoint_upload = StarletteUploadFile(
            filename=endpoint_name,
            file=io.BytesIO(endpoint_bytes),
        )
        simpress_upload = StarletteUploadFile(
            filename=simpress_name,
            file=io.BytesIO(simpress_bytes),
        )
        result = await update_inventory_files(
            endpoint_upload,
            simpress_upload,
            services=_inventory_update_services,
            automatic_import_id=batch_id,
        )
    except HTTPException as error:
        raise _as_domain_error(error) from error
    await asyncio.to_thread(usage_metrics.record_comparison, result)
    try:
        await asyncio.to_thread(prune_expired_histories)
    except Exception:
        logger.exception(
            "Falha na retenção após a atualização automática confirmada (lote=%s).",
            batch_id,
        )
    return result




async def previsualizar_planilhas(endpoint_file: UploadFile, details_file: UploadFile):
    return await preview_inventory_files(endpoint_file, details_file, services=_inventory_update_services)


@app.post("/api/planilhas/previsualizar")
async def previsualizar_planilhas_endpoint(
    endpoint_file: UploadFile = File(...),
    details_file: UploadFile = File(...),
):
    return await previsualizar_planilhas(endpoint_file, details_file)


async def atualizar_planilhas(
    endpoint_file: UploadFile,
    details_file: UploadFile,
    preview_token: Optional[str] = None,
):
    result = await update_inventory_files(
        endpoint_file, details_file, services=_inventory_update_services,
        idempotency_key=idempotency_key.get(), preview_token=preview_token,
    )
    await asyncio.to_thread(usage_metrics.record_comparison, result)
    try:
        await asyncio.to_thread(prune_expired_histories)
    except Exception:
        event_id = (result.get("comparacao") or {}).get("evento_id") if isinstance(result, dict) else None
        logger.exception("Falha na retenção após a atualização confirmada (evento=%s).", event_id or "desconhecido")
    return result


@app.post("/api/planilhas/atualizar")
async def atualizar_planilhas_endpoint(
    endpoint_file: UploadFile = File(...),
    details_file: UploadFile = File(...),
    preview_token: Optional[str] = Form(None, min_length=66, max_length=90),
):
    return await atualizar_planilhas(endpoint_file, details_file, preview_token=preview_token)



def _normalizar_termo_busca(valor: Any) -> str:
    return normalize_search_term(valor)


def _categoria_equipamento(tipo: Any, modelo: Any) -> str:
    """Agrupa os rótulos do inventário nas categorias usadas pela pesquisa."""
    return categorize_equipment(tipo, modelo, _normalizar_termo_busca)


def _refresh_equipment_search_index_locked() -> None:
    """Precompute normalized searchable values once per published cache generation."""
    inventory_state.refresh_search_index_locked(_inventory_state_services)


def _buscar_equipamentos_por_termo(q: str) -> List[Dict[str, Any]]:
    """Busca por série, modelo, apelido, cidade ou categoria e prioriza série."""
    term = _normalizar_termo_busca(q)
    if not term:
        return []


    current_index = inventory_state.current_search_index(_inventory_state_services)
    return search_normalized_equipment_index(current_index, term)


def _listar_equipamentos_indexados() -> List[Dict[str, Any]]:
    """Todos os equipamentos da base na ordem do índice de busca."""
    return [entry[2] for entry in inventory_state.current_search_index(_inventory_state_services)]


def _get_stats_snapshot() -> Dict[str, Any]:
    export_counts = count_export_changes()
    export_status = get_export_status()
    with STATE_LOCK:
        stats_snapshot = deepcopy(inventory_state.stats_cache)
        equipment_count = len(inventory_state.equipamentos_cache)
        unit_count = len(inventory_state.unidades_cache)
    return {
        **stats_snapshot,
        "total_equipamentos": equipment_count,
        "base_importada": equipment_count > 0,
        "exportacao_disponivel": export_status["available"],
        "exportacao_motivo": export_status["reason"],
        "total_apelidos_modificados": export_counts["apelidos"],
        "total_id_clientes_modificados": export_counts["id_clientes"],
        "total_alteracoes_exportaveis": export_counts["total"],
        "total_unidades": unit_count,
    }


def _unit_cache_snapshot() -> List[Dict[str, Any]]:
    with STATE_LOCK:
        return deepcopy(inventory_state.unidades_cache)






system_router, _system_exports = create_system_router(SystemRouteServices(
    get_readiness_status=lambda: get_startup_readiness(),
    get_transfer_form_defaults=lambda: load_origin_contact(TRANSFER_ORIGIN_CONTACT_FILE),
))
get_readiness_status = _system_exports.get_readiness_status
get_transfer_form_defaults = _system_exports.get_transfer_form_defaults
app.include_router(system_router)

units_router, _units_exports = create_units_router(UnitRouteServices(
    format_cep=format_cep,
    get_details_files=get_details_files,
    endpoint_upload_file=lambda: ENDPOINT_UPLOAD_FILE,
    modified_aliases_file=lambda: MODIFIED_APELIDOS_FILE,
    custom_units_file=lambda: CUSTOM_UNIDADES_FILE,
    export_model_file=lambda: EXPORT_MODEL_FILE,
    transaction_root=lambda: TRANSACTION_ROOT,
    get_units_snapshot=_unit_cache_snapshot,
    coerce_unit_id=_coerce_unit_id,
    unit_identity_key=_unit_identity_key,
    local_json_corruption_error=LocalJsonCorruptionError,
    load_database=lambda *args, **kwargs: load_database(*args, **kwargs),
    publish_database_state=publish_database_state,
))
_public_unit_list = _units_exports.public_unit_list
salvar_unidade = _units_exports.salvar_unidade
get_unidades = _units_exports.get_unidades
app.include_router(units_router)

equipment_search_router, _equipment_search_exports = create_equipment_search_router(EquipmentSearchServices(
    search=_buscar_equipamentos_por_termo,
    normalize=_normalizar_termo_busca,
    categories=_CATEGORIAS_EQUIPAMENTO,
    get_stats=_get_stats_snapshot,
    list_all=_listar_equipamentos_indexados,
))
buscar_seriais = _equipment_search_exports.buscar_seriais
buscar_seriais_expandida = _equipment_search_exports.buscar_seriais_expandida
get_stats = _equipment_search_exports.get_stats
app.include_router(equipment_search_router)

transfer_router, gerar_planilha = create_transfer_router(get_local_transfer_template)
app.include_router(transfer_router)

transfer_history_router, _transfer_history_exports = create_transfer_history_router(TransferHistoryServices(
    history_file=lambda: TRANSFER_HISTORY_FILE,
    transaction_root=lambda: TRANSACTION_ROOT,
))
listar_historico_transferencias = _transfer_history_exports.listar_historico_transferencias
registrar_historico_transferencia = _transfer_history_exports.registrar_historico_transferencia
app.include_router(transfer_history_router)


def _dashboard_equipment_snapshot() -> List[Dict[str, Any]]:
    with STATE_LOCK:
        return [
            {"serial": item.get("serial"), "tipo": item.get("tipo"), "modelo": item.get("modelo"), "retirada_cidade": item.get("retirada_cidade")}
            for item in inventory_state.equipamentos_cache.values()
        ]


def _comparison_history_signature() -> Optional[tuple]:
    try:
        stat = COMPARISON_HISTORY_FILE.stat()
    except OSError:
        return None
    return (str(COMPARISON_HISTORY_FILE.resolve(strict=False)), stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns)


# The panel's absence list of the newest event, kept while the history file signature
# and the event id match, so each refresh of the panel does not reparse the whole history.
_DASHBOARD_DETAILS_CACHE: Optional[tuple] = None
_DASHBOARD_DETAILS_CACHE_LOCK = threading.Lock()


def _dashboard_event_details(event_id: str) -> Optional[Dict[str, Any]]:
    global _DASHBOARD_DETAILS_CACHE
    signature = _comparison_history_signature()
    with _DASHBOARD_DETAILS_CACHE_LOCK:
        cached = _DASHBOARD_DETAILS_CACHE
    if signature is not None and cached is not None and cached[:2] == (signature, event_id):
        return deepcopy(cached[2])
    details = None
    for item in load_comparison_history():
        if isinstance(item, dict) and str(item.get("id") or "") == event_id and isinstance(item.get("comparacao"), dict):
            details = {"ausentes_no_endpoint": item["comparacao"].get("ausentes_no_endpoint")}
            break
    # A write between the two stats may have been read halfway; cache only a stable file.
    if signature is not None and signature == _comparison_history_signature():
        with _DASHBOARD_DETAILS_CACHE_LOCK:
            _DASHBOARD_DETAILS_CACHE = (signature, event_id, deepcopy(details))
    return details


def _dashboard_comparison_events() -> List[Dict[str, Any]]:
    eventos = get_historico_comparacoes()["eventos"]
    if not eventos:
        return eventos
    # The list endpoint omits the detailed comparison; the panel only reads the newest event,
    # so only that event receives the endpoint absence list needed for the notebook/MacBook count.
    details = _dashboard_event_details(str(eventos[0].get("id") or ""))
    if details is not None:
        return [{**eventos[0], "comparacao": details}, *eventos[1:]]
    return eventos


def _dashboard_ticket_report() -> Optional[Dict[str, Any]]:
    # Read-only; a missing file hides the stale tickets alert and an unreadable one makes it unavailable.
    if not TICKET_REPORT_FILE.exists():
        return None
    data = json.loads(TICKET_REPORT_FILE.read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else None


dashboard_router, _dashboard_exports = create_dashboard_router(DashboardServices(
    get_equipment=_dashboard_equipment_snapshot,
    get_stats=_get_stats_snapshot,
    get_comparison_events=_dashboard_comparison_events,
    categories=_CATEGORIAS_EQUIPAMENTO,
    categorize=_categoria_equipamento,
    get_usage_summary=usage_metrics.summary,
    set_manual_minutes=usage_metrics.set_manual_minutes,
    get_comparison_history=lambda: load_comparison_history(),
    get_details_file=lambda: DETAILS_UPLOAD_FILE,
    stale_base_days=STALE_BASE_ALERT_DAYS,
    get_ticket_report=_dashboard_ticket_report,
    stale_ticket_days=STALE_TICKET_ALERT_DAYS,
    display_comparison=comparison_history_for_display,
))
obter_painel = _dashboard_exports.obter_painel
app.include_router(dashboard_router)


def _check_imported_export():
    if get_primary_equipment_file() is not None:
        count_export_changes(locked=True)
        status = get_export_status()
        return None if status["available"] else status["reason"]
    return None  # A model can be configured before the first inventory pair.


local_setup_router, importar_modelo_local = create_local_setup_router(
    model_path=lambda kind: INPUT_DIR / "_modelos" / "transferencia.xlsx" if kind == "transferencia" else EXPORT_MODEL_FILE,
    transaction_root=lambda: TRANSACTION_ROOT,
    backups_directory=lambda: UPLOAD_BACKUPS_DIR,
    max_upload_bytes=lambda: MAX_UPLOAD_BYTES,
    parser_admission=lambda: WORKBOOK_PARSE_ADMISSION,
    export_dependency_paths=lambda: [*get_details_files(), ENDPOINT_UPLOAD_FILE, MODIFIED_APELIDOS_FILE],
    validate_export=_check_imported_export,
    origin_contact_file=lambda: TRANSFER_ORIGIN_CONTACT_FILE,
)
app.include_router(local_setup_router)

inventory_history_router, _inventory_history_exports = create_inventory_history_router(InventoryHistoryServices(
    manual_aliases_by_serial=_manual_aliases_by_serial,
    build_history_annotation_record=lambda *args, **kwargs: build_history_annotation_record(*args, **kwargs),
    comparison_annotation_item_id=comparison_annotation_item_id,
    comparison_history_path=lambda: COMPARISON_HISTORY_FILE,
    comparison_history_for_display=comparison_history_for_display,
    details_upload_path=lambda: DETAILS_UPLOAD_FILE,
    endpoint_upload_path=lambda: ENDPOINT_UPLOAD_FILE,
    get_equipment_snapshot=_history_equipment_snapshot,
    get_equipment_template_file=get_equipamentos_template_file,
    history_annotation_record_path=history_annotation_record_path,
    history_annotation_row_fingerprint=history_annotation_row_fingerprint,
    history_annotations_directory=lambda: HISTORY_ANNOTATIONS_DIR,
    history_annotations_sync_enabled=lambda: HISTORY_ANNOTATIONS_SYNC_ENABLED,
    input_directory=lambda: INPUT_DIR,
    atomic_file_transaction=lambda *args, **kwargs: AtomicFileTransaction(*args, **kwargs),
    inventory_state_transaction=inventory_state_transaction,
    load_comparison_history=load_comparison_history,
    load_database=lambda *args, **kwargs: load_database(*args, **kwargs),
    load_history_annotations=lambda directory, **kwargs: load_history_annotations(directory, **kwargs),
    load_modified_apelidos=load_modified_apelidos,
    local_annotation_outbox_directory=lambda: LOCAL_HISTORY_ANNOTATION_OUTBOX_DIR,
    max_upload_bytes=MAX_UPLOAD_BYTES,
    modified_aliases_path=lambda: MODIFIED_APELIDOS_FILE,
    normalize_reference_value=normalize_reference_value,
    publish_database_state=publish_database_state,
    save_history_annotation=lambda *args, **kwargs: save_history_annotation(*args, **kwargs),
    transaction_root=lambda: TRANSACTION_ROOT,
    upload_backups_directory=lambda: UPLOAD_BACKUPS_DIR,
    validate_uploaded_workbook=lambda *args, **kwargs: validate_uploaded_workbook(*args, **kwargs),
    logged_users_for_series=lambda series: user_changes.logged_users_for_series(_alias_change_services, series)["usuarios"],
))
flush_pending_history_annotations = _inventory_history_exports.flush_pending_history_annotations
migrate_legacy_history_annotations = _inventory_history_exports.migrate_legacy_history_annotations
_build_comparison_restore_context = _inventory_history_exports.build_comparison_restore_context
_comparison_event_has_restorable_backup = _inventory_history_exports.comparison_event_has_restorable_backup
_comparison_event_type = _inventory_history_exports.comparison_event_type
_comparison_restore_backup_paths = _inventory_history_exports.comparison_restore_backup_paths
_current_inventory_snapshot = _inventory_history_exports.current_inventory_snapshot
_sha256_file = _inventory_history_exports.sha256_file
_validated_comparison_backup_files = _inventory_history_exports.validated_comparison_backup_files
excluir_comparacao_historica = _inventory_history_exports.excluir_comparacao_historica
exportar_historico_comparacoes_excel = _inventory_history_exports.exportar_historico_comparacoes_excel
get_comparacao_historica = _inventory_history_exports.get_comparacao_historica
get_historico_comparacoes = _inventory_history_exports.get_historico_comparacoes
salvar_anotacao_comparacao = _inventory_history_exports.salvar_anotacao_comparacao
salvar_anotacao_comparacao_historica = _inventory_history_exports.salvar_anotacao_comparacao_historica
resolver_divergencia_apelido_manual = _inventory_history_exports.resolver_divergencia_apelido_manual
restaurar_comparacao_historica = _inventory_history_exports.restaurar_comparacao_historica
visualizar_restauracao_comparacao = _inventory_history_exports.visualizar_restauracao_comparacao
app.include_router(inventory_history_router)

operations_router = create_operations_router(
    load_comparison_history=load_comparison_history,
    comparison_history_for_display=_inventory_history_exports.history_annotation_view,
    ticket_report_path=lambda: TICKET_REPORT_FILE,
    transaction_root=lambda: TRANSACTION_ROOT,
    history_annotations_directory=lambda: HISTORY_ANNOTATIONS_DIR,
    local_annotation_outbox_directory=lambda: LOCAL_HISTORY_ANNOTATION_OUTBOX_DIR,
    history_annotations_sync_enabled=lambda: HISTORY_ANNOTATIONS_SYNC_ENABLED,
    load_history_annotation_records=lambda directory, **kwargs: load_history_annotations(directory, **kwargs),
    record_ticket_reconciliation=usage_metrics.record_ticket_reconciliation,
)
app.include_router(operations_router)


def _monthly_report_inventory() -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    with STATE_LOCK:
        return deepcopy(list(inventory_state.equipamentos_cache.values())), deepcopy(inventory_state.stats_cache)


_monthly_report_services = MonthlyReportServices(
    inventory_snapshot=_monthly_report_inventory,
    load_comparison_history=load_comparison_history,
    comparison_history_for_display=comparison_history_for_display,
    ticket_report_path=lambda: TICKET_REPORT_FILE,
    read_usage_metrics=usage_metrics.read,
)
app.include_router(create_monthly_report_router(_monthly_report_services))


def _generate_previous_month_report_best_effort() -> None:
    """Save last month's report when missing; failures wait for the next check."""
    if get_transaction_recovery_block() is not None:
        return
    try:
        generate_previous_month_report(_monthly_report_services, MONTHLY_REPORTS_DIR, TRANSACTION_ROOT)
    except Exception as error:
        logger.warning("O relatório mensal automático não pôde ser gerado agora (erro=%s); nova tentativa em breve.", type(error).__name__)


async def _monthly_report_worker() -> None:
    # The first check waits for the startup imports; later ones also cover month turns while running.
    await asyncio.sleep(MONTHLY_REPORT_FIRST_CHECK_SECONDS)
    while True:
        await asyncio.to_thread(_generate_previous_month_report_best_effort)
        await asyncio.sleep(MONTHLY_REPORT_CHECK_SECONDS)

# Serve static files and frontend
STATIC_DIR.mkdir(parents=True, exist_ok=True)
app.mount("/static", RevalidatingStaticFiles(directory=str(STATIC_DIR)), name="static")

@app.get("/favicon.ico", include_in_schema=False)
def serve_favicon():
    return FileResponse(STATIC_DIR / "favicon.svg", media_type="image/svg+xml")

@app.get("/")
def read_root():
    index_file = STATIC_DIR / "index.html"
    if index_file.exists():
        return frontend_response(STATIC_DIR)
    return {
        "mensagem": "API Simpress Transferência ativa. O arquivo static/index.html está sendo inicializado.",
        "docs": "/docs"
    }

if __name__ == "__main__":
    import uvicorn
    log_file = configure_logging(LOG_DIR)
    try:
        equipment_source: Path | str | None = get_primary_equipment_file()
    except FileBusyError:
        equipment_source = "INDISPONÍVEL no momento (arquivo em uso)"
    print(f"[INFO] Diretório do projeto: {BASE_DIR}")
    print(f"[INFO] Modelo de transferência: {TEMPLATE_FILE} ({'OK' if TEMPLATE_FILE.is_file() else 'AUSENTE'})")
    print(f"[INFO] Fonte EquipamentosDetalhes: {equipment_source or 'AUSENTE'}")
    print(f"[INFO] Porta local: {SERVER_PORT}")
    print(f"[INFO] Log local: {log_file}")
    # Access lines would carry ?q= searches (serials, hostnames) to the console.
    uvicorn.run(app, host="127.0.0.1", port=SERVER_PORT, proxy_headers=False, access_log=False)
