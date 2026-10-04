"""Atomic file operations and cell-scoped XLSX package validation."""

from __future__ import annotations

import hashlib
import json
import math
import os
import posixpath
import re
import shutil
import tempfile
import threading
import time
import warnings
import zipfile
from contextlib import ExitStack, contextmanager
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, Iterator, Optional
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import openpyxl
from app_config import MAX_XLSX_ARCHIVE_ENTRIES, MAX_XLSX_UNCOMPRESSED_BYTES
from xlsx_safety import validate_xlsx


class FileBusyError(RuntimeError):
    """A destination could not be locked or replaced after bounded retries."""


class FileTransactionError(RuntimeError):
    """A file transaction failed without silently claiming success."""


class TransactionConflictError(FileTransactionError):
    """Recovery found a destination that no longer matches either expected state."""


class WorkbookSafetyError(FileTransactionError):
    """A workbook contains features that this openpyxl flow cannot preserve safely."""


_PROCESS_LOCKS: Dict[str, threading.RLock] = {}
_PROCESS_LOCKS_GUARD = threading.Lock()
_REPLACE_ATTEMPTS = 4
_RETRY_DELAYS = (0.1, 0.25, 0.5)
_MAX_XLSX_ENTRIES = MAX_XLSX_ARCHIVE_ENTRIES
_MAX_XLSX_UNCOMPRESSED_BYTES = MAX_XLSX_UNCOMPRESSED_BYTES

_UNSUPPORTED_PART_PREFIXES = (
    "xl/activex/",
    "xl/ctrlprops/",
    "xl/embeddings/",
    "xl/externallinks/",
    "xl/model/",
    "xl/pivotcache/",
    "xl/pivottables/",
    "xl/querytables/",
    "xl/slicers/",
    "xl/threadedcomments/",
    "xl/timelinecaches/",
    "xl/webextensions/",
    "customxml/",
)
_UNSUPPORTED_PART_NAMES = {"xl/vbaproject.bin", "xl/persons/person.xml"}


def _canonical(path: Path) -> str:
    return os.path.normcase(str(path.resolve(strict=False)))


def _lockfile(path: Path) -> Path:
    return path.with_name(f".{path.name}.lock")


def _acquire_os_lock(handle: Any, deadline: float, label: str) -> None:
    if os.name == "nt":
        import msvcrt

        handle.seek(0, os.SEEK_END)
        if handle.tell() == 0:
            handle.write(b"\0")
            handle.flush()
        while True:
            try:
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                return
            except OSError as error:
                if time.monotonic() >= deadline:
                    raise FileBusyError(f"O arquivo {label} está sendo usado por outra operação local.") from error
                time.sleep(0.05)
    else:
        import fcntl

        while True:
            try:
                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                return
            except OSError as error:
                if time.monotonic() >= deadline:
                    raise FileBusyError(f"O arquivo {label} está sendo usado por outra operação local.") from error
                time.sleep(0.05)


def _release_os_lock(handle: Any) -> None:
    if os.name == "nt":
        import msvcrt

        handle.seek(0)
        msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
    else:
        import fcntl

        fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


@contextmanager
def lock_paths(paths: Iterable[Path], timeout: float = 15.0) -> Iterator[None]:
    """Acquire in-process and cross-process locks in a stable order."""
    ordered = sorted({_canonical(Path(path)): Path(path).resolve(strict=False) for path in paths}.items())
    deadline = time.monotonic() + timeout
    acquired_threads = []
    handles = []
    try:
        for key, _ in ordered:
            with _PROCESS_LOCKS_GUARD:
                process_lock = _PROCESS_LOCKS.setdefault(key, threading.RLock())
            remaining = max(0.0, deadline - time.monotonic())
            if not process_lock.acquire(timeout=remaining):
                raise FileBusyError("A planilha ou configuração está ocupada por outra operação local.")
            acquired_threads.append(process_lock)

        for _, target in ordered:
            target.parent.mkdir(parents=True, exist_ok=True)
            lock_path = _lockfile(target)
            handle = open(lock_path, "a+b")
            handles.append(handle)
            _acquire_os_lock(handle, deadline, target.name)

        yield
    finally:
        for handle in reversed(handles):
            try:
                _release_os_lock(handle)
            except OSError:
                pass
            handle.close()
        for process_lock in reversed(acquired_threads):
            process_lock.release()


def _hash_file(path: Path) -> Optional[str]:
    if not path.exists():
        return None
    digest = hashlib.sha256()
    with open(path, "rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_temp(target: Path, writer: Callable[[Path], None], validator: Optional[Callable[[Path], None]]) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temp_name = tempfile.mkstemp(
        prefix=f".{target.name}.", suffix=".tmp", dir=str(target.parent)
    )
    os.close(descriptor)
    temp_path = Path(temp_name)
    try:
        writer(temp_path)
        with open(temp_path, "r+b") as staged:
            staged.flush()
            os.fsync(staged.fileno())
        if validator:
            validator(temp_path)
        return temp_path
    except PermissionError as error:
        temp_path.unlink(missing_ok=True)
        raise FileBusyError(
            f"Não foi possível gravar o temporário de '{target.name}'. Verifique se o arquivo ou a pasta está aberto ou sendo sincronizado."
        ) from error
    except OSError as error:
        temp_path.unlink(missing_ok=True)
        raise FileTransactionError(f"Falha ao preparar '{target.name}': {error.strerror or error}.") from error
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise


def _replace_with_retries(source: Path, destination: Path) -> None:
    for attempt in range(_REPLACE_ATTEMPTS):
        try:
            os.replace(source, destination)
            _fsync_directory(destination.parent)
            return
        except PermissionError as error:
            if attempt == _REPLACE_ATTEMPTS - 1:
                raise FileBusyError(
                    f"Não foi possível substituir '{destination.name}'. Feche a planilha no Excel e tente novamente."
                ) from error
            time.sleep(_RETRY_DELAYS[min(attempt, len(_RETRY_DELAYS) - 1)])
        except OSError as error:
            if attempt == _REPLACE_ATTEMPTS - 1:
                raise FileTransactionError(f"Falha ao substituir '{destination.name}': {error.strerror or error}.") from error
            time.sleep(_RETRY_DELAYS[min(attempt, len(_RETRY_DELAYS) - 1)])


def _fsync_directory(directory: Path) -> None:
    if os.name == "nt":
        return
    descriptor = None
    try:
        descriptor = os.open(directory, os.O_RDONLY)
        os.fsync(descriptor)
    except OSError:
        pass
    finally:
        if descriptor is not None:
            os.close(descriptor)


def atomic_write_bytes(path: Path, data: bytes, validator: Optional[Callable[[Path], None]] = None) -> None:
    """Atomically replace one file from a validated same-directory temporary."""
    target = Path(path).resolve(strict=False)
    with lock_paths([target]):
        temp_path = _write_temp(target, lambda temp: temp.write_bytes(data), validator)
        try:
            _replace_with_retries(temp_path, target)
        finally:
            temp_path.unlink(missing_ok=True)


def atomic_write_json(path: Path, value: Any) -> None:
    payload = json.dumps(value, ensure_ascii=False, indent=2).encode("utf-8")
    atomic_write_bytes(Path(path), payload, lambda temp: json.loads(temp.read_text(encoding="utf-8")))


def atomic_copy_file(source: Path, destination: Path, validator: Optional[Callable[[Path], None]] = None) -> None:
    source = Path(source).resolve(strict=False)
    destination = Path(destination).resolve(strict=False)
    with lock_paths([source, destination]):
        def copy_to(temp: Path) -> None:
            shutil.copyfile(source, temp)

        temp_path = _write_temp(destination, copy_to, validator)
        try:
            _replace_with_retries(temp_path, destination)
        finally:
            temp_path.unlink(missing_ok=True)


def _safe_target(root: Path, relative_path: str) -> Path:
    candidate = (root / relative_path).resolve(strict=False)
    try:
        candidate.relative_to(root.resolve(strict=False))
    except ValueError as error:
        raise TransactionConflictError("O registro de transação contém um caminho fora do projeto.") from error
    return candidate


def _persist_manifest(manifest_path: Path, manifest: Dict[str, Any]) -> None:
    payload = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
    temp = _write_temp(manifest_path, lambda path: path.write_bytes(payload), None)
    try:
        _replace_with_retries(temp, manifest_path)
    finally:
        temp.unlink(missing_ok=True)


def _restore_entry(root: Path, transaction_dir: Path, entry: Dict[str, Any]) -> None:
    target = _safe_target(root, entry["target"])
    current_hash = _hash_file(target)
    original_hash = entry.get("original_hash")
    expected_hash = entry.get("new_hash")

    if current_hash == original_hash:
        return
    if current_hash != expected_hash:
        raise TransactionConflictError(
            f"A recuperação foi interrompida porque '{target.name}' não corresponde ao estado esperado; o arquivo foi preservado."
        )

    if original_hash is None:
        target.unlink(missing_ok=True)
        return

    backup = _safe_target(transaction_dir, entry["backup"])
    if _hash_file(backup) != original_hash:
        raise TransactionConflictError(f"A cópia de recuperação de '{target.name}' falhou na validação.")
    temp = _write_temp(target, lambda path: shutil.copyfile(backup, path), None)
    try:
        if _hash_file(temp) != original_hash:
            raise TransactionConflictError(f"A restauração temporária de '{target.name}' falhou na validação.")
        _replace_with_retries(temp, target)
    finally:
        temp.unlink(missing_ok=True)


def _recover_transaction(root: Path, transaction_dir: Path) -> None:
    manifest_path = transaction_dir / "manifest.json"
    if not manifest_path.exists():
        shutil.rmtree(transaction_dir, ignore_errors=True)
        return
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("version") != 1 or not isinstance(manifest.get("entries"), list):
            raise ValueError("versão ou lista de arquivos inválida")
    except Exception as error:
        raise TransactionConflictError(f"O registro de recuperação '{manifest_path.name}' está inválido: {error}.") from error

    if manifest.get("phase") == "committed":
        shutil.rmtree(transaction_dir, ignore_errors=True)
        return

    entries = manifest["entries"]
    for entry in reversed(entries):
        _restore_entry(root, transaction_dir, entry)
    shutil.rmtree(transaction_dir, ignore_errors=True)


def _recover_pending_locked(root: Path) -> None:
    transaction_root = root / ".transactions"
    if not transaction_root.is_dir():
        return
    for transaction_dir in sorted(transaction_root.iterdir()):
        if transaction_dir.is_dir():
            _recover_transaction(root, transaction_dir)


def recover_pending_transactions(root: Path) -> None:
    """Recover uncommitted transactions; safe to call repeatedly."""
    root = Path(root).resolve(strict=False)
    if not (root / ".transactions").is_dir():
        return
    with lock_paths([root / ".atomic-transactions"]):
        _recover_pending_locked(root)


class AtomicFileTransaction:
    """Journaled multi-file transaction with hash-checked, idempotent rollback."""

    def __init__(self, root: Path, targets: Iterable[Path]):
        self.root = Path(root).resolve(strict=False)
        self.targets = sorted({Path(item).resolve(strict=False) for item in targets}, key=_canonical)
        self.transaction_root = self.root / ".transactions"
        self.transaction_id = f"{time.strftime('%Y%m%d-%H%M%S')}-{os.getpid()}-{threading.get_ident()}-{os.urandom(4).hex()}"
        self.transaction_dir = self.transaction_root / self.transaction_id
        self.manifest_path = self.transaction_dir / "manifest.json"
        self.manifest: Dict[str, Any] = {"version": 1, "phase": "prepared", "entries": []}
        self._stack = ExitStack()
        self._entered = False
        self._committed = False
        self._closed = False

    def __enter__(self) -> "AtomicFileTransaction":
        transaction_dir_created = False
        try:
            self.root.mkdir(parents=True, exist_ok=True)
            self._stack.enter_context(lock_paths([self.root / ".atomic-transactions"]))
            _recover_pending_locked(self.root)
            if self.targets:
                self._stack.enter_context(lock_paths(self.targets))
            self.transaction_dir.mkdir(parents=True, exist_ok=False)
            transaction_dir_created = True
            (self.transaction_dir / "backups").mkdir()
            self._entered = True
            return self
        except BaseException as error:
            self._closed = True
            if transaction_dir_created:
                try:
                    shutil.rmtree(self.transaction_dir)
                except BaseException as cleanup_error:
                    error.add_note(f"A limpeza do diretório parcial da transação também falhou: {cleanup_error!r}")
            try:
                self._stack.close()
            except BaseException as cleanup_error:
                error.add_note(f"A liberação dos recursos da transação também falhou: {cleanup_error!r}")
            raise

    def add_target(self, target: Path) -> None:
        """Add and lock a destination discovered while the transaction-wide lock is held."""
        if not self._entered or self.manifest["entries"]:
            raise FileTransactionError("Novos arquivos só podem ser incluídos antes do primeiro estágio.")
        resolved = Path(target).resolve(strict=False)
        if resolved in self.targets:
            return
        self._stack.enter_context(lock_paths([resolved]))
        self.targets.append(resolved)

    def _target_relative(self, target: Path) -> str:
        try:
            return str(Path(target).resolve(strict=False).relative_to(self.root)).replace("\\", "/")
        except ValueError as error:
            raise FileTransactionError("Todos os arquivos da transação precisam ficar dentro da pasta do projeto.") from error

    def _find_entry(self, target: Path) -> Optional[Dict[str, Any]]:
        relative = self._target_relative(target)
        return next((item for item in self.manifest["entries"] if item["target"] == relative), None)

    def _stage(self, target: Path, writer: Callable[[Path], None], validator: Optional[Callable[[Path], None]], delete: bool) -> None:
        if not self._entered or self._closed:
            raise FileTransactionError("A transação não está aberta.")
        target = Path(target).resolve(strict=False)
        if target not in self.targets:
            raise FileTransactionError(f"'{target.name}' não foi incluído no bloqueio da transação.")
        if self._find_entry(target):
            raise FileTransactionError(f"'{target.name}' foi incluído mais de uma vez na mesma transação.")

        original_hash = _hash_file(target)
        index = len(self.manifest["entries"])
        backup_relative = f"backups/{index:04d}.bin"
        backup_path = self.transaction_dir / backup_relative
        if original_hash is not None:
            shutil.copyfile(target, backup_path)
            with open(backup_path, "r+b") as backup_file:
                backup_file.flush()
                os.fsync(backup_file.fileno())
            if _hash_file(backup_path) != original_hash:
                raise FileTransactionError(f"A cópia de segurança de '{target.name}' não passou na validação.")

        stage_path = None
        new_hash = None
        if not delete:
            stage_path = _write_temp(target, writer, validator)
            new_hash = _hash_file(stage_path)

        entry = {
            "target": self._target_relative(target),
            "stage": self._target_relative(stage_path) if stage_path else None,
            "backup": backup_relative if original_hash is not None else None,
            "original_hash": original_hash,
            "new_hash": new_hash,
        }
        self.manifest["entries"].append(entry)
        try:
            _persist_manifest(self.manifest_path, self.manifest)
        except Exception:
            if stage_path:
                stage_path.unlink(missing_ok=True)
            raise

    def stage_bytes(self, target: Path, data: bytes, validator: Optional[Callable[[Path], None]] = None) -> None:
        self._stage(target, lambda temp: temp.write_bytes(data), validator, delete=False)

    def stage_file(
        self,
        target: Path,
        writer: Callable[[Path], None],
        validator: Optional[Callable[[Path], None]] = None,
    ) -> None:
        self._stage(target, writer, validator, delete=False)

    def stage_delete(self, target: Path) -> None:
        self._stage(target, lambda _: None, None, delete=True)

    def update_json(self, target: Path, updater: Callable[[Any], Any], default: Any) -> Any:
        target = Path(target).resolve(strict=False)
        current = default
        if target.exists():
            try:
                current = json.loads(target.read_text(encoding="utf-8"))
            except Exception as error:
                raise FileTransactionError(f"Não foi possível ler '{target.name}' para atualizá-lo com segurança.") from error
        updated = updater(current)
        payload = json.dumps(updated, ensure_ascii=False, indent=2).encode("utf-8")
        self.stage_bytes(target, payload, lambda temp: json.loads(temp.read_text(encoding="utf-8")))
        return updated

    def _apply_entry(self, entry: Dict[str, Any]) -> None:
        target = _safe_target(self.root, entry["target"])
        current_hash = _hash_file(target)
        original_hash = entry.get("original_hash")
        new_hash = entry.get("new_hash")
        if current_hash == new_hash:
            return
        if current_hash != original_hash:
            raise TransactionConflictError(
                f"'{target.name}' mudou depois da preparação; a operação foi cancelada sem sobrescrever o arquivo."
            )

        if new_hash is None:
            if target.exists():
                for attempt in range(_REPLACE_ATTEMPTS):
                    try:
                        target.unlink()
                        break
                    except PermissionError as error:
                        if attempt == _REPLACE_ATTEMPTS - 1:
                            raise FileBusyError(f"Feche '{target.name}' no Excel e tente novamente.") from error
                        time.sleep(_RETRY_DELAYS[min(attempt, len(_RETRY_DELAYS) - 1)])
                _fsync_directory(target.parent)
            return

        stage = _safe_target(self.root, entry["stage"])
        if _hash_file(stage) != new_hash:
            raise TransactionConflictError(f"O temporário de '{target.name}' não passou na validação de integridade.")
        _replace_with_retries(stage, target)
        if _hash_file(target) != new_hash:
            raise TransactionConflictError(f"'{target.name}' não corresponde ao conteúdo preparado após a substituição.")

    def apply(self) -> None:
        if not self._entered:
            raise FileTransactionError("A transação não está aberta.")
        self.manifest["phase"] = "applying"
        _persist_manifest(self.manifest_path, self.manifest)
        for index, entry in enumerate(self.manifest["entries"]):
            self._apply_entry(entry)
            self.manifest["applied_through"] = index
            _persist_manifest(self.manifest_path, self.manifest)

    def commit(self) -> None:
        self.apply()
        for entry in self.manifest["entries"]:
            target = _safe_target(self.root, entry["target"])
            if _hash_file(target) != entry.get("new_hash"):
                raise TransactionConflictError(f"A confirmação de '{target.name}' falhou na verificação final.")
        self.manifest["phase"] = "committed"
        _persist_manifest(self.manifest_path, self.manifest)
        self._committed = True
        shutil.rmtree(self.transaction_dir, ignore_errors=True)

    def rollback(self) -> None:
        if self.manifest_path.exists():
            for entry in reversed(self.manifest["entries"]):
                _restore_entry(self.root, self.transaction_dir, entry)
        shutil.rmtree(self.transaction_dir, ignore_errors=True)

    def close_after_interruption(self) -> None:
        """Release test/process resources while leaving the durable journal for startup recovery."""
        self._closed = True
        self._stack.close()

    def __exit__(self, exc_type: Any, exc: Any, traceback: Any) -> bool:
        rollback_error = None
        try:
            if not self._committed and not self._closed:
                try:
                    self.rollback()
                except Exception as error:
                    rollback_error = error
        finally:
            self._closed = True
            self._stack.close()
        if rollback_error:
            if exc is not None:
                raise TransactionConflictError(f"A gravação falhou e a recuperação precisa de conferência: {rollback_error}") from exc
            raise rollback_error
        return False


def _cell_signature(cell: Any) -> tuple:
    hyperlink = None
    if cell.hyperlink is not None:
        hyperlink = (cell.hyperlink.target, cell.hyperlink.location, cell.hyperlink.tooltip, cell.hyperlink.display)
    comment = None
    if cell.comment is not None:
        comment = (cell.comment.text, cell.comment.author)
    return (cell.value, cell.data_type, cell.style_id, hyperlink, comment)


def _dimension_signature(dimensions: Any) -> tuple:
    result = []
    for key, item in sorted(dimensions.items(), key=lambda pair: str(pair[0])):
        result.append((
            str(key), getattr(item, "min", None), getattr(item, "max", None),
            getattr(item, "width", None), getattr(item, "height", None),
            getattr(item, "hidden", None), getattr(item, "bestFit", None),
            getattr(item, "outlineLevel", None), getattr(item, "collapsed", None),
            getattr(item, "style", None), getattr(item, "style_id", None),
            getattr(item, "thickTop", None), getattr(item, "thickBottom", None),
            getattr(item, "customHeight", None),
        ))
    return tuple(result)


def _conditional_formatting_signature(conditional_formatting: Any) -> tuple:
    return tuple(
        (repr(condition), tuple(repr(rule) for rule in rules))
        for condition, rules in conditional_formatting._cf_rules.items()
    )


def _properties_signature(properties: Any) -> tuple:
    """Exclude timestamps synthesized by openpyxl when absent from the package."""
    names = getattr(type(properties), "__elements__", ())
    return tuple(
        (name, repr(getattr(properties, name, None)))
        for name in names
        if name not in {"created", "modified"}
    )


def workbook_snapshot(workbook: Any) -> Dict[str, Any]:
    """Capture cell contents and workbook layout for round-trip comparisons."""
    worksheets = {}
    for ws in workbook.worksheets:
        cells = {
            cell.coordinate: _cell_signature(cell)
            for cell in ws._cells.values()
        }
        tables = tuple(sorted(
            (name, ws.tables[name].ref, ws.tables[name].displayName)
            for name in ws.tables.keys()
        ))
        validations = tuple(sorted((item.type, item.formula1, item.formula2, item.sqref.__str__()) for item in ws.data_validations.dataValidation))
        worksheets[ws.title] = {
            "cells": cells,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "dimension": ws.calculate_dimension(),
            "sheet_state": ws.sheet_state,
            "merged": tuple(sorted(str(item) for item in ws.merged_cells.ranges)),
            "row_dimensions": _dimension_signature(ws.row_dimensions),
            "column_dimensions": _dimension_signature(ws.column_dimensions),
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "auto_filter": ws.auto_filter.ref,
            "auto_filter_details": repr(ws.auto_filter),
            "tables": tables,
            "validations": validations,
            "conditional_formatting": _conditional_formatting_signature(ws.conditional_formatting),
            "chart_count": len(ws._charts),
            "image_count": len(ws._images),
            "protection": repr(ws.protection),
            "sheet_format": (
                ws.sheet_format.defaultRowHeight, ws.sheet_format.defaultColWidth,
                ws.sheet_format.zeroHeight, ws.sheet_format.baseColWidth,
            ),
            "print_area": str(ws.print_area),
            "print_title_rows": ws.print_title_rows,
            "print_title_cols": ws.print_title_cols,
            "page_setup": repr(ws.page_setup),
            "page_margins": repr(ws.page_margins),
            "sheet_properties": repr(ws.sheet_properties),
            "views": repr(ws.views),
            "print_options": repr(ws.print_options),
        }
    defined_names = tuple(sorted((str(name), str(item.attr_text)) for name, item in workbook.defined_names.items()))
    return {
        "sheetnames": tuple(workbook.sheetnames),
        "active_sheet": workbook.active.title if workbook.active else None,
        "epoch": str(workbook.epoch),
        "defined_names": defined_names,
        "calculation": repr(workbook.calculation),
        "properties": _properties_signature(workbook.properties),
        "style_catalog": tuple(repr(style) for style in workbook._cell_styles),
        "named_styles": tuple(repr(style) for style in workbook._named_styles),
        "worksheets": worksheets,
    }


_XLSX_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_XLSX_DOCUMENT_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_XLSX_PACKAGE_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_CELL_TAG_RE = re.compile(rb"<(?P<name>(?:[A-Za-z_][A-Za-z0-9_.-]*:)?c)(?=[\s/>])")
_CELL_REFERENCE_RE = re.compile(rb"(?<![A-Za-z0-9_.:-])r\s*=\s*([\"'])(.*?)\1")
_CELL_ATTRIBUTE_RE = re.compile(rb"\s+t\s*=\s*(?:\"[^\"]*\"|'[^']*')")
_ILLEGAL_XML_TEXT_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def _worksheet_part_paths(archive: zipfile.ZipFile) -> Dict[str, str]:
    """Map worksheet titles to their package paths without rewriting workbook XML."""
    try:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        relationships = {
            item.attrib["Id"]: item
            for item in relationships_root.findall(f"{_XLSX_PACKAGE_REL_NS}Relationship")
        }
        package_names = {item.filename for item in archive.infolist()}
        result = {}
        sheets = workbook_root.find(f"{_XLSX_MAIN_NS}sheets")
        if sheets is None:
            raise ValueError("A pasta de trabalho não contém abas.")
        for sheet in sheets.findall(f"{_XLSX_MAIN_NS}sheet"):
            title = sheet.attrib["name"]
            relationship_id = sheet.attrib[f"{_XLSX_DOCUMENT_REL_NS}id"]
            relationship = relationships[relationship_id]
            if relationship.attrib.get("TargetMode", "Internal").lower() == "external":
                raise ValueError("Uma aba aponta para um recurso externo.")
            target = relationship.attrib["Target"].replace("\\", "/")
            part_path = (
                posixpath.normpath(target.lstrip("/"))
                if target.startswith("/")
                else posixpath.normpath(posixpath.join("xl", target))
            )
            if part_path.startswith("../") or part_path not in package_names:
                raise ValueError("O pacote não contém a parte de uma aba declarada.")
            result[title] = part_path
        return result
    except (KeyError, OSError, zipfile.BadZipFile, ET.ParseError, ValueError) as error:
        raise WorkbookSafetyError("Não foi possível localizar as abas dentro do arquivo Excel.") from error


def _xml_tag_end(xml: bytes, start: int) -> int:
    """Return the byte after an XML start tag, accounting for quoted attributes."""
    quote = None
    for position in range(start, len(xml)):
        byte = xml[position]
        if quote is not None:
            if byte == quote:
                quote = None
        elif byte in (ord('"'), ord("'")):
            quote = byte
        elif byte == ord(">"):
            return position + 1
    raise WorkbookSafetyError("Uma aba contém XML incompleto e não pôde ser preservada.")


def _cell_ranges(xml: bytes) -> Dict[str, tuple]:
    """Find worksheet cell elements while leaving all other XML bytes untouched."""
    ranges = {}
    cursor = 0
    while True:
        match = _CELL_TAG_RE.search(xml, cursor)
        if match is None:
            break
        opening_end = _xml_tag_end(xml, match.start())
        opening = xml[match.start():opening_end]
        reference = _CELL_REFERENCE_RE.search(opening)
        if reference is None:
            raise WorkbookSafetyError("Uma célula da aba não tem referência válida.")
        address = reference.group(2).decode("ascii", errors="strict").upper()
        if opening[:-1].rstrip().endswith(b"/"):
            element_end = opening_end
        else:
            closing_tag = b"</" + match.group("name") + b">"
            close_start = xml.find(closing_tag, opening_end)
            if close_start < 0:
                raise WorkbookSafetyError("Uma célula da aba contém XML incompleto.")
            element_end = close_start + len(closing_tag)
        if address in ranges:
            raise WorkbookSafetyError(f"A aba contém a célula '{address}' mais de uma vez.")
        ranges[address] = (match.start(), element_end, opening_end, match.group("name"))
        cursor = element_end
    return ranges


def _replacement_cell_xml(opening: bytes, tag_name: bytes, value: Any, existing_content: bytes) -> bytes:
    """Encode one authorized cell value without changing its other attributes."""
    if re.search(rb"<(?:[A-Za-z_][A-Za-z0-9_.-]*:)?extLst(?=[\s/>])", existing_content):
        raise WorkbookSafetyError("A célula autorizada contém metadados avançados que não podem ser substituídos.")
    attributes = opening[1 + len(tag_name):-1]
    self_closing = attributes.rstrip().endswith(b"/")
    if self_closing:
        attributes = attributes.rstrip()[:-1]
    attributes = _CELL_ATTRIBUTE_RE.sub(b"", attributes).rstrip()

    if isinstance(value, str):
        if _ILLEGAL_XML_TEXT_RE.search(value):
            raise WorkbookSafetyError("O valor autorizado contém caracteres que o Excel não aceita.")
        encoded = escape(value).replace("\r", "&#13;").encode("utf-8")
        preserve_space = value != value.strip()
        space_attribute = b' xml:space="preserve"' if preserve_space else b""
        cell_type = b' t="inlineStr"'
        content = b"<is><t" + space_attribute + b">" + encoded + b"</t></is>"
    elif isinstance(value, bool):
        cell_type = b' t="b"'
        content = b"<v>" + (b"1" if value else b"0") + b"</v>"
    elif isinstance(value, int):
        cell_type = b""
        content = b"<v>" + str(value).encode("ascii") + b"</v>"
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise WorkbookSafetyError("O valor numérico autorizado não é finito.")
        cell_type = b""
        content = b"<v>" + repr(value).encode("ascii") + b"</v>"
    elif value is None:
        cell_type = b""
        content = b""
    else:
        raise WorkbookSafetyError("O tipo de valor autorizado não pode ser gravado com segurança.")

    return b"<" + tag_name + attributes + cell_type + b">" + content + b"</" + tag_name + b">"


def _patch_worksheet_cells(xml: bytes, updates: Dict[str, Any]) -> bytes:
    """Replace selected existing cell values and retain all other sheet XML bytes."""
    ranges = _cell_ranges(xml)
    patches = []
    for address, value in updates.items():
        normalized_address = str(address).upper()
        cell_range = ranges.get(normalized_address)
        if cell_range is None:
            raise WorkbookSafetyError(f"A célula autorizada '{normalized_address}' não existe no modelo.")
        start, end, opening_end, tag_name = cell_range
        opening = xml[start:opening_end]
        self_closing = opening[:-1].rstrip().endswith(b"/")
        content_end = end if self_closing else end - len(b"</" + tag_name + b">")
        existing_content = b"" if self_closing else xml[opening_end:content_end]
        replacement = _replacement_cell_xml(opening, tag_name, value, existing_content)
        patches.append((start, end, replacement))

    patched = xml
    for start, end, replacement in sorted(patches, reverse=True):
        patched = patched[:start] + replacement + patched[end:]
    return patched


def _write_package_with_cell_updates(
    source: Path,
    destination: Path,
    expected_values: Dict[tuple, Any],
) -> set:
    """Copy the original XLSX package and patch only authorized existing cells."""
    updates_by_sheet: Dict[str, Dict[str, Any]] = {}
    for key, value in expected_values.items():
        if not isinstance(key, tuple) or len(key) != 2:
            raise WorkbookSafetyError("A lista de alterações autorizadas da planilha é inválida.")
        sheet_name, address = key
        updates_by_sheet.setdefault(str(sheet_name), {})[str(address)] = value

    try:
        with zipfile.ZipFile(source, "r") as original:
            infos = original.infolist()
            names = [info.filename for info in infos]
            if len(names) != len(set(names)):
                raise WorkbookSafetyError("O arquivo Excel contém partes duplicadas no pacote.")
            worksheet_paths = _worksheet_part_paths(original)
            updated_paths = {}
            for sheet_name, updates in updates_by_sheet.items():
                part_path = worksheet_paths.get(sheet_name)
                if part_path is None:
                    raise WorkbookSafetyError(f"A aba '{sheet_name}' não existe no modelo original.")
                updated_paths[part_path] = _patch_worksheet_cells(original.read(part_path), updates)

            with zipfile.ZipFile(destination, "w") as output:
                output.comment = original.comment
                for info in infos:
                    content = updated_paths.get(info.filename)
                    if content is None:
                        content = original.read(info.filename)
                    output.writestr(info, content)
            return set(updated_paths)
    except WorkbookSafetyError:
        raise
    except (OSError, zipfile.BadZipFile, KeyError, RuntimeError) as error:
        raise WorkbookSafetyError("Não foi possível preservar a estrutura original do arquivo Excel.") from error


def _changed_expected_values(original_snapshot: Dict[str, Any], expected_values: Dict[tuple, Any]) -> Dict[tuple, Any]:
    """Skip authorized assignments that already match the original cell value."""
    changed = {}
    for (sheet_name, coordinate), value in expected_values.items():
        sheet = original_snapshot.get("worksheets", {}).get(sheet_name, {})
        original_cell = sheet.get("cells", {}).get(str(coordinate).upper())
        original_value = original_cell[0] if original_cell is not None else None
        if type(original_value) is not type(value) or original_value != value:
            changed[(sheet_name, coordinate)] = value
    return changed


def _updated_worksheet_parts(source: Path, expected_values: Dict[tuple, Any]) -> set:
    """Return package entries that may differ because a cell was explicitly updated."""
    if not expected_values:
        return set()
    try:
        with zipfile.ZipFile(source) as archive:
            paths = _worksheet_part_paths(archive)
    except (OSError, zipfile.BadZipFile) as error:
        raise WorkbookSafetyError("Não foi possível validar as abas da planilha temporária.") from error
    result = set()
    for sheet_name, _ in expected_values:
        part_path = paths.get(sheet_name)
        if part_path is None:
            raise WorkbookSafetyError(f"A aba '{sheet_name}' não existe na planilha temporária.")
        result.add(part_path)
    return result


def _expected_cell_data_type(value: Any) -> str:
    if isinstance(value, str):
        return "s"
    if isinstance(value, bool):
        return "b"
    if isinstance(value, (int, float)) or value is None:
        return "n"
    raise WorkbookSafetyError("O tipo de valor autorizado não pode ser validado com segurança.")


_EMPTY_CELL_DATA_TYPES = {"n", "s", "inlineStr"}


def _cell_values_match(actual: Any, expected: Any) -> bool:
    """Treat an empty string and a blank cell as the same Excel value."""
    return actual == expected or (
        actual in (None, "") and expected in (None, "")
    )


def _cell_data_type_matches(actual: tuple, expected_value: Any) -> bool:
    """Allow Excel's equivalent blank representations without accepting other types."""
    if actual[1] == _expected_cell_data_type(expected_value):
        return True
    return (
        actual[0] in (None, "")
        and expected_value in (None, "")
        and actual[1] in _EMPTY_CELL_DATA_TYPES
    )


def _cell_metadata_matches(actual: tuple, prior: tuple) -> bool:
    """Compare cell metadata, normalizing only empty-string storage differences."""
    if actual[2:] != prior[2:]:
        return False
    if actual[1] == prior[1]:
        return True
    return (
        actual[0] in (None, "")
        and prior[0] in (None, "")
        and actual[1] in _EMPTY_CELL_DATA_TYPES
        and prior[1] in _EMPTY_CELL_DATA_TYPES
    )


_VML_SHAPE_TAG = "{urn:schemas-microsoft-com:vml}shape"
_EXCEL_VML_CLIENT_DATA_TAG = "{urn:schemas-microsoft-com:office:excel}ClientData"


def _vml_part_contains_only_comment_notes(content: bytes) -> bool:
    """Return true only for legacy VML parts made exclusively of comment notes."""
    try:
        root = ET.fromstring(content)
    except (ET.ParseError, ValueError):
        return False

    def local_name(element: ET.Element) -> str:
        return element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""

    shapes = [element for element in root.iter() if local_name(element) == "shape"]
    client_data = [element for element in root.iter() if local_name(element) == "ClientData"]
    if not shapes or len(shapes) != len(client_data):
        return False

    for shape in shapes:
        if shape.tag != _VML_SHAPE_TAG:
            return False
        shape_data = [element for element in shape.iter() if local_name(element) == "ClientData"]
        if (
            len(shape_data) != 1
            or shape_data[0].tag != _EXCEL_VML_CLIENT_DATA_TAG
            or shape_data[0].attrib.get("ObjectType") != "Note"
        ):
            return False
    return True


def _package_features(
    source: Path,
    *,
    allow_threaded_comments: bool = False,
    allow_legacy_comments: bool = False,
    allow_custom_xml_passthrough: bool = False,
) -> None:
    try:
        with zipfile.ZipFile(source) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > _MAX_XLSX_ENTRIES:
                raise WorkbookSafetyError("A planilha tem uma estrutura ZIP inválida ou grande demais para validar.")
            if sum(entry.file_size for entry in entries) > _MAX_XLSX_UNCOMPRESSED_BYTES:
                raise WorkbookSafetyError("O conteúdo descompactado da planilha excede o limite de validação.")
            names = {item.filename for item in entries}
            lowered = {name.lower() for name in names}
            original_names = {name.lower(): name for name in names}
            for name in lowered:
                if allow_threaded_comments and (
                    name.startswith("xl/threadedcomments/") or name == "xl/persons/person.xml"
                ):
                    continue
                if allow_custom_xml_passthrough and name.startswith("customxml/"):
                    continue
                if name in _UNSUPPORTED_PART_NAMES or name.startswith(_UNSUPPORTED_PART_PREFIXES):
                    raise WorkbookSafetyError(f"A planilha contém o recurso '{name}' que esta gravação não preserva com segurança.")
                if name.startswith("xl/drawings/") and "vml" in name:
                    if (
                        allow_legacy_comments
                        and name.endswith(".vml")
                        and _vml_part_contains_only_comment_notes(archive.read(original_names[name]))
                    ):
                        continue
                    raise WorkbookSafetyError("A planilha contém comentários ou desenhos VML que esta gravação não preserva com segurança.")
            for name in names:
                if name.lower().startswith("xl/drawings/") and name.lower().endswith(".xml"):
                    content = archive.read(name)
                    if re.search(rb"<(?:[A-Za-z0-9_]+:)?sp(?:\s|>)", content) or b"AlternateContent" in content or b"<extLst" in content:
                        raise WorkbookSafetyError("A planilha contém formas ou extensões de desenho que esta gravação não preserva com segurança.")
    except WorkbookSafetyError:
        raise
    except (OSError, zipfile.BadZipFile, ValueError) as error:
        raise WorkbookSafetyError("A planilha não pôde ser lida para validar seus recursos.") from error


def load_workbook_safely(
    source: Any,
    *,
    allow_threaded_comments: bool = False,
    allow_legacy_comments: bool = False,
    allow_custom_xml_passthrough: bool = False,
) -> Any:
    """Load formulas/rich text and reject unsupported package features or warnings."""
    temporary_handle = None
    if isinstance(source, (str, os.PathLike, Path)):
        source_path = Path(source)
        validate_xlsx(source_path)
        _package_features(
            source_path,
            allow_threaded_comments=allow_threaded_comments,
            allow_legacy_comments=allow_legacy_comments,
            allow_custom_xml_passthrough=allow_custom_xml_passthrough,
        )
        if source_path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            # openpyxl checks filename extensions even for a valid XLSX archive;
            # staged files use a neutral .tmp suffix until commit.
            temporary_handle = source_path.open("rb")
            source = temporary_handle
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        try:
            workbook = openpyxl.load_workbook(source, data_only=False, keep_links=True, rich_text=True)
        finally:
            if temporary_handle is not None:
                temporary_handle.close()
    unsupported = [
        item for item in caught
        if re.search(r"unsupported|not supported|will be removed|extension|drawing", str(item.message), re.I)
        and not (
            allow_threaded_comments
            and str(item.message).strip() == "Unknown extension is not supported and will be removed"
        )
    ]
    if unsupported:
        workbook.close()
        raise WorkbookSafetyError(f"A biblioteca Excel avisou que um recurso não seria preservado: {unsupported[0].message}")
    return workbook


def validate_workbook_roundtrip(
    original_snapshot: Dict[str, Any],
    candidate_path: Path,
    expected_values: Dict[tuple, Any],
    *,
    original_package_parts: Optional[Dict[str, str]] = None,
    original_package_entries: Optional[Dict[str, str]] = None,
    allow_dimension_growth: bool = False,
    allow_threaded_comments: bool = False,
    allow_legacy_comments: bool = False,
    allow_custom_xml_passthrough: bool = False,
) -> None:
    """Reopen the temporary workbook and ensure only explicitly allowed values changed."""
    _package_features(
        candidate_path,
        allow_threaded_comments=allow_threaded_comments,
        allow_legacy_comments=allow_legacy_comments,
        allow_custom_xml_passthrough=allow_custom_xml_passthrough,
    )
    if original_package_entries is not None:
        candidate_entries = _package_hashes(candidate_path)
        if set(candidate_entries) != set(original_package_entries):
            raise WorkbookSafetyError("A planilha temporária alterou as partes do arquivo fora das células autorizadas.")
        allowed_parts = _updated_worksheet_parts(candidate_path, expected_values)
        changed_parts = {
            name for name, digest in original_package_entries.items()
            if candidate_entries[name] != digest
        }
        if changed_parts - allowed_parts:
            raise WorkbookSafetyError("A planilha temporária alterou recursos fora das células autorizadas.")
    if original_package_parts is not None:
        if _preserved_package_parts(candidate_path) != original_package_parts:
            raise WorkbookSafetyError("A planilha temporária alterou um desenho, gráfico ou mídia fora das células autorizadas.")
    workbook = load_workbook_safely(
        candidate_path,
        allow_threaded_comments=allow_threaded_comments,
        allow_legacy_comments=allow_legacy_comments,
        allow_custom_xml_passthrough=allow_custom_xml_passthrough,
    )
    try:
        candidate = workbook_snapshot(workbook)
    finally:
        workbook.close()

    if candidate["sheetnames"] != original_snapshot["sheetnames"]:
        raise WorkbookSafetyError("A planilha temporária alterou as abas existentes.")
    if candidate["active_sheet"] != original_snapshot["active_sheet"] or candidate["epoch"] != original_snapshot["epoch"]:
        raise WorkbookSafetyError("A planilha temporária alterou a estrutura do arquivo.")
    if candidate["defined_names"] != original_snapshot["defined_names"]:
        raise WorkbookSafetyError("A planilha temporária alterou fórmulas ou nomes definidos.")
    if any(candidate[key] != original_snapshot[key] for key in ("calculation", "properties", "style_catalog", "named_styles")):
        raise WorkbookSafetyError("A planilha temporária alterou as propriedades do arquivo.")

    expected_by_sheet: Dict[str, Dict[str, Any]] = {}
    for (sheet_name, coordinate), value in expected_values.items():
        expected_by_sheet.setdefault(sheet_name, {})[coordinate] = value

    for sheet_name, before in original_snapshot["worksheets"].items():
        after = candidate["worksheets"][sheet_name]
        allowed = expected_by_sheet.get(sheet_name, {})
        structural_keys = (
            "sheet_state", "merged", "row_dimensions", "column_dimensions", "freeze_panes",
            "auto_filter", "auto_filter_details", "tables", "validations", "conditional_formatting",
            "chart_count", "image_count", "protection", "sheet_format",
            "print_area", "print_title_rows", "print_title_cols", "page_setup", "page_margins",
            "sheet_properties", "views", "print_options",
        )
        for key in structural_keys:
            if before[key] != after[key]:
                raise WorkbookSafetyError(f"A planilha temporária alterou a estrutura '{key}' da aba '{sheet_name}'.")
        if not allow_dimension_growth and any(before[key] != after[key] for key in ("max_row", "max_column", "dimension")):
            raise WorkbookSafetyError(f"A planilha temporária alterou as dimensões da aba '{sheet_name}'.")

        before_cells = before["cells"]
        after_cells = after["cells"]
        for coordinate in set(before_cells) | set(after_cells) | set(allowed):
            prior = before_cells.get(coordinate)
            actual = after_cells.get(coordinate)
            expected = allowed.get(coordinate, prior[0] if prior else None)
            if actual is None:
                raise WorkbookSafetyError(f"A célula '{sheet_name}!{coordinate}' desapareceu na planilha temporária.")
            value_matches = _cell_values_match(actual[0], expected)
            if not value_matches:
                raise WorkbookSafetyError(f"O valor da célula '{sheet_name}!{coordinate}' não corresponde ao esperado.")
            if prior is not None:
                if coordinate in allowed:
                    metadata_matches = (
                        _cell_data_type_matches(actual, allowed[coordinate])
                        and actual[2:] == prior[2:]
                    )
                else:
                    metadata_matches = _cell_metadata_matches(actual, prior)
                if not metadata_matches:
                    raise WorkbookSafetyError(f"A célula '{sheet_name}!{coordinate}' alterou fórmula, estilo ou recurso associado.")
            elif coordinate not in allowed:
                raise WorkbookSafetyError(f"A planilha temporária criou a célula não autorizada '{sheet_name}!{coordinate}'.")
            elif not _cell_data_type_matches(actual, allowed[coordinate]):
                raise WorkbookSafetyError(f"A célula nova '{sheet_name}!{coordinate}' tem um tipo inesperado.")


def serialize_workbook_validated(
    source: Path,
    workbook: Any,
    original_snapshot: Dict[str, Any],
    expected_values: Dict[tuple, Any],
    *,
    allow_dimension_growth: bool = False,
    allow_threaded_comments: bool = False,
    allow_legacy_comments: bool = False,
    allow_custom_xml_passthrough: bool = False,
) -> bytes:
    """Save to a same-volume temporary, reopen, validate, and return its bytes."""
    source = Path(source).resolve(strict=False)
    original_entries = _package_hashes(source)

    def write(temp: Path) -> None:
        changed_values = _changed_expected_values(original_snapshot, expected_values)
        if changed_values:
            _write_package_with_cell_updates(
                source, temp, changed_values,
            )
        else:
            shutil.copyfile(source, temp)

    temp = _write_temp(
        source,
        write,
        lambda path: validate_workbook_roundtrip(
            original_snapshot, path, expected_values,
            original_package_entries=original_entries,
            allow_dimension_growth=allow_dimension_growth,
            allow_threaded_comments=allow_threaded_comments,
            allow_legacy_comments=allow_legacy_comments,
            allow_custom_xml_passthrough=allow_custom_xml_passthrough,
        ),
    )
    try:
        return temp.read_bytes()
    finally:
        temp.unlink(missing_ok=True)


def stage_workbook_update(
    transaction: AtomicFileTransaction,
    source: Path,
    update: Callable[[Any], Dict[tuple, Any]],
    *,
    allow_dimension_growth: bool = False,
) -> bool:
    """Edit a workbook into a same-volume transaction temp and verify before replacement."""
    source = Path(source).resolve(strict=False)
    if not source.is_file():
        return False
    workbook = load_workbook_safely(source)
    try:
        before = workbook_snapshot(workbook)
        original_parts = _preserved_package_parts(source)
        expected_values = update(workbook)
        if expected_values is None:
            return False

        def write(temp: Path) -> None:
            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter("always")
                workbook.save(temp)
            unsupported = [
                item for item in caught
                if re.search(r"unsupported|not supported|will be removed|extension|drawing", str(item.message), re.I)
            ]
            if unsupported:
                raise WorkbookSafetyError(
                    f"A planilha contém um recurso que não pode ser preservado: {unsupported[0].message}"
                )

        transaction.stage_file(
            source,
            write,
            lambda temp: validate_workbook_roundtrip(
                before, temp, expected_values, original_package_parts=original_parts,
                allow_dimension_growth=allow_dimension_growth
            ),
        )
    finally:
        workbook.close()
    return True


def _preserved_package_parts(source: Path) -> Dict[str, str]:
    """Fingerprint drawing, chart, and media parts for ordinary workbook saves."""
    prefixes = ("xl/drawings/", "xl/charts/", "xl/media/")
    result = {}
    try:
        with zipfile.ZipFile(source) as archive:
            for name in archive.namelist():
                if name.lower().startswith(prefixes):
                    result[name] = hashlib.sha256(archive.read(name)).hexdigest()
    except (OSError, zipfile.BadZipFile) as error:
        raise WorkbookSafetyError("Não foi possível conferir desenhos, gráficos e imagens da planilha.") from error
    return result


def _package_hashes(source: Path) -> Dict[str, str]:
    """Fingerprint every XLSX package entry for the download-only cell patch."""
    result = {}
    try:
        with zipfile.ZipFile(source) as archive:
            for info in archive.infolist():
                if info.filename in result:
                    raise WorkbookSafetyError("O arquivo Excel contém partes duplicadas no pacote.")
                result[info.filename] = hashlib.sha256(archive.read(info.filename)).hexdigest()
    except WorkbookSafetyError:
        raise
    except (OSError, zipfile.BadZipFile) as error:
        raise WorkbookSafetyError("Não foi possível conferir a estrutura interna da planilha.") from error
    return result
