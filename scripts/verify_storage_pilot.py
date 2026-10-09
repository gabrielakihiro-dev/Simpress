"""Exercise stage 6 with synthetic XLSX/JSON in a NEW disposable database.

Only this execution's database, role and individually journaled object keys are
removed. It refuses to adopt an existing database/login. No real files are read.
"""

import argparse
import io
import json
from pathlib import Path
import sys
import tempfile
import zipfile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import openpyxl
from fastapi.testclient import TestClient
from persistence.config import PersistenceSettings,PersistenceError
from persistence.postgres import MARKER,PostgresBackend,canonical
from persistence.backup import create,verify
from persistence.storage import LocalStorage,BUCKET
from scripts.run_postgres_pilot import checked_stack,admin_sql,create_login,drop_login
from scripts import supabase_local as local
import main as domain

DATABASE = "simpress_etapa6_validation"


def xlsx(headers, rows, *, native=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Equipamentos"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    meta = book.create_sheet("Metadados")
    meta["A1"] = "=1+1"
    output = io.BytesIO()
    book.save(output)
    book.close()
    if native:
        with zipfile.ZipFile(output,"a") as package:
            package.writestr("customXml/synthetic.xml",b"<test>preserve me</test>")
    return output.getvalue()


def inventory(*serials):
    headers = ["NumeroSerie","Apelido","Modelo","Tipo","Cidade","Endereco","UF","CEP","IdCliente","CNPJ Instalacao","CodigoContrato"]
    headers += ["CampoSintetico"+str(n) for n in range(domain.DETAILS_COMPLETE_SOURCE_COLUMN_COUNT-len(headers))]
    rows = [[s,"BASE", "Modelo sintético","Notebook","Cidade sintética","Rua sintética","PR","00000000","0001","00000000000000",str(domain.SUPPORTED_CONTRACT)]+[""]*(len(headers)-11) for s in serials]
    return {"detalhes":{"name":"EquipamentosDetalhes_sinteticos.xlsx","bytes":xlsx(headers,rows)},
            "endpoint":{"name":"InvAll_sintetico.xlsx","bytes":xlsx(["NumeroSerie","Nome do Computador"],[[s,"SYNTHETIC-HOST"] for s in serials])}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--local-disposable",action="store_true",required=True)
    parser.parse_args()
    cli,docker = checked_stack()
    status = json.loads(local.run([*cli,"status","--output","json","--agent","no"]))
    storage = LocalStorage(status["SERVICE_ROLE_KEY"])
    bucket_created = storage.not_found(storage.request("GET","bucket/"+BUCKET))
    created,login,backend = False,False,None
    assertions = 0
    def ok(condition,label):
        nonlocal assertions
        if not condition:
            raise AssertionError(label)
        assertions += 1
        print("PASS:",label,flush=True)
    def sql(statement):
        return admin_sql(docker,statement,database=DATABASE)
    def conflict(function,label,code=409):
        try:
            function()
            raise AssertionError(label)
        except PersistenceError as error:
            ok(error.status_code == code,label)
    try:
        admin_sql(docker,f"create database {DATABASE} template template0;")
        created = True
        sql("create schema auth; create table auth.users(id uuid primary key);")
        for path in sorted((ROOT/"supabase/migrations").glob("*.sql")):
            sql(path.read_text(encoding="utf-8"))
        sql("comment on schema simpress is '"+MARKER+"';")
        password = create_login(docker)
        login = True
        settings = PersistenceSettings("postgres",f"postgresql://simpress_piloto_login:{password}@127.0.0.1:54322/{DATABASE}",2,status["SERVICE_ROLE_KEY"])
        backend = PostgresBackend(settings,domain)
        backend.open()
        files = backend.files
        staged = inventory("SYNTHETIC-A")
        first = files.update(staged,"first-import")
        ok(len(backend.state()["equipamentos_cache"]) == 1,"validated source activates one batch")
        revision = backend.state()["revision"]
        ok(first == files.update(staged,"first-import") and backend.state()["revision"] == revision,"upload retry replays committed result without a new batch")
        conflict(lambda:files.update(inventory("SYNTHETIC-B"),"first-import"),"upload key cannot be reused for different bytes")
        with backend.transaction() as conn:
            source = conn.execute("select * from simpress.objetos_arquivo where chave like '%%/detalhes-%%'").fetchone()
            ok(storage.get(source) == staged["detalhes"]["bytes"],"original XLSX preserved byte for byte")
            ok(conn.execute("select count(*) as n from simpress.linhas_fonte where jsonb_array_length(celulas)>0").fetchone()["n"] == 2,"source row provenance and typed cells preserved")
        anonymous = storage.client.get("object/"+BUCKET+"/"+source["chave"],headers={"Authorization":"Bearer "+status["ANON_KEY"],"apikey":status["ANON_KEY"]})
        ok(anonymous.status_code != 200,"anonymous cannot download private object")
        public = storage.client.get("object/public/"+BUCKET+"/"+source["chave"])
        ok(public.status_code != 200,"public object URL cannot download private object")
        with patch.object(domain,"PERSISTENCE_SETTINGS",settings),patch.object(domain,"_POSTGRES_BACKEND",backend):
            client = TestClient(domain.app,base_url="http://127.0.0.1:8000",client=("127.0.0.1",1234))
            try:
                history = client.get("/api/historico-comparacoes")
                ok(history.status_code == 200 and len(history.json()["eventos"]) == 1,"HTTP reads SQL comparison history")
                event_id = "http-annotation-fixture"
                files.annotations("history_fixture.json",canonical({"historico":[{"id":event_id,
                    "horario":first["comparacao"]["horario_comparacao"],
                    "comparacao":{"sem_acesso_14_dias":[{"serie":"SYNTHETIC-A","motivo":"Fixture sintética"}]}}]}).encode(),"http-fixture")
                view = client.get("/api/historico-comparacoes/"+event_id).json()["evento"]
                category = "sem_acesso_14_dias"
                row = view["comparacao"][category][0]
                url = "/api/historico-comparacoes/"+event_id+"/anotacoes/"+row["anotacao_id"]
                ok(client.put(url,json={"anotacao":""}).status_code == 200,"empty note accepted over HTTP")
                ok(client.put(url,json={"anotacao":"a"*2000}).status_code == 200,"2000-character note accepted over HTTP")
                ok(client.put(url,json={"anotacao":"a"*2001}).status_code == 422,"2001-character note rejected over HTTP")
                ok(client.put(url+"other",json={"anotacao":"x"}).status_code == 404,"cross-row note reference rejected")
                download = client.get("/api/arquivos/"+str(source["id"]))
                ok(download.status_code == 200 and download.content == staged["detalhes"]["bytes"],"backend download preserves original bytes")
                exported = client.get("/api/historico-comparacoes/exportar/excel",params={"data":view["horario"][:10]})
                ok(exported.status_code == 200 and b"PK" == exported.content[:2],"existing history Excel export works with SQL notes")
                model = xlsx(["NumeroSerie","Apelido","IdCliente"],[["SYNTHETIC-A","BASE","0001"]],native=True)
                response = client.post("/api/piloto/modelos/exportacao",files={"upload":("Modelo_sintetico.xlsx",model)})
                ok(response.status_code == 200,"versioned private export template uploaded through backend")
                item = backend.state()["equipamentos_cache"]["SYNTHETIC-A"]
                edited = client.post("/api/equipamento/salvar",json={"serial":"SYNTHETIC-A","modelo":"EDITADO","idCliente":"0009","revisao":item["_revisao"],"novo":False})
                ok(edited.status_code == 200,"logical equipment edits committed before backup")
                ok(client.post("/api/equipamento/atualizar-apelido",json={"serial":"SYNTHETIC-A","apelido":"MANUAL"}).status_code == 200,"manual alias committed before backup")
                exported = client.get("/api/exportar-equipamentos")
                ok(exported.status_code == 200,"equipment export uses active SQL values")
                with zipfile.ZipFile(io.BytesIO(exported.content)) as package:
                    ok(package.read("customXml/synthetic.xml") == b"<test>preserve me</test>","native template XML survives export")
                workbook = openpyxl.load_workbook(io.BytesIO(exported.content))
                ok(workbook.active["B2"].value == "MANUAL" and workbook.active["C2"].value == "0009" and workbook["Metadados"]["A1"].value == "=1+1","exported alias/client values and unrelated formula preserved")
                workbook.close()
                sys.path.insert(0,str(ROOT/"tests"))
                from test_integrity_workflows import make_transfer_template,make_transfer_payload
                with tempfile.TemporaryDirectory(prefix="simpress-transfer-fixture-") as fixture_root:
                    template_path = make_transfer_template(Path(fixture_root)/"transfer.xlsx")
                    transfer_content = template_path.read_bytes()
                response = client.post("/api/piloto/modelos/transferencia",files={"upload":("Transferencia_sintetica.xlsx",transfer_content)})
                ok(response.status_code == 200,"private transfer template registered without reading operational template")
                response = client.post("/api/gerar-planilha",json=make_transfer_payload(serial="SYNTHETIC-A").model_dump())
                ok(response.status_code == 200,"transfer download generated with existing form contract")
                with zipfile.ZipFile(io.BytesIO(response.content)) as package,zipfile.ZipFile(io.BytesIO(transfer_content)) as original:
                    part = "xl/threadedComments/threadedComment1.xml"
                    ok(package.read(part) == original.read(part),"native threaded comments preserved in transfer download")
                second = files.update(inventory("SYNTHETIC-B"),"second-import")
                eid = second["comparacao"]["evento_id"]
                context = files.restore_context(eid)
                backend.execute("unit",{"id":None,"cidade":"Outra cidade","endereco":"Rua","uf":"PR","cep":"00000-000","cnpj":"00000000000000"},"invalidate-preview")
                conflict(lambda:files.restore(eid,context["preview_token"],"stale-restore"),"stale restore preview rejected")
                context = files.restore_context(eid)
                response = client.post("/api/historico-comparacoes/"+eid+"/restaurar",
                    json={"preview_token":context["preview_token"]},headers={"Idempotency-Key":"restore-logical"})
                ok(response.status_code == 200,"logical restoration uses existing HTTP contract")
                restored = response.json()
                ok(restored["evento"]["total_equipamentos_carregados"] == 1 and restored["evento"]["desfazer_disponivel"],"restore response includes loaded count and undo availability")
                current = backend.state()["equipamentos_cache"]["SYNTHETIC-A"]
                ok(current["modelo"] == "EDITADO" and current["idCliente"] == "0009" and current["apelido"] == "MANUAL","restore recovers logical edits and aliases beyond original XLSX")
                ok(not any(u["cidade"] == "Outra cidade" for u in backend.state()["unidades_cache"]),"restore recovers catalog visibility from its logical snapshot")
                ok(restored == files.restore(eid,context["preview_token"],"restore-logical"),"restore retry replays result after global revision changes")
            finally:
                client.close()
        before = backend.state()["revision"]
        with patch.object(files,"activate_upload",side_effect=RuntimeError("synthetic precommit failure")):
            try:
                files.update(inventory("SYNTHETIC-C"),"interrupted")
                raise AssertionError("failure injection ignored")
            except RuntimeError:
                pass
        ok(backend.state()["revision"] == before and "SYNTHETIC-A" in backend.state()["equipamentos_cache"],"failure after upload leaves active logical state unchanged")
        with backend.transaction() as conn:
            op = conn.execute("select id from simpress.operacoes where chave_idempotencia='files_update:interrupted'").fetchone()["id"]
        recovered = files.recover(op)
        ok(recovered["sucesso"] and "SYNTHETIC-C" in backend.state()["equipamentos_cache"],"journaled immutable uploads recover without original request files")
        conflict(lambda:files.cancel(op),"confirmed/referenced objects cannot be deleted")
        # Exact carryover with numeric distinctions and latest blank/conflict rules.
        for number,note,label in ((1,"inteiro","int"),(1.0,"decimal","float")):
            event = {"id":"legacy-"+label,"horario":"invalid-legacy-date","comparacao":{"sem_acesso_14_dias":[{"serie":"X","value":number,"anotacao":note}]}}
            files.annotations(label+".json",canonical({"historico":[event]}).encode(),label)
        with backend.transaction(write=True) as conn:
            event = {"id":"new-exact","comparacao":{"sem_acesso_14_dias":[{"serie":"X","value":1}]}}
            files.history.insert(conn,event,reapply=True)
            notes = conn.execute("select texto from simpress.anotacoes where evento_id='new-exact'").fetchall()
            ok([n["texto"] for n in notes] == ["inteiro"],"annotation carryover distinguishes canonical 1 from 1.0")
            files.history.insert(conn,{"id":"new-blank","comparacao":{"sem_acesso_14_dias":[{"serie":"X","value":1,"anotacao":""}]}})
            files.history.insert(conn,{"id":"after-blank","comparacao":{"sem_acesso_14_dias":[{"serie":"X","value":1}]}},reapply=True)
            ok(not conn.execute("select * from simpress.anotacoes where evento_id='after-blank'").fetchall(),"latest empty annotation blocks resurrection of older note")
            files.history.insert(conn,{"id":"duplicates","comparacao":{"sem_acesso_14_dias":[{"serie":"Y","anotacao":"one"},{"serie":"Y","anotacao":"two"}]}})
            files.history.insert(conn,{"id":"after-conflict","comparacao":{"sem_acesso_14_dias":[{"serie":"Y"}]}},reapply=True)
            ok(not conn.execute("select * from simpress.anotacoes where evento_id='after-conflict'").fetchall(),"conflicting duplicate notes do not carry over")
        partial = {"evento_id":"external-legacy","item_id":"row-stable","categoria":"sem_acesso_14_dias","anotacao":"legada","horario":"invalid-date","atualizado_em":"invalid-date"}
        files.annotations("partial.json",canonical({"anotacoes":[partial]}).encode(),"partial")
        view = next(e for e in files.history.load() if e["id"] == "external-legacy")
        ok(view["comparacao"]["sem_acesso_14_dias"][0]["anotacao"] == "legada","annotation accessible without local event history; invalid date preserved")
        files.history.mutate("note",{"event":"external-legacy","item":"row-stable","note":"editada"},"edit-partial")
        conflicting = files.annotations("old.json",canonical({"anotacoes":[{**partial,"anotacao":"antiga"}]}).encode(),"conflicting-legacy")
        ok(conflicting["conflitos"] == 1,"legacy import preserves newer annotation and journals conflict")
        files.annotations("deleted.json",canonical({"historico":[{"id":"deletion-fixture","comparacao":{
            "sem_acesso_14_dias":[{"serie":"DELETEONLY","anotacao":"original"}]}}]}).encode(),"deleted-fixture")
        deleted_view = next(e for e in files.history.load() if e["id"] == "deletion-fixture")
        deleted_item = deleted_view["comparacao"]["sem_acesso_14_dias"][0]["anotacao_id"]
        files.history.mutate("delete",{"event":"deletion-fixture"},"delete-fixture")
        deleted_import = files.annotations("deleted-retry.json",canonical({"anotacoes":[{
            "evento_id":"deletion-fixture","item_id":deleted_item,"categoria":"sem_acesso_14_dias",
            "anotacao":"ressuscitada","atualizado_em":"2100-01-01T00:00:00+00:00"}]}).encode(),"deleted-retry")
        ok(deleted_import["conflitos"] == 1 and not any(e["id"] == "deletion-fixture" for e in files.history.load()),
           "legacy reimport cannot resurrect annotations on a deleted event")
        files.annotations("retention.json",canonical({"historico":[{"id":"expired-empty","horario":"2000-01-01T00:00:00+00:00","comparacao":{}},
            {"id":"expired-noted","horario":"2000-01-01T00:00:00+00:00","comparacao":{"sem_acesso_14_dias":[{"serie":"OLD","anotacao":"protegida"}]}},
            {"id":"invalid-date","horario":"invalid-date","comparacao":{}}]}).encode(),"retention-fixture")
        ok(files.prune()["ocultados"] == 1 and any(e["id"] == "expired-noted" for e in files.history.load())
           and any(e["id"] == "invalid-date" for e in files.history.load()),"60-day visibility retention protects notes and invalid dates without deleting Storage")
        with patch.object(files,"activate_upload",side_effect=RuntimeError("synthetic orphan")):
            try:
                files.update(inventory("ORPHAN"),"orphan")
            except RuntimeError:
                pass
        with backend.transaction() as conn:
            orphan = conn.execute("select id from simpress.operacoes where chave_idempotencia='files_update:orphan'").fetchone()["id"]
        actual_delete = files.storage.delete
        def lost_delete_reply(obj):
            actual_delete(obj)
            raise PersistenceError(503,"synthetic lost DELETE reply")
        with patch.object(files.storage,"delete",side_effect=lost_delete_reply):
            conflict(lambda:files.cancel(orphan),"lost delete reply leaves durable deletion intent",503)
        ok(files.cancel(orphan)["sucesso"],"delete retry safely handles already absent orphan")
        actual_finish = files.finish
        committed_input = inventory("COMMITTED")
        def lost_commit_reply(*args):
            actual_finish(*args)
            raise PersistenceError(503,"synthetic lost COMMIT reply")
        with patch.object(files,"finish",side_effect=lost_commit_reply):
            conflict(lambda:files.update(committed_input,"lost-commit"),"lost COMMIT reply is reported without reverting the committed batch",503)
        committed_revision = backend.state()["revision"]
        replay = files.update(committed_input,"lost-commit")
        ok(replay["sucesso"] and backend.state()["revision"] == committed_revision,"retry after lost COMMIT reply does not create another batch")
        files.update(inventory("DUPLICATE","DUPLICATE"),"duplicate-source")
        with backend.transaction() as conn:
            count = conn.execute("select count(*) as n from simpress.equipamentos_lote where lote_id=(select lote_ativo_id from simpress.estado_aplicacao where id=1)").fetchone()["n"]
        ok(count == 2,"duplicate official serial rows are preserved without deduplication")
        with tempfile.TemporaryDirectory(prefix="simpress-backup-fixture-") as root:
            migration_root = Path(root)/"supabase/migrations"
            migration_root.mkdir(parents=True)
            for migration in (ROOT/"supabase/migrations").glob("*.sql"):
                (migration_root/migration.name).write_bytes(migration.read_bytes())
            path = create(files,Path(root))
            manifest = verify(path)
            ok(manifest["revision"] == backend.state()["revision"] and not manifest["auth_included"],"domain/Storage archive verified at a consistent revision without credentials")
            ok(len(manifest["migrations"]) == len(list((ROOT/"supabase/migrations").glob("*.sql"))),"backup includes all versioned migrations")
            with zipfile.ZipFile(path,"a") as archive:
                archive.writestr("unlisted.txt","synthetic tamper")
            conflict(lambda:verify(path),"tampered backup rejected")
        ok("SERVICE_ROLE_KEY" not in canonical(recovered) and status["SERVICE_ROLE_KEY"] not in canonical(recovered),"responses do not expose Storage credentials")
        print(f"Validated {assertions} stage 6 assertions.",flush=True)
    finally:
        if backend:
            backend.close()
        if created:
            descriptors = json.loads(sql("select coalesce(json_agg(o),'[]') from simpress.objetos_arquivo o;"))
            for obj in descriptors:
                storage.delete(obj)
            admin_sql(docker,f"drop database {DATABASE} with (force);")
        if login:
            drop_login(docker)
        if bucket_created:
            response = storage.request("DELETE","bucket/"+BUCKET)
            if response.status_code not in (200,204) and not storage.not_found(response):
                raise RuntimeError("Synthetic bucket cleanup incomplete")
        storage.close()


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        message = str(error) if isinstance(error,(AssertionError,PersistenceError)) else type(error).__name__
        print("Stage 6 validation failed:",message,file=sys.stderr)
        sys.exit(1)
